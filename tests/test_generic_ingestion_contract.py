"""Generic ingestion contracts built from synthetic workbooks only."""

from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from backend.routers import pipeline
from core.entity_engine import import_file_to_entities
from core.unified_integration import combine_entity_sources, export_unified_workbook


def _write_sheet(path: Path, sheet_name: str, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_builtin_template_is_layout_only_and_never_seeds_employee_data(monkeypatch) -> None:
    called = False

    def fake_import(_path: str):
        nonlocal called
        called = True
        return {"employee_profile": [{"工号": "OLD", "姓名": "旧模板人员"}]}

    monkeypatch.setattr(pipeline, "import_template_to_entities", fake_import)

    entities, logs = pipeline._load_template_base_entities(None, "builtin.xlsx")

    assert called is False
    assert entities == {}
    assert any("仅作为输出版式" in log for log in logs)


def test_integration_rejects_sources_when_no_person_can_be_recognized() -> None:
    with pytest.raises(HTTPException) as exc_info:
        pipeline._ensure_recognized_people({"source_person_count": 0, "output_person_count": 0})

    assert exc_info.value.status_code == 400
    assert "未识别到有效人员" in str(exc_info.value.detail)


def test_roster_confirmation_depends_on_semantic_content_not_filename() -> None:
    assert pipeline._record_confirms_current_roster(
        "attendance_record",
        {"工号": "X-1", "姓名": "任意人员", "实际出勤天数": 20},
    )
    assert pipeline._record_confirms_current_roster(
        "social_security",
        {"身份证号": "110101199001011234", "姓名": "任意人员", "个人养老": 800},
    )
    assert not pipeline._record_confirms_current_roster(
        "employee_profile",
        {"工号": "X-1", "姓名": "历史人员", "公司": "旧公司"},
    )


def test_common_identity_aliases_apply_to_every_business_entity(tmp_path: Path) -> None:
    source = tmp_path / "unknown-tax-source.xlsx"
    _write_sheet(source, "unknown-sheet", [
        ["应扣税额", "员工编号", "累计应纳税所得额", "人员姓名"],
        [678.9, "G-9001", 99000, "通用验证员"],
    ])

    parsed = import_file_to_entities(str(source), source.name)

    assert parsed["tax_record"][0]["工号"] == "G-9001"
    assert parsed["tax_record"][0]["姓名"] == "通用验证员"
    assert parsed["tax_record"][0]["本次应扣税额"] == pytest.approx(678.9)


def test_generic_effective_date_is_preserved_for_time_based_resolution(tmp_path: Path) -> None:
    source = tmp_path / "department-change.xlsx"
    _write_sheet(source, "员工变化", [
        ["员工编号", "人员姓名", "所属部门", "生效日期"],
        ["G-9001", "通用验证员", "新部门", "2026-07-20"],
    ])

    parsed = import_file_to_entities(str(source), source.name)

    assert parsed["employee_profile"][0]["_effective_date"] == "2026-07-20"


def test_same_name_cannot_bridge_different_strong_identities() -> None:
    prepared = combine_entity_sources({}, {
        "employee_profile": [
            {"工号": "G-9001", "姓名": "通用验证员", "公司": "虚构通用公司"},
            {"身份证号": "110101199001011234", "姓名": "通用验证员", "公司": "虚构通用公司"},
        ],
        "salary_detail": [
            {"工号": "G-9001", "姓名": "通用验证员", "月基本薪资": 18888.88},
        ],
        "social_security": [
            {"身份证号": "110101199001011234", "姓名": "通用验证员", "个人养老": 900},
        ],
    })

    assert prepared["source_person_count"] == 2
    assert len(prepared["entities"]["employee_profile"]) == 2
    assert prepared["entities"]["salary_detail"][0]["_person_key"] != prepared["entities"]["social_security"][0]["_person_key"]
    assert any(issue["issue_type"] == "ambiguous_person" for issue in prepared["issues"])


def test_builtin_layout_source_sheets_are_cleared_without_removing_layout(tmp_path: Path) -> None:
    source = tmp_path / "builtin-copy.xlsx"
    workbook = Workbook()
    payroll = workbook.active
    payroll.title = "工资核算"
    payroll["A1"] = "模板标题"
    attendance = workbook.create_sheet("考勤")
    attendance.append(["工号", "姓名", "出勤天数"])
    attendance.append(["OLD-001", "旧模板人员", 21.75])
    tax = workbook.create_sheet("个税导出核对")
    tax.append(["工号", "姓名", "应扣税额"])
    tax.append(["OLD-001", "旧模板人员", 999.99])
    workbook.save(source)
    workbook.close()

    cleared = pipeline._clear_builtin_layout_source_data(str(source))

    assert set(cleared) == {"考勤", "个税导出核对"}
    result = load_workbook(source, data_only=False)
    assert result["工资核算"]["A1"].value == "模板标题"
    assert "考勤" in result.sheetnames
    assert "个税导出核对" in result.sheetnames
    assert result["考勤"]["A2"].value is None
    assert result["个税导出核对"]["B2"].value is None
    result.close()


def test_arbitrary_filenames_sheet_names_headers_and_column_order_export_to_template(tmp_path: Path) -> None:
    roster = tmp_path / "alpha-001.xlsx"
    attendance = tmp_path / "beta-002.xlsx"
    tax = tmp_path / "gamma-003.xlsx"
    output = tmp_path / "result.xlsx"

    _write_sheet(roster, "任意数据页", [
        ["这是一行无关标题"],
        [None], [None], [None], [None], [None], [None], [None], [None], [None], [None],
        ["银行卡号", "基本工资", "所属公司", "人员姓名", "员工编号"],
        ["622200001", 12345.67, "新公司甲", "测试甲", "N-001"],
    ])
    _write_sheet(attendance, "A区明细", [
        ["加班工资", "出勤天数", "员工编号", "人员姓名"],
        [321.5, 21.5, "N-001", "测试甲"],
    ])
    _write_sheet(tax, "B区明细", [
        ["累计应纳税所得额", "应扣税额", "姓名", "工号"],
        [80000, 456.78, "测试甲", "N-001"],
    ])

    entities: dict[str, list[dict]] = {}
    for path in (roster, attendance, tax):
        parsed = import_file_to_entities(str(path), path.name)
        for entity_id, records in parsed.items():
            if entity_id != "_log":
                entities.setdefault(entity_id, []).extend(records)

    prepared = combine_entity_sources({}, entities)
    assert prepared["source_person_count"] == 1
    export_unified_workbook(
        pipeline.BUILTIN_TEMPLATE_PATH,
        prepared["entities"],
        str(output),
    )

    workbook = load_workbook(output, data_only=False)
    payroll = workbook["工资核算"]
    assert payroll["A4"].value == "N-001"
    assert payroll["B4"].value == "测试甲"
    assert payroll["D4"].value == "新公司甲"
    assert payroll["Z4"].value == "=ROUND(X4*Y4,-1)"
    assert isinstance(payroll["AI4"].value, str) and payroll["AI4"].value.startswith("=")
    assert isinstance(payroll["AT4"].value, str) and payroll["AT4"].value.startswith("=")
    assert isinstance(payroll["BQ4"].value, str) and payroll["BQ4"].value.startswith("=")
    assert isinstance(payroll["CB4"].value, str) and payroll["CB4"].value.startswith("=")
    assert isinstance(payroll["BA4"].value, str) and payroll["BA4"].value.startswith("=")
    workbook.close()


def test_arbitrary_social_security_sheet_exports_without_supplier_specific_names(tmp_path: Path) -> None:
    source = tmp_path / "vendor-neutral-77.xlsx"
    output = tmp_path / "social-result.xlsx"
    _write_sheet(source, "随机页签", [
        ["说明：列顺序由供应商自行决定"],
        ["个人住房12%", "证件号码", "单位养老16%", "人员姓名", "个人养老8%", "单位名称"],
        [360, "110101199001011234", 1600, "通用测试乙", 800, "新公司乙"],
    ])

    parsed = import_file_to_entities(str(source), source.name)
    assert parsed["social_security"][0]["身份证号"] == "110101199001011234"
    prepared = combine_entity_sources({}, {
        key: value for key, value in parsed.items() if key != "_log"
    })
    export_unified_workbook(
        pipeline.BUILTIN_TEMPLATE_PATH,
        prepared["entities"],
        str(output),
    )

    workbook = load_workbook(output, data_only=False)
    payroll = workbook["工资核算"]
    assert payroll["B4"].value == "通用测试乙"
    assert isinstance(payroll["BB4"].value, str) and payroll["BB4"].value.startswith("=")
    assert isinstance(payroll["BD4"].value, str) and payroll["BD4"].value.startswith("=")
    assert isinstance(payroll["BF4"].value, str) and payroll["BF4"].value.startswith("=")
    assert isinstance(payroll["BI4"].value, str) and payroll["BI4"].value.startswith("=")
    workbook.close()
