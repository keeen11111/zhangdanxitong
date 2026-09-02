from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest


def _source_paths(tmp_path: Path) -> dict[str, str]:
    names = [
        "薪资数据-科园-7月薪资（8.14发薪）.xlsx",
        "益药科园2026.07.xlsx",
        "鹤安长泰2026.07.xlsx",
        "202608_税款计算_工资薪金所得-益药科园.xls",
        "202608_税款计算_劳务报酬-益药科园.xls",
        "202608_税款计算_工资薪金所得-鹤安长泰.xls",
        "202608_税款计算_劳务报酬-鹤安长泰.xls",
    ]
    values: dict[str, str] = {}
    for index, name in enumerate(names):
        path = tmp_path / f"normalized-{index}.xlsx"
        path.write_bytes(b"PK\x03\x04")
        values[name] = str(path)
    return values


def test_detect_keyuan_batch_requires_the_complete_source_set(tmp_path: Path) -> None:
    from backend.keyuan_workflow import detect_keyuan_batch

    sources = _source_paths(tmp_path)
    batch = detect_keyuan_batch(
        "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        sources,
    )

    assert batch is not None
    assert batch.payroll_period == "2026.07"
    assert batch.payment_period == "2026.08"
    assert batch.output_filename == "202608（所属月202607）-北京科园-鹤安-大药房工资核算总表-v2.xlsx"
    assert batch.salary_source.name.startswith("normalized-")
    assert len(batch.social_sources) == 2
    assert len(batch.tax_sources) == 4

    sources.pop("202608_税款计算_劳务报酬-鹤安长泰.xls")
    assert detect_keyuan_batch(
        "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        sources,
    ) is None


def test_detect_keyuan_batch_does_not_claim_other_months_or_templates(tmp_path: Path) -> None:
    from backend.keyuan_workflow import detect_keyuan_batch

    sources = _source_paths(tmp_path)
    assert detect_keyuan_batch("通用工资总表.xlsx", sources) is None
    sources["202609_税款计算_工资薪金所得-益药科园.xls"] = sources.pop(
        "202608_税款计算_工资薪金所得-益药科园.xls"
    )
    assert detect_keyuan_batch(
        "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        sources,
    ) is None


def test_keyuan_batch_requires_the_project_to_use_the_payroll_period(tmp_path: Path) -> None:
    from backend.keyuan_workflow import detect_keyuan_batch

    batch = detect_keyuan_batch(
        "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        _source_paths(tmp_path),
    )

    assert batch is not None
    assert batch.matches_project_month("2026.07") is True
    assert batch.matches_project_month("2026.08") is False
    assert batch.matches_project_month("2026.09") is False


def test_complete_keyuan_batch_is_not_diverted_to_model_by_stale_project_month(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.agent_execution as execution
    from core.document_agent.orchestrator import ToolRegistry

    sources = _source_paths(tmp_path)
    master = tmp_path / "master.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook.save(master)
    workbook.close()
    output = tmp_path / "draft.xlsx"

    class FakeBatch:
        payroll_period = "2026.07"

        def matches_project_month(self, _month: str) -> bool:
            return False

    def fake_execute(*_args: object, **kwargs: object) -> dict[str, object]:
        path = Path(kwargs["output_path"])
        book = openpyxl.Workbook()
        book.active.title = "工资核算"
        book["工资核算"].insert_cols(1)
        book["工资核算"]["A2"] = "班制"
        for name in ("台账", "个税导出核对", "OA请款及审批"):
            book.create_sheet(name)
        book.save(path)
        book.close()
        return {"status": "passed", "change_count": 0, "changes": [], "issues": []}

    monkeypatch.setattr(execution, "execute_keyuan_batch", fake_execute)
    monkeypatch.setattr(execution.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(execution.ModelConfig, "fallback_from_env", lambda: None)
    monkeypatch.setattr(execution, "DATA_DIR", tmp_path)
    run = {
        "run_id": "x" * 32,
        "master_file": "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        "salary_month": "2026.09",
        "_master_path": str(master),
        "_source_paths": sources,
        "workbook_updates": [],
        "messages": [],
        "model_plan": {"steps": []},
    }
    monkeypatch.setattr(execution, "detect_keyuan_batch", lambda *_args, **_kwargs: FakeBatch())
    execution.execute_model_plan(
        run, draft=output, registry=ToolRegistry(), read_schemas=[], materials=[], rules={},
        save=lambda _run: None, emit=lambda *_args, **_kwargs: None,
    )
    assert run["status"] == "completed"


def test_personnel_change_writer_adds_new_rows_without_duplicates() -> None:
    from scripts.keyuan_python_executor import _apply_personnel_changes

    target_book = openpyxl.Workbook()
    target = target_book.active
    target.title = "人员异动"
    target["A79"] = "内部调动:"
    target["A80"] = "序号"
    target["B80"] = "姓 名"
    target["C80"] = "员工编号"
    target["D80"] = "调出部门"
    target["F80"] = "原职务"
    target["G80"] = "调入部门"
    target["I80"] = "现职务"
    target["J80"] = "生效日期"
    target["A81"] = 1
    target["B81"] = "旧员工"
    target["C81"] = "001"
    target["A100"] = "跨公司调动："
    target["A101"] = "序号"
    target["B101"] = "姓 名"
    target["C101"] = "员工编号"
    target["D101"] = "调入/调出"
    target["E101"] = "调出公司/部门"
    target["G101"] = "调入公司/部门"
    target["H101"] = "现职务"
    target["I101"] = "生效日期"
    target["A102"] = 1
    target["B102"] = "旧跨公司"
    target["C102"] = "002"
    target["A150"] = "劳动合同变更："
    target["A151"] = "序号"
    target["B151"] = "姓 名"
    target["A152"] = None

    source_book = openpyxl.Workbook()
    source = source_book.active
    source.title = "入离职、转岗、转正、其他"
    source.append(["内部调动"])
    source.append(["序号", "姓 名", "员工编号", "调出部门", "原组别", "原职务", "调入部门", "现组别", "现职务", "生效日期"])
    source.append([1, "新员工", None, "原部门", None, "原岗位", "新部门", None, "新岗位", "2026-07-01"])
    source.append([None] * 10)
    source.append(["夸公司调动"])
    source.append(["姓名", "新工号", "原公司", "新公司", "新成本中心/部门", "入职时间", "转岗时间", "岗位"])
    source.append(["新跨公司", "003", "旧公司", "新公司", "新部门", None, "2026-06-01", "新岗位"])
    source.append([None] * 8)
    source.append(["其他"])
    source.append(["姓名", "备注"])
    source.append(["新备注", "产假回岗"])

    changes: list[dict[str, object]] = []
    issues: list[dict[str, str]] = []
    _apply_personnel_changes(target, source, changes, issues, "salary.xlsx")

    assert target["B82"].value == "新员工"
    assert target["C82"].value is None
    assert target["I82"].value == "新岗位"
    assert target["B103"].value == "新跨公司"
    assert target["D103"].value == "旧公司"
    assert target["E103"].value == "新部门"
    assert target["G103"].value == "新公司"
    assert target["B148"].value == "新备注"
    assert not issues

    second_changes: list[dict[str, object]] = []
    _apply_personnel_changes(target, source, second_changes, issues, "salary.xlsx")
    assert not second_changes


def test_execute_keyuan_batch_stages_normalized_inputs_and_returns_audited_output(
    tmp_path: Path,
) -> None:
    from backend.keyuan_workflow import detect_keyuan_batch, execute_keyuan_batch
    import scripts.keyuan_python_executor as executor

    sources = _source_paths(tmp_path)
    master = tmp_path / "normalized-master.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook.save(master)
    workbook.close()
    batch = detect_keyuan_batch(
        "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        sources,
    )
    assert batch is not None
    output = tmp_path / "result" / "draft.xlsx"

    def fake_process(master_path: Path, salary_path: Path, social_paths: object, tax_paths: object, output_path: Path) -> dict[str, object]:
        sample_dir = master_path.parent
        assert (sample_dir / batch.master_filename).is_file()
        assert (sample_dir / batch.salary_source_name).is_file()
        assert all((sample_dir / name).is_file() for name, _path in batch.social_sources)
        tax_dir = sample_dir / "回复：202608科园、鹤安应发工资的全部附件20260811"
        assert all((tax_dir / name).is_file() for name, _path in batch.tax_sources)
        workbook = openpyxl.load_workbook(master_path)
        workbook["工资核算"].insert_cols(1)
        workbook["工资核算"]["A2"] = "班制"
        workbook["工资核算"].cell(1, 99).value = "末列"
        workbook.create_sheet("台账")
        workbook.create_sheet("个税导出核对")
        workbook.create_sheet("OA请款及审批")
        workbook.save(output_path)
        workbook.close()
        return {"status": "passed", "changes": [{"sheet": "工资核算", "cell": "A2", "before": None, "after": "班制"}], "issues": []}

    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(executor, "process_keyuan_python", fake_process)

    result = execute_keyuan_batch(batch, master_path=master, source_paths=sources, output_path=output, work_root=tmp_path / "work")
    monkeypatch.undo()

    assert result["status"] == "passed"
    assert result["change_count"] == 1
    assert result["download_filename"] == batch.output_filename
    assert result["structural"] == {
        "sheet_names": True,
        "payroll_columns": True,
        "payroll_header": True,
        "formula_references": True,
    }
    assert output.is_file()
    assert master.is_file()


def test_complete_keyuan_run_uses_fixed_executor_without_model(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.agent_execution as execution
    from core.document_agent.orchestrator import ToolRegistry

    output = tmp_path / "Agent草稿.xlsx"
    captured: dict[str, object] = {}

    def fake_executor(batch: object, **kwargs: object) -> dict[str, object]:
        captured["batch"] = batch
        output_path = Path(kwargs["output_path"])
        workbook = openpyxl.Workbook()
        workbook.active.title = "工资核算"
        workbook["工资核算"]["A2"] = "班制"
        workbook["工资核算"].cell(1, 99).value = "末列"
        for name in ("台账", "个税导出核对", "OA请款及审批"):
            workbook.create_sheet(name)
        workbook.save(output_path)
        workbook.close()
        return {"status": "passed", "change_count": 2, "changes": [{"sheet": "工资核算", "cell": "A2"}], "issues": [], "structural": {}}

    monkeypatch.setattr(execution, "execute_keyuan_batch", fake_executor)
    monkeypatch.setattr(execution.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(execution.ModelConfig, "fallback_from_env", lambda: None)
    monkeypatch.setattr(execution, "DATA_DIR", tmp_path)

    source_paths = _source_paths(tmp_path)
    master = tmp_path / "master.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook.save(master)
    workbook.close()
    run = {
        "run_id": "r" * 32,
        "master_file": "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        "salary_month": "2026.07",
        "_master_path": str(master),
        "_source_paths": source_paths,
        "workbook_updates": [],
        "messages": [],
        "model_plan": {"steps": []},
    }

    execution.execute_model_plan(
        run,
        draft=output,
        registry=ToolRegistry(),
        read_schemas=[],
        materials=[],
        rules={},
        save=lambda _run: None,
        emit=lambda *_args, **_kwargs: None,
    )

    assert captured["batch"] is not None
    assert run["keyuan_workflow"]["status"] == "passed"
    assert run["status"] == "completed"
    assert run["validation"]["status"] == "passed"
    assert output.is_file()


def test_background_complete_keyuan_batch_skips_model_plan_generation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A complete batch starts the fixed executor even when no model is configured."""
    import backend.routers.agent as agent

    run = {
        "run_id": "z" * 32,
        "project_id": "project-1",
        "tenant_id": "tenant-a",
        "status": "planning",
        "salary_month": "2026.07",
        "master_file": "202607（所属月202606)-北京科园-鹤安-大药房工资核算总表-v2.xlsx",
        "_source_paths": {},
        "events": [],
        "messages": [],
        "workbook_updates": [],
        "plan_confirmation": {"required": True, "confirmed": False},
        "month_confirmation": {"required": False, "confirmed": False},
    }

    class FakeBatch:
        payroll_period = "2026.07"

        def matches_project_month(self, _month: str) -> bool:
            return True

    execute_calls: list[str] = []
    monkeypatch.setattr(agent, "_load_run", lambda *_args, **_kwargs: run)
    monkeypatch.setattr(agent, "_save_run", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(agent, "_material_context", lambda *_args, **_kwargs: [])
    monkeypatch.setattr(agent, "find_basic_salary_source", lambda _paths: None)
    monkeypatch.setattr(agent, "detect_keyuan_batch", lambda *_args, **_kwargs: FakeBatch())
    monkeypatch.setattr(agent, "SessionLocal", lambda: SimpleNamespace(close=lambda: None))
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(agent.ModelConfig, "fallback_from_env", lambda: None)
    monkeypatch.setattr(agent, "confirm_agent_plan", lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("model plan must be skipped")))

    def fake_execute(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        execute_calls.append("fixed")
        run["execution_result"] = {"status": "completed", "code": None, "content": "done"}
        run["status"] = "completed"
        return run

    monkeypatch.setattr(agent, "execute_agent_run", fake_execute)
    monkeypatch.setattr(agent, "_finalize_agent_output", lambda _run: None)

    agent._run_agent_workflow(run["run_id"], run["tenant_id"])

    assert execute_calls == ["fixed"]
    assert run["model_plan"]["model"] == "fixed-python-keyuan"
    assert run["plan_confirmation"] == {"required": True, "confirmed": True}
