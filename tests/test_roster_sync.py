from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
import pytest


DETAIL_HEADERS = [
    "序号", "姓名", "薪资所属月", "社保所属月", "公积金所属月", "单位社保合计",
    "残保金企业", "个人社保合计", "基本公积金企业", "基本公积金个人",
    "基本工资", "岗位津贴", "企龄津贴", "差旅津贴", "绩效基数", "实际绩效",
    "绩效调整", "排班绩效", "加班费", "翻班津贴", "餐费补贴", "奖金",
    "节日福利", "高温费", "税前调整", "病假扣款", "事假扣款", "最低薪资补偿",
    "旷工扣款", "其他扣款", "缺勤扣款", "独生子女", "应发工资", "税后调整",
    "经济补偿金", "应税工资", "个调税", "实发工资", "管理费", "代收代付", "合计",
]


def _create_master(path: Path) -> None:
    workbook = openpyxl.Workbook()
    detail = workbook.active
    detail.title = "明细"
    detail.append(["结算月：", "2026-06", "所属月：", "2026-06"])
    detail.append(DETAIL_HEADERS)
    common = {
        "薪资所属月": "2026-06", "社保所属月": "2026-06", "公积金所属月": "2026-06",
        "单位社保合计": 100, "残保金企业": 10, "个人社保合计": 20,
        "基本公积金企业": 30, "基本公积金个人": 30, "企龄津贴": 0,
        "差旅津贴": 0, "最低薪资补偿": 0, "旷工扣款": 0, "其他扣款": 0,
        "独生子女": 0, "税后调整": 0, "经济补偿金": 0, "个调税": 5,
        "管理费": 100,
    }
    for sequence, name in enumerate(("保留甲", "应删除", "保留乙"), start=1):
        values = dict(common, 序号=sequence, 姓名=name, 基本工资=999, 岗位津贴=999,
                      绩效基数=999, 实际绩效=999, 绩效调整=999, 排班绩效=999,
                      加班费=999, 翻班津贴=999, 餐费补贴=999, 奖金=999,
                      节日福利=999, 高温费=999, 税前调整=999, 病假扣款=999,
                      事假扣款=999, 缺勤扣款=999)
        detail.append([values.get(header) for header in DETAIL_HEADERS])
    detail.append(["*合计*", "79"] + [None] * (len(DETAIL_HEADERS) - 2))

    notice = workbook.create_sheet("付款通知书")
    notice.append(["费用所属期间：", "2026-06"])
    notice.append(["项目", None, None, None, "人数", "金额"])
    for label in (
        "单位社保合计", "个人社保合计", "社保小计", "残保金企业",
        "基本公积金企业", "基本公积金个人", "公积金小计", "应发工资",
        "应税工资", "个调税", "实发工资", "代收代付",
    ):
        notice.append([label, None, None, None, "79" if label in {"社保小计", "残保金企业", "公积金小计"} else 3, 0])
    notice.append(["管理费：", None, None, None, 0])
    notice.append(["本期总计应付款：", None, None, None, 0])
    notice.append(["人民币大写：", None, None, None, "旧金额"])
    workbook.save(path)
    workbook.close()


def _create_source(path: Path, *, include_unknown: bool = False) -> None:
    workbook = openpyxl.Workbook()
    driver = workbook.active
    driver.title = "司机"
    driver.append(["员工姓名", "基本工资"])
    driver.append(["司机员", 99999])
    outsource = workbook.create_sheet("外包")
    outsource.append(["员工姓名", "基本工资"])
    outsource.append(["外包员", 88888])
    dispatch = workbook.create_sheet("派遣")
    dispatch.append([
        "员工姓名", "基本工资", "岗位津贴", "绩效基数", "实际绩效", "绩效调整",
        "排班绩效", "加班工资", "翻班津贴", "餐饮津贴", "奖金", "节日福利",
        "高温费", "税前调整", "病假扣款", "事假扣款", "缺勤扣款", "税后调整",
    ])
    dispatch.append(["保留乙", 3000, 300, 500, 400, None, 20, 30, 40, 50, 60, 70, 80, 90, 10, 20, 30, 5])
    dispatch.append(["陌生人" if include_unknown else "保留甲", 2000, 200, 400, 300, 10, None, None, None, None, None, None, None, None, None, None, None, None])
    workbook.save(path)
    workbook.close()


def _row_by_name(sheet: openpyxl.worksheet.worksheet.Worksheet, name: str) -> int:
    for row in range(3, sheet.max_row + 1):
        if sheet.cell(row, 2).value == name:
            return row
    raise AssertionError(f"missing {name}")


def test_sync_master_to_source_roster_uses_only_the_selected_sheet(tmp_path: Path) -> None:
    from backend.roster_sync import sync_master_to_source_roster

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _create_master(master)
    _create_source(source)

    result = sync_master_to_source_roster(
        master_path=master,
        source_path=source,
        output_path=output,
        source_sheet="派遣",
        target_sheet="明细",
        target_period="2026-07",
    )

    workbook = openpyxl.load_workbook(output, data_only=False)
    detail = workbook["明细"]
    headers = {cell.value: cell.column for cell in detail[2] if cell.value}
    assert [detail.cell(row, 2).value for row in range(3, 5)] == ["保留甲", "保留乙"]
    assert detail.cell(5, 1).value == "*合计*"
    assert result["retained_names"] == ["保留甲", "保留乙"]
    assert result["removed_names"] == ["应删除"]
    assert result["employee_count"] == 2
    assert "司机员" not in result["retained_names"]
    assert "外包员" not in result["retained_names"]

    row_a = _row_by_name(detail, "保留甲")
    row_b = _row_by_name(detail, "保留乙")
    assert detail.cell(row_a, headers["基本工资"]).value == 2000
    assert detail.cell(row_b, headers["加班费"]).value == 30
    assert detail.cell(row_a, headers["加班费"]).value == 0
    assert detail.cell(row_a, headers["绩效调整"]).value == 10
    assert detail.cell(row_a, headers["薪资所属月"]).value == "2026-07"
    assert detail.cell(row_b, headers["社保所属月"]).value == "2026-07"
    assert detail.cell(row_b, headers["公积金所属月"]).value == "2026-07"

    expected_gross_b = 3000 + 300 + 400 + 0 + 20 + 30 + 40 + 50 + 60 + 70 + 80 + 90 - 10 - 20 - 30
    assert detail.cell(row_b, headers["应发工资"]).value == expected_gross_b
    assert detail.cell(row_b, headers["应税工资"]).value == expected_gross_b - 20 - 30
    assert detail.cell(row_b, headers["实发工资"]).value == expected_gross_b - 20 - 30 - 5 + 5
    assert detail.cell(row_b, headers["代收代付"]).value == expected_gross_b + 100 + 10 + 30
    assert detail.cell(row_b, headers["合计"]).value == expected_gross_b + 100 + 10 + 30 + 100

    notice = workbook["付款通知书"]
    assert notice["B1"].value == "2026-07"
    assert notice["E3"].value == 2
    assert notice["E5"].value == "79"
    assert notice["E15"].value == sum(detail.cell(row, headers["管理费"]).value for row in (3, 4))
    assert notice["E16"].value == detail.cell(5, headers["合计"]).value
    assert isinstance(notice["E17"].value, str) and notice["E17"].value.endswith(("整", "分"))
    workbook.close()


def test_sync_preserves_and_repairs_formulas_after_removing_non_roster_rows(
    tmp_path: Path,
) -> None:
    from backend.roster_sync import sync_master_to_source_roster

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _create_master(master)
    _create_source(source)
    workbook = openpyxl.load_workbook(master)
    detail = workbook["明细"]
    gross_column = DETAIL_HEADERS.index("应发工资") + 1
    total_column = DETAIL_HEADERS.index("合计") + 1
    for row in range(3, 6):
        detail.cell(row, gross_column).value = f"=K{row}+L{row}"
    detail.cell(6, gross_column).value = f"=SUM(AG3:AG5)"
    workbook["付款通知书"]["E16"] = f"='明细'!AO6"
    summary = workbook.create_sheet("汇总")
    summary["A1"] = "=SUM('明细'!K3:K5)"
    workbook.save(master)
    workbook.close()

    sync_master_to_source_roster(
        master_path=master,
        source_path=source,
        output_path=output,
        source_sheet="派遣",
        target_sheet="明细",
        target_period="2026-07",
    )

    updated = openpyxl.load_workbook(output, data_only=False)
    try:
        detail = updated["明细"]
        assert detail.cell(3, gross_column).value == "=K3+L3"
        assert detail.cell(4, gross_column).value == "=K4+L4"
        assert detail.cell(5, gross_column).value == "=SUM(AG3:AG4)"
        assert updated["付款通知书"]["E16"].value == "='明细'!AO5"
        assert updated["汇总"]["A1"].value == "=SUM('明细'!K3:K4)"
        assert detail.cell(5, total_column).value is not None
    finally:
        updated.close()


def test_sync_repairs_template_objects_after_removing_non_roster_rows(
    tmp_path: Path,
) -> None:
    from backend.roster_sync import sync_master_to_source_roster

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _create_master(master)
    _create_source(source)
    workbook = openpyxl.load_workbook(master)
    detail = workbook["明细"]
    detail.add_table(Table(displayName="PayrollTable", ref="A2:AO6"))
    detail.auto_filter.ref = "A2:AO6"
    validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
    validation.add("K3:K5")
    detail.add_data_validation(validation)
    detail.conditional_formatting.add(
        "K3:K5", CellIsRule(operator="greaterThan", formula=["0"]),
    )
    workbook.defined_names.add(DefinedName(
        "PayrollRows", attr_text="'明细'!$A$3:$AO$5",
    ))
    chart = BarChart()
    chart.add_data(Reference(detail, min_col=11, min_row=2, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(detail, min_col=2, min_row=3, max_row=5))
    detail.add_chart(chart, "AQ2")
    workbook.save(master)
    workbook.close()
    original_bytes = master.read_bytes()

    sync_master_to_source_roster(
        master_path=master,
        source_path=source,
        output_path=output,
        source_sheet="派遣",
        target_sheet="明细",
        target_period="2026-07",
    )

    assert master.read_bytes() == original_bytes
    updated = openpyxl.load_workbook(output, data_only=False)
    try:
        detail = updated["明细"]
        assert detail.tables["PayrollTable"].ref == "A2:AO5"
        assert detail.auto_filter.ref == "A2:AO5"
        assert str(detail.data_validations.dataValidation[0].sqref) == "K3:K4"
        conditional_ranges = [str(item.sqref) for item in detail.conditional_formatting]
        assert conditional_ranges == ["K3:K4"]
        assert updated.defined_names["PayrollRows"].attr_text == "'明细'!$A$3:$AO$4"
        series = detail._charts[0].series[0]
        assert series.val.numRef.f == "'明细'!$K$3:$K$4"
        assert series.cat.numRef.f == "'明细'!$B$3:$B$4"
    finally:
        updated.close()


def test_sync_rejects_a_source_person_missing_from_the_master_without_output(tmp_path: Path) -> None:
    from backend.roster_sync import RosterSyncError, sync_master_to_source_roster

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _create_master(master)
    _create_source(source, include_unknown=True)

    with pytest.raises(RosterSyncError, match="陌生人"):
        sync_master_to_source_roster(
            master_path=master,
            source_path=source,
            output_path=output,
            source_sheet="派遣",
            target_sheet="明细",
            target_period="2026-07",
        )

    assert not output.exists()


def test_sync_rejects_overwriting_a_formula_without_output(tmp_path: Path) -> None:
    from backend.roster_sync import RosterSyncError, sync_master_to_source_roster

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _create_master(master)
    _create_source(source)
    workbook = openpyxl.load_workbook(master)
    detail = workbook["明细"]
    basic_salary_column = DETAIL_HEADERS.index("基本工资") + 1
    detail.cell(3, basic_salary_column).value = "=1+1"
    workbook.save(master)
    workbook.close()

    with pytest.raises(RosterSyncError, match="公式"):
        sync_master_to_source_roster(
            master_path=master,
            source_path=source,
            output_path=output,
            source_sheet="派遣",
            target_sheet="明细",
            target_period="2026-07",
        )

    assert not output.exists()


def test_agent_completes_an_explicit_roster_scoped_run_without_a_model_write_loop(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.agent_execution as execution
    from core.document_agent.orchestrator import ToolRegistry

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "Agent草稿.xlsx"
    _create_master(master)
    _create_source(source)
    monkeypatch.setattr(execution.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(execution.ModelConfig, "fallback_from_env", lambda: None)
    run = {
        "run_id": "s" * 32,
        "master_file": master.name,
        "salary_month": "2026.07",
        "instruction": "只处理变更表格里面的派遣sheet的人员数据，更新在总表里，但是其他sheet的人员数据完全不保留，全部去掉，不显示。只要派遣sheet的人员",
        "_master_path": str(master),
        "_source_paths": {source.name: str(source)},
        "file_manifest": [],
        "workbook_updates": [],
        "messages": [],
        "model_plan": {"steps": []},
    }

    execution.execute_model_plan(
        run,
        draft=output,
        registry=ToolRegistry(),
        read_schemas=[],
        materials=[],
        rules={},
        save=lambda _run: None,
        emit=lambda *_args, **_kwargs: None,
    )

    assert output.is_file()
    assert run["status"] == "awaiting_review"
    assert run["execution_result"]["status"] == "execution_incomplete"
    assert run["execution_result"]["code"] == "VALIDATION_FAILED"
    assert run["roster_sync"]["employee_count"] == 2
    assert run["roster_sync"]["source_sheet"] == "派遣"
    assert run["workbook_updates"]


def test_deterministic_roster_sync_executes_skill_validations_after_the_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.agent_execution as execution
    from core.document_agent.orchestrator import ToolRegistry

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    output = tmp_path / "Agent草稿.xlsx"
    _create_master(master)
    _create_source(source)
    monkeypatch.setattr(execution.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(execution.ModelConfig, "fallback_from_env", lambda: None)
    validation_calls: list[str] = []
    emitted: list[tuple[str, dict]] = []
    registry = ToolRegistry()

    def validate_workbook() -> dict:
        validation_calls.append("validate_workbook")
        return {
            "readable": True,
            "unresolved_item_count": 0,
            "formula_errors": [],
            "summary_range_errors": [],
            "duplicate_identities": [],
            "identity_errors": [],
            "can_publish": True,
        }

    def validate_with_officecli() -> dict:
        validation_calls.append("validate_with_officecli")
        return {
            "valid": True,
            "recalculated": True,
            "formula_errors": [],
            "unevaluated_formulas": [],
        }

    registry.register("validate_workbook", validate_workbook)
    registry.register("validate_with_officecli", validate_with_officecli)
    run = {
        "run_id": "v" * 32,
        "master_file": master.name,
        "salary_month": "2026.07",
        "instruction": "只保留变更表格派遣 sheet 的人员，其他人员全部删除",
        "_master_path": str(master),
        "_source_paths": {source.name: str(source)},
        "file_manifest": [],
        "workbook_updates": [],
        "messages": [],
        "model_plan": {"steps": []},
    }

    execution.execute_model_plan(
        run,
        draft=output,
        registry=registry,
        read_schemas=[],
        materials=[],
        rules={},
        save=lambda _run: None,
        emit=lambda _run, event_type, payload, **_kwargs: emitted.append((event_type, payload)),
    )

    assert validation_calls == ["validate_workbook", "validate_with_officecli"]
    checkpoint = run["execution_checkpoint"]
    assert checkpoint["required_validation_tools"] == [
        "validate_workbook", "validate_with_officecli",
    ]
    assert checkpoint["completed_steps"] == [
        "validate_workbook", "validate_with_officecli",
    ]
    tool_results = [payload for event_type, payload in emitted if event_type == "tool_result"]
    assert [payload["name"] for payload in tool_results] == [
        "source_sheet_roster_sync", "validate_workbook", "validate_with_officecli",
    ]
    assert tool_results[0]["write_count"] > 0
    assert all(payload["status"] == "succeeded" for payload in tool_results)
    assert run["execution_result"]["status"] == "completed"
