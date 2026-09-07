import asyncio
from concurrent.futures import ThreadPoolExecutor
import json
from datetime import datetime, timezone
from types import SimpleNamespace
from typing import Any

import openpyxl

import backend.routers.agent as agent
from backend.routers.agent import _append_event, AgentMessageIn


def test_run_message_context_passes_manual_excerpt_to_model(monkeypatch) -> None:
    monkeypatch.setattr(
        agent,
        "_material_context",
        lambda _run: [{"filename": "操作手册.docx", "kind": "manual", "excerpt": "奖金按J列填写"}],
    )

    # workflow.started_at marks a run whose execution has begun; that keeps
    # the task-level status prompt (short, read-only) instead of alignment.
    messages = agent._run_message_messages({
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "conversation": [],
        "summary": {},
        "workflow": {"started_at": "2026-01-01T00:00:00+00:00"},
    })
    context = json.loads(messages[1]["content"])
    assert context["materials"][0]["excerpt"] == "奖金按J列填写"
    assert context["materials"][0]["text_length"] == len("奖金按J列填写")
    assert "最多回答3句、200字" in messages[0]["content"]
    assert "不要主动说明任何文档或手册的读取状态" in messages[0]["content"]


def test_run_message_alignment_mode_restates_intent_with_workbook_headers(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])
    source = tmp_path / "更新表.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "更新"
    sheet.append(["姓名", "奖金"])
    sheet.append(["张三", 100])
    workbook.save(source)

    messages = agent._run_message_messages({
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "conversation": [{"role": "user", "content": "奖金按更新表更新，缺人的先留着"}],
        "summary": {},
        "master_file": "总表.xlsx",
        "_master_path": str(tmp_path / "缺失.xlsx"),
        "_source_paths": {"更新表.xlsx": str(source)},
    })
    assert "需求对齐助手" in messages[0]["content"]
    assert "我的理解" in messages[0]["content"]
    context = json.loads(messages[1]["content"])
    # 读取失败的总表不进入上下文，但来源文件的结构（含表头）必须进入。
    assert [book["filename"] for book in context["workbooks"]] == ["更新表.xlsx"]
    assert context["workbooks"][0]["sheets"][0]["header_row"] == ["姓名", "奖金"]


def test_save_run_uses_independent_temporary_files_for_concurrent_updates(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "b" * 32, "tenant_id": "tenant-a", "project_id": "project-a",
        "events": [],
    }

    def save_many(_: int) -> None:
        for _ in range(20):
            agent._save_run(run)

    with ThreadPoolExecutor(max_workers=8) as executor:
        list(executor.map(save_many, range(8)))

    saved = agent._run_path("tenant-a", "b" * 32)
    assert json.loads(saved.read_text(encoding="utf-8"))["run_id"] == "b" * 32
    assert not list(tmp_path.rglob("*.tmp"))


def test_save_run_retries_transient_windows_file_lock(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "c" * 32, "tenant_id": "tenant-a", "project_id": "project-a",
        "events": [],
    }
    original_replace = type(tmp_path).replace
    attempts = 0

    def flaky_replace(self, target):
        nonlocal attempts
        if self.suffix == ".tmp" and attempts < 2:
            attempts += 1
            raise PermissionError("[WinError 32] file is in use")
        return original_replace(self, target)

    monkeypatch.setattr(type(tmp_path), "replace", flaky_replace)
    monkeypatch.setattr(agent.time, "sleep", lambda _seconds: None)

    agent._save_run(run)

    assert attempts == 2
    assert json.loads(agent._run_path("tenant-a", "c" * 32).read_text(encoding="utf-8"))["run_id"] == "c" * 32


def test_save_run_survives_a_longer_windows_lock_window(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "d" * 32, "tenant_id": "tenant-a", "project_id": "project-a",
        "events": [],
    }
    original_replace = type(tmp_path).replace
    attempts = 0

    def locked_for_a_few_seconds(self, target):
        nonlocal attempts
        if self.suffix == ".tmp" and attempts < 7:
            attempts += 1
            raise PermissionError("[WinError 5] file is in use")
        return original_replace(self, target)

    monkeypatch.setattr(type(tmp_path), "replace", locked_for_a_few_seconds)
    monkeypatch.setattr(agent.time, "sleep", lambda _seconds: None)

    agent._save_run(run)

    assert attempts == 7
    assert json.loads(agent._run_path("tenant-a", "d" * 32).read_text(encoding="utf-8"))["run_id"] == "d" * 32


def test_append_event_is_revisioned_and_does_not_mutate_previous_payload() -> None:
    run = {"run_id": "run-1", "events": [{"revision": 1, "type": "run_started", "payload": {"ok": True}}]}
    event = _append_event(run, "user_message", {"content": "确认更新"}, item_id="item-1")

    assert event["revision"] == 2
    assert event["run_id"] == "run-1"
    assert event["type"] == "user_message"
    assert event["item_id"] == "item-1"
    assert run["events"][0]["payload"] == {"ok": True}
    datetime.fromisoformat(event["at"].replace("Z", "+00:00"))


def test_append_event_redacts_sensitive_payload_fields() -> None:
    run = {"run_id": "run-2", "events": []}
    event = _append_event(run, "model_response", {"content": "ok", "api_key": "secret", "prompt": "private"})
    assert "api_key" not in event["payload"]
    assert "prompt" not in event["payload"]


def test_append_event_accepts_safe_progress_updates() -> None:
    run = {"run_id": "run-3", "events": []}
    event = _append_event(
        run,
        "progress",
        {"stage": "person", "label": "正在处理第 1 / 2 人", "current": 1, "total": 2},
        item_id="item-1",
    )

    assert event["type"] == "progress"
    assert event["payload"]["current"] == 1
    assert event["item_id"] == "item-1"


def test_list_agent_events_tolerates_malformed_revisions(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "d" * 32,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "events": [
            {"event_id": "bad", "revision": "not-a-number", "type": "progress", "payload": {}},
            {"event_id": "ok", "revision": 3, "type": "progress", "payload": {"label": "继续"}},
            {"event_id": "bool", "revision": True, "type": "progress", "payload": {}},
        ],
    }
    agent._save_run(run)

    result = agent.list_agent_events(
        "d" * 32,
        after_revision=0,
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    assert [event["event_id"] for event in result["events"]] == ["ok"]
    assert result["revision"] == 3


def test_load_run_retries_transient_file_lock(monkeypatch) -> None:
    class LockedRunPath:
        def __init__(self) -> None:
            self.attempts = 0

        def is_file(self) -> bool:
            return True

        def read_text(self, *, encoding: str) -> str:
            self.attempts += 1
            if self.attempts < 3:
                raise PermissionError("[WinError 32] file is in use")
            return json.dumps({"run_id": "e" * 32, "tenant_id": "tenant-a"})

    path = LockedRunPath()
    monkeypatch.setattr(agent, "_run_path", lambda *_args: path)
    monkeypatch.setattr(agent.time, "sleep", lambda _seconds: None)

    run = agent._load_run("e" * 32, SimpleNamespace(tenant_id="tenant-a"))

    assert run["run_id"] == "e" * 32
    assert path.attempts == 3


def test_stale_processing_checkpoint_becomes_resumable(monkeypatch) -> None:
    stale = {
        "run_id": "f" * 32,
        "status": "processing",
        "updated_at": "2020-01-01T00:00:00+00:00",
        "events": [],
    }
    monkeypatch.setattr(agent, "_ACTIVE_RUN_IDS", set())
    assert agent._recover_stale_processing(stale) is True
    assert stale["status"] == "execution_incomplete"
    assert stale["code"] == "WORKFLOW_INTERRUPTED"
    assert stale["events"][-1]["payload"]["failure_type"] == "stale_worker"


def test_stale_active_worker_checkpoint_stops_processing(monkeypatch) -> None:
    """事件断流超过租期时，即使 worker 标记仍在，也不能永久 processing。"""
    run_id = "ab" * 16
    stale = {
        "run_id": run_id,
        "status": "processing",
        "updated_at": "2020-01-01T00:00:00+00:00",
        "events": [],
    }
    monkeypatch.setattr(agent, "_ACTIVE_RUN_IDS", {run_id})

    try:
        assert agent._recover_stale_processing(stale) is True
        assert stale["status"] == "execution_incomplete"
        assert stale["code"] == "WORKER_UNRESPONSIVE"
        assert agent.run_stop_requested(run_id) is True
    finally:
        agent.clear_run_stop(run_id)


def test_sse_frame_has_an_explicit_type_and_redacts_sensitive_values() -> None:
    frame = agent._sse_frame("message.delta", {"content": "你好", "api_key": "secret"}, event_id="12")

    assert frame.startswith("id: 12\nevent: message.delta\n")
    payload = json.loads(frame.split("data: ", 1)[1])
    assert payload == {"content": "你好"}
    assert "secret" not in frame


def test_streamed_run_message_persists_the_user_and_agent_messages(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    monkeypatch.setattr(agent, "_stream_orchestrate_run_message", lambda _run: iter(["这是", "流式回复"]))
    run = {
        "run_id": "a" * 32, "tenant_id": "tenant-a", "project_id": "project-a",
        "conversation": [], "events": [],
    }
    agent._save_run(run)

    response = agent.stream_agent_run_message(
        "a" * 32, AgentMessageIn(message="请说明当前状态"),
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    async def read_body() -> str:
        chunks: list[str] = []
        async for chunk in response.body_iterator:
            chunks.append(chunk.decode("utf-8") if isinstance(chunk, bytes) else chunk)
        return "".join(chunks)

    body = asyncio.run(read_body())
    assert "event: message.accepted" in body
    assert "event: message.delta" in body
    assert "这是流式回复" in body
    saved = agent._load_run("a" * 32, SimpleNamespace(tenant_id="tenant-a"))
    assert [message["role"] for message in saved["conversation"]] == ["user", "agent"]


def test_run_event_stream_replays_after_revision_and_finishes_at_terminal_state(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "b" * 32, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "completed", "events": [],
    }
    _append_event(run, "run_started", {"label": "开始"})
    _append_event(run, "run_completed", {"label": "完成"})
    agent._save_run(run)

    response = agent.stream_agent_events(
        "b" * 32, after_revision=1,
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    async def read_body() -> str:
        chunks: list[str] = []
        async for chunk in response.body_iterator:
            chunks.append(chunk.decode("utf-8") if isinstance(chunk, bytes) else chunk)
        return "".join(chunks)

    body = asyncio.run(read_body())
    assert "id: 2" in body
    assert "event: run_completed" in body
    assert "event: run_started" not in body


def test_run_event_stream_recovers_stale_processing_run(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "9" * 32,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "processing",
        "updated_at": "2020-01-01T00:00:00+00:00",
        "events": [],
    }
    path = agent._run_path("tenant-a", run["run_id"])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(run), encoding="utf-8")
    response = agent.stream_agent_events(
        "9" * 32,
        after_revision=0,
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    async def read_body() -> str:
        chunks: list[str] = []
        async for chunk in response.body_iterator:
            chunks.append(chunk.decode("utf-8") if isinstance(chunk, bytes) else chunk)
        return "".join(chunks)

    body = asyncio.run(read_body())
    assert "WORKFLOW_INTERRUPTED" in body
    saved = agent._load_run("9" * 32, SimpleNamespace(tenant_id="tenant-a"))
    assert saved["status"] == "execution_incomplete"


def test_result_summary_uses_physical_workbook_updates(tmp_path, monkeypatch) -> None:
    import openpyxl

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    output = tmp_path / "updated.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资"
    workbook.save(output)
    workbook.close()
    run = {
        "run_id": "c" * 32, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "completed", "draft_filename": output.name,
        "validation": {"status": "structurally_valid"},
        "result": {"filename": output.name, "sha256": agent.file_digest(output), "change_count": 1},
        "workbook_updates": [{
            "target_sheet": "工资", "target_cell": "K1", "old_value": 21, "new_value": 23,
            "source_file": "考勤.xlsx", "source_sheet": "7月", "source_cells": ["E11"],
        }],
        "items": [], "events": [],
    }
    agent._save_run(run)

    result = agent.get_agent_result(
        "c" * 32, page=1, page_size=50,
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    assert result["available"] is True
    assert result["change_count"] == 1
    assert result["changes"][0]["target_cell"] == "K1"
    assert result["can_download"] is True


def test_background_worker_claim_is_idempotent() -> None:
    run_id = "worker-test"
    assert agent._claim_run_worker(run_id) is True
    assert agent._claim_run_worker(run_id) is False
    agent._release_run_worker(run_id)
    assert agent._claim_run_worker(run_id) is True
    agent._release_run_worker(run_id)


def test_processing_checkpoint_is_marked_stale_after_worker_restart() -> None:
    from datetime import timedelta

    old = datetime.now(timezone.utc) - timedelta(seconds=agent.PROCESSING_STALE_SECONDS + 1)
    assert agent._processing_is_stale({"status": "processing", "updated_at": old.isoformat()}) is True
    assert agent._processing_is_stale({"status": "processing", "updated_at": datetime.now(timezone.utc).isoformat()}) is False
    assert agent._processing_is_stale({"status": "completed", "updated_at": old.isoformat()}) is False


def test_saving_a_run_refreshes_the_processing_heartbeat(tmp_path, monkeypatch) -> None:
    from datetime import timedelta

    run_id = "c3" * 16
    old = datetime.now(timezone.utc) - timedelta(seconds=agent.PROCESSING_STALE_SECONDS + 1)
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "processing",
        "updated_at": old.isoformat(),
        "events": [],
    })

    saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert agent._processing_is_stale(saved) is False


def test_process_endpoint_recovers_a_stale_processing_checkpoint(tmp_path, monkeypatch) -> None:
    run_id = "f" * 32
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    started: list[tuple] = []

    class FakeThread:
        def __init__(self, *, target, args, daemon):
            started.append((target, args, daemon))

        def start(self):
            return None

    monkeypatch.setattr(agent, "Thread", FakeThread)
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "processing",
        "events": [],
    })
    monkeypatch.setattr(agent, "_processing_is_stale", lambda _run: True)
    agent._ACTIVE_RUN_IDS.discard(run_id)
    try:
        agent.process_agent_run(run_id, user=SimpleNamespace(tenant_id="tenant-a"))
        saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
        assert len(started) == 1
        assert any(event.get("payload", {}).get("stage") == "stale_recovery" for event in saved["events"])
        assert saved["workflow"]["stage"] == "queued"
    finally:
        agent._release_run_worker(run_id)


def test_process_endpoint_clears_a_previous_interruption_code_when_resuming(tmp_path, monkeypatch) -> None:
    run_id = "a1" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)

    class FakeThread:
        def __init__(self, **_kwargs):
            return None

        def start(self):
            return None

    monkeypatch.setattr(agent, "Thread", FakeThread)
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "execution_incomplete",
        "code": "WORKFLOW_INTERRUPTED",
        "detail": "后台处理意外中断，已保留完成的写入和事件；可直接续跑",
        "events": [],
    })
    agent._ACTIVE_RUN_IDS.discard(run_id)
    try:
        agent.process_agent_run(run_id, user=SimpleNamespace(tenant_id="tenant-a"))
        saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
        assert saved["status"] == "processing"
        assert "code" not in saved
        assert saved["detail"] == "后台 Agent 已接管处理，页面关闭后仍会继续"
    finally:
        agent._release_run_worker(run_id)


def test_execute_endpoint_allows_manual_retry_of_a_zero_write_progress_loop(tmp_path, monkeypatch) -> None:
    """A historical read-only loop can be retried, while preserving fail-closed blocks."""

    run_id = "a2" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)

    class FakeThread:
        def __init__(self, **_kwargs):
            return None

        def start(self):
            return None

    monkeypatch.setattr(agent, "Thread", FakeThread)
    monkeypatch.setattr(agent, "_require_confirmed_plan", lambda _run: None)
    monkeypatch.setattr(agent, "_verify_plan_files", lambda _run: None)
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "blocked",
        "code": "NO_PROGRESS_DETECTED",
        "workbook_updates": None,
        "plan_confirmation": {"required": False, "confirmed": True},
        "month_confirmation": {"required": False, "confirmed": True},
        "events": [],
    })
    agent._ACTIVE_RUN_IDS.discard(run_id)
    try:
        resumed = agent.execute_agent_run(
            run_id, user=SimpleNamespace(tenant_id="tenant-a"), db=SimpleNamespace(),
        )
        assert resumed["status"] == "processing"
        assert any(event["type"] == "run_started" for event in resumed["events"])
    finally:
        agent._release_run_worker(run_id)


def test_process_endpoint_resets_previous_round_plan_for_next_round(tmp_path, monkeypatch) -> None:
    run_id = "b2" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)

    class FakeThread:
        def __init__(self, **_kwargs):
            return None

        def start(self):
            return None

    monkeypatch.setattr(agent, "Thread", FakeThread)
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "completed",
        "draft_filename": "draft.xlsx",
        "workflow": {"started_at": "2026-01-01T00:00:00+00:00"},
        "model_plan": {"summary": "上一轮计划", "steps": ["上一轮步骤"], "questions": [], "model": "m"},
        "_plan_evidence_digest": "digest-old",
        "events": [],
    })
    agent._ACTIVE_RUN_IDS.discard(run_id)
    try:
        agent.process_agent_run(
            run_id,
            agent.AgentProcessIn(instruction="把奖金改为按新表更新"),
            user=SimpleNamespace(tenant_id="tenant-a"),
        )
        saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
        # 下一轮必须丢弃上一轮计划，让工作流按新的对齐对话重新规划。
        assert saved["model_plan"] is None
        assert "_plan_evidence_digest" not in saved
        assert saved["workflow"]["next_round"] is True
        assert any(event.get("payload", {}).get("stage") == "next_round" for event in saved["events"])
        assert any(
            message.get("content") == "把奖金改为按新表更新"
            for message in saved.get("conversation", [])
            if message.get("role") == "user"
        )
    finally:
        agent._release_run_worker(run_id)


def test_alignment_phase_rule_message_stays_in_conversation(tmp_path, monkeypatch) -> None:
    run_id = "c3" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: SimpleNamespace())
    monkeypatch.setattr(
        agent,
        "OpenAICompatibleProvider",
        lambda _config: SimpleNamespace(complete=lambda **_kwargs: SimpleNamespace(content="我的理解：先对齐，再执行。")),
    )
    processed: list[tuple] = []

    def _record_if_processed(*args: Any, **_kwargs: Any) -> None:
        processed.append(args)

    monkeypatch.setattr(agent, "process_agent_run", _record_if_processed)
    monkeypatch.setattr(agent, "parse_supported_workbook_rule", lambda _message: {"kind": "duty_roster_name_count"})
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "planning",
        "events": [],
        "conversation": [],
    })

    result = agent.message_agent_run(
        run_id,
        AgentMessageIn(message="把值班表出现次数填到总表"),
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    # 对齐阶段（执行未开始）：规则式消息进入对话让模型复述理解，
    # 不能直接触发固定流程。
    assert processed == []
    assert result["message"]["role"] == "agent"
    saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert [message["role"] for message in saved["conversation"]] == ["user", "agent"]


def test_change_report_groups_updates_and_announces_once() -> None:
    run = {
        "items": [
            {"status": "auto_applied"},
            {"status": "resolved"},
            {"status": "pending"},
        ],
        "workbook_updates": [
            {"sheet": "配送员值班费", "cell": "D2", "before": None, "after": 0, "rule": "值班次数"},
            {"sheet": "配送员值班费", "cell": "D3", "before": 1, "after": 0, "rule": "值班次数"},
            {"sheet": "工资核算", "cell": "F5", "before": 100, "after": 200, "rule": "奖金"},
        ],
        "conversation": [],
    }
    report = agent._build_change_report(run, run["workbook_updates"])
    assert "共写入 3 处改动" in report
    assert "配送员值班费（值班次数）：2 处" in report
    assert "工资核算（奖金）：1 处" in report
    assert "自动写入 1 项" in report
    assert "经确认写入 1 项" in report
    assert "另有 1 项待你确认" in report

    # 同一版本结果只播报一次；结果变化后再次播报。
    agent._announce_change_report(run, "sha-1")
    assert len(run["conversation"]) == 1
    agent._announce_change_report(run, "sha-1")
    assert len(run["conversation"]) == 1
    agent._announce_change_report(run, "sha-2")
    assert len(run["conversation"]) == 2
    assert all(message.get("kind") == "change_report" for message in run["conversation"])
    assert len(run["change_reports"]) == 2


def test_background_worker_records_a_sanitized_interruption_reason(tmp_path, monkeypatch) -> None:
    run_id = "b2" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    monkeypatch.setattr(agent, "_material_context", lambda _run: (_ for _ in ()).throw(RuntimeError("internal secret")))
    agent._save_run({
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "planning",
        "events": [],
    })

    agent._run_agent_workflow(run_id, "tenant-a")

    saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert saved["status"] == "execution_incomplete"
    failure = saved["events"][-1]
    assert failure["type"] == "run_failed"
    assert failure["payload"]["code"] == "WORKFLOW_INTERRUPTED"
    assert failure["payload"]["failure_type"] == "RuntimeError"
    assert "secret" not in json.dumps(failure, ensure_ascii=False)


def test_background_workflow_reaches_a_downloadable_terminal_result(tmp_path, monkeypatch) -> None:
    import openpyxl

    run_id = "d" * 32
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    run = {
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "planning", "instruction": "更新", "events": [], "items": [],
        "model_plan": {"summary": "计划", "steps": ["写入"], "questions": []},
        "plan_confirmation": {"required": True, "confirmed": True},
        "month_confirmation": {"required": False, "confirmed": True},
    }
    agent._save_run(run)

    def fake_execute(current_run_id, user, db):
        current = agent._load_run(current_run_id, user)
        output = tmp_path / "updated.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = 2
        workbook.save(output)
        workbook.close()
        current.update({
            "status": "awaiting_review", "draft_filename": output.name,
            "execution_result": {"status": "completed", "content": "完成"},
            "workbook_updates": [{"target_sheet": "Sheet", "target_cell": "A1", "old_value": 1, "new_value": 2}],
        })
        agent._save_run(current)
        return agent._public_run(current)

    monkeypatch.setattr(agent, "_execute_agent_run_sync", fake_execute)
    agent._run_agent_workflow(run_id, "tenant-a")

    completed = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert completed["status"] == "completed"
    assert completed["validation"]["status"] == "structurally_valid"
    assert completed["result"]["change_count"] == 1
    assert any(event["type"] == "run_completed" for event in completed["events"])
    progress = [event["payload"] for event in completed["events"] if event["type"] == "progress"]
    assert {item["stage"] for item in progress}.isdisjoint({"materials"})
    assert any(
        item["stage"] == "file_analysis" and "跳过解析" in item["label"]
        for item in progress
    )


def test_background_workflow_stops_after_a_zero_write_segment_until_user_resumes(tmp_path, monkeypatch) -> None:
    import openpyxl

    run_id = "e" * 32
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    agent._save_run({
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "planning", "instruction": "更新", "events": [], "items": [],
        "model_plan": {"summary": "计划", "steps": ["写入"], "questions": []},
        "plan_confirmation": {"required": True, "confirmed": True},
        "month_confirmation": {"required": False, "confirmed": True},
    })
    calls = 0

    def fake_execute(current_run_id, user, db):
        nonlocal calls
        calls += 1
        current = agent._load_run(current_run_id, user)
        current.update({
            "status": "execution_incomplete",
            "code": "NO_WRITES_PERFORMED",
            "execution_result": {"status": "execution_incomplete", "code": "NO_WRITES_PERFORMED"},
        })
        agent._save_run(current)
        return agent._public_run(current)

    monkeypatch.setattr(agent, "_execute_agent_run_sync", fake_execute)
    agent._run_agent_workflow(run_id, "tenant-a")

    completed = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert calls == 1
    assert completed["status"] == "execution_incomplete"
    assert completed["code"] == "NO_WRITES_PERFORMED"
    assert not any(event.get("type") == "run_completed" for event in completed["events"])
    assert not any(event.get("payload", {}).get("stage") == "segment_resume" for event in completed["events"])
    assert any(event.get("payload", {}).get("stage") == "resumable" for event in completed["events"])


def test_background_workflow_keeps_explicit_target_month_when_filename_month_differs(
    tmp_path, monkeypatch,
) -> None:
    run_id = "f" * 32
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: object())
    monkeypatch.setattr(agent.ModelConfig, "fallback_from_env", lambda: None)
    agent._save_run({
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "planning", "instruction": "生成五月数据", "events": [], "items": [],
        "salary_month": "2026.05", "target_salary_month": "2026.05",
        "month_authority": "explicit",
        "model_plan": {"summary": "计划", "steps": ["写入"], "questions": []},
        "plan_confirmation": {"required": True, "confirmed": True},
        "month_confirmation": {
            "required": True, "confirmed": False,
            "configured_month": "2026.05", "filename_month": "2026.04",
        },
    })
    observed_months: list[str] = []

    def fake_execute(current_run_id, user, db):
        current = agent._load_run(current_run_id, user)
        observed_months.append(str(current.get("salary_month")))
        current.update({
            "status": "execution_incomplete", "code": "VALIDATION_FAILED",
            "execution_result": {"status": "execution_incomplete", "code": "VALIDATION_FAILED"},
        })
        agent._save_run(current)
        return agent._public_run(current)

    monkeypatch.setattr(agent, "_execute_agent_run_sync", fake_execute)
    agent._run_agent_workflow(run_id, "tenant-a")

    completed = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert observed_months == ["2026.05"]
    assert completed["salary_month"] == "2026.05"
    assert completed["target_salary_month"] == "2026.05"


def test_background_workflow_preserves_month_question_without_resolved_target(
    tmp_path, monkeypatch,
) -> None:
    run_id = "a" * 32
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: object())
    monkeypatch.setattr(agent.ModelConfig, "fallback_from_env", lambda: None)
    agent._save_run({
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "planning", "instruction": "处理这批文件", "events": [], "items": [],
        "month_authority": "ambiguous_instruction", "draft_filename": "checkpoint.xlsx",
        "model_plan": {
            "summary": "计划", "steps": ["写入"],
            "questions": ["目标处理月份是2026年4月还是2026年5月？"],
        },
        "plan_confirmation": {"required": True, "confirmed": False},
        "month_confirmation": {"required": True, "confirmed": False},
    })
    execution_calls: list[str] = []
    monkeypatch.setattr(
        agent, "_execute_agent_run_sync",
        lambda current_run_id, *_args: execution_calls.append(current_run_id),
    )

    agent._run_agent_workflow(run_id, "tenant-a")

    blocked = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert blocked["status"] == "awaiting_review"
    assert blocked["model_plan"]["questions"] == ["目标处理月份是2026年4月还是2026年5月？"]
    assert execution_calls == []


def test_finalize_rejects_a_readable_draft_without_data_writes(tmp_path, monkeypatch) -> None:
    run_id = "zero-write-finalize"
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    output = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "模板"
    workbook.save(output)
    workbook.close()

    run = {
        "run_id": run_id,
        "tenant_id": "tenant-a",
        "project_id": "project-a",
        "status": "awaiting_review",
        "draft_filename": output.name,
        "events": [],
        "items": [],
        "workbook_updates": [],
        "execution_result": {"status": "completed", "content": "已完成"},
    }

    agent._finalize_agent_output(run)

    assert run["status"] == "execution_incomplete"
    assert run["code"] == "NO_WRITES_PERFORMED"
    assert run["validation"]["status"] == "not_verified"


def test_draft_acceptance_rejects_a_total_formula_that_omits_an_employee(
    tmp_path,
) -> None:
    output = tmp_path / "bad-summary-range.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["工号", "姓名", "应发工资"])
    sheet.append(["E001", "甲", 100])
    sheet.append(["E002", "乙", 200])
    sheet.append(["合计", None, "=SUM(C2:C2)"])
    workbook.save(output)
    workbook.close()

    acceptance = agent._inspect_draft_acceptance(output)

    assert acceptance["passed"] is False
    assert acceptance["summary_range_errors"] == [{
        "sheet": "明细", "cell": "C4", "formula": "=SUM(C2:C2)",
        "expected_rows": "2:3", "actual_rows": "2:2",
    }]


def test_draft_acceptance_avoids_read_only_random_cell_access(
    tmp_path,
    monkeypatch,
) -> None:
    output = tmp_path / "acceptance-performance.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["工号", "姓名", "应发工资"])
    sheet.append(["E001", "甲", 100])
    sheet.append(["合计", None, "=SUM(C2:C2)"])
    workbook.save(output)
    workbook.close()
    real_load_workbook = agent.openpyxl.load_workbook
    load_options: list[dict] = []

    def recording_load_workbook(*args, **kwargs):
        load_options.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(agent.openpyxl, "load_workbook", recording_load_workbook)

    agent._inspect_draft_acceptance(output)

    assert len(load_options) == 2
    assert all(options.get("read_only") is False for options in load_options)


def test_roster_acceptance_avoids_read_only_random_cell_access(
    tmp_path,
    monkeypatch,
) -> None:
    output = tmp_path / "roster-acceptance-performance.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["序号", "姓名"])
    sheet.append([1, "甲"])
    sheet.append([2, "乙"])
    sheet.append(["*合计*", None])
    workbook.save(output)
    workbook.close()
    real_load_workbook = agent.openpyxl.load_workbook
    load_options: list[dict] = []

    def recording_load_workbook(*args, **kwargs):
        load_options.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(agent.openpyxl, "load_workbook", recording_load_workbook)
    result = agent._verify_roster_acceptance({
        "instruction": "以派遣名单为准",
        "roster_sync": {
            "retained_names": ["甲", "乙"],
            "target_sheet": "明细",
        },
    }, output)

    assert result["passed"] is True
    assert load_options == [{"read_only": False, "data_only": False}]


def test_draft_acceptance_accepts_split_ranges_that_cover_every_employee(
    tmp_path,
) -> None:
    output = tmp_path / "split-summary-range.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["工号", "姓名", "应发工资"])
    sheet.append(["E001", "甲", 100])
    sheet.append(["E002", "乙", 200])
    sheet.append(["合计", None, "=SUM(C2:C2)+SUM(C3:C3)"])
    workbook.save(output)
    workbook.close()

    acceptance = agent._inspect_draft_acceptance(output)

    assert acceptance["summary_range_errors"] == []
    assert acceptance["passed"] is True


def test_draft_acceptance_rejects_a_total_formula_that_includes_itself(
    tmp_path,
) -> None:
    output = tmp_path / "self-referencing-summary.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["工号", "姓名", "应发工资"])
    sheet.append(["E001", "甲", 100])
    sheet.append(["E002", "乙", 200])
    sheet.append(["合计", None, "=SUM(C:C)"])
    workbook.save(output)
    workbook.close()

    acceptance = agent._inspect_draft_acceptance(output)

    assert acceptance["passed"] is False
    assert acceptance["summary_range_errors"][0]["cell"] == "C4"
    assert acceptance["summary_range_errors"][0]["reason"] == "summary_range_includes_total"


def test_draft_acceptance_rejects_a_formula_that_references_a_missing_sheet(
    tmp_path,
) -> None:
    output = tmp_path / "missing-sheet-reference.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "='不存在的表'!A1"
    workbook.save(output)
    workbook.close()

    acceptance = agent._inspect_draft_acceptance(output)

    assert acceptance["passed"] is False
    assert acceptance["formula_errors"] == [{
        "sheet": "Sheet", "cell": "A1", "error": "missing_sheet_reference",
        "referenced_sheet": "不存在的表",
    }]


def _cross_month_finalize_run(tmp_path, monkeypatch):
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    baseline = tmp_path / "april-accepted.xlsx"
    output = tmp_path / "may-draft.xlsx"
    for path in (baseline, output):
        workbook = openpyxl.Workbook()
        workbook.active.title = "汇总"
        workbook.active["A2"] = "5月" if path == output else "4月"
        workbook.active["B2"] = 100 if path == output else 50
        if path == output:
            workbook.active["A3"] = "4月"
            workbook.active["B3"] = 50
            workbook.active["A4"] = "3月"
            workbook.active["B4"] = 40
        workbook.save(path)
        workbook.close()
    baseline_sha = agent.file_digest(baseline)
    run = {
        "run_id": "month-finalize", "project_id": "project-a",
        "status": "awaiting_review", "draft_filename": output.name,
        "salary_month": "2026.05", "target_salary_month": "2026.05",
        "instruction": "新增5月并更新本月数据", "events": [], "items": [],
        "requires_month_roll_forward": True,
        "baseline": {"salary_month": "2026.04", "sha256": baseline_sha},
        "month_roll_forward": {
            "sheet": "汇总", "current_row": 2, "history_row": 3,
            "new_month": "5月", "target_salary_month": "2026.05",
            "baseline_sha256": baseline_sha,
            "history_values": {"A3": "4月", "B3": 50},
            "history_snapshot": {
                "A3": "4月", "B3": 50,
                "A4": "3月", "B4": 40,
            },
        },
        "_immutable_inputs": [{
            "role": "previous_accepted_result", "path": str(baseline), "sha256": baseline_sha,
        }],
        "workbook_updates": [
            {"rule": "roll_forward_month", "sheet": "汇总", "inserted_row": 2, "history_row": 3},
            {"target_sheet": "汇总", "target_cell": "B2", "old_value": 50, "new_value": 100},
        ],
        "execution_result": {"status": "completed", "content": "完成"},
        "plan_confirmation": {"required": True, "confirmed": True},
        "model_plan": {"steps": ["滚月", "写入", "校验"], "questions": []},
    }
    return run, baseline, output


def test_finalize_accepts_cross_month_run_only_when_history_and_inputs_are_unchanged(
    tmp_path, monkeypatch,
) -> None:
    run, _baseline, _output = _cross_month_finalize_run(tmp_path, monkeypatch)

    agent._finalize_agent_output(run)

    assert run["status"] == "completed"
    assert run["machine_acceptance"]["target_month_present"] is True
    assert run["machine_acceptance"]["historical_months_unchanged"] is True
    assert run["machine_acceptance"]["immutable_inputs_unchanged"] is True


def test_finalize_rejects_cross_month_run_with_only_a_month_row_insert(
    tmp_path, monkeypatch,
) -> None:
    run, _baseline, _output = _cross_month_finalize_run(tmp_path, monkeypatch)
    run["workbook_updates"] = [run["workbook_updates"][0]]

    agent._finalize_agent_output(run)

    assert run["status"] == "execution_incomplete"
    assert run["code"] == "NO_WRITES_PERFORMED"
    assert run["machine_acceptance"]["writes_traceable"] is False


def test_finalize_rejects_realistic_incomplete_model_report_after_row_insert(
    tmp_path, monkeypatch,
) -> None:
    """A model report saying it could not write must never publish a partial row insert."""
    output = tmp_path / "partial.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "姓名"
    workbook.active["A2"] = "占位"
    workbook.save(output)
    workbook.close()
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    run = {
        "run_id": "partial-row-insert",
        "project_id": "project-a",
        "status": "awaiting_review",
        "draft_filename": output.name,
        "salary_month": "2026.04",
        "target_salary_month": "2026.04",
        "instruction": "把新增人员写入总表",
        "events": [],
        "items": [],
        "workbook_updates": [{
            "sheet": "Sheet",
            "source_row": 2,
            "inserted_row": 3,
            "copied_columns": 10,
        }],
        "execution_result": {
            "status": "completed",
            "content": (
                "已插入占位行，但未能完成业务字段写入。"
                "无法可靠写入其余工作表，必须报告阻塞缺口。"
                "结论：本任务未完成。"
            ),
        },
        "plan_confirmation": {"required": True, "confirmed": True},
        "model_plan": {"steps": ["插入人员", "写入业务字段", "校验"], "questions": []},
        "validation": {"status": "not_verified"},
    }

    agent._finalize_agent_output(run)

    assert run["status"] == "execution_incomplete"
    assert run["code"] == "INCOMPLETE_MODEL_RESPONSE"


def test_finalize_rejects_cross_month_run_when_frozen_history_was_changed(
    tmp_path, monkeypatch,
) -> None:
    run, _baseline, output = _cross_month_finalize_run(tmp_path, monkeypatch)
    workbook = openpyxl.load_workbook(output)
    workbook["汇总"]["B3"] = 999
    workbook.save(output)
    workbook.close()

    agent._finalize_agent_output(run)

    assert run["status"] == "execution_incomplete"
    assert run["code"] == "MONTH_HISTORY_CHANGED"


def test_finalize_rejects_cross_month_run_when_older_history_was_changed(
    tmp_path, monkeypatch,
) -> None:
    run, _baseline, output = _cross_month_finalize_run(tmp_path, monkeypatch)
    workbook = openpyxl.load_workbook(output)
    workbook["汇总"]["B4"] = 999
    workbook.save(output)
    workbook.close()

    agent._finalize_agent_output(run)

    assert run["status"] == "execution_incomplete"
    assert run["code"] == "MONTH_HISTORY_CHANGED"


def test_finalize_rejects_cross_month_run_when_prior_accepted_result_changed(
    tmp_path, monkeypatch,
) -> None:
    run, baseline, _output = _cross_month_finalize_run(tmp_path, monkeypatch)
    baseline.write_bytes(b"changed")

    agent._finalize_agent_output(run)

    assert run["status"] == "execution_incomplete"
    assert run["code"] == "IMMUTABLE_INPUT_CHANGED"


def test_finalize_requires_every_skill_validation_tool(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    output = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "已更新"
    workbook.save(output)
    workbook.close()
    run = {
        "run_id": "validation-finalize", "project_id": "project-a",
        "status": "awaiting_review", "draft_filename": output.name,
        "instruction": "更新数据", "events": [
            {"type": "model_request", "payload": {"stage": "writing", "operation": "model_call"}},
            {"type": "tool_result", "payload": {"name": "validate_workbook", "status": "succeeded"}},
        ],
        "items": [],
        "workbook_updates": [{"target_sheet": "Sheet", "target_cell": "A1", "old_value": "旧", "new_value": "已更新"}],
        "execution_result": {"status": "completed", "content": "完成"},
        "execution_checkpoint": {
            "current_stage": "completed",
            "required_validation_tools": ["validate_workbook", "validate_with_officecli"],
        },
        "plan_confirmation": {"required": True, "confirmed": True},
        "model_plan": {"steps": ["写入"], "questions": []},
    }

    agent._finalize_agent_output(run)

    assert run["status"] == "awaiting_review"
    assert run["code"] == "VALIDATION_TOOL_NOT_EXECUTED"
    assert run["validation"]["missing_validation_tools"] == ["validate_with_officecli"]


def test_finalize_requires_skill_validations_for_deterministic_writes(
    tmp_path, monkeypatch,
) -> None:
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    output = tmp_path / "deterministic-draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "已更新"
    workbook.save(output)
    workbook.close()
    run = {
        "run_id": "deterministic-validation-finalize",
        "project_id": "project-a",
        "status": "awaiting_review",
        "draft_filename": output.name,
        "instruction": "按来源名单更新数据",
        "events": [{
            "type": "tool_result",
            "payload": {
                "name": "source_sheet_roster_sync",
                "status": "succeeded",
                "write_count": 1,
            },
        }],
        "items": [],
        "workbook_updates": [{
            "target_sheet": "Sheet",
            "target_cell": "A1",
            "old_value": "旧",
            "new_value": "已更新",
        }],
        "execution_result": {"status": "completed", "content": "完成"},
        "execution_checkpoint": {
            "current_stage": "completed",
            "required_validation_tools": ["validate_workbook", "validate_with_officecli"],
        },
        "plan_confirmation": {"required": True, "confirmed": True},
        "model_plan": {"steps": ["名单同步"], "questions": []},
    }

    agent._finalize_agent_output(run)

    assert run["status"] == "awaiting_review"
    assert run["code"] == "VALIDATION_TOOL_NOT_EXECUTED"
    assert run["validation"]["missing_validation_tools"] == [
        "validate_with_officecli", "validate_workbook",
    ]


def test_finalize_ignores_validations_that_ran_before_the_last_write(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    output = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "已更新"
    workbook.save(output)
    workbook.close()
    run = {
        "run_id": "stale-validation-finalize", "project_id": "project-a",
        "status": "awaiting_review", "draft_filename": output.name,
        "instruction": "更新数据", "events": [
            {"type": "model_request", "payload": {"stage": "writing", "operation": "model_call"}},
            {"type": "tool_result", "payload": {
                "name": "validate_workbook", "status": "succeeded",
            }},
            {"type": "tool_result", "payload": {
                "name": "validate_with_officecli", "status": "succeeded",
            }},
            {"type": "tool_result", "payload": {
                "name": "apply_source_cells", "status": "succeeded", "write_count": 1,
            }},
        ],
        "items": [],
        "workbook_updates": [{
            "target_sheet": "Sheet", "target_cell": "A1",
            "old_value": "旧", "new_value": "已更新",
        }],
        "execution_result": {"status": "completed", "content": "完成"},
        "execution_checkpoint": {
            "current_stage": "completed",
            "required_validation_tools": ["validate_workbook", "validate_with_officecli"],
        },
        "plan_confirmation": {"required": True, "confirmed": True},
        "model_plan": {"steps": ["写入"], "questions": []},
    }

    agent._finalize_agent_output(run)

    assert run["status"] == "awaiting_review"
    assert run["code"] == "VALIDATION_TOOL_NOT_EXECUTED"
    assert run["validation"]["missing_validation_tools"] == [
        "validate_with_officecli", "validate_workbook",
    ]


def test_finalize_rejects_explicit_roster_scope_without_roster_evidence(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    output = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "已更新"
    workbook.save(output)
    workbook.close()
    run = {
        "run_id": "roster-finalize", "project_id": "project-a",
        "status": "awaiting_review", "draft_filename": output.name,
        "instruction": "以派遣名单为准，其他人员全部删除", "events": [], "items": [],
        "workbook_updates": [{"target_sheet": "Sheet", "target_cell": "A1", "old_value": "旧", "new_value": "已更新"}],
        "execution_result": {"status": "completed", "content": "完成"},
        "plan_confirmation": {"required": True, "confirmed": True},
        "model_plan": {"steps": ["名单同步"], "questions": []},
    }

    agent._finalize_agent_output(run)

    assert run["status"] == "awaiting_review"
    assert run["code"] == "ROSTER_VALIDATION_FAILED"


def test_background_workflow_does_not_auto_resume_after_turn_limit(tmp_path, monkeypatch) -> None:
    import openpyxl

    run_id = "f1" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    agent._save_run({
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "planning", "instruction": "更新", "events": [], "items": [],
        "model_plan": {"summary": "计划", "steps": ["写入"], "questions": []},
        "plan_confirmation": {"required": True, "confirmed": True},
        "month_confirmation": {"required": False, "confirmed": True},
    })
    calls = 0

    def fake_execute(current_run_id, user, db):
        nonlocal calls
        calls += 1
        current = agent._load_run(current_run_id, user)
        if calls <= 7:
            current.update({
                "status": "execution_incomplete",
                "code": "MAX_TURNS_EXCEEDED",
                "execution_result": {"status": "execution_incomplete", "code": "MAX_TURNS_EXCEEDED"},
            })
        else:
            output = tmp_path / "continued-after-many-segments.xlsx"
            workbook = openpyxl.Workbook()
            workbook.save(output)
            workbook.close()
            current.update({
                "status": "awaiting_review", "draft_filename": output.name,
                "execution_result": {"status": "completed", "content": "完成"},
                "workbook_updates": [],
            })
        agent._save_run(current)
        return agent._public_run(current)

    monkeypatch.setattr(agent, "_execute_agent_run_sync", fake_execute)
    agent._run_agent_workflow(run_id, "tenant-a")

    completed = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert calls == 1
    assert completed["status"] == "execution_incomplete"
    assert completed["code"] == "MAX_TURNS_EXCEEDED"
    assert not any(event.get("payload", {}).get("stage") == "segment_resume" for event in completed["events"])


def test_background_workflow_waits_for_user_after_model_provider_error(tmp_path, monkeypatch) -> None:
    run_id = "e2" * 16
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / filename)
    monkeypatch.setattr(agent, "_material_context", lambda _run: [])

    class FakeDb:
        def close(self):
            return None

    monkeypatch.setattr(agent, "SessionLocal", FakeDb)
    agent._save_run({
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-a",
        "status": "planning", "instruction": "更新", "events": [], "items": [],
        "model_plan": {"summary": "计划", "steps": ["写入"], "questions": []},
        "plan_confirmation": {"required": True, "confirmed": True},
        "month_confirmation": {"required": False, "confirmed": True},
    })
    calls = 0

    def fake_execute(current_run_id, user, db):
        nonlocal calls
        calls += 1
        current = agent._load_run(current_run_id, user)
        if calls == 1:
            current.update({
                "status": "execution_incomplete",
                "code": "MODEL_PROVIDER_ERROR",
                "execution_result": {"status": "execution_incomplete", "code": "MODEL_PROVIDER_ERROR"},
            })
        else:
            current.update({
                "status": "completed",
                "execution_result": {"status": "completed", "content": "完成"},
                "workbook_updates": [{"sheet": "工资核算", "cell": "A1"}],
            })
        agent._save_run(current)
        return agent._public_run(current)

    monkeypatch.setattr(agent, "_execute_agent_run_sync", fake_execute)
    agent._run_agent_workflow(run_id, "tenant-a")

    completed = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert calls == 1
    assert completed["status"] == "execution_incomplete"
    assert completed["code"] == "MODEL_PROVIDER_ERROR"
    assert not any(event.get("payload", {}).get("stage") == "segment_resume" for event in completed["events"])
