from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from typing import Any

import openpyxl
from openpyxl.styles import Font
import pytest

from core.document_agent.orchestrator import ToolExecutionError
from core.document_agent.workbook_session import WorkbookSession


def _formula_session(tmp_path: Path, formula: Any = "=A21/21") -> SimpleNamespace:
    master = tmp_path / "原始总表.xlsx"
    source = tmp_path / "来源.xlsx"
    draft = tmp_path / "公式参数草稿.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "工资核算"
    sheet["J2"] = formula
    sheet["J2"].number_format = "#,##0.00"
    sheet["J2"].font = Font(bold=True)
    sheet["K2"] = "=B21/21"
    sheet["A21"] = 2100
    workbook.create_sheet("历史记录")["A1"] = "历史数据不变"
    workbook.save(master)
    workbook.close()
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 2000
    workbook.save(source)
    workbook.close()
    originals = {master: master.read_bytes(), source: source.read_bytes()}
    session = WorkbookSession(master, draft, {source.name: source})
    session.prepare_copy()
    return SimpleNamespace(session=session, draft=draft, originals=originals)


def _change(**overrides: Any) -> dict[str, Any]:
    return {
        "target_sheet": "工资核算", "target_cell": "J2", "expected_formula": "=A21/21",
        "old_divisor": 21, "new_divisor": 23, **overrides,
    }


@pytest.mark.parametrize("formula,expected", [
    ("=A21/21", "=A21/23"),
    ('=A21/21+210/210+SUM(A1:A21)/21+IF(A1="/21",21,0)',
     '=A21/23+210/210+SUM(A1:A21)/23+IF(A1="/21",21,0)'),
    ('=IF(A1="say ""/21""",A21/21,A2/210)', '=IF(A1="say ""/21""",A21/23,A2/210)'),
    ("=ROUND(A21 / 21,2)+A2/210+A3/21.5", "=ROUND(A21 / 23,2)+A2/210+A3/21.5"),
])
def test_divisor_changes_only_matching_numeric_tokens_after_division(
    tmp_path: Path, formula: str, expected: str,
) -> None:
    case = _formula_session(tmp_path, formula)

    updates = case.session.apply_formula_divisors([_change(expected_formula=formula)])

    assert isinstance(updates, list)
    assert len(updates) == 1
    assert updates[0]["old_value"] == formula
    assert updates[0]["new_value"] == expected
    assert updates[0]["target_sheet"] == "工资核算"
    assert updates[0]["target_cell"] == "J2"
    workbook = openpyxl.load_workbook(case.draft, data_only=False)
    try:
        sheet = workbook["工资核算"]
        assert sheet["J2"].value == expected
        assert sheet["J2"].data_type == "f"
        assert sheet["J2"].number_format == "#,##0.00"
        assert sheet["J2"].font.bold is True
        assert sheet["K2"].value == "=B21/21"
        assert sheet["A21"].value == 2100
        assert workbook["历史记录"]["A1"].value == "历史数据不变"
    finally:
        workbook.close()
    assert {path: path.read_bytes() for path in case.originals} == case.originals


@pytest.mark.parametrize("formula", [
    "=A21+21+210", "=A1/210", "=A1/21.5", "=A1/A21", '=IF(A1="/21",21,210)',
])
def test_formula_without_matching_divisor_is_rejected_without_writes(tmp_path: Path, formula: str) -> None:
    case = _formula_session(tmp_path, formula)
    before = case.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        case.session.apply_formula_divisors([_change(expected_formula=formula)])

    assert case.draft.read_bytes() == before
    assert {path: path.read_bytes() for path in case.originals} == case.originals


def test_non_formula_target_cannot_be_replaced_by_a_formula(tmp_path: Path) -> None:
    case = _formula_session(tmp_path, 2100)
    before = case.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        case.session.apply_formula_divisors([_change()])

    assert case.draft.read_bytes() == before
    assert {path: path.read_bytes() for path in case.originals} == case.originals


@pytest.mark.parametrize("overrides", [
    pytest.param({"expected_formula": "=A21/21+1"}, id="stale-formula"),
    pytest.param({"new_divisor": 0}, id="zero-new-divisor"),
    pytest.param({"new_divisor": "23"}, id="string-new-divisor"),
    pytest.param({"old_divisor": "21"}, id="string-old-divisor"),
    pytest.param({"new_divisor": True}, id="boolean-new-divisor"),
    pytest.param({"old_divisor": True}, id="boolean-old-divisor"),
    pytest.param({"new_divisor": float("nan")}, id="nan-new-divisor"),
    pytest.param({"new_divisor": float("inf")}, id="infinite-new-divisor"),
    pytest.param({"expected_formula": 21}, id="non-string-expected-formula"),
    pytest.param({"target_sheet": "不存在"}, id="unknown-target-sheet"),
    pytest.param({"target_cell": "J0"}, id="invalid-target-cell"),
    pytest.param({"target_cell": "J2:K2"}, id="range-not-single-cell"),
    pytest.param({"formula": "=arbitrary_replacement()"}, id="unknown-parameter"),
])
def test_invalid_or_stale_formula_change_is_rejected_atomically(
    tmp_path: Path, overrides: dict[str, Any],
) -> None:
    case = _formula_session(tmp_path)
    before = case.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        case.session.apply_formula_divisors([_change(**overrides)])

    assert case.draft.read_bytes() == before
    assert {path: path.read_bytes() for path in case.originals} == case.originals


def test_duplicate_target_in_one_formula_batch_is_rejected(tmp_path: Path) -> None:
    case = _formula_session(tmp_path)
    before = case.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        case.session.apply_formula_divisors([_change(), _change()])

    assert case.draft.read_bytes() == before
    assert {path: path.read_bytes() for path in case.originals} == case.originals


@pytest.mark.parametrize("second_change", [
    _change(target_cell="K2", expected_formula="=B21/21", old_divisor=22),
    _change(target_cell="K2", expected_formula="=B21/21+1"),
])
def test_later_formula_failure_does_not_persist_earlier_valid_change(
    tmp_path: Path, second_change: dict[str, Any],
) -> None:
    case = _formula_session(tmp_path)
    before = case.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        case.session.apply_formula_divisors([_change(), second_change])

    assert case.draft.read_bytes() == before
    assert {path: path.read_bytes() for path in case.originals} == case.originals
