from __future__ import annotations

from pathlib import Path
from datetime import datetime
import json

import openpyxl
import pytest


def test_basic_processor_updates_explicit_payroll_sources_and_returns_a_change_log(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    from scripts.keyuan_basic_processor import process_keyuan_basics

    master_path = tmp_path / "master.xlsx"
    master = openpyxl.Workbook()
    payroll = master.active
    payroll.title = "工资核算"
    payroll.cell(4, 2).value = "张三"
    attendance = master.create_sheet("考勤")
    attendance.cell(2, 2).value = "张三"
    duty = master.create_sheet("配送员值班费")
    duty.cell(2, 2).value = "张三"
    adjustment = master.create_sheet("其他调差累计")
    master.save(master_path)
    master.close()

    source_path = tmp_path / "薪资数据-科园-7月薪资.xlsx"
    source = openpyxl.Workbook()
    bonus = source.active
    bonus.title = "奖金-7月"
    bonus.cell(4, 4).value = "张三"
    bonus.cell(4, 10).value = 100
    bonus.cell(4, 11).value = 20
    bonus.cell(4, 12).value = 30
    source_attendance = source.create_sheet("考勤-7月")
    source_attendance.cell(9, 1).value = "张三"
    source_attendance.cell(9, 7).value = 21
    source_attendance.cell(9, 17).value = 2
    source_attendance.cell(9, 30).value = 3
    source_attendance.cell(9, 32).value = 4
    source_attendance.cell(9, 33).value = 5
    source_attendance.cell(9, 34).value = 6
    source_duty = source.create_sheet("值班")
    source_duty.cell(3, 3).value = "张三"
    source_duty.cell(4, 3).value = "张三"
    source_adjustment = source.create_sheet("补发补扣")
    source_adjustment.cell(1, 1).value = "张三"
    source_adjustment.cell(1, 2).value = "加班补发 775.98元"
    source.save(source_path)
    source.close()

    output_path = tmp_path / "processed.xlsx"
    result = process_keyuan_basics(master_path, source_path, output_path)

    assert output_path.is_file()
    assert result["change_count"] >= 10
    assert {change["rule"] for change in result["changes"]} >= {"奖金 J+K/L", "考勤字段", "值班次数", "补发补扣"}
    output = openpyxl.load_workbook(output_path, data_only=False)
    assert output["工资核算"].cell(4, 30).value == 120
    assert output["工资核算"].cell(4, 31).value == 30
    assert output["考勤"].cell(2, 7).value == 21
    assert output["考勤"].cell(2, 9).value == 7
    assert output["考勤"].cell(2, 16).value == 6
    assert output["配送员值班费"].cell(2, 4).value == 2
    assert [output["其他调差累计"].cell(2, column).value for column in range(1, 4)] == ["张三", 775.98, "加班补发 775.98元"]
    output.close()

    from scripts.keyuan_basic_processor import main

    cli_output = tmp_path / "cli-processed.xlsx"
    assert main([str(master_path), str(source_path), str(cli_output)]) == 0
    assert json.loads(capsys.readouterr().out)["change_count"] >= 10


def test_basic_processor_never_overwrites_the_salary_source(tmp_path: Path) -> None:
    from scripts.keyuan_basic_processor import process_keyuan_basics

    master_path = tmp_path / "master.xlsx"
    master = openpyxl.Workbook()
    master.active.title = "工资核算"
    master.create_sheet("考勤")
    master.create_sheet("配送员值班费")
    master.create_sheet("其他调差累计")
    master.save(master_path)
    master.close()

    source_path = tmp_path / "salary.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "奖金-7月"
    for sheet in ("考勤-7月", "值班", "补发补扣"):
        source.create_sheet(sheet)
    source.save(source_path)
    original = source_path.read_bytes()
    source.close()

    with pytest.raises(ValueError, match="独立草稿"):
        process_keyuan_basics(master_path, source_path, source_path)

    assert source_path.read_bytes() == original


def test_basic_processor_accepts_a_different_month_sheet_suffix(tmp_path: Path) -> None:
    from scripts.keyuan_basic_processor import process_keyuan_basics

    master_path = tmp_path / "master.xlsx"
    master = openpyxl.Workbook()
    payroll = master.active
    payroll.title = "工资核算"
    payroll.cell(4, 2).value = "张三"
    attendance = master.create_sheet("考勤")
    attendance.cell(2, 2).value = "张三"
    duty = master.create_sheet("配送员值班费")
    duty.cell(2, 2).value = "张三"
    master.create_sheet("其他调差累计")
    master.save(master_path)
    master.close()

    source_path = tmp_path / "202608薪资更新.xlsx"
    source = openpyxl.Workbook()
    bonus = source.active
    bonus.title = "奖金-8月"
    bonus.cell(4, 4).value = "张三"
    bonus.cell(4, 10).value = 100
    source_attendance = source.create_sheet("考勤-8月")
    source_attendance.cell(9, 1).value = "张三"
    source_attendance.cell(9, 7).value = 22
    source.create_sheet("值班")
    source.create_sheet("补发补扣")
    source.save(source_path)
    source.close()

    result = process_keyuan_basics(master_path, source_path, tmp_path / "processed.xlsx")

    assert result["status"] in {"passed", "needs_review"}
    output = openpyxl.load_workbook(tmp_path / "processed.xlsx", data_only=False)
    assert output["工资核算"].cell(4, 30).value == 100
    assert output["考勤"].cell(2, 7).value == 22
    output.close()


def test_basic_processor_rolls_oa_history_down_and_adds_july(tmp_path: Path) -> None:
    from scripts.keyuan_basic_processor import process_keyuan_basics

    master_path = tmp_path / "master.xlsx"
    master = openpyxl.Workbook()
    master.active.title = "工资核算"
    master.create_sheet("考勤")
    master.create_sheet("配送员值班费")
    master.create_sheet("其他调差累计")
    oa = master.create_sheet("OA请款及审批")
    for row, month in enumerate(("6月", "5月", "4月"), start=8):
        oa.cell(row, 1).value = month
        oa.cell(row, 2).value = f"value-{month}"
    # The current-month row uses formulas; history must snapshot its value
    # instead of inheriting the formula and following the new July total.
    oa["C8"] = "=C5"
    oa["C5"] = 100
    master.save(master_path)
    master.close()

    source_path = tmp_path / "salary.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "奖金-7月"
    source.create_sheet("考勤-7月")
    source.create_sheet("值班")
    source.create_sheet("补发补扣")
    source.save(source_path)
    source.close()

    output_path = tmp_path / "processed.xlsx"
    result = process_keyuan_basics(master_path, source_path, output_path)

    workbook = openpyxl.load_workbook(output_path, data_only=False)
    oa = workbook["OA请款及审批"]
    assert oa["A8"].value == "7月"
    assert oa["A9"].value == "6月"
    assert oa["A10"].value == "5月"
    assert oa["A11"].value == "4月"
    assert oa["B9"].value == "value-6月"
    assert oa["C8"].value == "=C5"
    assert oa["C9"].value != "=C5"
    assert any(change["rule"] == "OA新增7月" for change in result["changes"])
    workbook.close()


def test_basic_processor_merges_all_mapped_salary_sheets(tmp_path: Path) -> None:
    from scripts.keyuan_basic_processor import process_keyuan_basics

    master_path = tmp_path / "master.xlsx"
    master = openpyxl.Workbook()
    payroll = master.active
    payroll.title = "工资核算"
    payroll.cell(4, 2).value = "李文慧"
    payroll.cell(4, 17).value = "3000/年"
    attendance = master.create_sheet("考勤")
    attendance.cell(2, 2).value = "李楠"
    attendance.cell(2, 11).value = 9
    attendance.cell(2, 17).value = "常白班"
    duty = master.create_sheet("配送员值班费")
    duty.cell(2, 2).value = "李文慧"
    adjustment = master.create_sheet("其他调差累计")
    personnel = master.create_sheet("人员异动")
    personnel.append(["入职情况："])
    personnel.append(["序号", "姓 名", "员工工号", "部门", "组别", "职 务", "人员类别", "入职日期", "试用结束日期"])
    personnel.append([1, "李楠", "002056", "北京益药科园大药房有限公司", "北京门店", "收银员", "合同制人员", datetime(2026, 5, 12), datetime(2026, 11, 11)])
    personnel.append([None] * 9)
    personnel.append(["离职情况:"])
    personnel.append(["序号", "姓名", "员工工号", "部门", "组别", "职务", "人员类别", "入职日期", "离职日期", "离职原因"])
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append(["内部调动:"])
    personnel.append(["序号", "姓 名", "员工编号", "调出部门", "原组别", "原职务", "调入部门", "现组别", "现职务", "生效日期"])
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    personnel.append([None] * 10)
    heat = master.create_sheet("防暑降温费")
    heat.cell(2, 2).value = "李楠"
    heat.cell(2, 10).value = "旧标准"
    master.save(master_path)
    master.close()

    source_path = tmp_path / "薪资数据-科园-7月薪资.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "奖金-7月"
    source.create_sheet("考勤-7月")
    source.create_sheet("值班")
    source.create_sheet("补发补扣")
    transfer = source.create_sheet("入离职、转岗、转正、其他")
    transfer.append(["入职"])
    transfer.append(["姓名", "新工号", "新公司", "新成本中心/部门", "入职时间", "试用期结束时间", "岗位"])
    transfer.append(["李楠", "002228", "北京益药科园大药房有限公司", "北京大药房", datetime(2026, 5, 12), datetime(2026, 11, 11), "收银员"])
    transfer.append([None])
    transfer.append(["内部调动"])
    transfer.append(["序号", "姓 名", "员工编号", "调出部门", "原组别", "原职务", "调入部门", "现组别", "现职务", "生效日期", "月薪"])
    transfer.append([1, "李文慧", None, "客服组", None, "患者服务负责人", "客服组", None, "CAR-T执业药师", datetime(2026, 7, 1), 1050])
    high_temp = source.create_sheet("高温费")
    high_temp.append(["人员", "高温费（元/年）"])
    high_temp.append(["李楠", 1200])
    benefits = source.create_sheet("年节福利")
    benefits.append(["人员", "年节费（元/年）"])
    benefits.append(["李文慧", 0])
    work_cycle = source.create_sheet("科园做一休一人员")
    work_cycle.append(["李楠"])
    ignored = source.create_sheet("迟到早退旷工")
    ignored.append(["迟到"])
    ignored.append(["姓名", "迟到次数（求和）", None, None, None, "豁免"])
    ignored.append(["李楠", 3, None, None, None, "豁免"])
    source.save(source_path)
    source.close()

    output_path = tmp_path / "processed.xlsx"
    result = process_keyuan_basics(master_path, source_path, output_path)

    output = openpyxl.load_workbook(output_path, data_only=False)
    assert output["人员异动"]["C3"].value == "002228"
    assert output["工资核算"]["Y4"].value == 1050
    assert output["工资核算"]["AA4"].value == 1050
    assert output["工资核算"]["Q4"].value == "0/年"
    assert output["考勤"]["Q2"].value == "上一休一"
    assert output["考勤"]["K2"].value == 9
    assert output["防暑降温费"]["J2"].value == 1200
    assert any(change["rule"] in {"入职", "内部调动", "离职", "跨公司调动", "其他"} for change in result["changes"])
    output.close()
