from datetime import datetime
import json
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import quote

import openpyxl
import pytest
from fastapi import HTTPException

import backend.routers.financial_workbooks as financial_workbooks
from backend.main import app
from backend.models import Project, UploadFile
from backend.routers.projects import _ALLOWED_TYPES, _validate_declared_workbook_role


def test_generic_financial_master_role_accepts_a_non_payroll_workbook() -> None:
    assert "financial_master" in _ALLOWED_TYPES
    assert "financial_source" in _ALLOWED_TYPES

    _validate_declared_workbook_role(
        "financial_master",
        ["应收账款明细", "客户对账"],
    )


def test_payroll_template_role_still_requires_the_payroll_sheet() -> None:
    try:
        _validate_declared_workbook_role("template", ["应收账款明细"])
    except Exception as error:
        assert "工资核算" in str(error)
    else:
        raise AssertionError("工资模板角色不应接受非工资总表")


def test_generic_financial_integration_routes_are_publicly_registered() -> None:
    paths = set(app.openapi()["paths"])

    assert "/api/financial-workbooks/{project_id}/integrations" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/latest" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/progress" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/latest/download" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/latest/review/download" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/latest/draft/download" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/latest/release" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/published" in paths
    assert "/api/financial-workbooks/{project_id}/integrations/published/{version_id}/download" in paths
    assert "/api/financial-workbooks/work-queue" in paths


def test_financial_progress_write_retries_windows_file_lock(tmp_path, monkeypatch) -> None:
    target = tmp_path / "progress.json"
    actual_replace = financial_workbooks.os.replace
    attempts = 0

    def replace_with_transient_lock(source, destination):
        nonlocal attempts
        attempts += 1
        if attempts < 3:
            raise PermissionError(5, "Access is denied", str(destination))
        actual_replace(source, destination)

    monkeypatch.setattr(financial_workbooks.os, "replace", replace_with_transient_lock)

    financial_workbooks._write_json_atomically(target, {"status": "processing"})

    assert attempts == 3
    assert json.loads(target.read_text(encoding="utf-8")) == {"status": "processing"}


def test_financial_progress_lock_does_not_block_integration(monkeypatch) -> None:
    monkeypatch.setattr(
        financial_workbooks,
        "_write_json_atomically",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(PermissionError(5, "Access is denied")),
    )

    progress = financial_workbooks._write_progress(
        "project-1",
        status="processing",
        stages=[{"key": "prepare", "label": "准备文件", "status": "running"}],
        detail="正在准备文件",
    )

    assert progress["status"] == "processing"
    assert progress["detail"] == "正在准备文件"


class _FakeQuery:
    def __init__(self, records):
        self.records = records

    def filter(self, *_criteria):
        return self

    def first(self):
        return self.records[0] if self.records else None

    def all(self):
        return self.records


class _FakeDb:
    def __init__(self, project, files):
        self.project = project
        self.files = files

    def query(self, model):
        return _FakeQuery([self.project] if model is Project else self.files)


def test_generic_integration_excludes_legacy_payroll_source_files() -> None:
    master = SimpleNamespace(id="master", file_type="financial_master")
    financial_source = SimpleNamespace(id="financial", file_type="financial_source")
    payroll_source = SimpleNamespace(id="payroll", file_type="source")

    selected_master, selected_sources = financial_workbooks._select_integration_files(
        "project-id",
        _FakeDb(SimpleNamespace(id="project-id"), [master, financial_source, payroll_source]),
    )

    assert selected_master.id == "master"
    assert [file.id for file in selected_sources] == ["financial"]


def test_financial_work_queue_prioritizes_manual_review_and_lists_real_actions(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(financial_workbooks, "_META_DIR", tmp_path / "metadata")
    review_dir = tmp_path / "metadata" / "review-project"
    older_review_dir = tmp_path / "metadata" / "older-review-project"
    release_dir = tmp_path / "metadata" / "release-project"
    published_dir = tmp_path / "metadata" / "published-project"
    review_dir.mkdir(parents=True)
    older_review_dir.mkdir(parents=True)
    release_dir.mkdir(parents=True)
    published_dir.mkdir(parents=True)
    (review_dir / "latest.json").write_text(json.dumps({
        "status": "review_required", "issues": [{}, {}], "completed_at": "2026-08-24T10:00:00+00:00",
    }), encoding="utf-8")
    (older_review_dir / "latest.json").write_text(json.dumps({
        "status": "review_required", "issues": [{}], "completed_at": "2026-08-23T10:00:00+00:00",
    }), encoding="utf-8")
    (release_dir / "latest.json").write_text(json.dumps({
        "status": "ready_for_release", "issues": [], "completed_at": "2026-08-24T11:00:00+00:00",
    }), encoding="utf-8")
    (published_dir / "latest.json").write_text(json.dumps({
        "status": "published", "issues": [], "completed_at": "2026-08-24T12:00:00+00:00",
    }), encoding="utf-8")
    projects = [
        SimpleNamespace(id="release-project", name="8月应付", salary_month="2026.08"),
        SimpleNamespace(id="published-project", name="7月应付", salary_month="2026.07"),
        SimpleNamespace(id="older-review-project", name="7月应收", salary_month="2026.07"),
        SimpleNamespace(id="review-project", name="8月应收", salary_month="2026.08"),
    ]

    queue = financial_workbooks._build_financial_work_queue(projects)

    assert [item["project_id"] for item in queue] == ["review-project", "older-review-project", "release-project"]
    assert queue[0]["action_label"] == "查看 2 项问题并修正来源"
    assert queue[2]["action_label"] == "确认发布正式账单"


def test_financial_work_queue_marks_changed_input_files_for_reintegration(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(financial_workbooks, "_META_DIR", tmp_path / "metadata")
    project_id = "stale-project"
    result_dir = tmp_path / "metadata" / project_id
    result_dir.mkdir(parents=True)
    (result_dir / "latest.json").write_text(json.dumps({
        "status": "ready_for_release", "issues": [], "completed_at": "2026-08-24T10:00:00+00:00",
        "input_signature": "outdated-signature",
    }), encoding="utf-8")
    files = [
        SimpleNamespace(id="master", file_type="financial_master", stored_path="stale/master.xlsx", created_at=datetime(2026, 8, 24)),
        SimpleNamespace(id="source", file_type="financial_source", stored_path="stale/source.xlsx", created_at=datetime(2026, 8, 24)),
    ]

    queue = financial_workbooks._build_financial_work_queue(
        [SimpleNamespace(id=project_id, name="8月费用", salary_month="2026.08")],
        {project_id: files},
    )

    assert queue == [{
        "project_id": project_id,
        "project_name": "8月费用",
        "salary_month": "2026.08",
        "status": "blocked",
        "issue_count": 0,
        "action_label": "文件已变化，请重新整合",
        "completed_at": "2026-08-24T10:00:00+00:00",
    }]


def _write_workbook(path: Path, sheet_name: str, rows: list[list[object]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_generic_integration_updates_a_non_payroll_master_and_returns_review_items(tmp_path, monkeypatch) -> None:
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    project_id = "financial-project"
    (upload_root / project_id).mkdir(parents=True)
    master_path = upload_root / project_id / "master.xlsx"
    source_path = upload_root / project_id / "source.xlsx"
    _write_workbook(master_path, "客户应收台账", [["单据号", "客户", "金额"], ["AR-1", "甲公司", 100]])
    _write_workbook(source_path, "本月应收更新", [["单据编号", "客户", "金额"], ["AR-1", "甲公司", 120], ["AR-2", "乙公司", 80]])

    project = SimpleNamespace(id=project_id, owner=SimpleNamespace(tenant_id="tenant-a"))
    now = datetime(2026, 8, 24)
    master = SimpleNamespace(
        id="master", file_type="financial_master", stored_path=f"{project_id}/master.xlsx",
        original_name="甲公司应收总表.xlsx", created_at=now,
    )
    source = SimpleNamespace(
        id="source", file_type="financial_source", stored_path=f"{project_id}/source.xlsx",
        original_name="8月应收更新.xlsx", created_at=now,
    )
    monkeypatch.setattr(financial_workbooks, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(financial_workbooks, "EXPORT_DIR", str(export_root))
    monkeypatch.setattr(financial_workbooks, "_META_DIR", tmp_path / "metadata")

    result = financial_workbooks.create_financial_workbook_integration(
        project_id,
        user=SimpleNamespace(tenant_id="tenant-a"),
        db=_FakeDb(project, [master, source]),
    )

    assert result.auto_update_count == 1
    assert result.source_file_count == 1
    assert result.review_filename
    assert result.status == "review_required"
    assert result.release_checks == {
        "can_release": False,
        "issue_count": 1,
        "matched_sheet_count": 1,
        "auto_update_count": 1,
    }
    assert [issue["issue_type"] for issue in result.issues] == ["unknown_record"]
    assert result.updates == [{
        "target_sheet": "客户应收台账",
        "target_cell": "C2",
        "old_value": 100,
        "new_value": 120,
        "source_file": "source.xlsx",
        "source_sheet": "本月应收更新",
        "source_row": 2,
    }]
    output = openpyxl.load_workbook(export_root / project_id / result.filename, data_only=False)
    assert output["客户应收台账"]["C2"].value == 120
    output.close()

    review = openpyxl.load_workbook(export_root / project_id / result.review_filename, data_only=False)
    assert review["客户应收台账"]["C2"].comment is not None
    assert "原值：100" in review["客户应收台账"]["C2"].comment.text
    assert "新值：120" in review["客户应收台账"]["C2"].comment.text
    assert "修改记录" in review.sheetnames
    review.close()

    progress = financial_workbooks.get_financial_workbook_integration_progress(
        project_id,
        user=SimpleNamespace(tenant_id="tenant-a"),
        db=_FakeDb(project, [master, source]),
    )
    assert progress["status"] == "completed_with_issues"
    assert all(stage["status"] == "completed" for stage in progress["stages"])


def test_completed_financial_result_is_downloadable_without_manual_release(tmp_path, monkeypatch) -> None:
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    project_id = "release-project"
    (upload_root / project_id).mkdir(parents=True)
    _write_workbook(
        upload_root / project_id / "master.xlsx",
        "费用明细",
        [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]],
    )
    _write_workbook(
        upload_root / project_id / "source.xlsx",
        "费用调整",
        [["单据编号", "费用项目", "金额"], ["BX-001", "运输费", 120]],
    )
    project = SimpleNamespace(id=project_id, owner=SimpleNamespace(tenant_id="tenant-a"))
    now = datetime(2026, 8, 24)
    master = SimpleNamespace(id="master", file_type="financial_master", stored_path=f"{project_id}/master.xlsx", original_name="总表.xlsx", created_at=now)
    source = SimpleNamespace(id="source", file_type="financial_source", stored_path=f"{project_id}/source.xlsx", original_name="更新.xlsx", created_at=now)
    db = _FakeDb(project, [master, source])
    user = SimpleNamespace(id="user-1", name="复核员", tenant_id="tenant-a")
    monkeypatch.setattr(financial_workbooks, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(financial_workbooks, "EXPORT_DIR", str(export_root))
    monkeypatch.setattr(financial_workbooks, "_META_DIR", tmp_path / "metadata")

    result = financial_workbooks.create_financial_workbook_integration(project_id, user=user, db=db)

    assert result.status == "ready_for_release"
    response = financial_workbooks.download_latest_financial_workbook_integration(project_id, user=user, db=db)
    assert response.headers["content-disposition"].startswith("attachment;")
    assert quote("正式稿") in response.headers["content-disposition"]

    review_response = financial_workbooks.download_latest_financial_workbook_review(project_id, user=user, db=db)
    assert review_response.headers["content-disposition"].startswith("attachment;")
    assert quote("修订稿") in review_response.headers["content-disposition"]


def test_published_financial_workbook_is_archived_with_publisher_and_survives_input_changes(tmp_path, monkeypatch) -> None:
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    project_id = "version-project"
    (upload_root / project_id).mkdir(parents=True)
    _write_workbook(
        upload_root / project_id / "master.xlsx",
        "应付账款",
        [["单据号", "供应商", "金额"], ["AP-001", "物流公司", 100]],
    )
    _write_workbook(
        upload_root / project_id / "source.xlsx",
        "应付调整",
        [["单据编号", "供应商", "金额"], ["AP-001", "物流公司", 120]],
    )
    project = SimpleNamespace(id=project_id, owner=SimpleNamespace(tenant_id="tenant-a"))
    created_at = datetime(2026, 8, 24)
    master = SimpleNamespace(id="master", file_type="financial_master", stored_path=f"{project_id}/master.xlsx", original_name="总表.xlsx", created_at=created_at)
    source = SimpleNamespace(id="source", file_type="financial_source", stored_path=f"{project_id}/source.xlsx", original_name="调整.xlsx", created_at=created_at)
    db = _FakeDb(project, [master, source])
    user = SimpleNamespace(id="user-1", name="复核员", email="reviewer@example.com", tenant_id="tenant-a")
    monkeypatch.setattr(financial_workbooks, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(financial_workbooks, "EXPORT_DIR", str(export_root))
    monkeypatch.setattr(financial_workbooks, "_META_DIR", tmp_path / "metadata")

    financial_workbooks.create_financial_workbook_integration(project_id, user=user, db=db)
    released = financial_workbooks.release_latest_financial_workbook_integration(project_id, user=user, db=db)

    assert released.published_version is not None
    assert released.published_version.version_id
    assert released.published_version.published_by == {"id": "user-1", "name": "复核员"}
    assert len(released.published_version.sha256) == 64
    versions = financial_workbooks.list_published_financial_workbook_versions(project_id, user=user, db=db)
    assert [version.version_id for version in versions] == [released.published_version.version_id]

    source.created_at = datetime(2026, 8, 25)
    with pytest.raises(HTTPException, match="已变化"):
        financial_workbooks.get_latest_financial_workbook_integration(project_id, user=user, db=db)

    archived = financial_workbooks.download_published_financial_workbook_version(
        project_id,
        released.published_version.version_id,
        user=user,
        db=db,
    )
    assert archived.headers["content-disposition"].startswith("attachment;")

    (export_root / project_id / "published" / released.published_version.filename).write_bytes(b"tampered")
    with pytest.raises(HTTPException, match="校验不一致"):
        financial_workbooks.download_published_financial_workbook_version(
            project_id,
            released.published_version.version_id,
            user=user,
            db=db,
        )
