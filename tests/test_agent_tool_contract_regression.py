from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from backend.routers.agent import _build_run_tool_registry, _model_tool_schemas
from backend.agent_execution import _deduplicate_tool_schemas
from core.document_agent.contracts import ToolCall, ToolResult
from core.document_agent.model import ModelConfig, ModelProviderError, ModelResponse, OpenAICompatibleProvider
from core.document_agent.orchestrator import ModelOrchestrator, ToolRegistry
from core.document_agent.workbook_session import month_roll_forward_schema


def test_execution_tool_contract_has_unique_names_after_composition() -> None:
    """The provider must never receive duplicate function names.

    The generic execution phase composes the router's read contract with its
    write/validation additions.  Validation tools are present in both sets;
    this regression catches the HTTP 400 raised before the model can run.
    """
    read_schemas = [
        schema for schema in _model_tool_schemas()
        if schema["function"]["name"] in {"validate_workbook", "validate_with_officecli"}
    ]
    composed = _deduplicate_tool_schemas([
        *read_schemas,
        *read_schemas,
        {
            "type": "function",
            "function": {
                "name": "apply_source_cells",
                "description": "受控写入",
                "parameters": {"type": "object", "properties": {}, "required": [], "additionalProperties": False},
            },
        },
    ])

    names = [schema["function"]["name"] for schema in composed]
    assert len(names) == len(set(names))
    assert names.count("validate_workbook") == 1
    assert names.count("validate_with_officecli") == 1
    assert "apply_source_cells" in names


def test_execution_tool_contract_rejects_conflicting_duplicate_definition() -> None:
    first = {
        "type": "function",
        "function": {
            "name": "validate_workbook",
            "description": "first",
            "parameters": {"type": "object", "properties": {}, "required": []},
        },
    }
    conflicting = {
        "type": "function",
        "function": {
            "name": "validate_workbook",
            "description": "different",
            "parameters": {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
        },
    }

    with pytest.raises(ValueError, match="重复工具名称且定义冲突"):
        _deduplicate_tool_schemas([first, conflicting])


@pytest.mark.parametrize("name", ["prepare_workbook_copy", "apply_source_cells", "roll_forward_month"])
def test_new_workbook_write_tool_names_are_validated_and_executable(name: str) -> None:
    registry = ToolRegistry()
    registry.register(name, lambda: {"executed": name})

    result = registry.execute(ToolCall(call_id="new-write-tool", name=name))

    assert result.name == name
    assert result.status == "succeeded"
    assert result.ok is True
    assert result.output == {"executed": name}
    assert ToolResult.model_validate_json(result.model_dump_json()) == result


def test_month_roll_forward_schema_never_accepts_model_supplied_history_values() -> None:
    parameters = month_roll_forward_schema()["function"]["parameters"]
    change = parameters["properties"]["change"]

    assert "history_values" not in change["properties"]
    assert set(change["required"]) == {
        "sheet", "current_row", "expected_current_month", "new_month",
        "start_column", "end_column",
    }


def test_cross_month_run_rejects_data_writes_until_roll_forward_succeeds(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.agent_execution as execution

    master = tmp_path / "april.xlsx"
    source = tmp_path / "may-change.xlsx"
    draft = tmp_path / "may-draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "汇总"
    workbook.active["A2"] = "4月"
    workbook.active["B2"] = 0
    workbook.active["A3"] = "3月"
    workbook.active["B3"] = 50
    workbook.save(master)
    workbook.close()
    workbook = openpyxl.Workbook()
    workbook.active.title = "变更"
    workbook.active["B2"] = 100
    workbook.save(source)
    workbook.close()
    draft.write_bytes(master.read_bytes())
    run = {
        "run_id": "month-run", "master_file": master.name,
        "salary_month": "2026.05", "target_salary_month": "2026.05",
        "instruction": (
            "表页—配送员值班费，填D列值班天数，把D列的原有的天数清除，"
            "根据“5月薪资数据”中“值班”表页值班人员名字出现的次数来填。"
        ),
        "_master_path": str(master),
        "_source_paths": {source.name: str(source)}, "file_manifest": [],
        "workbook_updates": [], "messages": [], "model_plan": {"steps": []},
        "requires_month_roll_forward": True,
        "baseline": {"salary_month": "2026.04", "sha256": "baseline-sha"},
    }
    registry = ToolRegistry()
    monkeypatch.setattr(execution.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(execution.ModelConfig, "fallback_from_env", lambda: None)
    execution.execute_model_plan(
        run, draft=draft, registry=registry, read_schemas=[], materials=[], rules={},
        save=lambda _run: None, emit=lambda *_args, **_kwargs: None,
    )
    before = draft.read_bytes()

    rejected = registry.execute(ToolCall(
        call_id="write-before-roll", name="apply_source_cells", arguments={"changes": [{
            "source_file": source.name, "source_sheet": "变更", "source_cells": ["B2"],
            "target_sheet": "汇总", "target_cell": "B2", "expected_value": 0,
            "aggregation": "copy",
        }]},
    ))

    assert rejected.status == "failed"
    assert "先执行月份滚动" in str(rejected.error)
    assert draft.read_bytes() == before

    rolled = registry.execute(ToolCall(
        call_id="roll", name="roll_forward_month", arguments={"change": {
            "sheet": "汇总", "current_row": 2, "expected_current_month": "4月",
            "new_month": "5月", "start_column": 1, "end_column": 2,
        }},
    ))
    rejected_delete = registry.execute(ToolCall(
        call_id="delete-current-month", name="delete_rows", arguments={"change": {
            "sheet": "汇总", "rows": [2], "expected_cells": {"A2": "5月"},
        }},
    ))
    rejected_older_delete = registry.execute(ToolCall(
        call_id="delete-older-month", name="delete_rows", arguments={"change": {
            "sheet": "汇总", "rows": [4], "expected_cells": {"A4": "3月"},
        }},
    ))
    rejected_older_insert = registry.execute(ToolCall(
        call_id="insert-in-history", name="insert_and_copy_row", arguments={"change": {
            "sheet": "汇总", "source_row": 3, "insert_before_row": 4,
            "expected_source_cells": {"A3": "4月"}, "copy_max_column": 2,
        }},
    ))
    rejected_older_write = registry.execute(ToolCall(
        call_id="write-older-month", name="apply_source_cells", arguments={"changes": [{
            "source_file": source.name, "source_sheet": "变更", "source_cells": ["B2"],
            "target_sheet": "汇总", "target_cell": "B4", "expected_value": 50,
            "aggregation": "copy",
        }]},
    ))
    written = registry.execute(ToolCall(
        call_id="write-after-roll", name="apply_source_cells", arguments={"changes": [{
            "source_file": source.name, "source_sheet": "变更", "source_cells": ["B2"],
            "target_sheet": "汇总", "target_cell": "B2", "expected_value": 0,
            "aggregation": "copy",
        }]},
    ))

    assert rolled.status == "succeeded"
    assert rejected_delete.status == "failed"
    assert "历史月份" in str(rejected_delete.error)
    assert rejected_older_delete.status == "failed"
    assert rejected_older_insert.status == "failed"
    assert rejected_older_write.status == "failed"
    assert written.status == "succeeded"
    check = openpyxl.load_workbook(draft, data_only=False)
    try:
        assert check["汇总"]["A2"].value == "5月"
        assert check["汇总"]["B2"].value == 100
        assert check["汇总"]["A3"].value == "4月"
        assert check["汇总"]["B3"].value == 0
        assert check["汇总"]["A4"].value == "3月"
        assert check["汇总"]["B4"].value == 50
    finally:
        check.close()


def test_cross_month_roll_forward_executes_deferred_roster_rule(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A deferred deterministic roster instruction must run after month roll."""
    import backend.agent_execution as execution

    master = tmp_path / "april.xlsx"
    source = tmp_path / "may-change.xlsx"
    draft = tmp_path / "may-draft.xlsx"
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary["A2"] = "4月"
    summary["B2"] = 10
    summary["A3"] = "3月"
    summary["B3"] = 9
    roster = workbook.create_sheet("明细")
    roster.append(["序号", "姓名"])
    roster.append([1, "旧人员"])
    roster.append([2, "新人员"])
    roster.append(["*合计*", "79"])
    workbook.save(master)
    workbook.close()
    workbook = openpyxl.Workbook()
    source_sheet = workbook.active
    source_sheet.title = "派遣"
    source_sheet.append(["员工姓名"])
    source_sheet.append(["新人员"])
    workbook.save(source)
    workbook.close()

    class Config:
        max_output_tokens = 512

        def model_copy(self, *, update):
            return self

    class Provider:
        def __init__(self, _config):
            self.turn = 0

        def complete(self, *, messages, tools):
            self.turn += 1
            if self.turn == 1:
                from core.document_agent.contracts import ToolCall
                from core.document_agent.model import ModelResponse
                return ModelResponse(tool_calls=[ToolCall(
                    call_id="roll-forward", name="roll_forward_month", arguments={
                        "change": {
                            "sheet": "汇总", "current_row": 2,
                            "expected_current_month": "4月", "new_month": "5月",
                            "start_column": 1, "end_column": 2,
                        },
                    },
                )])
            from core.document_agent.model import ModelResponse
            return ModelResponse(content="已完成")

    monkeypatch.setattr(execution.ModelConfig, "from_env", lambda: Config())
    monkeypatch.setattr(execution.ModelConfig, "fallback_from_env", lambda: None)
    monkeypatch.setattr(execution, "OpenAICompatibleProvider", Provider)
    monkeypatch.setattr(execution, "select_project_skills", lambda _run: {
        "routes": [{"id": "workbook-update", "required_validation_tools": []}],
        "selected_routes": [], "runtime_routes": [], "reference_only_routes": [],
        "required_validation_tools": [],
    })

    run = {
        "run_id": "deferred-roster", "master_file": master.name,
        "salary_month": "2026.05", "target_salary_month": "2026.05",
        "instruction": "只处理来源文件里的派遣 sheet，最终名单以派遣为准，不保留其他人员",
        "_master_path": str(master), "_source_paths": {source.name: str(source)},
        "file_manifest": [], "source_files": [source.name], "workbook_updates": [],
        "messages": [], "model_plan": {"steps": []},
        "requires_month_roll_forward": True,
        "baseline": {"salary_month": "2026.04", "sha256": "baseline-sha"},
    }
    registry = execution.ToolRegistry()

    execution.execute_model_plan(
        run, draft=draft, registry=registry, read_schemas=[], materials=[], rules={},
        save=lambda _run: None, emit=lambda *_args, **_kwargs: None,
    )

    assert run["month_roll_forward"]["new_month"] == "5月"
    assert run["roster_sync"]["retained_names"] == ["新人员"]
    assert run["roster_sync"]["change_count"] > 0
    assert run["execution_result"]["status"] == "completed"
    assert run["status"] == "awaiting_review"
    updated = openpyxl.load_workbook(draft, data_only=False)
    try:
        detail = updated["明细"]
        assert [detail.cell(row, 2).value for row in range(2, detail.max_row + 1)] == ["新人员", "79"]
    finally:
        updated.close()


@pytest.mark.parametrize("tool_name", ["apply_cell_changes", "rollback_work_item"])
@pytest.mark.parametrize("target_cell", ["B3", "B4"])
def test_route_registry_never_mutates_the_frozen_history_row(
    tool_name: str, target_cell: str, tmp_path: Path,
) -> None:
    draft = tmp_path / "may-draft.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet.append(["月份", "金额"])
    sheet.append(["5月", 200])
    sheet.append(["4月", 100])
    sheet.append(["3月", 100])
    workbook.save(draft)
    workbook.close()
    run = {
        "run_id": "month-run", "project_id": "project-1",
        "salary_month": "2026.05", "draft_filename": draft.name,
        "month_roll_forward": {"sheet": "汇总", "history_row": 3},
        "workbook_updates": [],
        "items": [{
            "id": "history-item", "status": "needs_review",
            "issue_type": "agent_review", "risk_level": "low",
            "target_sheet": "汇总", "target_cell": target_cell,
            "current_value": 100, "candidate_values": [999],
            "applied_history": [{"old_value": 90, "new_value": 100}],
        }],
    }
    registry = _build_run_tool_registry(run, workbook_path=draft)
    arguments = (
        {"item_id": "history-item", "value": 999}
        if tool_name == "apply_cell_changes"
        else {"item_id": "history-item"}
    )

    result = registry.execute(ToolCall(
        call_id=f"frozen-{tool_name}", name=tool_name, arguments=arguments,
    ))

    assert result.status == "failed"
    assert "历史月份行已冻结" in str(result.error)
    check = openpyxl.load_workbook(draft, data_only=False)
    try:
        assert check["汇总"][target_cell].value == 100
    finally:
        check.close()


@pytest.mark.parametrize(
    "schema", _model_tool_schemas(), ids=lambda schema: schema["function"]["name"],
)
def test_every_advertised_tool_survives_typed_registry_round_trip(
    schema: dict[str, Any],
) -> None:
    name = schema["function"]["name"]
    registry = ToolRegistry()
    registry.register(name, lambda **arguments: {"received": arguments})

    call = ToolCall(call_id=f"call-{name}", name=name, arguments={})
    result = registry.execute(call)

    assert result.status == "succeeded"
    assert result.ok is True
    assert result.call_id == call.call_id
    assert result.name == name
    assert result.output == {"received": {}}
    assert ToolResult.model_validate_json(result.model_dump_json()) == result


def test_officecli_validation_tool_is_exposed_without_command_arguments() -> None:
    schema = next(
        item for item in _model_tool_schemas()
        if item["function"]["name"] == "validate_with_officecli"
    )
    assert schema["function"]["parameters"] == {
        "type": "object",
        "properties": {},
        "required": [],
        "additionalProperties": False,
    }
    assert "公式重算" in schema["function"]["description"]
    assert "当前独立草稿" in schema["function"]["description"]


def test_officecli_validation_is_scoped_to_the_current_draft(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    draft = tmp_path / "draft.xlsx"
    draft.write_bytes(b"placeholder")
    captured: list[tuple[list[str], dict[str, Any]]] = []

    class Completed:
        def __init__(self, stdout: str) -> None:
            self.returncode = 0
            self.stdout = stdout
            self.stderr = ""

    def fake_run(args: list[str], **kwargs: Any) -> Completed:
        captured.append((args, kwargs))
        if args[1] == "validate":
            return Completed("workbook is valid")
        return Completed(json.dumps({
            "success": True,
            "data": {
                "matches": 1,
                "results": [{
                    "path": "/Sheet/A3", "text": "3",
                    "format": {
                        "formula": "SUM(A1:A2)", "computedValue": "3", "evaluated": True,
                    },
                }],
            },
        }))

    monkeypatch.setattr("backend.routers.agent.subprocess.run", fake_run)
    registry = _build_run_tool_registry({
        "run_id": "run-1", "project_id": "project-1", "draft_filename": draft.name,
        "items": [],
    }, workbook_path=draft)

    result = registry.execute(ToolCall(
        call_id="officecli", name="validate_with_officecli", arguments={},
    ))

    assert result.ok is True
    assert result.output["valid"] is True
    assert result.output["recalculated"] is True
    assert result.output["formula_count"] == 1
    assert [call[0][1] for call in captured] == ["validate", "query"]
    assert all(str(draft) in call[0] for call in captured)
    assert all(call[1]["shell"] is False for call in captured)
    assert all(call[1]["encoding"] == "utf-8" for call in captured)


def test_officecli_nonzero_exit_is_a_failed_validation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    draft = tmp_path / "draft.xlsx"
    draft.write_bytes(b"placeholder")

    class Completed:
        returncode = 1
        stdout = "formula reference error"
        stderr = ""

    monkeypatch.setattr("backend.routers.agent.subprocess.run", lambda *_args, **_kwargs: Completed())
    registry = _build_run_tool_registry({
        "run_id": "run-1", "project_id": "project-1", "draft_filename": draft.name,
        "items": [],
    }, workbook_path=draft)

    result = registry.execute(ToolCall(
        call_id="officecli-invalid", name="validate_with_officecli", arguments={},
    ))

    assert result.status == "failed"
    assert "未通过" in str(result.error)


@pytest.mark.parametrize(
    ("formula_result", "expected_error"),
    [
        (
            {"path": "/Sheet/A6", "text": "#DIV/0!", "format": {
                "formula": "1/0", "computedValue": "#DIV/0!", "evaluated": True,
            }},
            "#DIV/0!",
        ),
        (
            {"path": "/Sheet/A5", "text": "#OCLI_NOTEVAL!", "format": {
                "formula": "UNSUPPORTED(A1)", "evaluated": False,
            }},
            "无法重算",
        ),
    ],
)
def test_officecli_validation_rejects_formula_errors_or_unevaluated_formulas(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    formula_result: dict[str, Any],
    expected_error: str,
) -> None:
    draft = tmp_path / "draft.xlsx"
    draft.write_bytes(b"placeholder")

    class Completed:
        returncode = 0
        stderr = ""

        def __init__(self, stdout: str) -> None:
            self.stdout = stdout

    def fake_run(args: list[str], **_kwargs: Any) -> Completed:
        if args[1] == "validate":
            return Completed("workbook is valid")
        return Completed(json.dumps({
            "success": True,
            "data": {"matches": 1, "results": [formula_result]},
        }))

    monkeypatch.setattr("backend.routers.agent.subprocess.run", fake_run)
    registry = _build_run_tool_registry({
        "run_id": "run-1", "project_id": "project-1", "draft_filename": draft.name,
        "items": [],
    }, workbook_path=draft)

    result = registry.execute(ToolCall(
        call_id="officecli-formula-error", name="validate_with_officecli", arguments={},
    ))

    assert result.status == "failed"
    assert expected_error in str(result.error)


def test_officecli_validation_allows_preexisting_baseline_findings(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    baseline = tmp_path / "baseline.xlsx"
    draft = tmp_path / "draft.xlsx"
    baseline.write_bytes(b"baseline")
    draft.write_bytes(b"draft")
    schema_payload = json.dumps({
        "success": False,
        "warnings": [
            {"message": "Found 1 validation error(s):", "code": "warning"},
            {"message": "[Schema] WPS extension is not declared.", "code": "warning"},
            {"message": "Path: /x:worksheet[1]/x:autoFilter[1]", "code": "warning"},
            {"message": "Part: /xl/worksheets/sheet1.xml", "code": "warning"},
        ],
    })
    formula_payload = json.dumps({
        "success": True,
        "data": {
            "matches": 1,
            "results": [{
                "path": "/Sheet/A3", "text": "#N/A",
                "format": {
                    "formula": "NA()", "computedValue": "#N/A", "evaluated": True,
                },
            }],
        },
    })

    class Completed:
        stderr = ""

        def __init__(self, *, returncode: int, stdout: str) -> None:
            self.returncode = returncode
            self.stdout = stdout

    def fake_run(args: list[str], **_kwargs: Any) -> Completed:
        if args[1] == "validate":
            return Completed(returncode=1, stdout=schema_payload)
        return Completed(returncode=0, stdout=formula_payload)

    monkeypatch.setattr("backend.routers.agent.subprocess.run", fake_run)
    registry = _build_run_tool_registry({
        "run_id": "run-1", "project_id": "project-1", "draft_filename": draft.name,
        "_master_path": str(baseline), "items": [],
    }, workbook_path=draft)

    result = registry.execute(ToolCall(
        call_id="officecli-baseline", name="validate_with_officecli", arguments={},
    ))

    assert result.status == "succeeded"
    assert result.output["valid"] is True
    assert result.output["recalculated"] is True
    assert result.output["baseline_schema_issue_count"] == 1
    assert result.output["baseline_formula_error_count"] == 1
    assert result.output["new_schema_errors"] == []
    assert result.output["formula_errors"] == []


@pytest.mark.parametrize("name", ["inspect_source_file", "read_source_range"])
def test_advertised_source_tools_read_the_registered_run_source(
    name: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "tenant-a" / "奖金来源.xlsx"
    source.parent.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "奖金名单"
    sheet.append(["工号", "奖金"])
    sheet.append(["E001", 2000])
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr("backend.routers.agent.UPLOAD_DIR", str(tmp_path))
    run = {
        "run_id": "run-1", "project_id": "project-1",
        "source_files": [source.name], "_source_paths": {source.name: str(source)},
    }
    registry = _build_run_tool_registry(run, workbook_path=tmp_path / "draft.xlsx")
    schema = next(item for item in _model_tool_schemas() if item["function"]["name"] == name)
    arguments = {"filename": source.name}
    if name == "read_source_range":
        arguments.update({"sheet": "奖金名单", "range": "A1:B2"})
    assert set(arguments) == set(schema["function"]["parameters"]["required"])

    result = registry.execute(ToolCall(call_id="source-call", name=name, arguments=arguments))

    assert result.ok is True, result.error
    assert result.status == "succeeded"
    assert result.output["filename"] == source.name
    if name == "inspect_source_file":
        assert result.output["sheets"] == [{
            "name": "奖金名单", "max_row": 2, "max_column": 2,
            "sample": [["工号", "奖金"], ["E001", 2000]],
        }]
    else:
        assert result.output["values"] == [["工号", "奖金"], ["E001", 2000]]
    assert result.observation == result.output


@pytest.mark.parametrize("name", ["inspect_source_file", "read_source_range"])
def test_responses_provider_accepts_advertised_source_function_calls(
    name: str, monkeypatch: pytest.MonkeyPatch,
) -> None:
    provider = OpenAICompatibleProvider(ModelConfig(
        provider="openai_compatible", base_url="https://model.example/v1",
        api_key="test-secret", model="finance-model", api_style="responses",
    ))
    arguments = {"filename": "奖金来源.xlsx"}
    if name == "read_source_range":
        arguments.update({"sheet": "奖金名单", "range": "A1:B2"})
    captured: dict[str, Any] = {}

    def fake_request(request: Any, **_kwargs: Any) -> _Response:
        captured["payload"] = json.loads(request.data)
        return _Response({
            "id": "response-source", "status": "completed",
            "output": [{
                "type": "function_call", "call_id": "source-call", "name": name,
                "arguments": json.dumps(arguments, ensure_ascii=False),
            }],
        })

    monkeypatch.setattr("urllib.request.urlopen", fake_request)
    response = provider.complete(
        messages=[{"role": "user", "content": "核对本次运行的奖金来源"}],
        tools=_model_tool_schemas(),
    )

    assert name in {tool["name"] for tool in captured["payload"]["tools"]}
    assert [call.model_dump() for call in response.tool_calls] == [{
        "call_id": "source-call", "name": name, "arguments": arguments,
    }]


def test_multi_tool_response_remains_one_assistant_turn_before_all_results() -> None:
    calls = [
        ToolCall(call_id="call-inspect", name="inspect_workbook"),
        ToolCall(call_id="call-read", name="read_range", arguments={"sheet": "奖金名单", "range": "A1:B2"}),
    ]
    snapshots: list[list[dict[str, Any]]] = []

    class Provider:
        def complete(self, *, messages: Any, tools: Any) -> ModelResponse:
            snapshots.append(deepcopy(messages))
            if len(snapshots) == 1:
                return ModelResponse(content="同时检查结构和奖金来源。", tool_calls=calls)
            return ModelResponse(content="已核对来源。")

    registry = ToolRegistry()
    registry.register("inspect_workbook", lambda: {"sheets": ["奖金名单"]})
    registry.register("read_range", lambda **_arguments: {"values": [["E001", 2000]]})
    original_messages = [{"role": "user", "content": "核对奖金"}]

    result = ModelOrchestrator(provider=Provider(), registry=registry).run(
        run_id="run-1", messages=original_messages, tools=_model_tool_schemas(),
    )

    assert result.status == "completed"
    assert original_messages == [{"role": "user", "content": "核对奖金"}]
    follow_up = snapshots[1]
    assert [message["role"] for message in follow_up] == ["user", "assistant", "tool", "tool"]
    assert follow_up[1]["content"] == "同时检查结构和奖金来源。"
    assert [call["id"] for call in follow_up[1]["tool_calls"]] == ["call-inspect", "call-read"]
    assert [message["tool_call_id"] for message in follow_up[2:]] == ["call-inspect", "call-read"]
    follow_up_results = [json.loads(message["content"]) for message in follow_up[2:]]
    assert follow_up_results[0] == {"sheets": ["奖金名单"]}
    assert follow_up_results[1]["checkpoint_id"]
    assert [cell["value"] for cell in follow_up_results[1]["current_cells"]] == ["E001", 2000]
    assert "values" not in follow_up_results[1]
    assert [event.type for event in result.events] == [
        "model_request", "model_response", "tool_call", "tool_result", "tool_call", "tool_result",
        "model_request", "model_response",
    ]


def test_malformed_tool_arguments_are_retried_with_same_provider() -> None:
    """A transient provider formatting error must not abort the workbook run."""
    calls: list[list[dict[str, Any]]] = []

    class Provider:
        def complete(self, *, messages: Any, tools: Any) -> ModelResponse:
            calls.append(deepcopy(messages))
            if len(calls) == 1:
                raise ModelProviderError("模型工具调用参数无效")
            return ModelResponse(content="已继续完成核对。")

    result = ModelOrchestrator(
        provider=Provider(),
        registry=ToolRegistry(),
        fallback_providers=[object()],
    ).run(
        run_id="run-tool-retry",
        messages=[{"role": "user", "content": "继续处理"}],
        tools=_model_tool_schemas(),
    )

    assert result.status == "completed"
    assert result.content == "已继续完成核对。"
    assert len(calls) == 2
    assert calls[1][-1]["role"] == "user"
    assert "不是合法 JSON" in calls[1][-1]["content"]
    assert [event.payload.get("stage") for event in result.events if event.type == "model_request"] == [
        "model_call", "model_call",
    ]


class _Response:
    def __init__(self, payload: dict[str, Any]):
        self._body = json.dumps(payload).encode("utf-8")

    def read(self) -> bytes:
        return self._body

    def __enter__(self) -> _Response:
        return self

    def __exit__(self, *_args: object) -> None:
        return None
