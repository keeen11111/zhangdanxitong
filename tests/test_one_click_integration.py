"""一次上传与一键整合的确定性规则测试。"""

import base64
import json
import re
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import openpyxl
import pandas as pd
import pytest
from fastapi import HTTPException
from openpyxl.drawing.image import Image
from openpyxl.worksheet.formula import ArrayFormula

import backend.routers.pipeline as pipeline_module
from backend.routers.pipeline import (
    _build_export_preview,
    _choose_template_candidate,
    _copy_source_workbooks,
    _copy_novel_source_sheets,
    _freeze_external_workbook_formulas,
    _confirm_all_pending_diff,
    _content_disposition,
    _ensure_manual_issues_export,
    _load_current_export_meta,
    _project_source_signature,
    _normalize_export_filename,
    _extract_master_salary_month,
    _resolve_effective_salary_month,
    _resolve_manual_issue,
    _resolve_layout_template_path,
    _sheet_content_signature,
    _update_duty_roster_counts,
    _write_social_contributions_to_ledger,
)
from core.entity_engine import import_file_to_entities
from backend.routers.projects import (
    _classify_uploaded_workbook,
    _normalize_excel_content,
    _read_upload_content,
    _validate_declared_workbook_role,
    _validate_ooxml_archive,
    _workbook_binary_kind,
    _workbook_dimensions,
    analyze_workbook_structure,
    upload_file,
)


TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9WlEPyAAAAAASUVORK5CYII="
)


def _add_tiny_image(sheet, tmp_path, filename: str = "tiny.png") -> None:
    image_path = tmp_path / filename
    image_path.write_bytes(TINY_PNG)
    sheet.add_image(Image(str(image_path)), "A1")


def _add_cached_value_to_formula_cell(path, sheet_xml_name: str, coordinate: str, value: str) -> None:
    """Add an Excel cached result to a formula cell for a realistic source fixture."""
    temporary_path = path.with_suffix(".cached.xlsx")
    with ZipFile(path, "r") as source_archive, ZipFile(temporary_path, "w") as target_archive:
        for item in source_archive.infolist():
            payload = source_archive.read(item.filename)
            if item.filename == sheet_xml_name:
                text = payload.decode("utf-8")
                pattern = re.compile(
                    rf'(<c r="{re.escape(coordinate)}"[^>]*>.*?<f>.*?</f>)(?:<v>.*?</v>)?(</c>)',
                    re.DOTALL,
                )
                text, count = pattern.subn(rf"\1<v>{value}</v>\2", text, count=1)
                assert count == 1
                payload = text.encode("utf-8")
            target_archive.writestr(item, payload)
    temporary_path.replace(path)


def test_workbook_with_payroll_master_sheet_is_template() -> None:
    sheets = {
        "工资核算": pd.DataFrame({"工号": ["E001"], "姓名": ["测试员工"]}),
        "员工档案": pd.DataFrame({"工号": ["E001"]}),
    }

    result = _classify_uploaded_workbook("2026年工资主模板.xlsx", sheets)

    assert result == "template"


def test_regular_business_workbook_is_source() -> None:
    sheets = {
        "考勤明细": pd.DataFrame({"工号": ["E001"], "迟到次数": [0]}),
    }

    result = _classify_uploaded_workbook("6月考勤.xlsx", sheets)

    assert result == "source"


def test_image_only_workbook_is_accepted_as_source_data(tmp_path) -> None:
    workbook_path = tmp_path / "10-业务运营部配送团队通讯费.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "业务运营部配送团队通讯费"
    _add_tiny_image(workbook.active, tmp_path)
    workbook.save(workbook_path)
    content = workbook_path.read_bytes()
    sheets = pd.read_excel(BytesIO(content), sheet_name=None)

    assert _workbook_dimensions(sheets, content) == (0, 0)


def test_legacy_excel_content_is_detected_even_with_xlsx_suffix() -> None:
    legacy_ole_header = bytes.fromhex("D0 CF 11 E0 A1 B1 1A E1") + b"legacy"

    assert _workbook_binary_kind(legacy_ole_header) == "xls"
    assert _workbook_binary_kind(b"PK\x03\x04modern") == "xlsx"


def test_encrypted_ooxml_content_is_detected_as_xlsx_even_with_ole_container() -> None:
    encrypted_ooxml = (
        bytes.fromhex("D0 CF 11 E0 A1 B1 1A E1")
        + "EncryptionInfo".encode("utf-16le")
        + b"padding"
        + "EncryptedPackage".encode("utf-16le")
    )

    assert _workbook_binary_kind(encrypted_ooxml) == "xlsx-encrypted"


def test_encrypted_ooxml_normalization_uses_xlsx_source_extension(monkeypatch) -> None:
    encrypted_ooxml = (
        bytes.fromhex("D0 CF 11 E0 A1 B1 1A E1")
        + "EncryptionInfo".encode("utf-16le")
        + "EncryptedPackage".encode("utf-16le")
    )
    calls: list[str] = []

    def convert(content: bytes, source_extension: str = ".xls") -> bytes:
        calls.append(source_extension)
        workbook = BytesIO()
        openpyxl.Workbook().save(workbook)
        return workbook.getvalue()

    monkeypatch.setattr(
        "backend.routers.projects._convert_legacy_excel_with_desktop_excel",
        convert,
    )

    normalized, extension = _normalize_excel_content(encrypted_ooxml)

    assert normalized.startswith(b"PK\x03\x04")
    assert extension == ".xlsx"
    assert calls == [".xlsx"]


def test_ooxml_upload_does_not_invoke_desktop_conversion(monkeypatch) -> None:
    workbook_bytes = BytesIO()
    workbook = openpyxl.Workbook()
    workbook.save(workbook_bytes)
    content = workbook_bytes.getvalue()
    monkeypatch.setattr(
        "backend.routers.projects._convert_legacy_excel_with_desktop_excel",
        lambda _content: (_ for _ in ()).throw(AssertionError("unexpected conversion")),
    )

    normalized, extension = _normalize_excel_content(content)

    assert normalized == content
    assert extension == ".xlsx"


def test_upload_rejects_non_excel_suffix() -> None:
    upload = SimpleNamespace(filename="伪装文件.txt", file=BytesIO(b"PK\x03\x04data"))

    with pytest.raises(HTTPException, match="不是支持的 Excel 文件"):
        _read_upload_content(upload)


def test_upload_rejects_content_over_size_limit(monkeypatch) -> None:
    monkeypatch.setattr("backend.routers.projects._MAX_UPLOAD_BYTES", 4)
    upload = SimpleNamespace(filename="总表.xlsx", file=BytesIO(b"12345"))

    with pytest.raises(HTTPException, match="不能超过") as exc_info:
        _read_upload_content(upload)

    assert exc_info.value.status_code == 413


def test_single_upload_removes_stored_file_when_database_commit_fails(tmp_path, monkeypatch) -> None:
    workbook_bytes = BytesIO()
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook.active["A1"] = "工号"
    workbook.save(workbook_bytes)
    project = SimpleNamespace(id="project-id")

    class FailingDb:
        def query(self, _model):
            return self

        def filter(self, *_criteria):
            return self

        def first(self):
            return None

        def add(self, _record) -> None:
            pass

        def commit(self) -> None:
            raise RuntimeError("database unavailable")

        def rollback(self) -> None:
            pass

    monkeypatch.setattr("backend.routers.projects.UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr("backend.routers.projects._load_project_or_404", lambda *_args: project)
    upload = SimpleNamespace(filename="总表.xlsx", file=BytesIO(workbook_bytes.getvalue()))

    with pytest.raises(RuntimeError, match="database unavailable"):
        upload_file(
            project_id=project.id,
            file_type="template",
            sheet_name=None,
            file=upload,
            user=SimpleNamespace(id="user-id"),
            db=FailingDb(),
        )

    assert list(tmp_path.rglob("*.xlsx")) == []


def test_master_role_requires_the_payroll_sheet() -> None:
    with pytest.raises(HTTPException) as exc_info:
        _validate_declared_workbook_role("template", ["考勤", "奖金"])

    assert exc_info.value.status_code == 400
    assert "工资核算" in str(exc_info.value.detail)


def test_change_role_rejects_a_complete_master_workbook() -> None:
    with pytest.raises(HTTPException) as exc_info:
        _validate_declared_workbook_role(
            "source",
            ["工资核算", "工资汇总表", "工资条", "台账"],
        )

    assert exc_info.value.status_code == 400
    assert "总表" in str(exc_info.value.detail)


def test_ooxml_archive_rejects_excessive_uncompressed_content(monkeypatch) -> None:
    monkeypatch.setattr("backend.routers.projects._MAX_UNCOMPRESSED_WORKBOOK_BYTES", 4)
    content = BytesIO()
    with ZipFile(content, "w") as archive:
        archive.writestr("xl/workbook.xml", "12345")

    with pytest.raises(HTTPException, match="解压后内容过大"):
        _validate_ooxml_archive(content.getvalue())


def test_master_workbook_analysis_classifies_every_sheet(tmp_path) -> None:
    workbook_path = tmp_path / "总表.xlsx"
    workbook = openpyxl.Workbook()
    payroll = workbook.active
    payroll.title = "工资核算"
    payroll["A1"] = "姓名"
    payroll["B2"] = "=1+1"
    workbook.create_sheet("人员异动")["A1"] = "新工号"
    workbook.create_sheet("规则备忘")["A1"] = "说明"
    workbook.save(workbook_path)

    analysis = analyze_workbook_structure(str(workbook_path))

    assert analysis["sheet_count"] == 3
    assert {sheet["category"] for sheet in analysis["sheets"]} == {
        "工资核算主表",
        "人员异动",
        "规则政策",
    }
    assert analysis["formula_count"] == 1
    payroll_analysis = next(
        sheet for sheet in analysis["sheets"] if sheet["name"] == "工资核算"
    )
    assert payroll_analysis["content_types"] == ["数据", "公式"]


def test_project_source_signature_is_order_independent() -> None:
    first = SimpleNamespace(
        id="file-1",
        original_name="2-考勤.xlsx",
        stored_path="project/file-1.xlsx",
        file_type="source",
        row_count=20,
        col_count=4,
        created_at=None,
    )
    second = SimpleNamespace(
        id="file-2",
        original_name="10-奖金.xlsx",
        stored_path="project/file-2.xlsx",
        file_type="source",
        row_count=8,
        col_count=3,
        created_at=None,
    )

    assert _project_source_signature([first, second]) == _project_source_signature([second, first])


def test_sheet_name_does_not_determine_source_record_priority() -> None:
    assert not hasattr(pipeline_module, "_source_record_priority")


def test_master_salary_month_uses_explicit_owned_month_marker() -> None:
    assert (
        _extract_master_salary_month("202608（所属月202607）-北京科园工资核算总表.xlsx")
        == "2026.07"
    )


def test_project_month_conflict_with_master_is_allowed_and_uses_master_month() -> None:
    assert (
        _resolve_effective_salary_month(
            "2026.06",
            "202608（所属月202607）-北京科园工资核算总表.xlsx",
        )
        == "2026.07"
    )


def test_resolve_manual_issue_applies_recommended_value_or_keeps_current() -> None:
    final_data = {
        "entities": {
            "employee_profile": [{"工号": "E001", "姓名": "甲", "部门": "旧部门"}],
        },
        "issues": [{
            "issue_id": "issue-1",
            "issue_type": "department_transfer",
            "employee_id": "E001",
            "person_name": "甲",
            "target_field": "部门",
            "current_value": "旧部门",
            "proposed_value": "新部门",
            "status": "pending",
        }],
    }

    assert _resolve_manual_issue(final_data, "issue-1", "apply_proposed") == "新部门"
    assert final_data["entities"]["employee_profile"][0]["部门"] == "新部门"
    assert final_data["issues"][0]["status"] == "confirmed"

    final_data["issues"][0]["status"] = "pending"
    final_data["entities"]["employee_profile"][0]["部门"] = "旧部门"
    assert _resolve_manual_issue(final_data, "issue-1", "keep_current") == "旧部门"
    assert final_data["entities"]["employee_profile"][0]["部门"] == "旧部门"
    assert final_data["issues"][0]["status"] == "confirmed"


def test_duty_roster_counts_replace_old_master_values(tmp_path) -> None:
    output_path = tmp_path / "总表.xlsx"
    output = openpyxl.Workbook()
    target = output.active
    target.title = "配送员值班费"
    target.append(["序号", "值班人员姓名", "值班费标准", "值班天数", "值班费合计"])
    target.append([1, "马昭", 300, 4, "=ROUND(C2*D2,0)"])
    target.append([2, "高阳阳", 300, 3, "=ROUND(C3*D3,0)"])
    output.save(output_path)

    source_path = tmp_path / "7月变更.xlsx"
    source = openpyxl.Workbook()
    roster = source.active
    roster.title = "值班"
    roster.append(["业务运营部配送团队 7 月值班排班表"])
    roster.append(["序号", "值班日期", "值班人员"])
    for index, name in enumerate(["高阳阳", "马昭", "高阳阳", "马昭", "高阳阳", "马昭", "高阳阳", "马昭"], 1):
        roster.append([index, f"7.{index}", name])
    source.save(source_path)

    updated = _update_duty_roster_counts(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    assert updated == 2
    assert result["配送员值班费"]["D2"].value == 4
    assert result["配送员值班费"]["D3"].value == 4
    assert result["配送员值班费"]["E3"].value == "=ROUND(C3*D3,0)"


def test_attendance_import_combines_workday_and_rest_day_overtime_and_keeps_travel(tmp_path) -> None:
    source_path = tmp_path / "薪资数据-7月.xlsx"
    source = openpyxl.Workbook()
    attendance = source.active
    attendance.title = "考勤-7月"
    attendance.append([
        "姓名",
        "工号",
        "出勤天数",
        "工作日加班时长（加班费）(求和)",
        "公休日加班时长（加班费）(求和)",
        "节假日加班时长（加班费）(求和)",
        "出差",
    ])
    attendance.append(["甲", "E001", 23, 2, 4, 8, 3])
    source.save(source_path)

    parsed = import_file_to_entities(str(source_path), source_path.name)

    record = parsed["attendance_record"][0]
    assert record["普通加班时长"] == 6
    assert record["法定节日加班时长"] == 8
    assert record["出差天数"] == 3


def test_project_source_signature_changes_with_source_metadata() -> None:
    original = SimpleNamespace(
        id="file-1",
        original_name="考勤.xlsx",
        stored_path="project/file-1.xlsx",
        file_type="source",
        row_count=20,
        col_count=4,
        created_at=None,
    )
    reclassified = SimpleNamespace(**{**original.__dict__, "file_type": "template"})

    assert _project_source_signature([original]) != _project_source_signature([reclassified])


def test_sheet_content_signature_treats_empty_string_as_excel_blank() -> None:
    source = openpyxl.Workbook()
    source.active["G4"] = ""
    target = openpyxl.Workbook()

    assert _sheet_content_signature(source.active) == _sheet_content_signature(target.active)


def test_latest_export_is_rejected_after_source_files_change(tmp_path, monkeypatch) -> None:
    source = SimpleNamespace(
        id="file-1",
        original_name="考勤.xlsx",
        stored_path="project/file-1.xlsx",
        file_type="source",
        row_count=20,
        col_count=4,
        created_at=None,
    )

    class Query:
        def filter(self, *_args):
            return self

        def all(self):
            return [source]

    class Database:
        def query(self, *_args):
            return Query()

    monkeypatch.setattr(pipeline_module, "SESSION_DIR", str(tmp_path))
    (tmp_path / "project_export_meta.json").write_text(
        json.dumps({"filename": "old.xlsx", "source_signature": "outdated"}),
        encoding="utf-8",
    )

    with pytest.raises(HTTPException, match="来源文件已变化") as error:
        _load_current_export_meta("project", Database())

    assert error.value.status_code == 409


def test_manual_issues_export_is_created_for_legacy_metadata(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(pipeline_module, "EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setattr(pipeline_module, "SESSION_DIR", str(tmp_path / "sessions"))
    (tmp_path / "sessions").mkdir()
    meta = {
        "filename": "工资核算_2026.06.xlsx",
        "issues": [{
            "person_name": "李楠",
            "issue_type": "missing_value",
            "target_field": "月基本薪资",
            "message": "缺失",
            "source_files": ["考勤.xlsx"],
            "source_sheets": ["考勤"],
        }],
    }

    filename, path = _ensure_manual_issues_export("project", meta)

    workbook = openpyxl.load_workbook(path, data_only=False)
    assert filename.endswith(".xlsx")
    assert workbook["待人工处理"]["B2"].value == "李楠"


def test_manual_issues_export_excludes_confirmed_items_when_rebuilt(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(pipeline_module, "EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setattr(pipeline_module, "SESSION_DIR", str(tmp_path / "sessions"))
    (tmp_path / "sessions").mkdir()
    meta = {
        "filename": "工资核算_2026.06.xlsx",
        "issues_filename": "待人工处理.xlsx",
        "issues_path": "project/待人工处理.xlsx",
        "issues": [
            {
                "issue_id": "pending-1",
                "person_name": "待处理人员",
                "issue_type": "missing_value",
                "target_field": "月基本薪资",
                "message": "缺失",
                "status": "pending",
            },
            {
                "issue_id": "confirmed-1",
                "person_name": "已处理人员",
                "issue_type": "department_transfer",
                "target_field": "部门",
                "message": "已确认",
                "status": "confirmed",
            },
        ],
    }

    _, path = _ensure_manual_issues_export("project", meta)

    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook["待人工处理"]
    assert sheet.max_row == 2
    assert sheet["B2"].value == "待处理人员"
    workbook.close()


def test_one_click_confirmation_only_changes_pending_items() -> None:
    diff = {
        "hire": [
            {"工号": "E001", "status": "pending"},
            {"工号": "E002", "status": "ignored"},
        ],
        "attendance": {"items": [{"工号": "E003", "status": "pending"}]},
        "summary": {"hire": 2},
    }

    confirmed = _confirm_all_pending_diff(diff)

    assert confirmed == 2
    assert diff["hire"][0]["status"] == "confirmed"
    assert diff["hire"][1]["status"] == "ignored"
    assert diff["attendance"]["items"][0]["status"] == "confirmed"


def test_download_header_supports_chinese_filename() -> None:
    header = _content_disposition("工资核算_2026.06.xlsx")

    assert 'filename="payroll-result.xlsx"' in header
    assert "filename*=UTF-8''%E5%B7%A5%E8%B5%84%E6%A0%B8%E7%AE%97_2026.06.xlsx" in header
    header.encode("latin-1")


def test_export_preview_returns_non_empty_rows_from_generated_workbook(tmp_path) -> None:
    workbook_path = tmp_path / "工资核算.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "工资核算"
    sheet.append(["姓名", "应发工资"])
    sheet.append(["张三", 10000])
    workbook.save(workbook_path)

    preview = _build_export_preview(str(workbook_path))

    assert preview["sheet_name"] == "工资核算"
    assert preview["rows"] == [["姓名", "应发工资"], ["张三", 10000]]


def test_export_preview_prefers_payroll_sheet_in_full_master_workbook(tmp_path) -> None:
    workbook_path = tmp_path / "完整总表.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "规则备忘"
    payroll = workbook.create_sheet("工资核算")
    payroll.append(["姓名", "应发工资"])
    payroll.append(["李四", 12000])
    workbook.save(workbook_path)

    preview = _build_export_preview(str(workbook_path))

    assert preview["sheet_name"] == "工资核算"
    assert preview["rows"] == [["姓名", "应发工资"], ["李四", 12000]]


def test_one_click_selects_master_even_when_legacy_classifier_marks_every_file_template() -> None:
    files = [
        SimpleNamespace(id="main", original_name="大药房工资核算总表-v2.xlsx", file_type="template", row_count=150, col_count=80),
        SimpleNamespace(id="hean", original_name="鹤安长泰2026.06.xlsx", file_type="template", row_count=8, col_count=25),
        SimpleNamespace(id="keyuan", original_name="薪资数据-科园-6月薪资.xlsx", file_type="template", row_count=48, col_count=32),
        SimpleNamespace(id="yiyao", original_name="益药科园2026.06.xlsx", file_type="template", row_count=36, col_count=28),
    ]

    selected = _choose_template_candidate(files)

    assert selected.id == "main"


def test_one_click_only_selects_a_workbook_compatible_with_export() -> None:
    files = [
        SimpleNamespace(id="summary", original_name="03-工资汇总表.xlsx", file_type="template", row_count=200, col_count=90),
        SimpleNamespace(id="payroll", original_name="04-工资核算.xlsx", file_type="template", row_count=80, col_count=70),
    ]

    selected = _choose_template_candidate(files, compatible_file_ids={"payroll"})

    assert selected.id == "payroll"


def test_output_filename_is_safe_and_keeps_xlsx_suffix() -> None:
    assert _normalize_export_filename("  六月工资总表  ") == "六月工资总表.xlsx"
    assert _normalize_export_filename("六月:工资/总表.xlsx") == "六月工资总表.xlsx"
    assert _normalize_export_filename("../六月工资总表") == "六月工资总表.xlsx"


def test_copy_source_workbooks_adds_detail_sheets(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_path = tmp_path / "attendance.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "考勤"
    source.active["A1"] = "姓名"
    source.active["A2"] = "张三"
    source.save(source_path)

    copied = _copy_source_workbooks(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    assert copied == ["考勤"]
    assert result["考勤"]["A2"].value == "张三"


def test_copy_source_workbooks_freezes_external_formula_to_cached_value(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_path = tmp_path / "bonus.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "奖金-7月"
    source.active["A1"] = "人员"
    source.active["A2"] = "张三"
    source.active["B2"] = "='[missing.xlsx]附1-月度考核'!J4"
    source.save(source_path)
    _add_cached_value_to_formula_cell(source_path, "xl/worksheets/sheet1.xml", "B2", "950")

    copied = _copy_source_workbooks(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    assert copied == ["奖金-7月"]
    assert result["奖金-7月"]["B2"].value == 950


def test_copy_source_workbooks_rejects_external_formula_without_cached_value(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_path = tmp_path / "bonus.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "奖金-7月"
    source.active["B2"] = "='[missing.xlsx]附1-月度考核'!J4"
    source.save(source_path)

    with pytest.raises(ValueError, match="缓存结果"):
        _copy_source_workbooks(str(output_path), [str(source_path)])


def test_cached_external_formula_value_rejects_excel_error_cell() -> None:
    source = openpyxl.Workbook()
    source.active["A1"] = "='[missing.xlsx]明细'!A1"
    cached = openpyxl.Workbook()
    cached.active["A1"] = "#SPILL!"
    cached.active["A1"].data_type = "e"

    with pytest.raises(ValueError, match="缓存结果"):
        pipeline_module._cached_external_formula_value(source.active["A1"], cached.active["A1"])


def test_freeze_external_workbook_formulas_removes_all_external_formula_text(tmp_path) -> None:
    workbook_path = tmp_path / "result.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "奖金-7月"
    source.active["A1"] = "='[missing.xlsx]明细'!A1"
    source.active["A2"] = "='[missing.xlsx]明细'!A2"
    source.save(workbook_path)
    _add_cached_value_to_formula_cell(workbook_path, "xl/worksheets/sheet1.xml", "A1", "100")
    _add_cached_value_to_formula_cell(workbook_path, "xl/worksheets/sheet1.xml", "A2", "200")

    assert _freeze_external_workbook_formulas(str(workbook_path)) == 2

    result = openpyxl.load_workbook(workbook_path, data_only=False, keep_links=False)
    try:
        assert result["奖金-7月"]["A1"].value == 100
        assert result["奖金-7月"]["A2"].value == 200
    finally:
        result.close()


def test_freeze_external_workbook_formulas_uses_matching_fallback_cache(tmp_path) -> None:
    workbook_path = tmp_path / "result.xlsx"
    source_path = tmp_path / "master.xlsx"

    source = openpyxl.Workbook()
    source.active.title = "防暑降温费"
    source.active["G2"] = "='[previous.xlsx]防暑降温费'!G2"
    source.save(source_path)
    _add_cached_value_to_formula_cell(source_path, "xl/worksheets/sheet1.xml", "G2", "100")

    output = openpyxl.Workbook()
    output.active.title = "防暑降温费"
    output.active["G2"] = source.active["G2"].value
    output.save(workbook_path)

    assert pipeline_module._verify_unchanged_formula_signature(
        str(source_path),
        str(workbook_path),
    ) == 1
    assert _freeze_external_workbook_formulas(str(workbook_path), str(source_path)) == 1
    result = openpyxl.load_workbook(workbook_path, data_only=False, keep_links=False)
    try:
        assert result["防暑降温费"]["G2"].value == 100
    finally:
        result.close()


def test_social_billing_values_are_written_to_the_output_ledger(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    ledger = workbook.create_sheet("台账")
    ledger.append(["序号", "姓名"])
    ledger.append([
        "新工号", "姓名", "缴费公司-与劳动合同公司主体一致", "单位养老\n16%",
        "个人养老\n8%", "单位失业\n0.5%", "个人失业\n0.5%", "单位工伤\n0.3%、0.4%",
        "单位生育\n", "单位医疗9.8%", "个人医疗\n2%+3", "单位住房 12%", "个人住房 12%",
    ])
    ledger.append(["E001", "张三"])
    ledger.append(["E002", "李四"])
    workbook.save(output_path)

    updated = _write_social_contributions_to_ledger(
        str(output_path),
        {"social_security": [{
            "姓名": "张三",
            "缴费公司": "大药房",
            "公司养老": 1145.92,
            "个人养老": 572.96,
            "公司失业": 35.81,
            "个人失业": 35.81,
            "公司工伤": 14.32,
            "公司医疗": 701.88,
            "个人医疗": 146.24,
            "公司公积金": 860,
            "个人公积金": 860,
        }]},
    )

    result = openpyxl.load_workbook(output_path, data_only=False)["台账"]
    assert updated == 1
    assert result["C3"].value == "大药房"
    assert result["D3"].value == 1145.92
    assert result["E3"].value == 572.96
    assert result["K3"].value == 146.24
    assert result["L3"].value == 860
    assert result["M3"].value == 860
    assert result["D4"].value is None


def test_social_billing_uses_employee_id_when_the_person_name_changed(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    ledger = workbook.create_sheet("台账")
    ledger.append(["序号", "姓名"])
    ledger.append(["新工号", "姓名", "缴费公司-与劳动合同公司主体一致", "单位养老\n16%"])
    ledger.append(["E001", "旧姓名", "旧公司", 1])
    workbook.save(output_path)

    updated = _write_social_contributions_to_ledger(
        str(output_path),
        {
            "employee_profile": [{
                "工号": "E001",
                "姓名": "新姓名",
                "身份证号": "110101199001011234",
                "_person_key": "employee:E001",
            }],
            "social_security": [{
                "姓名": "新姓名",
                "身份证号": "110101199001011234",
                "缴费公司": "新公司",
                "公司养老": 1234.56,
                "_person_key": "employee:E001",
            }],
        },
    )

    result = openpyxl.load_workbook(output_path, data_only=False)["台账"]
    assert updated == 1
    assert result["A3"].value == "E001"
    assert result["B3"].value == "新姓名"
    assert result["C3"].value == "新公司"
    assert result["D3"].value == 1234.56


def test_copy_source_workbooks_preserves_image_only_sheet(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_path = tmp_path / "10-业务运营部配送团队通讯费.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "业务运营部配送团队通讯费"
    _add_tiny_image(source.active, tmp_path, "source.png")
    source.save(source_path)

    copied = _copy_source_workbooks(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    assert copied == ["业务运营部配送团队通讯费"]
    assert len(result["业务运营部配送团队通讯费"]._images) == 1


def test_copy_source_workbooks_replaces_a_changed_supporting_sheet(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    payslip = output.create_sheet("工资条")
    payslip["A1"] = "模板数据"
    output.save(output_path)

    source_path = tmp_path / "05-工资条.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "工资条"
    source.active["A1"] = "本月原始数据"
    source.save(source_path)

    copied = _copy_source_workbooks(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    archived_names = [name for name in result.sheetnames if name != "工资条" and "工资条" in name]
    assert archived_names == []
    assert result["工资条"]["A1"].value == "本月原始数据"
    assert copied == ["工资条"]


def test_copy_source_workbooks_protects_payroll_sheet_and_uses_original_filename_for_archive_title(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output["工资核算"]["A1"] = "总表权威数据"
    output.save(output_path)

    source_path = tmp_path / "3064853f32d8407286ada07e3ab1308d.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "工资核算"
    source.active["A1"] = "本月原始数据"
    source.save(source_path)

    copied = _copy_source_workbooks(
        str(output_path),
        [str(source_path)],
        source_names={str(source_path): "05-工资条.xlsx"},
    )

    result = openpyxl.load_workbook(output_path, data_only=False)
    assert copied == ["来源-05-工资条-工资核算"]
    assert result["工资核算"]["A1"].value == "总表权威数据"
    assert result[copied[0]]["A1"].value == "本月原始数据"


def test_copy_source_workbooks_keeps_both_sheets_when_titles_repeat(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_paths = []
    for filename, value in (("first.xlsx", "第一份"), ("second.xlsx", "第二份")):
        source_path = tmp_path / filename
        source = openpyxl.Workbook()
        source.active.title = "考勤"
        source.active["A1"] = value
        source.save(source_path)
        source_paths.append(str(source_path))

    copied = _copy_source_workbooks(str(output_path), source_paths)

    result = openpyxl.load_workbook(output_path, data_only=False)
    attendance_sheets = [sheet for sheet in result.worksheets if "考勤" in sheet.title]
    assert len(attendance_sheets) == 2
    assert {sheet["A1"].value for sheet in attendance_sheets} == {"第一份", "第二份"}
    assert len(copied) == 2


def test_copy_source_workbooks_verifies_array_formulas_by_content(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_path = tmp_path / "array-formula.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "计算明细"
    source.active["A1"] = ArrayFormula(ref="A1", text="=SUM(B1:B2)")
    source.active["B1"] = 1
    source.active["B2"] = 2
    source.save(source_path)

    copied = _copy_source_workbooks(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    formula = result["计算明细"]["A1"].value
    assert copied == ["计算明细"]
    assert isinstance(formula, ArrayFormula)
    assert formula.text == "=SUM(B1:B2)"


def test_copy_source_workbooks_refreshes_existing_detail_sheet_without_changing_order(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output["工资核算"]["A1"] = "=考勤!A2"
    attendance = output.create_sheet("考勤")
    attendance["A1"] = "姓名"
    attendance["A2"] = "旧员工"
    output.create_sheet("待人工处理")
    output.save(output_path)

    source_path = tmp_path / "attendance.xlsx"
    source = openpyxl.Workbook()
    source.active.title = "考勤"
    source.active["A1"] = "姓名"
    source.active["A2"] = "新员工"
    source.save(source_path)

    copied = _copy_source_workbooks(str(output_path), [str(source_path)])

    result = openpyxl.load_workbook(output_path, data_only=False)
    assert copied == ["考勤"]
    assert result.sheetnames == ["工资核算", "考勤", "待人工处理"]
    assert result["考勤"]["A2"].value == "新员工"
    assert result["工资核算"]["A1"].value == "=考勤!A2"


def test_copy_novel_source_sheets_maps_synonyms_and_only_adds_new_business_sheet(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.create_sheet("人员异动")
    output.save(output_path)

    source_path = tmp_path / "changes.xlsx"
    source = openpyxl.Workbook()
    personnel = source.active
    personnel.title = "入离职、转岗、转正"
    personnel.append(["工号", "姓名", "新部门", "生效日期"])
    personnel.append(["E001", "甲", "新部门", "2026-07-20"])
    novel = source.create_sheet("全新项目专项成本")
    novel.append(["项目编码", "项目名称", "专项成本"])
    novel.append(["P001", "新项目", 5000])
    source.save(source_path)

    result = _copy_novel_source_sheets(str(output_path), [str(source_path)])

    workbook = openpyxl.load_workbook(output_path, data_only=False)
    assert "入离职、转岗、转正" not in workbook.sheetnames
    assert "全新项目专项成本" in workbook.sheetnames
    assert workbook["全新项目专项成本"]["C2"].value == 5000
    assert result["copied_sheets"] == ["全新项目专项成本"]
    assert result["mapped_sheets"][0]["source_sheet"] == "入离职、转岗、转正"
    assert result["mapped_sheets"][0]["target_sheet"] == "人员异动"
    workbook.close()


def test_known_payroll_detail_sheet_updates_payroll_without_being_added(tmp_path) -> None:
    output_path = tmp_path / "result.xlsx"
    output = openpyxl.Workbook()
    output.active.title = "工资核算"
    output.save(output_path)

    source_path = tmp_path / "bonus.xlsx"
    source = openpyxl.Workbook()
    bonus = source.active
    bonus.title = "奖金-7月"
    bonus.append(["工号", "姓名", "绩效奖金"])
    bonus.append(["E001", "甲", 1000])
    source.save(source_path)

    result = _copy_novel_source_sheets(str(output_path), [str(source_path)])

    workbook = openpyxl.load_workbook(output_path, data_only=False)
    assert "奖金-7月" not in workbook.sheetnames
    assert result["copied_sheets"] == []
    assert result["mapped_sheets"][0]["target_sheet"] == "工资核算"
    workbook.close()


def test_export_updates_the_uploaded_master_instead_of_switching_to_builtin(tmp_path, monkeypatch) -> None:
    uploaded_template = tmp_path / "04-工资核算.xlsx"
    uploaded = openpyxl.Workbook()
    uploaded.active.title = "工资核算"
    uploaded.save(uploaded_template)

    full_master = tmp_path / "大药房工资核算总表-v2.xlsx"
    master = openpyxl.Workbook()
    master.active.title = "规则备忘"
    master.create_sheet("工资核算")
    master.create_sheet("工资条")
    master.save(full_master)
    monkeypatch.setattr(pipeline_module, "BUILTIN_TEMPLATE_PATH", str(full_master))

    selected = _resolve_layout_template_path(str(uploaded_template))

    assert selected == str(uploaded_template)
