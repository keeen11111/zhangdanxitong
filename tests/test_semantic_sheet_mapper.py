from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from core.sheet_mapper import apply_semantic_sheet_updates
from backend.routers.pipeline import _apply_semantic_source_updates


def _save_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_maps_differently_named_personnel_sheets_by_content(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "人员异动": [["工号", "姓名", "部门", "岗位"], ["E001", "张三", "A部门", "专员"]],
    })
    _save_workbook(update_path, {
        "人员变更": [["员工工号", "人员姓名", "变更后部门", "职位"], ["E001", "张三", "B部门", "主管"]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook["人员异动"]
    assert workbook.sheetnames == ["人员异动"]
    assert sheet["C2"].value == "B部门"
    assert sheet["D2"].value == "主管"
    assert result["auto_update_count"] == 2
    assert result["issues"] == []


def test_maps_unknown_non_personnel_sheet_using_its_business_key(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]],
    })
    _save_workbook(update_path, {
        "八月费用调整": [["费用项目", "单据编号", "金额"], ["运输费", "BX-001", 120]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用明细"]["C2"].value == 120
    assert result["auto_update_count"] == 1
    assert result["issues"] == []


def test_sends_ambiguous_target_sheet_to_manual_review(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "人员异动": [["工号", "姓名", "部门"], ["E001", "张三", "A部门"]],
        "人员档案": [["工号", "姓名", "部门"], ["E001", "张三", "A部门"]],
    })
    _save_workbook(update_path, {
        "人员变更": [["员工工号", "人员姓名", "变更后部门"], ["E001", "张三", "B部门"]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["人员异动"]["C2"].value == "A部门"
    assert workbook["人员档案"]["C2"].value == "A部门"
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "ambiguous_sheet"
    assert result["issues"][0]["source_sheets"] == ["人员变更"]


def test_uses_user_confirmed_target_sheet_for_an_ambiguous_source(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "人员异动": [["工号", "姓名", "部门"], ["E001", "张三", "A部门"]],
        "人员档案": [["工号", "姓名", "部门"], ["E001", "张三", "A部门"]],
    })
    _save_workbook(update_path, {
        "人员变更": [["员工工号", "人员姓名", "变更后部门"], ["E001", "张三", "B部门"]],
    })

    result = apply_semantic_sheet_updates(
        master_path,
        [update_path],
        output_path,
        sheet_overrides={("更新表.xlsx", "人员变更"): "人员档案"},
    )

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["人员异动"]["C2"].value == "A部门"
    assert workbook["人员档案"]["C2"].value == "B部门"
    assert result["auto_update_count"] == 1
    assert result["issues"] == []


def test_matching_policy_can_hold_back_a_target_field(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {"费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]]})
    _save_workbook(update_path, {"费用调整": [["单据编号", "费用项目", "金额"], ["BX-001", "运输费", 120]]})

    result = apply_semantic_sheet_updates(
        master_path,
        [update_path],
        output_path,
        matching_policy={"source_rules": [{"source_sheet": "费用调整", "target_field": "金额", "action": "review"}]},
    )

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用明细"]["C2"].value == 100
    assert result["auto_update_count"] == 0
    assert any(issue["issue_type"] == "configured_field_review" for issue in result["issues"])


def test_pipeline_wrapper_never_adds_unmatched_source_sheet(tmp_path: Path) -> None:
    output_path = tmp_path / "结果.xlsx"
    source_path = tmp_path / "更新表.xlsx"
    _save_workbook(output_path, {
        "费用明细": [["单据号", "金额"], ["BX-001", 100]],
    })
    _save_workbook(source_path, {
        "完全未知主题": [["其他编号", "说明"], ["X-1", "不应复制"]],
    })

    result = _apply_semantic_source_updates(str(output_path), [str(source_path)])

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["费用明细"]
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "unmatched_sheet"


def test_sends_unknown_record_to_manual_review_without_inserting_after_total(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [
            ["单据号", "项目名称", "金额"],
            ["BX-001", "运输费", 100],
            ["合计", None, "=SUM(C2:C2)"],
        ],
    })
    _save_workbook(update_path, {
        "费用调整": [["单据编号", "费用项目", "金额"], ["BX-002", "住宿费", 300]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook["费用明细"]
    assert sheet.max_row == 3
    assert sheet["A3"].value == "合计"
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "unknown_record"


def test_requires_a_stable_business_key_before_updating_a_matching_sheet(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用汇总": [["费用类别", "金额"], ["差旅", 100]],
    })
    _save_workbook(update_path, {
        "费用调整": [["费用类别", "金额"], ["差旅", 150]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用汇总"]["B2"].value == 100
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "missing_business_key"


def test_does_not_treat_a_generic_sequence_number_as_a_business_key(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用清单": [["编号", "项目名称", "金额"], ["1", "运输费", 100]],
    })
    _save_workbook(update_path, {
        "费用调整": [["编号", "项目名称", "金额"], ["1", "运输费", 150]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用清单"]["C2"].value == 100
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "missing_business_key"


def test_sends_duplicate_source_business_keys_to_manual_review(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]],
    })
    _save_workbook(update_path, {
        "费用调整": [
            ["单据编号", "费用项目", "金额"],
            ["BX-001", "运输费", 120],
            ["BX-001", "运输费", 120],
        ],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用明细"]["C2"].value == 100
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "duplicate_source_record"


def test_never_overwrites_a_formula_cell_with_a_source_value(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", "=100+0"]],
    })
    _save_workbook(update_path, {
        "费用调整": [["单据编号", "费用项目", "金额"], ["BX-001", "运输费", 120]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用明细"]["C2"].value == "=100+0"
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "formula_target_conflict"


def test_never_overwrites_an_array_formula_with_a_source_value(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    master = Workbook()
    target = master.active
    target.title = "费用明细"
    target.append(["单据号", "项目名称", "金额"])
    target.append(["BX-001", "运输费", 100])
    target["C2"] = ArrayFormula(ref="C2", text="=100+0")
    master.save(master_path)
    master.close()
    _save_workbook(update_path, {
        "费用调整": [["单据编号", "费用项目", "金额"], ["BX-001", "运输费", 120]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    value = load_workbook(output_path, data_only=False)["费用明细"]["C2"].value
    assert isinstance(value, ArrayFormula)
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "formula_target_conflict"


def test_does_not_copy_a_source_formula_into_the_master(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]],
    })
    _save_workbook(update_path, {
        "费用调整": [["单据编号", "费用项目", "金额"], ["BX-001", "运输费", "=1+1"]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["费用明细"]["C2"].value == 100
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "source_formula_conflict"


def test_requires_topic_evidence_beyond_only_document_number_and_amount(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "收款登记": [["单据号", "金额"], ["D-001", 100]],
    })
    _save_workbook(update_path, {
        "付款登记": [["单据号", "金额"], ["D-001", 900]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["收款登记"]["B2"].value == 100
    assert result["auto_update_count"] == 0
    assert result["issues"][0]["issue_type"] == "insufficient_topic_evidence"


def test_returns_an_audit_record_for_each_automatic_cell_update(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]],
    })
    _save_workbook(update_path, {
        "费用调整": [["单据编号", "费用项目", "金额"], ["BX-001", "运输费", 120]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    assert result["updates"] == [{
        "target_sheet": "费用明细",
        "target_cell": "C2",
        "old_value": 100,
        "new_value": 120,
        "source_file": "更新表.xlsx",
        "source_sheet": "费用调整",
        "source_row": 2,
    }]


def test_prefers_current_month_attendance_value_over_previous_month(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    june_path = tmp_path / "考勤-6月.xlsx"
    july_path = tmp_path / "考勤-7月.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "考勤": [["工号", "考勤项目", "实际出勤天数"], ["E001", "月度考勤", 0]],
    })
    _save_workbook(june_path, {
        "考勤-6月": [["工号", "考勤项目", "实际出勤天数"], ["E001", "月度考勤", 20]],
    })
    _save_workbook(july_path, {
        "考勤-7月": [["工号", "考勤项目", "实际出勤天数"], ["E001", "月度考勤", 22]],
    })

    result = apply_semantic_sheet_updates(
        master_path, [june_path, july_path], output_path, salary_month="2026.07"
    )

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["考勤"]["C2"].value == 22
    assert result["auto_update_count"] == 1
    assert result["issues"] == []


def test_uses_original_upload_names_to_identify_the_current_month(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    june_path = tmp_path / "upload-a.xlsx"
    july_path = tmp_path / "upload-b.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "考勤": [["工号", "考勤项目", "实际出勤天数"], ["E001", "月度考勤", 0]],
    })
    _save_workbook(june_path, {
        "考勤": [["工号", "考勤项目", "实际出勤天数"], ["E001", "月度考勤", 20]],
    })
    _save_workbook(july_path, {
        "考勤": [["工号", "考勤项目", "实际出勤天数"], ["E001", "月度考勤", 22]],
    })

    result = apply_semantic_sheet_updates(
        master_path,
        [june_path, july_path],
        output_path,
        salary_month="2026.07",
        source_names={str(june_path): "考勤-6月.xlsx", str(july_path): "考勤-7月.xlsx"},
    )

    assert load_workbook(output_path, data_only=False)["考勤"]["C2"].value == 22
    assert result["auto_update_count"] == 1
    assert result["issues"] == []


def test_sends_an_empty_source_sheet_to_manual_review_instead_of_ignoring_it(tmp_path: Path) -> None:
    master_path = tmp_path / "总表.xlsx"
    update_path = tmp_path / "更新表.xlsx"
    output_path = tmp_path / "结果.xlsx"
    _save_workbook(master_path, {
        "费用明细": [["单据号", "项目名称", "金额"], ["BX-001", "运输费", 100]],
    })
    _save_workbook(update_path, {
        "空白费用调整": [["单据编号", "费用项目", "金额"]],
    })

    result = apply_semantic_sheet_updates(master_path, [update_path], output_path)

    assert result["auto_update_count"] == 0
    assert result["issues"] == [{
        "issue_type": "empty_source_sheet",
        "message": "来源 Sheet 没有可处理的数据记录，未自动写入。",
        "source_files": ["更新表.xlsx"],
        "source_sheets": ["空白费用调整"],
    }]
