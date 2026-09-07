from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
import openpyxl
from fastapi import HTTPException

from backend.routers.agent import (
    AgentRunCreateIn,
    create_agent_run,
    _classify_message_decision,
    _detect_month_conflict,
    _issue_to_work_item,
    _extract_candidate_value,
    _public_model_status,
    _parse_model_reply,
    _parse_rule_candidates,
    _active_sheet_mapping_overrides,
    _rerun_after_sheet_mapping,
    _try_safe_auto_apply,
    _active_memory_decision,
    AgentApplyIn,
    AgentMessageIn,
    apply_agent_item,
    AgentRulePackageActivateIn,
    activate_agent_rule_package,
    download_accepted_agent_draft,
    message_agent_run,
)


class _Query:
    def __init__(self, first_value=None, all_values=None):
        self.first_value = first_value
        self.all_values = all_values or []

    def filter(self, *_args):
        return self

    def first(self):
        return self.first_value

    def all(self):
        return self.all_values


class _Db:
    def __init__(self, project):
        self.project = project

    def query(self, model):
        from backend.models import Project
        return _Query(self.project if model is Project else None, [])


class _DbWithFiles:
    def __init__(self, project, files):
        self.project = project
        self.files = files

    def query(self, model):
        from backend.models import Project, UploadFile
        if model is Project:
            return _Query(self.project, [])
        if model is UploadFile:
            return _Query(None, self.files)
        return _Query(None, [])


def test_agent_message_parser_only_accepts_explicit_business_choices() -> None:
    assert _classify_message_decision("确认更新到总表") == "apply_proposed"
    assert _classify_message_decision("本次不更新，保留原值") == "keep_current"
    assert _classify_message_decision("我觉得这个人应该处理一下") is None


def test_month_conflict_is_a_blocking_confirmation() -> None:
    conflict = _detect_month_conflict(
        "202606（所属月202605)-北京科园-鹤安-大药房工资核算总表.xlsx",
        "2026.06",
    )

    assert conflict["required"] is True
    assert conflict["filename_month"] == "2026.05"
    assert conflict["configured_month"] == "2026.06"


def test_explicit_target_month_is_not_overridden_by_master_ownership_month(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.routers.agent as agent

    upload_root = tmp_path / "uploads"
    project_dir = upload_root / "project-1"
    project_dir.mkdir(parents=True)
    master = project_dir / "master.xlsx"
    source = project_dir / "source.xlsx"
    for path in (master, source):
        workbook = openpyxl.Workbook()
        workbook.save(path)
        workbook.close()
    project = SimpleNamespace(
        id="project-1", name="长期项目", salary_month="2026.04",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="测试企业")),
    )
    files = [
        SimpleNamespace(
            id="master-1", file_type="financial_master",
            original_name="202605（所属月202604）工资总表.xlsx",
            stored_path="project-1/master.xlsx",
        ),
        SimpleNamespace(
            id="source-1", file_type="financial_source",
            original_name="2026年5月变更.xlsx", stored_path="project-1/source.xlsx",
        ),
    ]
    monkeypatch.setattr(agent, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")

    result = create_agent_run(
        AgentRunCreateIn(project_id="project-1", target_salary_month="2026.05"),
        user=SimpleNamespace(id="user-1", tenant_id="tenant-a"),
        db=_DbWithFiles(project, files),
    )

    assert result["salary_month"] == "2026.05"
    assert result["target_salary_month"] == "2026.05"
    assert result["month_confirmation"]["filename_month"] == "2026.04"
    assert result["month_confirmation"]["confirmed"] is True
    assert result["month_confirmation"]["required"] is False


def test_ambiguous_instruction_months_block_without_explicit_target(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.routers.agent as agent

    upload_root = tmp_path / "uploads"
    project_dir = upload_root / "project-1"
    project_dir.mkdir(parents=True)
    master = project_dir / "master.xlsx"
    source = project_dir / "source.xlsx"
    for path in (master, source):
        workbook = openpyxl.Workbook()
        workbook.save(path)
        workbook.close()
    project = SimpleNamespace(
        id="project-1", name="长期项目", salary_month="2026.04",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="测试企业")),
    )
    files = [
        SimpleNamespace(id="master-1", file_type="financial_master", original_name="202604总表.xlsx", stored_path="project-1/master.xlsx"),
        SimpleNamespace(id="source-1", file_type="financial_source", original_name="变更.xlsx", stored_path="project-1/source.xlsx"),
    ]
    monkeypatch.setattr(agent, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")

    result = create_agent_run(
        AgentRunCreateIn(project_id="project-1", instruction="把2026年4月总表更新为2026年5月数据"),
        user=SimpleNamespace(id="user-1", tenant_id="tenant-a"),
        db=_DbWithFiles(project, files),
    )

    assert result["status"] == "blocked"
    assert result["code"] == "TARGET_MONTH_REQUIRED"
    assert result["month_confirmation"]["required"] is True


def test_next_month_run_uses_latest_accepted_result_and_only_new_sources(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.routers.agent as agent

    upload_root = tmp_path / "uploads"
    project_dir = upload_root / "project-1"
    export_dir = tmp_path / "exports" / "project-1"
    project_dir.mkdir(parents=True)
    export_dir.mkdir(parents=True)
    paths = {
        "master": project_dir / "master.xlsx",
        "old": project_dir / "old.xlsx",
        "new": project_dir / "new.xlsx",
        "prior": export_dir / "Agent草稿_prior.xlsx",
    }
    for index, path in enumerate(paths.values(), start=1):
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = index
        workbook.save(path)
        workbook.close()
    monkeypatch.setattr(agent, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda project_id, filename: tmp_path / "exports" / project_id / filename)
    prior_run = {
        "run_id": "a" * 32, "tenant_id": "tenant-a", "project_id": "project-1",
        "salary_month": "2026.04", "status": "completed",
        "created_at": "2026-05-01T00:00:00+00:00", "updated_at": "2026-05-01T00:00:00+00:00",
        "draft_filename": paths["prior"].name,
        "result": {"filename": paths["prior"].name, "sha256": agent.file_digest(paths["prior"])},
        "validation": {"status": "structurally_valid"},
        "file_manifest": [{
            "id": "source-old", "filename": "4月变更.xlsx", "role": "financial_source",
            "sha256": agent.file_digest(paths["old"]),
        }],
        "items": [], "events": [],
    }
    agent._save_run(prior_run)
    project = SimpleNamespace(
        id="project-1", name="长期项目", salary_month="2026.04",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="测试企业")),
    )
    files = [
        SimpleNamespace(id="master-1", file_type="financial_master", original_name="202604总表.xlsx", stored_path="project-1/master.xlsx"),
        SimpleNamespace(id="source-old", file_type="financial_source", original_name="4月变更.xlsx", stored_path="project-1/old.xlsx"),
        SimpleNamespace(id="source-new", file_type="financial_source", original_name="5月变更.xlsx", stored_path="project-1/new.xlsx"),
    ]

    result = create_agent_run(
        AgentRunCreateIn(project_id="project-1", target_salary_month="2026.05"),
        user=SimpleNamespace(id="user-1", tenant_id="tenant-a"),
        db=_DbWithFiles(project, files),
    )

    saved = agent._load_run(result["run_id"], SimpleNamespace(tenant_id="tenant-a"))
    assert saved["_master_path"] == str(paths["prior"].resolve())
    assert saved["source_files"] == ["5月变更.xlsx"]
    assert saved["baseline"]["run_id"] == prior_run["run_id"]
    assert saved["baseline"]["salary_month"] == "2026.04"
    assert saved["requires_month_roll_forward"] is True


def test_legacy_incomplete_report_is_not_accepted_as_next_month_baseline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.routers.agent as agent

    run_dir = tmp_path / "runs" / "tenant-a"
    export_dir = tmp_path / "exports" / "project-1"
    run_dir.mkdir(parents=True)
    export_dir.mkdir(parents=True)
    output = export_dir / "partial.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "partial"
    workbook.save(output)
    workbook.close()
    partial = {
        "run_id": "legacy-partial",
        "tenant_id": "tenant-a",
        "project_id": "project-1",
        "salary_month": "2026.05",
        "status": "completed",
        "validation": {"status": "structurally_valid"},
        "result": {"filename": output.name, "sha256": agent.file_digest(output)},
        "execution_result": {
            "status": "completed",
            "content": "已插入占位行，但未能完成业务字段写入，任务未完成。",
        },
        "instruction": "把新增人员写入总表",
    }
    (run_dir / "legacy-partial.json").write_text(
        json.dumps(partial, ensure_ascii=False), encoding="utf-8",
    )
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(
        agent, "_result_path", lambda project_id, filename: tmp_path / "exports" / project_id / filename,
    )

    assert agent._latest_accepted_project_run("tenant-a", "project-1") is None
    public = agent._public_run(partial)
    assert public["status"] == "execution_incomplete"
    assert public["execution_result"]["code"] == "INCOMPLETE_MODEL_RESPONSE"


def test_third_month_uses_highest_accepted_month_and_never_replays_older_sources(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.routers.agent as agent

    upload_root = tmp_path / "uploads"
    project_dir = upload_root / "project-1"
    export_dir = tmp_path / "exports" / "project-1"
    project_dir.mkdir(parents=True)
    export_dir.mkdir(parents=True)
    paths = {
        "master": project_dir / "master.xlsx",
        "april_source": project_dir / "april.xlsx",
        "may_source": project_dir / "may.xlsx",
        "june_source": project_dir / "june.xlsx",
        "april_result": export_dir / "Agent草稿_april.xlsx",
        "may_result": export_dir / "Agent草稿_may.xlsx",
    }
    for index, path in enumerate(paths.values(), start=1):
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = index
        workbook.save(path)
        workbook.close()
    monkeypatch.setattr(agent, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(
        agent,
        "_result_path",
        lambda project_id, filename: tmp_path / "exports" / project_id / filename,
    )
    accepted_runs = [
        {
            "run_id": "a" * 32, "tenant_id": "tenant-a", "project_id": "project-1",
            "salary_month": "2026.04", "status": "completed",
            # Reopening an old run can make it newer by timestamp; it must not
            # become the baseline for the next salary month.
            "created_at": "2026-05-01T00:00:00+00:00",
            "updated_at": "2026-07-01T00:00:00+00:00",
            "result": {
                "filename": paths["april_result"].name,
                "sha256": agent.file_digest(paths["april_result"]),
            },
            "validation": {"status": "structurally_valid"},
            "file_manifest": [{
                "id": "source-april", "filename": "4月变更.xlsx",
                "role": "financial_source", "sha256": agent.file_digest(paths["april_source"]),
            }],
            "items": [], "events": [],
        },
        {
            "run_id": "b" * 32, "tenant_id": "tenant-a", "project_id": "project-1",
            "salary_month": "2026.05", "status": "completed",
            "created_at": "2026-06-01T00:00:00+00:00",
            "updated_at": "2026-06-01T00:00:00+00:00",
            "result": {
                "filename": paths["may_result"].name,
                "sha256": agent.file_digest(paths["may_result"]),
            },
            "validation": {"status": "structurally_valid"},
            "file_manifest": [{
                "id": "source-may", "filename": "5月变更.xlsx",
                "role": "financial_source", "sha256": agent.file_digest(paths["may_source"]),
            }],
            "items": [], "events": [],
        },
    ]
    for accepted in reversed(accepted_runs):
        agent._save_run(accepted)
    project = SimpleNamespace(
        id="project-1", name="长期项目", salary_month="2026.04",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="测试企业")),
    )
    files = [
        SimpleNamespace(id="master", file_type="financial_master", original_name="总表.xlsx", stored_path="project-1/master.xlsx"),
        SimpleNamespace(id="april", file_type="financial_source", original_name="4月变更.xlsx", stored_path="project-1/april.xlsx"),
        SimpleNamespace(id="may", file_type="financial_source", original_name="5月变更.xlsx", stored_path="project-1/may.xlsx"),
        SimpleNamespace(id="june", file_type="financial_source", original_name="6月变更.xlsx", stored_path="project-1/june.xlsx"),
    ]

    result = create_agent_run(
        AgentRunCreateIn(project_id="project-1", target_salary_month="2026.06"),
        user=SimpleNamespace(id="user-1", tenant_id="tenant-a"),
        db=_DbWithFiles(project, files),
    )

    saved = agent._load_run(result["run_id"], SimpleNamespace(tenant_id="tenant-a"))
    assert saved["baseline"]["run_id"] == "b" * 32
    assert saved["baseline"]["salary_month"] == "2026.05"
    assert saved["source_files"] == ["6月变更.xlsx"]

    duplicate_month = create_agent_run(
        AgentRunCreateIn(project_id="project-1", target_salary_month="2026.05"),
        user=SimpleNamespace(id="user-1", tenant_id="tenant-a"),
        db=_DbWithFiles(project, files),
    )
    assert duplicate_month["status"] == "blocked"
    assert duplicate_month["code"] == "TARGET_MONTH_REQUIRED"


def test_issue_conversion_does_not_expose_raw_model_prompt_or_credentials() -> None:
    issue = {
        "issue_type": "conflicting_value",
        "message": "多个来源值不同",
        "person_name": "张三",
        "target_sheet": "工资核算",
        "target_cell": "J12",
        "candidate_values": [2000, 3000],
        "source_files": ["来源.xlsx"],
        "source_sheets": ["奖金"],
    }

    item = _issue_to_work_item(issue, 0)

    assert item["id"]
    assert item["status"] == "needs_review"
    assert item["person_name"] == "张三"
    assert "prompt" not in item
    assert "password" not in item


def test_issue_conversion_has_stable_id_for_retry() -> None:
    issue = {"issue_type": "unknown_record", "message": "不存在", "source_files": []}
    first = _issue_to_work_item(issue, 0)
    second = _issue_to_work_item(issue, 0)
    assert first["id"] == second["id"]


def test_empty_source_issue_is_auto_skipped_and_keeps_audit_message() -> None:
    item = _issue_to_work_item({
        "issue_type": "empty_source_sheet",
        "source_sheets": ["补发补扣"],
        "candidate_target_sheets": ["工资核算"],
    }, 0)

    assert item["status"] == "skipped"
    assert item["decision"] == "skip"
    assert item["messages"][0]["role"] == "agent"


def test_issue_conversion_preserves_safe_source_evidence() -> None:
    item = _issue_to_work_item({
        "issue_type": "ambiguous_sheet",
        "candidate_target_sheets": ["工资核算", "台账"],
        "source_rows": [4, 5],
        "business_key": "E001",
        "risk_level": "high",
    }, 0)

    assert item["candidate_target_sheets"] == ["工资核算", "台账"]
    assert item["source_rows"] == [4, 5]
    assert item["business_key"] == "E001"
    assert item["risk_level"] == "high"


def test_safe_auto_apply_requires_single_candidate_and_writes_draft(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import openpyxl
    import backend.routers.agent as agent

    draft = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook["工资核算"]["J2"] = 0
    workbook.save(draft)
    workbook.close()
    monkeypatch.setattr(agent, "_result_path", lambda _project, _filename: draft)
    run = {"run_id": "run-1", "project_id": "p", "draft_filename": draft.name, "items": []}
    item = {
        "id": "item-1", "status": "needs_review", "issue_type": "review_required",
        "target_sheet": "工资核算", "target_cell": "J2", "candidate_values": [2000],
        "current_value": 0, "person_name": "李某", "original_issue": None,
    }

    assert _try_safe_auto_apply(run, item) is True
    assert item["status"] == "auto_applied"
    workbook = openpyxl.load_workbook(draft, data_only=False)
    assert workbook["工资核算"]["J2"].value == 2000
    workbook.close()


def test_safe_auto_apply_rejects_high_risk_even_with_one_candidate(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import openpyxl
    import backend.routers.agent as agent

    draft = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook["工资核算"]["J2"] = 0
    workbook.save(draft)
    workbook.close()
    monkeypatch.setattr(agent, "_result_path", lambda _project, _filename: draft)
    run = {"run_id": "run-1", "project_id": "p", "draft_filename": draft.name, "items": []}
    item = {
        "id": "item-1", "status": "needs_review", "issue_type": "review_required",
        "risk_level": "high", "target_sheet": "工资核算", "target_cell": "J2",
        "candidate_values": [2000], "current_value": 0, "person_name": "李某",
    }

    assert _try_safe_auto_apply(run, item) is False
    assert item["status"] == "needs_review"
    workbook = openpyxl.load_workbook(draft, data_only=False)
    assert workbook["工资核算"]["J2"].value == 0
    workbook.close()


def test_active_company_memory_reuses_only_matching_candidate(monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent.experience_store, "suggest_memory", lambda *_args, **_kwargs: [{
        "status": "active", "rule_id": "rule-1", "decision": "confirmed",
        "updates": {"action": "apply_proposed", "value": 2000},
    }])
    run = {"tenant_id": "tenant-a", "project_id": "p"}
    item = {
        "issue_type": "bonus_source", "risk_level": "low", "target_sheet": "工资核算",
        "source_sheets": ["奖金"], "candidate_values": [2000],
    }
    assert _active_memory_decision(run, item) == ("apply_proposed", 2000, "rule-1")

    item["candidate_values"] = [3000]
    assert _active_memory_decision(run, item) is None


def test_active_company_sheet_memory_is_reused_only_for_a_single_valid_target(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import openpyxl
    import backend.routers.agent as agent

    master = tmp_path / "master.xlsx"
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资核算"
    workbook.create_sheet("台账")
    workbook.save(master)
    workbook.close()
    workbook = openpyxl.Workbook()
    workbook.active.title = "奖金-7月"
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(agent.experience_store, "suggest_memory", lambda *_args, **_kwargs: [{
        "status": "active", "updates": {"target_sheet": "工资核算"}, "rule_id": "map-1",
    }])
    run = {
        "tenant_id": "tenant-1", "project_id": "project-1",
        "_source_paths": {"薪资数据.xlsx": str(source)},
    }
    assert _active_sheet_mapping_overrides(run, master) == {("薪资数据.xlsx", "奖金-7月"): "工资核算"}

    monkeypatch.setattr(agent.experience_store, "suggest_memory", lambda *_args, **_kwargs: [
        {"status": "active", "updates": {"target_sheet": "工资核算"}},
        {"status": "active", "updates": {"target_sheet": "台账"}},
    ])
    assert _active_sheet_mapping_overrides(run, master) == {}


def test_sheet_mapping_rerun_updates_existing_draft_and_resolves_issue(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import openpyxl
    import backend.routers.agent as agent

    upload_root = tmp_path / "uploads"
    project_dir = upload_root / "project-1"
    project_dir.mkdir(parents=True)
    source_path = project_dir / "source.xlsx"
    draft_path = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "人员异动"
    first.append(["工号", "姓名", "部门"])
    first.append(["E001", "张三", "A部门"])
    second = workbook.create_sheet("人员档案")
    second.append(["工号", "姓名", "部门"])
    second.append(["E001", "张三", "A部门"])
    workbook.save(draft_path)
    workbook.close()
    workbook = openpyxl.Workbook()
    workbook.active.title = "人员变更"
    workbook.active.append(["员工工号", "人员姓名", "变更后部门"])
    workbook.active.append(["E001", "张三", "B部门"])
    workbook.save(source_path)
    workbook.close()
    monkeypatch.setattr(agent, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(agent, "_result_path", lambda _project, _filename: draft_path)
    monkeypatch.setattr(agent, "_write_result_meta", lambda *_args: None)
    monkeypatch.setattr(agent.experience_store, "record_memory", lambda **_kwargs: None)

    class _FilesQuery:
        def all(self):
            return [
                SimpleNamespace(file_type="financial_master", stored_path="project-1/master.xlsx", original_name="master.xlsx"),
                SimpleNamespace(file_type="financial_source", stored_path="project-1/source.xlsx", original_name="source.xlsx"),
            ]

        def filter(self, *_args):
            return self

    class _FilesDb:
        def query(self, *_args):
            return _FilesQuery()

    run = {
        "run_id": "a" * 32, "tenant_id": "tenant-a", "project_id": "project-1",
        "salary_month": "2026.07", "draft_filename": draft_path.name, "rule_version": "v1",
        "sheet_mappings": [{"source_file": "source.xlsx", "source_sheet": "人员变更", "target_sheet": "人员档案"}],
        "items": [{
            "id": "item-1", "status": "needs_review", "revision": 1,
            "issue_type": "ambiguous_sheet", "source_files": ["source.xlsx"],
            "source_sheets": ["人员变更"], "candidate_target_sheets": ["人员异动", "人员档案"],
            "original_issue": {"issue_type": "ambiguous_sheet", "source_files": ["source.xlsx"], "source_sheets": ["人员变更"]},
        }],
        "integration_meta": {"filename": draft_path.name, "review_filename": ""},
    }
    result = _rerun_after_sheet_mapping(
        run, run["items"][0], "人员档案", "同表头且业务主题一致", SimpleNamespace(id="u1", tenant_id="tenant-a"), _FilesDb(),
    )
    assert result["applied"] is True
    assert run["items"][0]["status"] == "resolved"
    workbook = openpyxl.load_workbook(draft_path, data_only=False)
    assert workbook["人员档案"]["C2"].value == "B部门"
    assert workbook["人员异动"]["C2"].value == "A部门"
    workbook.close()
    assert not any(issue.get("issue_type") == "ambiguous_sheet" for issue in run["integration_meta"]["issues"])


def test_agent_run_is_blocked_with_actionable_detail_until_files_exist(tmp_path: Path, monkeypatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    project = SimpleNamespace(
        id="project-1", salary_month="2026.07",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="科园")),
    )
    user = SimpleNamespace(id="user-1", tenant_id="tenant-a")

    result = create_agent_run(
        AgentRunCreateIn(project_id="project-1"), user=user, db=_Db(project)
    )

    assert result["status"] == "blocked"
    assert "上传" in result["detail"]


def test_configured_named_sample_project_can_start_without_uploaded_files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_load_demo_for_project", lambda *_args: {
        "filename": "样本一_已更新.xlsx", "sha256": "digest", "_reference_path": str(tmp_path / "reference.xlsx"),
    })
    project = SimpleNamespace(
        id="project-1", name="样本一", salary_month="2026.07",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="演示企业")),
    )
    user = SimpleNamespace(id="user-1", tenant_id="tenant-a")

    result = create_agent_run(AgentRunCreateIn(project_id="project-1"), user=user, db=_Db(project))

    assert result["status"] == "planning"
    assert result["execution_mode"] == "demo"
    assert result["plan_confirmation"] == {"required": False, "confirmed": True}


def test_message_candidate_extraction_requires_an_offered_value() -> None:
    assert _extract_candidate_value("采用 2,000 元", [2000, 3000]) == 2000
    assert _extract_candidate_value("采用 2500 元", [2000, 3000]) is None


def test_public_model_status_never_exposes_secret(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PAYROLL_MODEL_PROVIDER", "openai_compatible")
    monkeypatch.setenv("PAYROLL_MODEL_BASE_URL", "https://dashscope.example/v1")
    monkeypatch.setenv("PAYROLL_MODEL_API_KEY", "secret-value")
    monkeypatch.setenv("PAYROLL_MODEL_NAME", "qwen3.8-max")
    monkeypatch.setenv("PAYROLL_MODEL_API_STYLE", "responses")
    monkeypatch.setenv("PAYROLL_MODEL_ENABLE_THINKING", "true")

    status = _public_model_status()

    assert status["configured"] is True
    assert status["model"] == "qwen3.8-max"
    assert status["api_style"] == "responses"
    assert status["enable_thinking"] is True
    assert "secret-value" not in repr(status)


def test_model_reply_parser_accepts_only_bounded_decision_shape() -> None:
    parsed = _parse_model_reply('{"reply":"建议使用 K 列","decision":"apply_proposed","value":800}')

    assert parsed["reply"] == "建议使用 K 列"
    assert parsed["decision"] == "apply_proposed"
    assert parsed["value"] == 800

    plain = _parse_model_reply("需要结合人员类别确认。")
    assert plain == {"reply": "需要结合人员类别确认。", "decision": None, "value": None}


def test_apply_is_idempotent_for_double_submit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    run = {
        "run_id": "a" * 32,
        "tenant_id": "tenant-a",
        "project_id": "project-1",
        "rule_version": "tenant-a:2026.07",
        "salary_month": "2026.07",
        "items": [{
            "id": "item-1", "status": "needs_review", "revision": 1,
            "candidate_values": [2000], "proposed_value": 2000,
            "target_sheet": "工资", "target_cell": "J2",
            "current_value": 0, "source_files": [], "source_sheets": [],
        }],
        "integration_meta": {"issues": [], "matches": [{"sheet": "工资"}]},
    }
    agent._save_run(run)
    monkeypatch.setattr(agent, "_apply_cell_value", lambda _run, item, value: item.update(applied_value=value))
    user = SimpleNamespace(id="user-1", tenant_id="tenant-a")

    first = apply_agent_item("a" * 32, "item-1", AgentApplyIn(action="apply_proposed", value=2000, expected_revision=1), user=user)
    second = apply_agent_item("a" * 32, "item-1", AgentApplyIn(action="apply_proposed", value=2000, expected_revision=1), user=user)

    assert first["items"][0]["status"] == "resolved"
    assert second["items"][0]["status"] == "resolved"


def test_manual_item_apply_commits_workbook_via_atomic_replace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Manual writes must use the same atomic boundary as model writes."""
    import backend.routers.agent as agent

    workbook_path = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资"
    workbook.active["J2"] = 0
    workbook.save(workbook_path)
    workbook.close()

    replacements: list[tuple[str, str]] = []
    original_replace = agent.os.replace

    def recording_replace(source: str | bytes | Path, target: str | bytes | Path) -> None:
        replacements.append((str(source), str(target)))
        original_replace(source, target)

    monkeypatch.setattr(agent, "_result_path", lambda *_args: workbook_path)
    monkeypatch.setattr(agent.os, "replace", recording_replace)
    run = {
        "project_id": "project-1", "draft_filename": workbook_path.name,
        "workbook_updates": [],
    }
    item = {
        "id": "item-1", "target_sheet": "工资", "target_cell": "J2",
        "current_value": 0,
    }

    agent._apply_cell_value(run, item, 2000)

    assert replacements and replacements[-1][1] == str(workbook_path)
    assert run["workbook_updates"] == [{
        "item_id": "item-1", "status": "applied",
        "target_sheet": "工资", "target_cell": "J2",
        "before": 0, "after": 2000,
        "old_value": 0, "new_value": 2000,
        "rule": "manual_item_apply",
    }]
    check = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        assert check["工资"]["J2"].value == 2000
    finally:
        check.close()


def test_retry_rolls_back_once_atomically_and_records_the_reverse_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import backend.routers.agent as agent

    run_id = "b" * 32
    workbook_path = tmp_path / "draft.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工资"
    workbook.active["J2"] = 2000
    workbook.save(workbook_path)
    workbook.close()
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda *_args: workbook_path)
    run = {
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-1",
        "draft_filename": workbook_path.name, "events": [], "workbook_updates": [],
        "items": [{
            "id": "item-1", "status": "resolved", "retry_count": 0,
            "target_sheet": "工资", "target_cell": "J2",
            "applied_value": 2000,
            "applied_history": [{"old_value": 0, "new_value": 2000}],
        }],
    }
    agent._save_run(run)

    result = agent.retry_agent_item(
        run_id, "item-1", user=SimpleNamespace(id="user-1", tenant_id="tenant-a"),
    )

    check = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        assert check["工资"]["J2"].value == 0
    finally:
        check.close()
    assert result["items"][0]["status"] == "needs_review"
    assert result["items"][0].get("applied_value") is None
    saved = agent._load_run(run_id, SimpleNamespace(tenant_id="tenant-a"))
    assert saved["items"][0]["applied_history"] == []
    assert saved["workbook_updates"][-1]["rule"] == "retry_rollback"


def test_rule_candidate_parser_rejects_untrusted_free_form_output() -> None:
    with pytest.raises(ValueError, match="规则 JSON"):
        _parse_rule_candidates("请执行 =SUM(A1:A2)")


def test_rule_package_context_separates_active_and_candidate_packages(monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "_load_material_index", lambda _tenant, _project: [
        {"kind": "rule_package", "material_id": "candidate", "project_id": "p", "filename": "candidate.json",
         "created_at": "2026-08-02", "rule_package": {"package_id": "candidate", "version": "v2", "status": "candidate", "rules": [{"rule_id": "new"}], "sources": []}},
        {"kind": "rule_package", "material_id": "active", "project_id": "p", "filename": "active.json",
         "created_at": "2026-08-01", "rule_package": {"package_id": "active", "version": "v1", "status": "active", "rules": [{"rule_id": "old"}], "sources": []}},
    ])

    context = agent._rule_package_context({"tenant_id": "tenant-a", "project_id": "p"})

    assert context["active"]["package_id"] == "active"
    assert [item["package_id"] for item in context["candidates"]] == ["candidate"]


def test_activate_rule_package_retires_previous_active_package(monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    records = [
        {"kind": "rule_package", "material_id": "candidate", "project_id": "p", "filename": "candidate.json",
         "created_at": "2026-08-02", "rule_package": {"package_id": "candidate", "version": "v2", "status": "candidate", "rules": [], "sources": []}},
        {"kind": "rule_package", "material_id": "active", "project_id": "p", "filename": "active.json",
         "created_at": "2026-08-01", "rule_package": {"package_id": "active", "version": "v1", "status": "active", "rules": [], "sources": []}},
    ]
    saved: list[dict] = []
    monkeypatch.setattr(agent, "_project_or_404", lambda _project, _user, _db: object())
    monkeypatch.setattr(agent, "_load_material_index", lambda _tenant, _project: records)
    monkeypatch.setattr(agent, "_save_material_index", lambda _tenant, _project, value: saved.extend(value))

    result = activate_agent_rule_package(
        "p", "candidate", AgentRulePackageActivateIn(confirm=True),
        user=SimpleNamespace(tenant_id="tenant-a"), db=None,
    )

    assert result["package"]["status"] == "active"
    assert {item["rule_package"]["package_id"]: item["rule_package"]["status"] for item in saved} == {
        "candidate": "active", "active": "retired",
    }


def _accepted_download_run(*, validation_status: str, draft_filename: str) -> dict:
    return {
        "run_id": "d" * 32,
        "tenant_id": "tenant-a",
        "project_id": "project-1",
        "validation": {"status": validation_status},
        "draft_filename": draft_filename,
    }


def test_accepted_draft_download_rejects_unvalidated_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    agent._save_run(_accepted_download_run(validation_status="not_verified", draft_filename="draft.xlsx"))

    with pytest.raises(HTTPException, match="尚未通过验收") as exc_info:
        download_accepted_agent_draft("d" * 32, user=SimpleNamespace(tenant_id="tenant-a"))

    assert getattr(exc_info.value, "status_code", None) == 409


def test_accepted_draft_download_returns_validated_workbook(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    workbook = tmp_path / "draft.xlsx"
    workbook.write_bytes(b"validated workbook")
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project_id, _filename: workbook)
    agent._save_run(_accepted_download_run(validation_status="passed", draft_filename=workbook.name))

    response = download_accepted_agent_draft("d" * 32, user=SimpleNamespace(tenant_id="tenant-a"))

    assert response.path == workbook
    assert response.filename == workbook.name


def test_accepted_draft_download_rejects_path_traversal_filename(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    agent._save_run(_accepted_download_run(validation_status="passed", draft_filename="../draft.xlsx"))

    with pytest.raises(HTTPException, match="文件无效") as exc_info:
        download_accepted_agent_draft("d" * 32, user=SimpleNamespace(tenant_id="tenant-a"))

    assert getattr(exc_info.value, "status_code", None) == 409


def test_run_conversation_works_without_an_open_review_item(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import backend.routers.agent as agent

    monkeypatch.setattr(agent, "RUN_DIR", tmp_path)
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: SimpleNamespace())
    monkeypatch.setattr(
        agent,
        "OpenAICompatibleProvider",
        lambda _config: SimpleNamespace(complete=lambda **_kwargs: SimpleNamespace(content="可以继续说明本次处理要求。")),
    )
    run = _accepted_download_run(validation_status="not_verified", draft_filename="draft.xlsx")
    run.update(status="awaiting_review", detail="当前批次仍有未完成核对", items=[])
    agent._save_run(run)

    result = message_agent_run(
        "d" * 32,
        AgentMessageIn(message="请说明还有哪些内容没完成"),
        user=SimpleNamespace(tenant_id="tenant-a"),
    )

    assert result["message"]["role"] == "agent"
    assert result["message"]["content"] == "可以继续说明本次处理要求。"
    saved = agent._load_run("d" * 32, SimpleNamespace(tenant_id="tenant-a"))
    assert [entry["role"] for entry in saved["conversation"]] == ["user", "agent"]
    assert saved["items"] == []
