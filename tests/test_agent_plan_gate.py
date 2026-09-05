from __future__ import annotations

from copy import deepcopy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from fastapi import HTTPException
import openpyxl
import pytest

import backend.routers.agent as agent
from backend.models import Project, UploadFile
from core.document_agent.contracts import ToolCall
from core.document_agent.model import ModelConfig, ModelResponse


class _Query:
    def __init__(self, first_value: Any = None, all_values: list[Any] | None = None):
        self.first_value = first_value
        self.all_values = all_values or []

    def filter(self, *_args: Any) -> _Query:
        return self

    def first(self) -> Any:
        return self.first_value

    def all(self) -> list[Any]:
        return self.all_values


class _Db:
    def __init__(self, project: Any, files: list[Any]):
        self.project = project
        self.files = files

    def query(self, model: Any) -> _Query:
        if model is Project:
            return _Query(first_value=self.project)
        if model is UploadFile:
            return _Query(all_values=self.files)
        raise AssertionError(f"Unexpected database model: {model}")


@pytest.fixture
def scenario(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> SimpleNamespace:
    upload_root = tmp_path / "uploads"
    project_dir = upload_root / "project-1"
    project_dir.mkdir(parents=True)
    files = []
    for filename, role, value in [
        ("202608（所属月202607）-北京科园-鹤安-大药房工资核算总表-v2.xlsx", "financial_master", 0),
        ("奖金来源.xlsx", "financial_source", 2000),
    ]:
        path = project_dir / filename
        workbook = openpyxl.Workbook()
        workbook.active.title = "工资核算"
        workbook.active.append(["工号", "姓名", "奖金"])
        workbook.active.append(["E001", "测试人员", value])
        workbook.save(path)
        workbook.close()
        files.append(SimpleNamespace(
            id=role, original_name=filename, file_type=role,
            stored_path=f"project-1/{filename}", file_size=path.stat().st_size,
        ))
    project = SimpleNamespace(
        id="project-1", salary_month="2026.07",
        owner=SimpleNamespace(tenant_id="tenant-a", tenant=SimpleNamespace(name="测试公司")),
    )
    user = SimpleNamespace(id="user-1", tenant_id="tenant-a")
    calls: list[str] = []

    def legacy_integration(*_args: Any, **_kwargs: Any) -> SimpleNamespace:
        calls.append("legacy_integration")
        return SimpleNamespace(model_dump=lambda: {
            "filename": "legacy-draft.xlsx", "issues": [], "updates": [], "matches": [],
        })

    def release(*_args: Any, **_kwargs: Any) -> SimpleNamespace:
        calls.append("release")
        return SimpleNamespace(model_dump=lambda: {"filename": "published.xlsx"})

    def model_complete(*_args: Any, **_kwargs: Any) -> Any:
        calls.append("model")
        raise AssertionError("An unconfirmed run must not invoke the model")

    monkeypatch.setattr(agent, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(agent, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(agent, "_result_path", lambda _project, filename: tmp_path / "results" / filename)
    monkeypatch.setattr(agent, "_load_material_index", lambda *_args: [])
    monkeypatch.setattr(agent.experience_store, "suggest_memory", lambda *_args, **_kwargs: [])
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(agent.OpenAICompatibleProvider, "complete", model_complete)
    monkeypatch.setattr(agent, "create_financial_workbook_integration", legacy_integration)
    monkeypatch.setattr(agent, "release_latest_financial_workbook_integration", release)
    return SimpleNamespace(
        root=tmp_path, user=user, db=_Db(project, files), calls=calls,
        originals={path: path.read_bytes() for path in project_dir.glob("*.xlsx")},
    )


@pytest.mark.parametrize("auto_publish", [False, True])
def test_create_run_only_inventories_files_without_integrating_or_creating_draft(
    scenario: SimpleNamespace, auto_publish: bool,
) -> None:
    result = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1", auto_publish=auto_publish),
        user=scenario.user, db=scenario.db,
    )

    assert scenario.calls == []
    assert not result.get("draft_filename")
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_new_run_returns_unconfirmed_plan_and_file_roles(scenario: SimpleNamespace) -> None:
    result = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )

    assert result["plan_confirmation"]["required"] is True
    assert result["plan_confirmation"]["confirmed"] is False
    assert {entry["filename"]: entry["role"] for entry in result["file_manifest"]} == {
        "202608（所属月202607）-北京科园-鹤安-大药房工资核算总表-v2.xlsx": "financial_master",
        "奖金来源.xlsx": "financial_source",
    }
    saved = agent._load_run(result["run_id"], scenario.user)
    assert saved["plan_confirmation"] == result["plan_confirmation"]


def test_keyuan_files_use_generic_agent_mode(scenario: SimpleNamespace) -> None:
    scenario.db.project.salary_month = "2026.09"
    scenario.db.files[0].original_name = "202608（所属月202607）-北京科园-鹤安-大药房工资核算总表-v2.xlsx"
    scenario.db.files[1].original_name = "薪资数据-科园-7月薪资（8.14发薪）.xlsx"

    result = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )

    assert result.get("execution_mode") is None
    assert "自动化流程" not in str(result.get("detail") or "")


def test_project_with_configured_demo_uses_deterministic_demo_mode_by_default(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    scenario.db.project.name = "北京"
    monkeypatch.setattr(agent, "_load_demo_for_project", lambda *_args: {
        "filename": "北京工资核算总表.xlsx",
        "sha256": "verified-digest",
        "_reference_path": str(scenario.root / "verified-result.xlsx"),
    })

    result = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )

    assert result["execution_mode"] == "demo"
    assert result["demo_reference"] == {
        "filename": "北京工资核算总表.xlsx",
        "sha256": "verified-digest",
    }
    assert result["plan_confirmation"] == {"required": False, "confirmed": True}


def test_beijing_showcase_can_start_without_uploaded_payroll_files(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    scenario.db.project.name = "北京"
    scenario.db.files = []
    monkeypatch.setattr(agent, "_load_demo_for_project", lambda *_args: {
        "filename": "待确定稿.xlsx",
        "sha256": "verified-digest",
        "_reference_path": str(scenario.root / "verified-result.xlsx"),
    })

    result = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )

    assert result["execution_mode"] == "demo"
    assert result["status"] == "planning"
    assert result["plan_confirmation"] == {"required": False, "confirmed": True}


def test_legacy_beijing_run_is_upgraded_to_demo_before_next_start(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """An existing Beijing run must not fall back to the real model path."""
    monkeypatch.setattr(agent, "_load_demo_for_project", lambda *_args: {
        "filename": "待确定稿.xlsx",
        "sha256": "verified-digest",
        "_reference_path": str(scenario.root / "verified-result.xlsx"),
    })
    run = {
        "run_id": "c" * 32,
        "tenant_id": "tenant-a",
        "project_id": "project-1",
        "project_name": "北京",
        "salary_month": "2026.09",
        "status": "completed",
        "workflow": {"started_at": "2026-09-04T00:00:00+00:00"},
        "draft_filename": "旧结果.xlsx",
        "plan_confirmation": {"required": True, "confirmed": True},
        "items": [],
        "events": [],
        "messages": [],
    }
    agent._save_run(run)

    assert agent._coerce_named_demo_run(run) is True
    assert run["execution_mode"] == "demo"
    assert run["demo_reference"]["filename"] == "待确定稿.xlsx"
    assert run["demo_reference_project_id"] == "project-1"
    assert run["plan_confirmation"]["confirmed"] is True


@pytest.mark.parametrize("action", ["execute", "publish"])
def test_unconfirmed_plan_blocks_execution_and_publication_even_with_no_pending_items(
    scenario: SimpleNamespace, action: str,
) -> None:
    run_id = "a" * 32
    agent._save_run({
        "run_id": run_id, "tenant_id": "tenant-a", "project_id": "project-1",
        "salary_month": "2026.07", "status": "ready_to_publish", "items": [],
        "plan_confirmation": {"required": True, "confirmed": False},
        "month_confirmation": {"required": False, "confirmed": True},
    })

    with pytest.raises(HTTPException) as raised:
        if action == "execute":
            agent.execute_agent_run(run_id, user=scenario.user)
        else:
            agent.publish_agent_run(
                run_id, agent.AgentPublishIn(confirm_month=True), user=scenario.user, db=scenario.db,
            )

    assert raised.value.status_code == 409
    assert scenario.calls == []
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def _prepared_run(scenario: SimpleNamespace, *, confirmed: bool = False) -> dict[str, Any]:
    master = next(path for path in scenario.originals if path.name.endswith("工资核算总表-v2.xlsx"))
    source = next(path for path in scenario.originals if path.name == "奖金来源.xlsx")
    run = {
        "run_id": "b" * 32, "tenant_id": "tenant-a", "project_id": "project-1",
        "salary_month": "2026.07", "instruction": "按手册核对奖金", "status": "awaiting_review",
        "master_file": master.name, "source_files": [source.name],
        "_master_path": str(master), "_source_paths": {source.name: str(source)},
        "file_manifest": [
            {"filename": path.name, "role": role, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
            for path, role in [(master, "financial_master"), (source, "financial_source")]
        ],
        "plan_confirmation": {"required": True, "confirmed": confirmed},
        "month_confirmation": {"required": False, "confirmed": True},
        "items": [], "events": [], "messages": [],
    }
    agent._save_run(run)
    return run


def _stub_plan_provider(monkeypatch: pytest.MonkeyPatch, plan: dict[str, Any]) -> list[dict[str, Any]]:
    captured: list[dict[str, Any]] = []
    config = ModelConfig(
        provider="openai_compatible", base_url="https://model.example/v1",
        api_key="test-secret", model="test-planner",
    )

    def complete(_self: Any, **kwargs: Any) -> ModelResponse:
        captured.append(kwargs)
        return ModelResponse(content=json.dumps(plan), request_id="plan-1")

    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: config)
    monkeypatch.setattr(agent.OpenAICompatibleProvider, "complete", complete)
    return captured


@pytest.mark.parametrize("instruction", [None, "E001奖金取来源表2000元，保留原始总表"])
def test_generating_plan_is_read_only_and_uses_user_answers_without_confirming(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch, instruction: str | None,
) -> None:
    run = _prepared_run(scenario)
    plan = {"summary": "核对并更新副本", "steps": ["核对E001奖金"], "questions": []}
    captured = _stub_plan_provider(monkeypatch, plan)
    result = agent.confirm_agent_plan(
        run["run_id"], agent.AgentPlanIn(confirm_plan=False, instruction=instruction), user=scenario.user,
    )

    assert {key: result["model_plan"][key] for key in plan} == plan
    assert result["plan_confirmation"]["confirmed"] is False
    assert not result.get("draft_filename")
    assert scenario.calls == []
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals
    context = json.loads(captured[0]["messages"][1]["content"])
    assert run["instruction"] in context["instruction"]
    if instruction:
        assert instruction in context["instruction"]
    saved = agent._load_run(run["run_id"], scenario.user)
    assert saved["model_plan"] == result["model_plan"]


def test_plan_confirmation_with_open_questions_still_requires_answers(
    scenario: SimpleNamespace,
) -> None:
    run = _prepared_run(scenario)
    run["model_plan"] = {"summary": "需要澄清", "steps": ["核对奖金"], "questions": ["同一人员两个奖金来源选哪个？"]}
    agent._save_run(run)

    with pytest.raises(HTTPException) as raised:
        agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=True), user=scenario.user)

    assert raised.value.status_code == 409
    assert scenario.calls == []
    assert agent._load_run(run["run_id"], scenario.user)["plan_confirmation"]["confirmed"] is False


def test_plan_confirmation_without_saved_plan_skips_plan_when_model_unavailable(
    scenario: SimpleNamespace,
) -> None:
    # 上次模型故障可能把运行留在“需要确认计划但没有已保存计划”的状态；
    # 确认请求此时降级为重新生成，模型仍不可用时跳过计划直接放行处理，
    # 不再用 409 卡住唯一的续跑入口。
    run = _prepared_run(scenario)

    result = agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=True), user=scenario.user)

    assert result["plan_confirmation"]["confirmed"] is True
    assert result["model_plan"]["model"] == "plan-skipped"
    assert result["model_plan"]["questions"] == []
    assert scenario.calls == []
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_explicit_confirmation_accepts_complete_plan_with_unchanged_inputs(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    run = _prepared_run(scenario)
    _stub_plan_provider(monkeypatch, {"summary": "核对奖金", "steps": ["核对E001奖金"], "questions": []})
    agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=False), user=scenario.user)
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: None)

    result = agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=True), user=scenario.user)

    assert result["plan_confirmation"]["confirmed"] is True
    assert scenario.calls == []
    assert not result.get("draft_filename")
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_plan_confirmation_keeps_only_selected_steps_and_records_custom_instruction(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    run = _prepared_run(scenario)
    steps = ["核对E001奖金", "核对E002考勤"]
    _stub_plan_provider(monkeypatch, {"summary": "核对并更新副本", "steps": steps, "questions": []})
    agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=False), user=scenario.user)

    result = agent.confirm_agent_plan(
        run["run_id"],
        agent.AgentPlanIn(
            confirm_plan=True,
            selected_steps=[steps[0]],
            instruction="奖金仅处理有来源依据的人员",
        ),
        user=scenario.user,
    )

    assert result["model_plan"]["steps"] == [steps[0]]
    saved = agent._load_run(run["run_id"], scenario.user)
    assert "奖金仅处理有来源依据的人员" in saved["instruction"]
    assert saved["events"][-1]["payload"]["selected_steps"] == [steps[0]]
    assert scenario.calls == []


@pytest.mark.parametrize("action", ["confirm", "execute"])
@pytest.mark.parametrize("filename", ["202608（所属月202607）-北京科园-鹤安-大药房工资核算总表-v2.xlsx", "奖金来源.xlsx"])
def test_changed_input_blocks_plan_confirmation_or_execution(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch, action: str, filename: str,
) -> None:
    run = _prepared_run(scenario)
    requests = _stub_plan_provider(monkeypatch, {"summary": "核对奖金", "steps": ["核对E001奖金"], "questions": []})
    agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=False), user=scenario.user)
    if action == "execute":
        agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=True), user=scenario.user)
    requests.clear()
    changed = next(path for path in scenario.originals if path.name == filename)
    workbook = openpyxl.load_workbook(changed)
    workbook.active["C2"] = 3000
    workbook.save(changed)
    workbook.close()
    current_files = {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")}

    with pytest.raises(HTTPException) as raised:
        if action == "confirm":
            agent.confirm_agent_plan(run["run_id"], agent.AgentPlanIn(confirm_plan=True), user=scenario.user)
        else:
            agent.execute_agent_run(run["run_id"], user=scenario.user)

    assert raised.value.status_code == 409
    assert "文件" in str(raised.value.detail)
    assert scenario.calls == []
    assert requests == []
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == current_files


def _create_and_confirm_plan(scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> str:
    _stub_plan_provider(monkeypatch, {
        "summary": "按来源更新独立副本，原件不动", "steps": ["核对E001身份，将来源C2奖金复制到总表C2"], "questions": [],
    })
    created = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )
    run_id = created["run_id"]
    agent.confirm_agent_plan(run_id, agent.AgentPlanIn(confirm_plan=False), user=scenario.user)
    confirmed = agent.confirm_agent_plan(run_id, agent.AgentPlanIn(confirm_plan=True), user=scenario.user)
    assert confirmed["plan_confirmation"]["confirmed"] is True
    assert not confirmed.get("draft_filename")
    return run_id


def test_agent_main_flow_uses_model_tools_to_write_only_draft_and_cannot_self_approve(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    run_id = _create_and_confirm_plan(scenario, monkeypatch)
    snapshots: list[dict[str, Any]] = []

    class ExecutionProvider:
        def __init__(self, _config: Any):
            pass

        def complete(self, **kwargs: Any) -> ModelResponse:
            snapshots.append(deepcopy(kwargs))
            if len(snapshots) == 1:
                return ModelResponse(tool_calls=[
                    ToolCall(call_id="copy", name="prepare_workbook_copy"),
                    ToolCall(call_id="target", name="read_range", arguments={"sheet": "工资核算", "range": "A2:C2"}),
                    ToolCall(call_id="source", name="read_source_range", arguments={
                        "filename": "奖金来源.xlsx", "sheet": "工资核算", "range": "A2:C2",
                    }),
                ])
            if len(snapshots) == 2:
                return ModelResponse(tool_calls=[ToolCall(call_id="write", name="apply_source_cells", arguments={
                    "changes": [{
                        "source_file": "奖金来源.xlsx", "source_sheet": "工资核算", "source_cells": ["C2"],
                        "target_sheet": "工资核算", "target_cell": "C2", "expected_value": 0, "aggregation": "copy",
                    }],
                })])
            return ModelResponse(content="已将E001奖金2000写入草稿，仍需验收。")

    monkeypatch.setattr("backend.agent_execution.OpenAICompatibleProvider", ExecutionProvider)

    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert scenario.calls == []
    assert len(snapshots) == 3
    advertised = {schema["function"]["name"] for schema in snapshots[0]["tools"]}
    assert {
        "prepare_workbook_copy",
        "read_range",
        "read_source_range",
        "apply_source_cells",
        "apply_formula_divisors",
        "insert_and_copy_row",
    } <= advertised
    read_results = {
        message["tool_call_id"]: json.loads(message["content"])
        for message in snapshots[1]["messages"] if message["role"] == "tool"
    }
    assert read_results["target"]["values"] == [["E001", "测试人员", 0]]
    assert read_results["source"]["values"] == [["E001", "测试人员", 2000]]
    draft = agent._result_path("project-1", result["draft_filename"])
    assert draft not in scenario.originals
    workbook = openpyxl.load_workbook(draft, data_only=False)
    try:
        assert workbook["工资核算"]["C2"].value == 2000
    finally:
        workbook.close()
    assert {path: path.read_bytes() for path in scenario.originals} == scenario.originals
    assert result["workbook_updates"] == [{
        "target_sheet": "工资核算", "target_cell": "C2", "old_value": 0, "new_value": 2000,
        "source_file": "奖金来源.xlsx", "source_sheet": "工资核算", "source_cells": ["C2"], "aggregation": "copy",
    }]
    tool_results = [event["payload"] for event in result["events"] if event["type"] == "tool_result"]
    assert [(item["name"], item["status"]) for item in tool_results] == [
        ("prepare_workbook_copy", "succeeded"), ("read_range", "succeeded"),
        ("read_source_range", "succeeded"), ("apply_source_cells", "succeeded"),
    ]
    assert result["validation"]["status"] == "not_verified"
    saved = agent._load_run(run_id, scenario.user)
    assert saved["workbook_updates"] == result["workbook_updates"]
    with pytest.raises(HTTPException) as raised:
        agent.publish_agent_run(run_id, agent.AgentPublishIn(), user=scenario.user, db=scenario.db)
    assert raised.value.status_code == 409
    assert scenario.calls == []


def test_matching_keyuan_source_leaves_basic_processor_to_model_tool(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """固定脚本已工具化：科园来源不再触发基础处理器自动执行。"""
    scenario.db.files[1].original_name = "薪资数据-科园-7月薪资.xlsx"
    run_id = _create_and_confirm_plan(scenario, monkeypatch)
    processor_calls: list[tuple[Path, Path, Path]] = []

    def fake_processor(master: Path, source: Path, output: Path) -> dict[str, Any]:
        processor_calls.append((master, source, output))
        return {"status": "passed", "changes": [], "issues": []}

    monkeypatch.setattr("scripts.keyuan_basic_processor.process_keyuan_basics", fake_processor)

    class ProseProvider:
        def __init__(self, _config: Any):
            pass

        def complete(self, **_kwargs: Any) -> ModelResponse:
            return ModelResponse(content="已了解需求，本轮尚无写入。")

    monkeypatch.setattr("backend.agent_execution.OpenAICompatibleProvider", ProseProvider)
    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    # 模型没有调用 run_basic_payroll_processor 工具时，固定脚本一次都不能跑
    assert processor_calls == []
    assert not result.get("basic_processor")
    assert not result.get("draft_filename")
    assert not result.get("workbook_updates")
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_basic_source_detection_accepts_generic_monthly_filename(tmp_path: Path) -> None:
    from backend.agent_execution import find_basic_salary_source

    source = tmp_path / "更新表-2026-08.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "奖金-8月"
    workbook.create_sheet("考勤-8月")
    workbook.save(source)
    workbook.close()

    detected = find_basic_salary_source({"更新表-2026-08.xlsx": str(source)})

    assert detected == ("更新表-2026-08.xlsx", source)


def test_basic_source_detection_prefers_salary_workbook_over_tax_attachment(tmp_path: Path) -> None:
    from backend.agent_execution import find_basic_salary_source

    # Upload normalization keeps the original .xls name but stores an .xlsx path.
    tax = tmp_path / "tax-upload.xlsx"
    tax_book = openpyxl.Workbook()
    tax_book.active.title = "个税导出核对"
    tax_book.save(tax)
    tax_book.close()

    salary = tmp_path / "薪资数据-科园-7月薪资（8.14发薪）.xlsx"
    salary_book = openpyxl.Workbook()
    salary_book.active.title = "奖金-7月"
    salary_book.create_sheet("考勤-7月")
    salary_book.create_sheet("值班")
    salary_book.save(salary)
    salary_book.close()

    detected = find_basic_salary_source({"202608_税款计算_工资薪金所得-益药科园.xls": str(tax), salary.name: str(salary)})

    assert detected == (salary.name, salary)


def test_model_configuration_gate_blocks_before_any_processor_run(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """未配置模型时直接阻断，固定脚本不作为前置步骤运行。"""
    scenario.db.files[1].original_name = "薪资数据-科园-8月薪资.xlsx"
    run_id = _create_and_confirm_plan(scenario, monkeypatch)
    calls: list[Path] = []

    def fake_processor(master: Path, source: Path, output: Path) -> dict[str, Any]:
        calls.append(output)
        return {"status": "passed", "changes": [], "issues": []}

    monkeypatch.setattr("scripts.keyuan_basic_processor.process_keyuan_basics", fake_processor)
    monkeypatch.setattr("backend.agent_execution.ModelConfig.from_env", lambda: None)
    monkeypatch.setattr("backend.agent_execution.ModelConfig.fallback_from_env", lambda: None)

    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert calls == []
    assert result["status"] == "blocked"
    assert "尚未配置模型服务" in result["detail"]
    assert "基础 Python 更新已完成" not in result["detail"]
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_background_workflow_blocks_on_missing_model_before_planning(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """固定脚本前置流程已移除：无模型配置时先阻断，不进入计划或执行。"""
    run = _prepared_run(scenario)
    run["status"] = "planning"
    calls: list[str] = []

    monkeypatch.setattr(agent, "_load_run", lambda *_args, **_kwargs: run)
    monkeypatch.setattr(agent, "_save_run", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(agent, "_material_context", lambda *_args, **_kwargs: [])
    monkeypatch.setattr(agent, "_rule_package_context", lambda *_args, **_kwargs: {})
    monkeypatch.setattr(agent, "SessionLocal", lambda: SimpleNamespace(close=lambda: None))
    monkeypatch.setattr(agent, "_release_run_worker", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: None)
    monkeypatch.setattr(agent.ModelConfig, "fallback_from_env", lambda: None)
    monkeypatch.setattr(
        agent,
        "confirm_agent_plan",
        lambda *_args, **_kwargs: calls.append("plan"),
    )

    agent._run_agent_workflow(run["run_id"], "tenant-a")

    assert calls == []
    assert run["status"] == "blocked"
    assert run["code"] == "MODEL_CONFIGURATION_REQUIRED"
    assert "尚未配置模型服务" in run["detail"]


def test_scope_restriction_detection_is_generic() -> None:
    """Any client's wording of a scope restriction must be detected."""
    from backend.agent_execution import user_scope_restriction

    samples = [
        "只处理派遣表页里的数据",
        "仅更新奖金Sheet，其他的不要动",
        "司机、外包不处理",
        "只要养老数据，其余别碰",
        "把派遣的人员整理进去，其他sheet不要动",
    ]
    for text in samples:
        run = {"instruction": text, "conversation": []}
        assert user_scope_restriction(run) is not None, text

    unrestricted = [
        {"instruction": "开始处理", "conversation": []},
        {"instruction": "按手册核对全部工资数据", "conversation": []},
        {"instruction": "处理所有sheet", "conversation": [
            {"role": "user", "content": "继续"},
        ]},
    ]
    for run in unrestricted:
        assert user_scope_restriction(run) is None, run["instruction"]

    # 对齐聊天里说的范围限制同样生效
    run = {"instruction": "开始处理", "conversation": [
        {"role": "assistant", "content": "我的理解是全量处理"},
        {"role": "user", "content": "不行，只处理派遣，其他不要动"},
    ]}
    assert user_scope_restriction(run) is not None


def _dispatch_tool_and_finish(monkeypatch: pytest.MonkeyPatch, tool_name: str) -> None:
    """让模拟模型第一轮调用指定工具，第二轮输出总结。"""
    from core.document_agent.contracts import ToolCall
    from core.document_agent.model import ModelConfig, ModelResponse

    config = ModelConfig(
        provider="openai_compatible", base_url="https://model.example/v1",
        api_key="test-secret", model="test-executor",
    )
    monkeypatch.setattr("backend.agent_execution.ModelConfig.from_env", lambda: config)
    monkeypatch.setattr("backend.agent_execution.ModelConfig.fallback_from_env", lambda: None)

    class DispatchProvider:
        def __init__(self, _config: Any):
            self.calls = 0

        def complete(self, **_kwargs: Any) -> ModelResponse:
            self.calls += 1
            if self.calls == 1:
                return ModelResponse(
                    tool_calls=[ToolCall(call_id="call-1", name=tool_name, arguments={})],
                )
            return ModelResponse(content="本轮执行完毕。")

    monkeypatch.setattr("backend.agent_execution.OpenAICompatibleProvider", DispatchProvider)


def test_scope_restriction_refuses_basic_processor_dispatch(scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> None:
    """用户限制了范围时，即使模型误调度基础处理器，也必须拒绝全量写入。"""
    run_id = _create_and_confirm_plan(scenario, monkeypatch)
    run = agent._load_run(run_id, scenario.user)
    run["instruction"] = "只处理派遣Sheet的数据，其他不要动"
    agent._save_run(run)
    processor_calls: list[Path] = []

    def fake_processor(master: Path, source: Path, output: Path) -> dict[str, Any]:
        processor_calls.append(output)
        return {"status": "passed", "changes": [], "issues": []}

    monkeypatch.setattr("scripts.keyuan_basic_processor.process_keyuan_basics", fake_processor)
    _dispatch_tool_and_finish(monkeypatch, "run_basic_payroll_processor")

    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert processor_calls == []
    assert result["basic_processor"]["status"] == "skipped"
    assert "范围" in result["basic_processor"]["reason"]
    assert not result.get("workbook_updates")
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_template_mismatch_reports_honest_failure_not_silent_skip(scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> None:
    """总表模板不匹配时，脚本诚实回报 failed 及原因，而不是替模型做适用性判断。"""
    run_id = _create_and_confirm_plan(scenario, monkeypatch)
    run = agent._load_run(run_id, scenario.user)
    run["master_file"] = "1-攀时（上海）高性能材料有限公司-攀时-派遣 2026-06 结算单.xlsx"
    agent._save_run(run)
    # 模板不匹配：scenario 总表只有"工资核算"Sheet，缺少脚本要求的考勤等Sheet，
    # process_keyuan_basics 会在写入前抛 KeyError，工具转为诚实 failed 结果。
    monkeypatch.setattr(
        "backend.agent_execution.find_basic_salary_source",
        lambda _paths: ("奖金来源.xlsx", scenario.root / "uploads" / "project-1" / "奖金来源.xlsx"),
    )
    _dispatch_tool_and_finish(monkeypatch, "run_basic_payroll_processor")

    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert result["basic_processor"]["status"] == "failed"
    assert "模板" in result["basic_processor"]["reason"]
    assert not result.get("workbook_updates")
    # 脚本失败前可能已复制草稿副本，但原始总表和来源文件必须一个字节都不变
    for path, content in scenario.originals.items():
        assert path.read_bytes() == content


def test_public_run_recovers_legacy_month_only_completion(scenario: SimpleNamespace) -> None:
    run = {
        "run_id": "legacy-mixed", "tenant_id": "tenant-a", "project_id": "project-1",
        "status": "completed", "items": [], "workbook_updates": [{"sheet": "工资核算", "cell": "AE4"}],
        "basic_processor": {"change_count": 1, "issue_count": 1, "issues": [{"detail": "目标单元格是公式，基础处理器跳过写入"}]},
        "execution_result": {"status": "completed", "code": "DETERMINISTIC_FORMULA_PRESERVED"},
    }
    assert agent._recover_incomplete_formula_completion(run) is True
    assert run["status"] == "execution_incomplete"
    assert run["code"] == "INCOMPLETE_DATA_PROCESSING"
    assert run["plan_confirmation"]["confirmed"] is True
    assert run["month_confirmation"]["confirmed"] is True


def test_model_prose_without_write_tool_calls_never_creates_or_approves_a_draft(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    run_id = _create_and_confirm_plan(scenario, monkeypatch)

    class ProseOnlyProvider:
        def __init__(self, _config: Any):
            pass

        def complete(self, **_kwargs: Any) -> ModelResponse:
            return ModelResponse(content="已全部写入并通过所有验收，可以发布。")

    monkeypatch.setattr("backend.agent_execution.OpenAICompatibleProvider", ProseOnlyProvider)

    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert not result.get("draft_filename")
    assert not result.get("workbook_updates")
    assert result["status"] == "execution_incomplete"
    assert result["execution_result"]["code"] == "NO_WRITES_PERFORMED"
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals
    assert result["validation"]["status"] == "not_verified"
    assert scenario.calls == []
    with pytest.raises(HTTPException) as raised:
        agent.publish_agent_run(run_id, agent.AgentPublishIn(), user=scenario.user, db=scenario.db)
    assert raised.value.status_code == 409


def test_empty_model_execution_is_resumable_with_an_actionable_detail(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    run_id = _create_and_confirm_plan(scenario, monkeypatch)

    class EmptyProvider:
        def __init__(self, _config: Any):
            pass

        def complete(self, **_kwargs: Any) -> ModelResponse:
            return ModelResponse()

    monkeypatch.setattr("backend.agent_execution.OpenAICompatibleProvider", EmptyProvider)

    result = agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert result["status"] == "execution_incomplete"
    assert result["execution_result"]["code"] == "EMPTY_MODEL_RESPONSE"
    assert "自动重试" in result["detail"]


def test_model_execution_context_includes_project_skill_policy(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Project-local skills are available to the runtime without full manuals."""
    run_id = _create_and_confirm_plan(scenario, monkeypatch)
    model_contexts: list[dict[str, Any]] = []

    class ContextProvider:
        def __init__(self, _config: Any):
            pass

        def complete(self, **kwargs: Any) -> ModelResponse:
            model_contexts.append(json.loads(kwargs["messages"][1]["content"]))
            return ModelResponse()

    monkeypatch.setattr("backend.agent_execution.OpenAICompatibleProvider", ContextProvider)

    agent._execute_agent_run_sync(run_id, scenario.user, scenario.db)

    assert model_contexts
    skill_context = model_contexts[0]["project_agent_skills"]
    assert "spreadsheets" in skill_context["installed_skills"]
    assert "先概览再读必要范围" in skill_context["execution_policy"]
    routes = {route["id"]: route for route in skill_context["task_routing"]}
    assert routes["workbook-update"]["runtime_mode"] == "registered_tools_only"


def test_public_run_recovers_legacy_empty_execution_as_resumable() -> None:
    result = agent._public_run({
        "run_id": "run-1",
        "status": "awaiting_review",
        "detail": "本轮未完成执行，请查看运行记录后重试；未发布正式结果",
        "items": [],
        "workbook_updates": [],
        "execution_result": {"status": "completed", "code": None, "content": ""},
        "validation": {"status": "not_verified"},
    })

    assert result["status"] == "execution_incomplete"
    assert result["execution_result"]["code"] == "EMPTY_MODEL_RESPONSE"
    assert "继续执行" in result["detail"]


def test_public_run_exposes_actionable_provider_error_detail() -> None:
    result = agent._public_run({
        "run_id": "run-1",
        "status": "execution_incomplete",
        "items": [],
        "workbook_updates": [],
        "execution_result": {
            "status": "execution_incomplete",
            "code": "MODEL_PROVIDER_ERROR",
            "content": "模型 API Key 或服务权限认证失败，请检查模型地址和密钥",
        },
    })

    assert "API Key" in result["detail"]
    assert "请继续执行" in result["detail"]


def _stub_ready_material(scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> None:
    """让任务带一份已读取全文的说明材料（模拟 docx 上传登记后的状态）。"""
    monkeypatch.setattr(agent, "_load_material_index", lambda *_args: [{
        "material_id": "m1", "project_id": "project-1", "filename": "1.docx",
        "kind": "manual", "status": "ready",
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "text": "只处理变更表派遣表页里的数据", "size_bytes": 10,
        "created_at": "2026-09-03T00:00:00+00:00",
    }])
    scenario.material_ready = True


def test_material_opener_pauses_first_start_for_alignment(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """读了说明材料后第一次“开始处理”不执行，先复述理解对齐颗粒度。"""
    _stub_ready_material(scenario, monkeypatch)
    opener_calls: list[dict[str, Any]] = []

    def opener_complete(_self: Any, **kwargs: Any) -> ModelResponse:
        opener_calls.append(kwargs)
        return ModelResponse(content="①我的理解：只处理变更表派遣表页 ②处理方式：按名单对齐 ③待确认：删除范围外人员？")

    config = ModelConfig(
        provider="openai_compatible", base_url="https://model.example/v1",
        api_key="test-secret", model="test-planner",
    )
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: config)
    monkeypatch.setattr(agent.OpenAICompatibleProvider, "complete", opener_complete)

    created = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )
    result = agent.process_agent_run(created["run_id"], None, user=scenario.user)

    # 第一次启动被对齐门槛拦截：不执行、不建草稿，只输出理解开场。
    # The alignment opener has finished its work and is waiting for the user.
    # Leaving it in ``planning`` makes the UI report an active analysis and
    # leaves SSE followers waiting forever even though no worker exists.
    assert result["status"] == "awaiting_review"
    assert result.get("alignment_opened") is True
    agent_messages = [m for m in result["conversation"] if m.get("role") == "agent"]
    assert len(agent_messages) == 1
    assert "我的理解" in agent_messages[0]["content"]
    assert opener_calls, "开场理解必须真实调用模型"
    assert scenario.calls == [], "对齐阶段不得触发集成执行"
    assert not result.get("draft_filename")
    assert {path: path.read_bytes() for path in scenario.root.rglob("*.xlsx")} == scenario.originals


def test_second_start_after_opener_proceeds_to_execution(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """开场对齐完成后，再次“开始处理”正常进入执行。"""
    _stub_ready_material(scenario, monkeypatch)
    config = ModelConfig(
        provider="openai_compatible", base_url="https://model.example/v1",
        api_key="test-secret", model="test-planner",
    )
    monkeypatch.setattr(agent.ModelConfig, "from_env", lambda: config)
    monkeypatch.setattr(agent.OpenAICompatibleProvider, "complete",
                        lambda _self, **_kwargs: ModelResponse(content="①我的理解：仅派遣 ②处理方式：对齐名单 ③待确认：无"))

    created = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )
    run_id = created["run_id"]
    first = agent.process_agent_run(run_id, None, user=scenario.user)
    assert first["alignment_opened"] is True

    workflow_calls: list[str] = []
    monkeypatch.setattr(agent, "_run_agent_workflow", lambda *args: workflow_calls.append(args[0]))
    second = agent.process_agent_run(run_id, None, user=scenario.user)

    assert second["status"] == "processing"
    assert workflow_calls == [run_id]


def test_explicit_start_bypasses_alignment_opener(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """按钮启动是完整执行意图，不应先消耗一次点击做对齐开场。"""
    _stub_ready_material(scenario, monkeypatch)
    created = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )
    run_id = created["run_id"]
    workflow_calls: list[str] = []
    monkeypatch.setattr(agent, "_alignment_opener_needed", lambda _run: True)
    monkeypatch.setattr(agent, "_open_alignment_conversation", lambda _run: (_ for _ in ()).throw(AssertionError("不应调用对齐开场")))
    monkeypatch.setattr(agent, "_run_agent_workflow", lambda *args: workflow_calls.append(args[0]))

    result = agent.process_agent_run(
        run_id, agent.AgentProcessIn(start_processing=True), user=scenario.user,
    )

    assert result["status"] == "processing"
    assert workflow_calls == [run_id]


def test_explicit_start_auto_confirms_unselected_plan_questions() -> None:
    run: dict[str, Any] = {"run_id": "auto-start", "instruction": "按规则处理", "plan_confirmation": {"confirmed": False}}
    plan = {"steps": ["核对奖金"], "questions": ["奖金来源有两份，按哪份？"]}

    assert agent._auto_confirm_plan_for_start(run, plan) is True
    assert plan["questions"] == []
    assert run["plan_confirmation"]["confirmed"] is True
    assert "未单独选择的计划事项按 Agent 建议处理" in run["instruction"]
    assert run["events"][0]["payload"]["stage"] == "plan_auto_confirmed"


def test_prior_agent_chat_skips_opener_gate(
    scenario: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """用户已在对话中与 Agent 对齐过（存在 Agent 回复）时不再强制开场。"""
    _stub_ready_material(scenario, monkeypatch)
    created = agent.create_agent_run(
        agent.AgentRunCreateIn(project_id="project-1"), user=scenario.user, db=scenario.db,
    )
    run_id = created["run_id"]
    run = agent._load_run(run_id, scenario.user)
    run.setdefault("conversation", []).append({
        "role": "agent", "content": "①我的理解：只处理派遣 ②处理方式：按名单 ③待确认：无",
        "at": "2026-09-03T00:00:00+00:00",
    })
    agent._save_run(run)

    workflow_calls: list[str] = []
    monkeypatch.setattr(agent, "_run_agent_workflow", lambda *args: workflow_calls.append(args[0]))
    result = agent.process_agent_run(run_id, None, user=scenario.user)

    assert result["status"] == "processing"
    assert not result.get("alignment_opened")
    assert workflow_calls == [run_id]
