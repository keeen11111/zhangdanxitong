from __future__ import annotations

from pathlib import Path

import openpyxl

from backend.natural_language_rules import apply_supported_workbook_rules, parse_supported_workbook_rule


INSTRUCTION = (
    "\u8868\u9875\u2014\u914d\u9001\u5458\u503c\u73ed\u8d39\uff0c\u586bD\u5217\u503c\u73ed\u5929\u6570\uff0c\u628aD\u5217\u7684\u539f\u6709\u7684\u5929\u6570\u6e05\u9664\uff0c"
    "\u6839\u636e\u201c1-\u85aa\u8d44\u6570\u636e-\u79d1\u56ed-7\u6708\u85aa\u8d44\uff088.14\u53d1\u85aa\uff09\u201d\u4e2d\u201c\u503c\u73ed\u201d\u8868\u9875\u503c\u73ed\u4eba\u5458\u540d\u5b57\u51fa\u73b0\u7684\u6b21\u6570\u6765\u586b\u3002"
)


def _workbooks(tmp_path: Path) -> tuple[Path, Path]:
    draft_path = tmp_path / "草稿.xlsx"
    draft = openpyxl.Workbook()
    target = draft.active
    target.title = "\u914d\u9001\u5458\u503c\u73ed\u8d39"
    target.append(["\u5e8f\u53f7", "\u503c\u73ed\u4eba\u5458\u59d3\u540d", "\u503c\u73ed\u8d39\u6807\u51c6", "\u503c\u73ed\u5929\u6570", "\u503c\u73ed\u8d39\u5408\u8ba1"])
    target.append([1, "\u9a6c\u662d", 300, 99, "=C2*D2"])
    target.append([2, "\u9ad8\u9633\u9633", 300, 88, "=C3*D3"])
    target.append([3, "\u672a\u6392\u73ed\u4eba\u5458", 300, 77, "=C4*D4"])
    draft.save(draft_path)
    draft.close()

    source_path = tmp_path / "\u85aa\u8d44\u6570\u636e-\u79d1\u56ed-7\u6708\u85aa\u8d44\uff088.14\u53d1\u85aa\uff09.xlsx"
    source = openpyxl.Workbook()
    roster = source.active
    roster.title = "\u503c\u73ed"
    roster.append(["\u5e8f\u53f7", "\u65e5\u671f", "\u503c\u73ed\u4eba\u5458"])
    roster.append([1, "7.1", "\u9ad8\u9633\u9633"])
    roster.append([2, "7.2", "\u9a6c\u662d"])
    roster.append([3, "7.3", "\u9ad8\u9633\u9633"])
    source.save(source_path)
    source.close()
    return draft_path, source_path


def test_parses_explicit_duty_count_instruction() -> None:
    assert parse_supported_workbook_rule(INSTRUCTION) == {
        "kind": "duty_roster_name_count",
        "target_sheet": "配送员值班费",
        "target_column": "D",
        "source_sheet": "值班",
        "source_hint": "1-薪资数据-科园-7月薪资（8.14发薪）",
    }


def test_parses_explicit_source_sheet_roster_scope() -> None:
    instruction = (
        "只处理变更文件里的派遣sheet，除此之外的人员数据不保留，不显示在生成的总表里。"
        "目标月份明确为2026.07，最终名单以派遣sheet为准；司机和外包sheet完全不处理。"
    )

    assert parse_supported_workbook_rule(instruction) == {
        "kind": "source_sheet_roster_sync",
        "target_sheet": "明细",
        "source_sheet": "派遣",
        "source_hint": "",
        "target_period": "2026-07",
    }


def test_parses_roster_scope_after_the_real_change_workbook_phrase() -> None:
    instruction = (
        "只处理变更表格里面的派遣sheet的人员数据，更新在总表里，但是其他sheet的人员数据完全不保留，"
        "全部去掉，不显示。只要派遣sheet的人员"
    )

    rule = parse_supported_workbook_rule(instruction)

    assert rule == {
        "kind": "source_sheet_roster_sync",
        "target_sheet": "明细",
        "source_sheet": "派遣",
        "source_hint": "",
    }


def test_replaces_duty_days_clears_unmatched_people_and_preserves_formulas(tmp_path: Path) -> None:
    draft_path, source_path = _workbooks(tmp_path)

    outcome = apply_supported_workbook_rules(
        instruction=INSTRUCTION,
        draft_path=draft_path,
        source_paths={source_path.name: source_path},
    )

    workbook = openpyxl.load_workbook(draft_path, data_only=False)
    sheet = workbook["\u914d\u9001\u5458\u503c\u73ed\u8d39"]
    assert outcome["results"][0]["status"] == "applied"
    assert sheet["D2"].value == 1
    assert sheet["D3"].value == 2
    assert sheet["D4"].value is None
    assert sheet["E2"].value == "=C2*D2"
    workbook.close()


def test_does_not_clear_any_value_when_source_cannot_be_uniquely_matched(tmp_path: Path) -> None:
    draft_path, source_path = _workbooks(tmp_path)
    second_source = tmp_path / "\u53e6\u4e00\u4e2a\u503c\u73ed\u6765\u6e90.xlsx"
    second_source.write_bytes(source_path.read_bytes())

    outcome = apply_supported_workbook_rules(
        instruction=INSTRUCTION.replace("\u201c1-\u85aa\u8d44\u6570\u636e-\u79d1\u56ed-7\u6708\u85aa\u8d44\uff088.14\u53d1\u85aa\uff09\u201d", ""),
        draft_path=draft_path,
        source_paths={source_path.name: source_path, second_source.name: second_source},
    )

    workbook = openpyxl.load_workbook(draft_path, data_only=False)
    assert outcome["results"][0]["status"] == "needs_review"
    assert workbook["\u914d\u9001\u5458\u503c\u73ed\u8d39"]["D2"].value == 99
    workbook.close()


def test_refuses_to_overwrite_formula_destination(tmp_path: Path) -> None:
    draft_path, source_path = _workbooks(tmp_path)
    workbook = openpyxl.load_workbook(draft_path)
    workbook["\u914d\u9001\u5458\u503c\u73ed\u8d39"]["D2"] = "=1+1"
    workbook.save(draft_path)
    workbook.close()

    outcome = apply_supported_workbook_rules(
        instruction=INSTRUCTION,
        draft_path=draft_path,
        source_paths={source_path.name: source_path},
    )

    workbook = openpyxl.load_workbook(draft_path, data_only=False)
    assert outcome["results"][0]["status"] == "needs_review"
    assert workbook["\u914d\u9001\u5458\u503c\u73ed\u8d39"]["D2"].value == "=1+1"
    workbook.close()
