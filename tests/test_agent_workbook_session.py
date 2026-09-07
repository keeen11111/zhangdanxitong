from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from typing import Any
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
import pytest

from core.document_agent.orchestrator import ToolExecutionError
from core.document_agent.workbook_session import WorkbookSession


@pytest.fixture
def workbooks(tmp_path: Path) -> SimpleNamespace:
    master = tmp_path / "原始总表.xlsx"
    source = tmp_path / "奖金来源.xlsx"
    draft = tmp_path / "独立草稿.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "工资核算"
    sheet["J2"] = 0
    sheet["K2"] = 100
    sheet["J3"] = "=SUM(J2:K2)"
    sheet["J2"].number_format = "#,##0.00"
    sheet["J2"].font = Font(bold=True, color="FF0000")
    workbook.create_sheet("历史记录")["A1"] = "保留历史数据"
    workbook.save(master)
    workbook.close()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "奖金明细"
    sheet["B2"] = 2000
    sheet["B3"] = 500
    sheet["B4"] = "=B2+B3"
    sheet["B5"] = "待核对"
    workbook.save(source)
    workbook.close()
    return SimpleNamespace(
        master=master, source=source, draft=draft,
        originals={master: master.read_bytes(), source: source.read_bytes()},
    )


def _session(workbooks: SimpleNamespace) -> WorkbookSession:
    return WorkbookSession(
        master=workbooks.master, draft=workbooks.draft,
        sources={workbooks.source.name: workbooks.source},
    )


def _change(**overrides: Any) -> dict[str, Any]:
    return {
        "source_file": "奖金来源.xlsx", "source_sheet": "奖金明细", "source_cells": ["B2"],
        "target_sheet": "工资核算", "target_cell": "J2", "expected_value": 0,
        "aggregation": "copy", **overrides,
    }


def _set_formula_cache(path: Path, coordinate: str, value: int | float) -> None:
    """Inject one numeric cached result into the first worksheet fixture."""
    with ZipFile(path, "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    worksheet_name = "xl/worksheets/sheet1.xml"
    root = ElementTree.fromstring(members[worksheet_name])
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cell = next(node for node in root.iter(f"{namespace}c") if node.get("r") == coordinate)
    cached = cell.find(f"{namespace}v")
    if cached is None:
        cached = ElementTree.SubElement(cell, f"{namespace}v")
    cached.text = str(value)
    members[worksheet_name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
    temporary = path.with_name(f"{path.stem}.cached.xlsx")
    with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    temporary.replace(path)


def test_prepare_copy_creates_independent_draft_without_touching_inputs(workbooks: SimpleNamespace) -> None:
    session = _session(workbooks)

    session.prepare_copy()

    assert workbooks.draft.is_file()
    assert workbooks.draft.read_bytes() == workbooks.originals[workbooks.master]
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


@pytest.mark.parametrize("input_name", ["master", "source"])
def test_prepare_copy_rejects_draft_that_is_any_input(workbooks: SimpleNamespace, input_name: str) -> None:
    with pytest.raises(ToolExecutionError):
        session = WorkbookSession(
            master=workbooks.master, draft=getattr(workbooks, input_name),
            sources={workbooks.source.name: workbooks.source},
        )
        session.prepare_copy()

    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_apply_source_cells_copies_and_sums_with_auditable_updates(workbooks: SimpleNamespace) -> None:
    session = _session(workbooks)
    session.prepare_copy()

    updates = session.apply_source_cells([
        _change(),
        _change(source_cells=["B2", "B3"], target_cell="K2", expected_value=100, aggregation="sum"),
    ])

    assert isinstance(updates, list)
    assert [{key: update[key] for key in (
        "old_value", "new_value", "target_sheet", "target_cell", "source_file", "source_sheet",
    )} for update in updates] == [
        {"old_value": 0, "new_value": 2000, "target_sheet": "工资核算", "target_cell": "J2",
         "source_file": "奖金来源.xlsx", "source_sheet": "奖金明细"},
        {"old_value": 100, "new_value": 2500, "target_sheet": "工资核算", "target_cell": "K2",
         "source_file": "奖金来源.xlsx", "source_sheet": "奖金明细"},
    ]
    workbook = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = workbook["工资核算"]
        assert sheet["J2"].value == 2000
        assert sheet["K2"].value == 2500
        assert sheet["J3"].value == "=SUM(J2:K2)"
        assert sheet["J2"].number_format == "#,##0.00"
        assert sheet["J2"].font.bold is True
        assert workbook["历史记录"]["A1"].value == "保留历史数据"
    finally:
        workbook.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


@pytest.mark.parametrize("change", [
    pytest.param(_change(target_cell="J3", expected_value="=SUM(J2:K2)"), id="target-formula"),
    pytest.param(_change(source_cells=["B4"]), id="source-formula-without-cache"),
    pytest.param(_change(source_file="未授权来源.xlsx"), id="unregistered-source"),
    pytest.param(_change(expected_value=999), id="stale-target-value"),
    pytest.param(_change(source_sheet="不存在的来源页"), id="unknown-source-sheet"),
    pytest.param(_change(target_sheet="不存在的目标页"), id="unknown-target-sheet"),
    pytest.param(_change(source_cells=["B0"]), id="invalid-source-cell"),
    pytest.param(_change(target_cell="J0"), id="invalid-target-cell"),
    pytest.param(_change(aggregation="execute"), id="unsupported-aggregation"),
    pytest.param(_change(source_cells=["B2", "B5"], aggregation="sum"), id="nonnumeric-sum"),
])
def test_invalid_change_is_sanitized_and_leaves_draft_and_inputs_unchanged(
    workbooks: SimpleNamespace, change: dict[str, Any],
) -> None:
    session = _session(workbooks)
    session.prepare_copy()
    draft_before = workbooks.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        session.apply_source_cells([change])

    assert workbooks.draft.read_bytes() == draft_before
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


@pytest.mark.parametrize("invalid_change", [
    _change(source_cells=["B4"], target_cell="K2", expected_value=100),
    _change(target_cell="K2", expected_value=999),
])
def test_batch_failure_never_persists_an_earlier_valid_change(
    workbooks: SimpleNamespace, invalid_change: dict[str, Any],
) -> None:
    session = _session(workbooks)
    session.prepare_copy()
    draft_before = workbooks.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        session.apply_source_cells([_change(), invalid_change])

    assert workbooks.draft.read_bytes() == draft_before
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_insert_and_copy_row_preserves_style_and_translates_formulas(workbooks: SimpleNamespace) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet["A2"] = "2026年6月"
        sheet["B2"] = 100
        sheet["C2"] = "=B2*2"
        sheet.row_dimensions[2].height = 24
        sheet["B2"].font = Font(bold=True, color="00FF00")
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    update = session.insert_and_copy_row({
        "sheet": "工资核算",
        "source_row": 2,
        "insert_before_row": 3,
        "expected_source_cells": {"A2": "2026年6月", "C2": "=B2*2"},
        "copy_max_column": 3,
    })

    assert update == {
        "sheet": "工资核算", "source_row": 2, "inserted_row": 3,
        "copied_columns": 3,
    }
    workbook = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = workbook["工资核算"]
        assert [sheet.cell(3, column).value for column in range(1, 4)] == ["2026年6月", 100, "=B3*2"]
        assert sheet.row_dimensions[3].height == 24
        assert sheet["B3"].font.bold is True
        assert sheet["B3"].font.color.rgb == "0000FF00"
    finally:
        workbook.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_insert_and_copy_row_updates_moved_and_summary_formula_references(
    workbooks: SimpleNamespace,
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet["A2"] = "模板人员"
        sheet["B2"] = 10
        sheet["C2"] = "=B2*2"
        sheet["A3"] = "原有人员"
        sheet["B3"] = 20
        sheet["C3"] = "=B3*2"
        sheet["D1"] = "=SUM(B1:B2)"
        summary = workbook.create_sheet("汇总")
        summary["A1"] = "=SUM('工资核算'!$B$2:$B$3)"
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    session.insert_and_copy_row({
        "sheet": "工资核算",
        "source_row": 2,
        "insert_before_row": 3,
        "expected_source_cells": {"A2": "模板人员", "C2": "=B2*2"},
        "copy_max_column": 3,
    })

    workbook = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = workbook["工资核算"]
        assert sheet["C3"].value == "=B3*2"
        assert sheet["C4"].value == "=B4*2"
        assert sheet["D1"].value == "=SUM(B1:B2)"
        assert workbook["汇总"]["A1"].value == "=SUM('工资核算'!$B$2:$B$4)"
    finally:
        workbook.close()


def test_insert_and_copy_row_repairs_template_objects(
    workbooks: SimpleNamespace,
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet.append(["月份", "金额", "倍数"])
        sheet.append(["6月", 100, "=B2*2"])
        sheet.append(["8月", 300, "=B3*2"])
        sheet.append(["合计", "=SUM(B2:B3)", "=SUM(C2:C3)"])
        sheet.add_table(Table(displayName="MonthlyTable", ref="A1:C4"))
        sheet.auto_filter.ref = "A1:C4"
        validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
        validation.add("B2:B3")
        sheet.add_data_validation(validation)
        sheet.conditional_formatting.add(
            "B2:B3", CellIsRule(operator="greaterThan", formula=["0"]),
        )
        workbook.defined_names.add(DefinedName(
            "MonthlyRows", attr_text="'工资核算'!$A$2:$C$3",
        ))
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
        sheet.add_chart(chart, "E2")
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    session.insert_and_copy_row({
        "sheet": "工资核算",
        "source_row": 2,
        "insert_before_row": 3,
        "expected_source_cells": {"A2": "6月", "C2": "=B2*2"},
        "copy_max_column": 3,
    })

    updated = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = updated["工资核算"]
        assert sheet.tables["MonthlyTable"].ref == "A1:C5"
        assert sheet.auto_filter.ref == "A1:C5"
        assert str(sheet.data_validations.dataValidation[0].sqref) == "B2:B4"
        assert [str(item.sqref) for item in sheet.conditional_formatting] == ["B2:B4"]
        assert updated.defined_names["MonthlyRows"].attr_text == "'工资核算'!$A$2:$C$4"
        series = sheet._charts[0].series[0]
        assert series.val.numRef.f == "'工资核算'!$B$2:$B$4"
        assert series.cat.numRef.f == "'工资核算'!$A$2:$A$4"
    finally:
        updated.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_formula_row_repair_preserves_digit_suffixed_function_names() -> None:
    inserted = WorkbookSession._shift_formula_references_for_row_insert(
        '=LOG10(B3)+"LOG10(B3)"',
        current_sheet="工资核算",
        target_sheet="工资核算",
        insert_before_row=3,
    )
    deleted = WorkbookSession._shift_formula_references_for_row_delete(
        "=LOG10(B3)",
        current_sheet="工资核算",
        target_sheet="工资核算",
        deleted_rows=[2],
    )

    assert inserted == '=LOG10(B4)+"LOG10(B3)"'
    assert deleted == "=LOG10(B2)"


def test_formula_row_repair_handles_a_direct_quoted_cross_sheet_reference() -> None:
    inserted = WorkbookSession._shift_formula_references_for_row_insert(
        "='明细'!AO6",
        current_sheet="付款通知书",
        target_sheet="明细",
        insert_before_row=4,
    )
    deleted = WorkbookSession._shift_formula_references_for_row_delete(
        "='明细'!AO6",
        current_sheet="付款通知书",
        target_sheet="明细",
        deleted_rows=[4],
    )

    assert inserted == "='明细'!AO7"
    assert deleted == "='明细'!AO5"


def test_personnel_row_insert_writes_identity_and_clears_template_business_values(
    workbooks: SimpleNamespace,
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet.append(["工号", "姓名", "基本工资", "应发工资"])
        sheet.append(["E001", "模板人员", 5000, "=C2"])
        sheet.append(["*合计*", None, "=SUM(C2:C2)", "=SUM(D2:D2)"])
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    update = session.insert_and_copy_row({
        "sheet": "工资核算", "source_row": 2, "insert_before_row": 3,
        "expected_source_cells": {"A2": "E001", "B2": "模板人员"},
        "copy_max_column": 4,
        "identity_values": {"A3": "E002", "B3": "李楠"},
    })

    assert update["identity_validated"] is True
    assert update["identity_cells"] == {"A3": "E002", "B3": "李楠"}
    draft = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = draft["工资核算"]
        assert [sheet["A3"].value, sheet["B3"].value] == ["E002", "李楠"]
        assert sheet["C3"].value is None
        assert sheet["D3"].value == "=C3"
        assert sheet["C4"].value == "=SUM(C2:C3)"
    finally:
        draft.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_personnel_row_insert_without_complete_identity_is_atomic(
    workbooks: SimpleNamespace,
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet.append(["工号", "姓名", "基本工资"])
        sheet.append(["E001", "模板人员", 5000])
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()
    before = workbooks.draft.read_bytes()

    with pytest.raises(ToolExecutionError, match="完整身份字段"):
        session.insert_and_copy_row({
            "sheet": "工资核算", "source_row": 2, "insert_before_row": 3,
            "expected_source_cells": {"A2": "E001", "B2": "模板人员"},
            "copy_max_column": 3,
            "identity_values": {"B3": "李楠"},
        })

    assert workbooks.draft.read_bytes() == before
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_delete_rows_repairs_moved_and_cross_sheet_summary_formulas(
    workbooks: SimpleNamespace,
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet.append(["姓名", "基本工资", "应发工资"])
        sheet.append(["甲", 100, "=B2*2"])
        sheet.append(["乙", 200, "=B3*2"])
        sheet.append(["丙", 300, "=B4*2"])
        sheet.append(["*合计*", "=SUM(B2:B4)", "=SUM(C2:C4)"])
        summary = workbook.create_sheet("汇总")
        summary["A1"] = "=SUM('工资核算'!B2:B4)"
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    session.delete_rows({
        "sheet": "工资核算", "rows": [3], "expected_cells": {"A3": "乙"},
    })

    draft = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = draft["工资核算"]
        assert sheet["A3"].value == "丙"
        assert sheet["C3"].value == "=B3*2"
        assert sheet["B4"].value == "=SUM(B2:B3)"
        assert sheet["C4"].value == "=SUM(C2:C3)"
        assert draft["汇总"]["A1"].value == "=SUM('工资核算'!B2:B3)"
    finally:
        draft.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_delete_rows_repairs_template_objects(
    workbooks: SimpleNamespace,
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook.remove(workbook["工资核算"])
        sheet = workbook.create_sheet("工资核算", 0)
        sheet.append(["姓名", "金额"])
        sheet.append(["甲", 100])
        sheet.append(["乙", 200])
        sheet.append(["丙", 300])
        sheet.append(["合计", "=SUM(B2:B4)"])
        sheet.add_table(Table(displayName="PayrollTable", ref="A1:B5"))
        sheet.auto_filter.ref = "A1:B5"
        validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
        validation.add("B2:B4")
        sheet.add_data_validation(validation)
        sheet.conditional_formatting.add(
            "B2:B4", CellIsRule(operator="greaterThan", formula=["0"]),
        )
        workbook.defined_names.add(DefinedName(
            "PayrollRows", attr_text="'工资核算'!$A$2:$B$4",
        ))
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
        sheet.add_chart(chart, "D2")
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    session.delete_rows({
        "sheet": "工资核算", "rows": [3], "expected_cells": {"A3": "乙"},
    })

    updated = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = updated["工资核算"]
        assert sheet.tables["PayrollTable"].ref == "A1:B4"
        assert sheet.auto_filter.ref == "A1:B4"
        assert str(sheet.data_validations.dataValidation[0].sqref) == "B2:B3"
        assert [str(item.sqref) for item in sheet.conditional_formatting] == ["B2:B3"]
        assert updated.defined_names["PayrollRows"].attr_text == "'工资核算'!$A$2:$B$3"
        series = sheet._charts[0].series[0]
        assert series.val.numRef.f == "'工资核算'!$B$2:$B$3"
        assert series.cat.numRef.f == "'工资核算'!$A$2:$A$3"
    finally:
        updated.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


@pytest.mark.parametrize("change", [
    {"sheet": "工资核算", "source_row": 2, "insert_before_row": 2,
     "expected_source_cells": {"A2": "2026年6月"}, "copy_max_column": 3},
    {"sheet": "工资核算", "source_row": 2, "insert_before_row": 3,
     "expected_source_cells": {"A2": "stale"}, "copy_max_column": 3},
    {"sheet": "不存在", "source_row": 2, "insert_before_row": 3,
     "expected_source_cells": {"A2": "2026年6月"}, "copy_max_column": 3},
])
def test_insert_and_copy_row_rejects_invalid_or_stale_requests_atomically(
    workbooks: SimpleNamespace, change: dict[str, Any],
) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        workbook["工资核算"]["A2"] = "2026年6月"
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()
    draft_before = workbooks.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        session.insert_and_copy_row(change)

    assert workbooks.draft.read_bytes() == draft_before
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_apply_verified_values_updates_non_formula_cells_without_changing_styles(workbooks: SimpleNamespace) -> None:
    session = _session(workbooks)
    session.prepare_copy()

    updates = session.apply_verified_values([
        {"target_sheet": "工资核算", "target_cell": "J2", "expected_value": 0, "new_value": None},
        {"target_sheet": "工资核算", "target_cell": "K2", "expected_value": 100, "new_value": 23},
    ])

    assert updates == [
        {"target_sheet": "工资核算", "target_cell": "J2", "old_value": 0, "new_value": None},
        {"target_sheet": "工资核算", "target_cell": "K2", "old_value": 100, "new_value": 23},
    ]
    workbook = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = workbook["工资核算"]
        assert sheet["J2"].value is None
        assert sheet["J2"].number_format == "#,##0.00"
        assert sheet["J2"].font.bold is True
        assert sheet["K2"].value == 23
        assert sheet["J3"].value == "=SUM(J2:K2)"
    finally:
        workbook.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


@pytest.mark.parametrize("changes", [
    [{"target_sheet": "工资核算", "target_cell": "J3", "expected_value": "=SUM(J2:K2)", "new_value": 1}],
    [{"target_sheet": "工资核算", "target_cell": "J2", "expected_value": 999, "new_value": 1}],
    [{"target_sheet": "工资核算", "target_cell": "J2", "expected_value": 0, "new_value": "=1+1"}],
    [
        {"target_sheet": "工资核算", "target_cell": "J2", "expected_value": 0, "new_value": 1},
        {"target_sheet": "工资核算", "target_cell": "J2", "expected_value": 0, "new_value": 2},
    ],
])
def test_apply_verified_values_rejects_unsafe_requests_atomically(
    workbooks: SimpleNamespace, changes: list[dict[str, Any]],
) -> None:
    session = _session(workbooks)
    session.prepare_copy()
    draft_before = workbooks.draft.read_bytes()

    with pytest.raises(ToolExecutionError):
        session.apply_verified_values(changes)

    assert workbooks.draft.read_bytes() == draft_before
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_roll_forward_summary_row_keeps_current_formulas_and_freezes_history(workbooks: SimpleNamespace) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        sheet = workbook["工资核算"]
        sheet["A2"] = "6月"
        sheet["B2"] = 40
        sheet["C2"] = "=B2*2"
        sheet["E4"] = "=C2"
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    _set_formula_cache(workbooks.master, "C2", 80)
    workbooks.originals[workbooks.master] = workbooks.master.read_bytes()
    session = _session(workbooks)
    session.prepare_copy()

    update = session.roll_forward_summary_row({
        "sheet": "工资核算", "current_row": 2, "expected_current_month": "6月", "new_month": "7月",
        "start_column": 1, "end_column": 3,
    })

    assert update["sheet"] == "工资核算"
    assert update["current_row"] == 2
    assert update["history_row"] == 3
    assert update["new_month"] == "7月"
    assert update["history_values"] == {"A3": "6月", "B3": 40, "C3": 80}
    assert update["history_snapshot"] == {
        "A3": "6月", "B3": 40, "C3": 80,
        "A4": None, "B4": None, "C4": None,
        "A5": None, "B5": None, "C5": None,
    }
    workbook = openpyxl.load_workbook(workbooks.draft, data_only=False)
    try:
        sheet = workbook["工资核算"]
        assert [sheet.cell(2, column).value for column in range(1, 4)] == ["7月", 40, "=B2*2"]
        assert [sheet.cell(3, column).value for column in range(1, 4)] == ["6月", 40, 80]
        assert sheet["E5"].value == "=C3"
    finally:
        workbook.close()
    assert {path: path.read_bytes() for path in workbooks.originals} == workbooks.originals


def test_roll_forward_summary_row_rejects_missing_history_values_atomically(workbooks: SimpleNamespace) -> None:
    workbook = openpyxl.load_workbook(workbooks.master)
    try:
        sheet = workbook["工资核算"]
        sheet["A2"] = "6月"
        sheet["C2"] = "=B2*2"
        workbook.save(workbooks.master)
    finally:
        workbook.close()
    session = _session(workbooks)
    session.prepare_copy()
    draft_before = workbooks.draft.read_bytes()

    with pytest.raises(ToolExecutionError, match="缓存"):
        session.roll_forward_summary_row({
            "sheet": "工资核算", "current_row": 2, "expected_current_month": "6月", "new_month": "7月",
            "start_column": 1, "end_column": 3,
        })

    assert workbooks.draft.read_bytes() == draft_before


def test_roll_forward_closes_formula_view_when_cached_view_cannot_open(
    workbooks: SimpleNamespace, monkeypatch: pytest.MonkeyPatch,
) -> None:
    session = _session(workbooks)
    session.prepare_copy()
    real_load_workbook = openpyxl.load_workbook
    formula_workbook = real_load_workbook(workbooks.draft, data_only=False)
    closed: list[bool] = []
    real_close = formula_workbook.close

    def tracked_close() -> None:
        closed.append(True)
        real_close()

    formula_workbook.close = tracked_close

    def fail_cached_view(_path: Path, *, data_only: bool):
        if data_only:
            raise OSError("cached view unavailable")
        return formula_workbook

    monkeypatch.setattr("core.document_agent.workbook_session.openpyxl.load_workbook", fail_cached_view)

    with pytest.raises(ToolExecutionError, match="无法读取"):
        session.roll_forward_summary_row({
            "sheet": "工资核算", "current_row": 2,
            "expected_current_month": None, "new_month": "7月",
            "start_column": 1, "end_column": 3,
        })

    assert closed == [True]
