"""统一总表的人数守恒与异常处理测试。"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from core.entity_engine import import_file_to_entities
from core.unified_integration import (
    _verify_unchanged_formula_signature,
    build_complete_roster,
    combine_entity_sources,
    create_review_workbook,
    discover_people_in_workbook,
    export_unified_workbook,
    write_manual_issues_workbook,
    _normalize_manual_issues,
    _matching_policy,
)


def test_complete_roster_sends_name_only_record_to_manual_review() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "公司": "甲公司"},
        ],
        "salary_detail": [
            {"工号": "E002", "姓名": "乙", "月基本薪资": 8000},
        ],
        "attendance_record": [
            {"姓名": "丙", "实际出勤天数": 21},
        ],
        "social_security": [
            {"姓名": "丁", "身份证号": "110101199001011234", "个人养老": 640},
        ],
        "tax_record": [
            {"工号": "E002", "姓名": "乙", "本次应扣税额": 80},
            {"姓名": "合计", "本次应扣税额": 80},
            {"姓名": "35", "本次应扣税额": 80},
        ],
    }

    result = build_complete_roster(entities)

    roster = result["entities"]["employee_profile"]
    assert result["source_person_count"] == 3
    assert result["output_person_count"] == 3
    assert {row["姓名"] for row in roster} == {"甲", "乙", "丁"}
    assert all(row.get("_person_key") for row in roster)
    assert any(
        issue["issue_type"] == "unmatched_person" and issue["person_name"] == "丙"
        for issue in result["issues"]
    )


def test_same_name_with_different_employee_ids_is_not_merged() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "王强"},
            {"工号": "E002", "姓名": "王强"},
        ],
        "tax_record": [
            {"姓名": "王强", "本次应扣税额": 100},
        ],
    }

    result = build_complete_roster(entities)

    roster = result["entities"]["employee_profile"]
    assert len(roster) == 2
    assert {row.get("工号") for row in roster if row.get("工号")} == {"E001", "E002"}
    assert any(issue["issue_type"] == "ambiguous_person" for issue in result["issues"])
    assert any(issue["issue_type"] == "unmatched_person" for issue in result["issues"])


def test_unified_export_excludes_name_only_record_without_manual_issue_sheet(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["AI2"] = "实际出勤天数"
    payroll["BF2"] = "个人公积金"
    summary = template.create_sheet("汇总")
    summary["A1"] = "=VLOOKUP(A2,[1]Sheet1!$A:$B,2,0)"
    summary["B1"] = "#N/A"
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "统一工资核算总表.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲"}],
        "salary_detail": [{"工号": "E002", "姓名": "乙", "月基本薪资": 8000}],
        "attendance_record": [{"姓名": "丙", "实际出勤天数": 21}],
        "social_security": [{"姓名": "丁", "身份证号": "110101199001011234", "个人公积金": 320}],
        "tax_record": [],
    }

    result = export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    payroll_out = exported["工资核算"]
    assert result["source_person_count"] == 3
    assert result["output_person_count"] == 3
    assert {payroll_out.cell(row=row, column=2).value for row in range(4, 7)} == {"甲", "乙", "丁"}
    row_by_name = {payroll_out.cell(row=row, column=2).value: row for row in range(4, 7)}
    assert payroll_out[f"BF{row_by_name['丁']}"].value == 320
    assert any(
        issue["issue_type"] == "unmatched_person" and issue["person_name"] == "丙"
        for issue in result["issues"]
    )
    assert "待人工处理" not in exported.sheetnames
    assert exported["汇总"]["A1"].value == "=VLOOKUP(A2,[1]Sheet1!$A:$B,2,0)"
    assert exported["汇总"]["B1"].value is None


def test_manual_issues_are_exported_to_a_standalone_workbook(tmp_path: Path) -> None:
    issues_path = tmp_path / "待人工处理.xlsx"
    write_manual_issues_workbook(str(issues_path), [{
        "issue_type": "missing_value",
        "person_name": "李楠",
        "target_field": "月基本薪资",
        "message": "月基本薪资缺失，统一总表对应单元格已留空。",
        "source_files": [
            "考勤.xlsx", "个税导出核对.xlsx", "防暑降温费.xlsx", "其他调差累计.xlsx",
            "工资核算.xlsx", "绩效考核.xlsx", "工资条.xlsx", "薪酬福利政策.xlsx",
        ],
        "source_sheets": ["考勤", "个税导出核对", "防暑降温费", "其他调差累计", "工资核算", "绩效考核", "工资条"],
    }])

    workbook = load_workbook(issues_path, data_only=False)
    sheet = workbook["待人工处理"]
    assert workbook.sheetnames == ["待人工处理"]
    assert sheet["A1"].value == "序号"
    assert sheet["C2"].value == "李楠"
    assert sheet["I2"].value == "待处理"
    assert sheet.row_dimensions[2].height > 52


def test_people_discovery_finds_valid_names_across_sheets(tmp_path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "不规则考勤"
    first["B3"] = "员工姓名"
    first["A3"] = "员工工号"
    first.append([])
    first["A4"] = "E001"
    first["B4"] = "甲"
    first["A5"] = "E002"
    first["B5"] = "乙"
    first["B6"] = "合计"
    second = workbook.create_sheet("临时补贴")
    second["C1"] = "姓名"
    second["C2"] = "丙"
    path = tmp_path / "source.xlsx"
    workbook.save(path)

    discovered = discover_people_in_workbook(str(path), "source.xlsx")

    assert {row["姓名"] for row in discovered} == {"甲", "乙", "丙"}
    assert {row.get("工号") for row in discovered if row.get("工号")} == {"E001", "E002"}


def test_people_discovery_ignores_section_labels_and_policy_sentences(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "姓名"
    for row, value in enumerate(
        ["序号", "无", "内部调动", "离职、退休", "试用期转正", "1、晚打卡15分钟内均不算迟到；", "张三"],
        start=2,
    ):
        sheet.cell(row=row, column=1, value=value)
    path = tmp_path / "人员异动.xlsx"
    workbook.save(path)

    discovered = discover_people_in_workbook(str(path), path.name)

    assert [row["姓名"] for row in discovered] == ["张三"]


def test_combining_template_and_sources_keeps_source_only_people() -> None:
    template_entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲"}],
        "salary_detail": [],
    }
    source_entities = {
        "employee_profile": [],
        "salary_detail": [{"工号": "E002", "姓名": "乙", "月基本薪资": 8000}],
        "social_security": [{"姓名": "丙", "身份证号": "110101199001011234"}],
    }

    result = combine_entity_sources(template_entities, source_entities)

    assert result["source_person_count"] == 3
    assert {row["姓名"] for row in result["entities"]["employee_profile"]} == {"甲", "乙", "丙"}


def test_authoritative_roster_only_adds_people_confirmed_by_multiple_sources() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "_roster_authoritative": True, "_source_file": "master.xlsx"},
            {"工号": "E002", "姓名": "乙", "_source_file": "attendance.xlsx"},
            {"姓名": "乙", "_source_file": "tax.xlsx"},
            {"工号": "H999", "姓名": "历史人员", "_source_file": "history.xlsx"},
        ],
        "salary_detail": [
            {"工号": "E001", "姓名": "甲", "月基本薪资": 8000, "_source_file": "salary.xlsx"},
            {"姓名": "乙", "绩效奖金": 500, "_source_file": "bonus.xlsx"},
            {"工号": "H999", "姓名": "历史人员", "月基本薪资": 6000, "_source_file": "history.xlsx"},
        ],
    }

    result = build_complete_roster(entities)

    assert result["source_person_count"] == 1
    assert {row["姓名"] for row in result["entities"]["employee_profile"]} == {"甲"}
    assert any(issue["issue_type"] == "unmatched_person" for issue in result["issues"])
    assert result["input_person_count"] == 3
    assert result["output_person_count"] == 1
    assert result["manual_only_person_count"] == 2
    assert result["coverage_gap_count"] == 0


def test_unique_name_and_context_can_auto_link_record_without_strong_identity() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "公司": "北京公司",
                "部门": "零售部",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
        ],
        "salary_detail": [
            {
                "姓名": "甲",
                "公司": "北京公司",
                "部门": "零售部",
                "月基本薪资": 8000,
                "_source_file": "薪资.xlsx",
            },
        ],
    }

    result = build_complete_roster(entities)

    assert result["output_person_count"] == 1
    assert result["entities"]["salary_detail"][0]["_person_key"] == "employee:E001"
    assert result["entities"]["salary_detail"][0]["_match_confidence"] >= 0.85
    assert "姓名" in result["entities"]["salary_detail"][0]["_match_basis"]
    assert not any(issue["issue_type"] == "unmatched_person" for issue in result["issues"])


def test_same_name_with_conflicting_strong_identity_is_never_auto_linked() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "公司": "北京公司",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
        ],
        "salary_detail": [
            {
                "工号": "E999",
                "姓名": "甲",
                "公司": "北京公司",
                "月基本薪资": 8000,
                "_source_file": "薪资.xlsx",
            },
        ],
    }

    result = build_complete_roster(entities)

    assert [(row.get("工号"), row.get("姓名")) for row in result["entities"]["employee_profile"]] == [("E001", "甲")]
    assert any(issue["issue_type"] == "identity_conflict_with_master" for issue in result["issues"])
    assert not any(row.get("_person_key") == "employee:E001" for row in result["entities"]["salary_detail"])


def test_same_name_with_multiple_authoritative_candidates_requires_manual_review() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "公司": "北京公司", "部门": "零售部", "_roster_authoritative": True},
            {"工号": "E002", "姓名": "甲", "公司": "北京公司", "部门": "运营部", "_roster_authoritative": True},
        ],
        "salary_detail": [{"姓名": "甲", "公司": "北京公司", "部门": "财务部", "月基本薪资": 8000}],
    }

    result = build_complete_roster(entities)

    assert result["output_person_count"] == 2
    assert any(issue["issue_type"] == "ambiguous_person" for issue in result["issues"])


def test_matching_policy_normalizes_weights_and_keeps_safe_threshold_floor() -> None:
    threshold, weights = _matching_policy({
        "auto_match_threshold": 0.2,
        "field_weights": {"姓名": 2, "公司": 1, "部门": 1, "岗位": 0},
    })

    assert threshold == 0.85
    assert round(sum(weights.values()), 6) == 1
    assert weights["姓名"] == 0.5


def test_matching_policy_can_force_a_source_sheet_to_manual_review() -> None:
    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲", "公司": "北京公司", "部门": "零售部", "_roster_authoritative": True}],
        "salary_detail": [{"姓名": "甲", "公司": "北京公司", "部门": "零售部", "_source_sheet": "临时表"}],
    }

    result = build_complete_roster(entities, salary_month="2026.08", matching_policy={
        "source_rules": [{"source_sheet": "临时表", "month": "2026.08", "action": "review"}],
    })

    assert any(issue["issue_type"] == "unmatched_person" for issue in result["issues"])
    assert "_person_key" not in result["entities"]["salary_detail"][0]


def test_matching_policy_can_force_a_target_field_to_manual_review() -> None:
    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲", "_roster_authoritative": True}],
        "salary_detail": [{"工号": "E001", "姓名": "甲", "绩效奖金": 500, "_source_sheet": "奖金"}],
    }

    result = build_complete_roster(entities, salary_month="2026.08", matching_policy={
        "source_rules": [{"source_sheet": "奖金", "target_field": "绩效奖金", "action": "review"}],
    })

    assert result["entities"]["salary_detail"][0]["绩效奖金"] is None
    assert any(issue["issue_type"] == "configured_field_review" for issue in result["issues"])


def test_department_transfer_overwrites_master_when_source_value_is_unique() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "原部门",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "新部门",
                "_source_priority": 30,
                "_source_file": "人员异动.xlsx",
            },
        ],
        "salary_detail": [],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["employee_profile"][0]["部门"] == "新部门"
    assert not any(item["issue_type"] == "department_transfer" for item in result["issues"])


def test_department_leaf_name_overwrites_master_when_source_value_is_unique() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "零售业务北区-北京市-支持组",
                "_roster_authoritative": True,
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "支持组",
                "_source_priority": 30,
            },
        ],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["employee_profile"][0]["部门"] == "支持组"
    assert not any(issue["issue_type"] == "department_transfer" for issue in result["issues"])


def test_conflicting_department_sources_require_manual_review() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "零售业务北区-北京市-支持组",
                "_roster_authoritative": True,
            },
            {"工号": "E001", "姓名": "甲", "部门": "支持组", "_source_priority": 30},
            {"工号": "E001", "姓名": "甲", "部门": "运营组", "_source_priority": 30},
        ],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["employee_profile"][0]["部门"] == "零售业务北区-北京市-支持组"
    assert any(issue["issue_type"] == "conflicting_value" for issue in result["issues"])


def test_department_containment_overwrites_master_when_source_value_is_unique() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "管理部",
                "_roster_authoritative": True,
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "大客户管理部",
                "_source_priority": 30,
            },
        ],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["employee_profile"][0]["部门"] == "大客户管理部"
    assert not any(issue["issue_type"] == "department_transfer" for issue in result["issues"])


def test_source_identity_conflict_with_master_is_excluded_for_manual_review() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "同名人员",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
            {
                "工号": "E999",
                "姓名": "同名人员",
                "_roster_confirmation": True,
                "_source_priority": 30,
                "_source_file": "变更.xlsx",
            },
        ],
        "salary_detail": [
            {"工号": "E999", "姓名": "同名人员", "绩效奖金": 500, "_source_priority": 30},
        ],
    }

    result = build_complete_roster(entities)

    assert [(row["工号"], row["姓名"]) for row in result["entities"]["employee_profile"]] == [
        ("E001", "同名人员")
    ]
    issue = next(issue for issue in result["issues"] if issue["issue_type"] == "identity_conflict_with_master")
    assert issue["source_files"] == ["变更.xlsx", "总表.xlsx"]


def test_conflicting_source_values_are_left_blank_for_manual_review() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "_roster_authoritative": True},
        ],
        "salary_detail": [
            {"工号": "E001", "姓名": "甲", "绩效奖金": 100, "_source_priority": 0},
            {"工号": "E001", "姓名": "甲", "绩效奖金": 200, "_source_priority": -10, "_source_sheet": "奖金-6月"},
            {"工号": "E001", "姓名": "甲", "绩效奖金": 300, "_source_priority": 30, "_source_sheet": "奖金-7月"},
        ],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["salary_detail"][0]["绩效奖金"] is None
    issue = next(
        issue
        for issue in result["issues"]
        if issue["issue_type"] == "conflicting_value"
        and issue["target_field"] == "绩效奖金"
    )
    assert issue["candidate_values"] == [100, 200, 300]


def test_current_month_attendance_and_bonus_override_previous_month_values() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "_roster_authoritative": True},
        ],
        "attendance_record": [
            {"工号": "E001", "姓名": "甲", "实际出勤天数": 20, "_source_sheet": "考勤-6月"},
            {"工号": "E001", "姓名": "甲", "实际出勤天数": 22, "_source_sheet": "考勤-7月"},
        ],
        "salary_detail": [
            {"工号": "E001", "姓名": "甲", "绩效奖金": 100, "_source_sheet": "奖金-6月"},
            {"工号": "E001", "姓名": "甲", "绩效奖金": 300, "_source_sheet": "奖金-7月"},
        ],
    }

    result = build_complete_roster(entities, salary_month="2026.07")

    assert result["entities"]["attendance_record"][0]["实际出勤天数"] == 22
    assert result["entities"]["salary_detail"][0]["绩效奖金"] == 300
    assert not any(
        issue["issue_type"] == "conflicting_value"
        and issue["target_field"] in {"实际出勤天数", "绩效奖金"}
        for issue in result["issues"]
    )


def test_current_month_attendance_value_also_resolves_profile_field_conflicts() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "人员类别": "合同制", "_source_sheet": "考勤-6月"},
            {"工号": "E001", "姓名": "甲", "人员类别": "劳务派遣", "_source_sheet": "考勤-7月"},
        ],
    }

    result = build_complete_roster(entities, salary_month="2026.07")

    assert result["entities"]["employee_profile"][0]["人员类别"] == "劳务派遣"
    assert not any(
        issue["issue_type"] == "conflicting_value" and issue["target_field"] == "人员类别"
        for issue in result["issues"]
    )


def test_sheet_month_takes_priority_over_a_current_month_package_filename() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001", "姓名": "甲", "人员类别": "合同制",
                "_source_file": "园区-7月薪资包.xlsx", "_source_sheet": "考勤-6月",
            },
            {
                "工号": "E001", "姓名": "甲", "人员类别": "劳务派遣",
                "_source_file": "园区-7月薪资包.xlsx", "_source_sheet": "考勤-7月",
            },
        ],
    }

    result = build_complete_roster(entities, salary_month="2026.07")

    assert result["entities"]["employee_profile"][0]["人员类别"] == "劳务派遣"
    assert not any(
        issue["issue_type"] == "conflicting_value" and issue["target_field"] == "人员类别"
        for issue in result["issues"]
    )


def test_conflicting_profile_values_are_left_blank_for_manual_review() -> None:
    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "银行卡号": "622200000000000001"},
            {"工号": "E001", "姓名": "甲", "银行卡号": "622200000000000002"},
        ],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["employee_profile"][0]["银行卡号"] is None
    assert any(
        issue["issue_type"] == "conflicting_value"
        and issue["target_field"] == "银行卡号"
        for issue in result["issues"]
    )


def test_source_row_with_only_a_name_is_not_linked_to_the_master() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
        ],
        "attendance_record": [
            {
                "姓名": "甲",
                "实际出勤天数": 21,
                "_source_file": "考勤.xlsx",
            },
        ],
    }

    result = build_complete_roster(entities)

    assert result["entities"]["attendance_record"][0].get("_person_key") is None
    assert any(issue["issue_type"] == "unmatched_person" for issue in result["issues"])


def test_export_does_not_rebuild_an_already_consolidated_roster(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲", "_person_key": "employee:E001"},
            {"工号": "E002", "姓名": "乙", "_person_key": "employee:E002"},
        ],
        "salary_detail": [
            {"_person_key": "employee:E001", "月基本薪资": 8000},
            {"姓名": "不应二次扩入的人", "月基本薪资": 1000},
        ],
    }

    result = export_unified_workbook(str(template_path), entities, str(output_path))

    assert result["source_person_count"] == 2
    assert result["output_person_count"] == 2


def test_export_keeps_existing_formula_when_no_replacement_value_exists(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["X4"] = 10000
    payroll["Y4"] = 0.8
    payroll["Z4"] = "=ROUND(X4*Y4,-1)"
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲"}],
        "salary_detail": [{"工号": "E001", "姓名": "甲"}],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["工资核算"]["Z4"].value == "=ROUND(X4*Y4,-1)"


def test_export_never_replaces_a_template_formula_with_imported_value(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["AI2"] = "实际出勤天数"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["AI4"] = "=SUMIF(考勤!A:A,A4,考勤!G:G)"
    attendance = template.create_sheet("考勤")
    attendance.append(["新工号", "姓名", None, None, None, None, "出勤天数"] + [None] * 15 + ["姓名", "出勤天数"])
    attendance.append(["E999", "乙", None, None, None, None, "=VLOOKUP(B2,W:X,2,0)"] + [None] * 15 + ["甲", 20])
    attendance["AA2"] = "=SUM(AE2:AF2)"
    attendance["AE2"] = 3
    attendance["AF2"] = 2
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲"}],
        "attendance_record": [{"工号": "E001", "姓名": "甲", "实际出勤天数": 0, "普通加班时长": 4}],
    }

    result = export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["工资核算"]["AI4"].value == "=SUMIF(考勤!A:A,A4,考勤!G:G)"
    assert exported["考勤"]["W2"].value == "甲"
    assert exported["考勤"]["X2"].value == 0
    assert exported["考勤"]["AA2"].value == "=SUM(AE2:AF2)"
    assert exported["考勤"]["AE2"].value == 4
    assert exported["考勤"]["AF2"].value == 0
    assert "公式完整性校验通过：原有 3 个公式全部保留" in result["log"]


def test_export_updates_inputs_that_drive_salary_formulas(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["X4"] = 9000
    payroll["Y4"] = 1
    payroll["Z4"] = "=ROUND(X4*Y4,-1)"
    payroll["AA4"] = "=X4-Z4"
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲"}],
        "salary_detail": [{"工号": "E001", "姓名": "甲", "月基本薪资": 8000, "月绩效标准": 2000}],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["工资核算"]["X4"].value == 10000
    assert exported["工资核算"]["Y4"].value == 0.8
    assert exported["工资核算"]["Z4"].value == "=ROUND(X4*Y4,-1)"
    assert exported["工资核算"]["AA4"].value == "=X4-Z4"


def test_formula_integrity_check_rejects_missing_or_changed_formulas(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["Z4"] = "=ROUND(X4*Y4,-1)"
    payroll["CL4"] = ArrayFormula(ref="CL4", text="=ROUND(MAX(CK4*{0.03;0.1},0),2)")
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    exported = load_workbook(template_path, data_only=False)
    exported["工资核算"]["Z4"] = 8000
    exported["工资核算"]["CL4"] = "=ROUND(MAX(CK4*0.03,0),2)"
    exported.save(output_path)
    exported.close()

    try:
        _verify_unchanged_formula_signature(str(template_path), str(output_path))
    except RuntimeError as error:
        message = str(error)
        assert "缺失 1 个" in message
        assert "被改写 1 个" in message
        assert "工资核算!Z4" in message
        assert "工资核算!CL4" in message
    else:
        raise AssertionError("公式缺失或改写时必须阻止导出")


def test_export_preserves_openpyxl_array_formula_objects(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["CL4"] = ArrayFormula(ref="CL4", text="=ROUND(MAX(CK4*{0.03;0.1},0),2)")
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "甲"}],
        "tax_record": [{"工号": "E001", "姓名": "甲", "累计应扣缴税额": 123}],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    formula = exported["工资核算"]["CL4"].value
    assert isinstance(formula, ArrayFormula)
    assert formula.text == "=ROUND(MAX(CK4*{0.03;0.1},0),2)"


def test_inserted_employee_gets_formula_from_same_column_even_if_last_row_is_static(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["BA4"] = "=SUM(Z4:AZ4)"
    payroll["A5"] = "E002"
    payroll["B5"] = "乙"
    payroll["BA5"] = 100
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲"},
            {"工号": "E002", "姓名": "乙"},
            {"工号": "E003", "姓名": "丙"},
        ],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["工资核算"]["BA6"].value == "=SUM(Z6:AZ6)"


def test_row_formula_is_not_mistaken_for_a_total_when_employee_range_expands(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["A5"] = "E002"
    payroll["B5"] = "乙"
    payroll["CI5"] = "=SUM(CC5:CH5)"
    payroll["CI6"] = "=SUM(CI4:CI5)"
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲"},
            {"工号": "E002", "姓名": "乙"},
            {"工号": "E003", "姓名": "丙"},
        ],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["工资核算"]["CI5"].value == "=SUM(CC5:CH5)"
    assert exported["工资核算"]["CI7"].value == "=SUM(CI4:CI6)"


def test_people_discovery_carries_current_roster_profile_fields(tmp_path: Path) -> None:
    workbook = Workbook()
    attendance = workbook.active
    attendance.title = "考勤"
    attendance.append(["新工号", "姓名", "成本中心/部门", "入职日期", "餐补标准"])
    attendance.append(["E002", "乙", "零售业务部", "2026-07-01", "500元/月"])
    path = tmp_path / "attendance.xlsx"
    workbook.save(path)

    discovered = discover_people_in_workbook(str(path), "attendance.xlsx")

    assert discovered == [{
        "工号": "E002",
        "姓名": "乙",
        "身份证号": None,
        "公司": None,
        "部门": "零售业务部",
        "入职日期": "2026-07-01",
        "餐补标准": "500元/月",
        "_source_file": "attendance.xlsx",
        "_source_sheet": "考勤",
        "_discovery_only": True,
    }]


def test_attendance_import_uses_primary_name_column_and_enriches_profile(tmp_path: Path) -> None:
    workbook = Workbook()
    attendance = workbook.active
    attendance.title = "考勤"
    attendance.append([
        "新工号", "姓名", "成本中心/部门", "本月状态", "入职日期", "餐补标准随工资发",
        "出勤天数", None, None, None, None, None, None, None, None, None, None, None,
        None, None, None, None, "姓名", "出勤天数",
    ])
    attendance.append([
        "E002", "乙", "零售业务部", "在职", "2026-07-01", "500元/月",
        14, None, None, None, None, None, None, None, None, None, None, None,
        None, None, None, None, "另一人员", 20,
    ])
    path = tmp_path / "17-考勤.xlsx"
    workbook.save(path)

    imported = import_file_to_entities(str(path), path.name)

    attendance_row = next(row for row in imported["attendance_record"] if row.get("工号") == "E002")
    profile_row = next(row for row in imported["employee_profile"] if row.get("工号") == "E002")
    salary_row = next(row for row in imported["salary_detail"] if row.get("工号") == "E002")
    assert attendance_row["姓名"] == "乙"
    assert attendance_row["实际出勤天数"] == 14
    assert profile_row["部门"] == "零售业务部"
    assert profile_row["入职日期"] == "2026-07-01"
    assert salary_row["餐补标准"] == "500元/月"


def test_export_refreshes_formula_driven_business_sheet_inputs(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "旧姓名"

    attendance = template.create_sheet("考勤")
    attendance.append(["新工号", "姓名", "成本中心/部门", "本月状态", "入职日期", "餐补标准随工资发", "出勤天数"])
    attendance.append(["E001", "旧姓名", "旧部门", "在职", "2020-01-01", "30元/天", "=VLOOKUP(B2,W:X,2,0)"])
    attendance["W1"] = "姓名"
    attendance["X1"] = "出勤天数"
    attendance["W2"] = "新姓名"

    heat = template.create_sheet("防暑降温费")
    heat.append(["新工号", "姓名", "成本中心/部门", "出勤天数", "班制", "防暑降温费标准"])
    heat.append(["E001", "旧姓名", "旧部门", "=VLOOKUP(B2,工资核算!B:AI,34,0)", "旧班制", "=VLOOKUP(B2,工资核算!B:R,17,0)"])

    tax_declaration = template.create_sheet("个税申报")
    tax_declaration.append(["工号", "姓名", "证照类型", "证照号码", "本期收入"])
    tax_declaration.append(["E001", "旧姓名", "居民身份证", "OLD-ID", "=VLOOKUP(B2,工资核算!B:BG,52,0)"])

    tax_check = template.create_sheet("个税导出核对")
    tax_check.append(["工号", "姓名", "证件类型", "证件号码", "税款所属期起", "税款所属期止", "所得项目", "本期收入"])
    tax_check.append(["E001", "旧姓名", "居民身份证", "OLD-ID", "2026-06-01", "2026-06-30", "正常工资薪金", 100])

    ledger = template.create_sheet("台账")
    ledger.append(list(range(1, 18)))
    ledger.append(["新工号", "姓名", "缴费公司-与劳动合同公司主体一致"])
    ledger.append(["E001", "旧姓名", "旧公司"])

    payslip = template.create_sheet("工资条")
    payslip["D1"] = "员工"
    payslip["D2"] = "旧姓名"
    payslip["K2"] = "=VLOOKUP(D2,工资核算!B:CN,10,0)"

    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{
            "工号": "E001",
            "姓名": "新姓名",
            "部门": "新部门",
            "本月状态": "在职",
            "入职日期": "2026-07-01",
            "身份证号": "110101199001011234",
        }],
        "salary_detail": [{"工号": "E001", "姓名": "新姓名", "餐补标准": "500元/月"}],
        "attendance_record": [{"工号": "E001", "姓名": "新姓名", "实际出勤天数": 20, "班制": "常白班"}],
        "tax_record": [{
            "工号": "E001",
            "姓名": "新姓名",
            "身份证号": "110101199001011234",
            "税款所属期起": "2026-07-01",
            "税款所属期止": "2026-07-31",
            "本期收入": 8888,
        }],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["考勤"]["B2"].value == "新姓名"
    assert exported["考勤"]["C2"].value == "新部门"
    assert exported["考勤"]["F2"].value == "500元/月"
    assert exported["考勤"]["G2"].value == "=VLOOKUP(B2,W:X,2,0)"
    assert exported["防暑降温费"]["B2"].value == "新姓名"
    assert exported["防暑降温费"]["E2"].value == "常白班"
    assert exported["防暑降温费"]["D2"].value.startswith("=VLOOKUP")
    assert exported["个税申报"]["B2"].value == "新姓名"
    assert exported["个税申报"]["D2"].value == "110101199001011234"
    assert exported["个税申报"]["E2"].value.startswith("=VLOOKUP")
    assert exported["个税导出核对"]["B2"].value == "新姓名"
    assert exported["个税导出核对"]["D2"].value == "110101199001011234"
    assert exported["个税导出核对"]["E2"].value == "2026-07-01"
    assert exported["个税导出核对"]["H2"].value == 8888
    assert exported["台账"]["B3"].value == "新姓名"
    assert exported["工资条"]["D2"].value == "新姓名"
    assert exported["工资条"]["K2"].value.startswith("=VLOOKUP")
    exported.close()


def test_export_refreshes_literal_heat_allowance_attendance_days(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "旧姓名"

    attendance = template.create_sheet("考勤")
    attendance.append(["新工号", "姓名"])
    attendance.append(["E001", "旧姓名"])

    heat = template.create_sheet("防暑降温费")
    heat.append(["新工号", "姓名", "成本中心/部门", "出勤天数", "班制"])
    heat.append(["E001", "旧姓名", "旧部门", 0, "旧班制"])

    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{"工号": "E001", "姓名": "新姓名", "部门": "新部门"}],
        "attendance_record": [{"工号": "E001", "姓名": "新姓名", "实际出勤天数": 20, "班制": "常白班"}],
    }

    export_unified_workbook(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    assert exported["防暑降温费"]["D2"].value == 20
    assert exported["防暑降温费"]["E2"].value == "常白班"
    exported.close()


def test_template_quirks_that_are_auto_handled_are_not_manual_issues(tmp_path: Path) -> None:
    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["A4"] = "E001"
    payroll["B4"] = "甲"
    payroll["CP4"] = "=VLOOKUP(B4,[1]Sheet1!$B:$R,13,0)"
    payroll["BQ46"] = "#N/A"
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"
    template.save(template_path)

    entities = {
        "employee_profile": [{
            "工号": "E001",
            "姓名": "甲",
            "身份证号": "110101199001011234",
        }],
        "salary_detail": [{"工号": "E001", "姓名": "甲", "月基本薪资": 8000}],
    }

    result = export_unified_workbook(str(template_path), entities, str(output_path))

    assert result["issue_count"] == 0


def test_latest_effective_record_wins_before_salary_month_cutoff() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "旧部门",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "第一部门",
                "_effective_date": "2026-07-05",
                "_source_file": "人员异动.xlsx",
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "最终部门",
                "_effective_date": "2026-07-20",
                "_source_file": "人员变更.xlsx",
            },
        ],
    }

    result = build_complete_roster(entities, salary_month="2026.07")

    assert result["entities"]["employee_profile"][0]["部门"] == "最终部门"
    assert not any(issue["issue_type"] == "department_transfer" for issue in result["issues"])


def test_future_and_same_day_conflicting_records_require_manual_review() -> None:
    entities = {
        "employee_profile": [
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "旧部门",
                "_roster_authoritative": True,
                "_source_file": "总表.xlsx",
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "甲部门",
                "_effective_date": "2026-07-20",
                "_source_file": "变更一.xlsx",
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "乙部门",
                "_effective_date": "2026-07-20",
                "_source_file": "变更二.xlsx",
            },
            {
                "工号": "E001",
                "姓名": "甲",
                "部门": "未来部门",
                "_effective_date": "2026-08-01",
                "_source_file": "未来变更.xlsx",
            },
        ],
    }

    result = build_complete_roster(entities, salary_month="2026.07")

    assert result["entities"]["employee_profile"][0]["部门"] is None
    issue_types = {issue["issue_type"] for issue in result["issues"]}
    assert "effective_date_conflict" in issue_types
    assert "future_effective_date" in issue_types


def test_manual_issues_workbook_sorts_simple_before_complex(tmp_path: Path) -> None:
    output_path = tmp_path / "待人工处理.xlsx"
    issues = [
        {
            "issue_type": "conflicting_identity",
            "person_name": "复杂人员",
            "target_field": "工号",
            "message": "身份存在冲突",
            "source_files": ["复杂.xlsx"],
            "source_sheets": ["人员"],
        },
        {
            "issue_type": "missing_value",
            "person_name": "简单人员",
            "target_field": "身份证号",
            "message": "字段缺失",
            "source_files": ["简单.xlsx"],
            "source_sheets": ["名单"],
        },
    ]

    write_manual_issues_workbook(str(output_path), issues)

    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook["待人工处理"]
    assert [cell.value for cell in sheet[1]] == [
        "序号", "事项编号", "人员姓名", "问题类型", "处理难度", "目标字段", "问题说明",
        "建议操作", "处理状态", "处理结果", "处理值", "处理备注", "来源文件", "来源工作表",
        "可选动作", "可复制内容",
    ]
    assert sheet["C2"].value == "简单人员"
    assert sheet["E2"].value == "简单"
    assert sheet["J2"].value == "待处理"
    assert len(sheet.data_validations.dataValidation) == 1
    assert sheet["C3"].value == "复杂人员"
    assert sheet["E3"].value == "复杂"
    workbook.close()


def test_actionable_manual_issue_includes_choices_and_copyable_patch() -> None:
    issues = _normalize_manual_issues([{
        "issue_type": "department_transfer",
        "person_name": "甲",
        "employee_id": "E001",
        "target_field": "部门",
        "current_value": "旧部门",
        "proposed_value": "新部门",
        "message": "部门发生变化",
        "source_files": ["人员异动.xlsx"],
        "source_sheets": ["入离职、转岗、转正"],
    }])

    issue = issues[0]
    assert issue["recommended_action"] == "apply_proposed"
    assert issue["action_options"] == [
        {"value": "apply_proposed", "label": "采用新值：新部门"},
        {"value": "keep_current", "label": "保留原值：旧部门"},
    ]
    assert issue["copy_text"] == "E001\t甲\t部门\t旧部门\t新部门"


def test_candidate_value_manual_issue_includes_candidates_in_copyable_patch() -> None:
    issue = _normalize_manual_issues([{
        "issue_type": "conflicting_value",
        "person_name": "甲",
        "employee_id": "E001",
        "target_field": "绩效奖金",
        "candidate_values": [100, 200, 300],
        "message": "来源值冲突",
    }])[0]

    assert issue["recommended_action"] == "copy_only"
    assert issue["copy_text"] == "E001\t甲\t绩效奖金\t100；200；300"


def test_review_workbook_annotates_changes_without_using_color_markers(tmp_path: Path) -> None:
    base_path = tmp_path / "总表.xlsx"
    final_path = tmp_path / "正式稿.xlsx"
    review_path = tmp_path / "修改稿.xlsx"

    base = Workbook()
    base_sheet = base.active
    base_sheet.title = "工资核算"
    base_sheet.append(["工号", "姓名", "月基本薪资"])
    base_sheet.append(["E001", "甲", 8000])
    base_sheet.append(["E003", "丙", 7000])
    base.save(base_path)
    base.close()

    final = Workbook()
    final_sheet = final.active
    final_sheet.title = "工资核算"
    final_sheet.append(["工号", "姓名", "月基本薪资"])
    final_sheet.append(["E001", "甲", 9000])
    final_sheet.append(["E002", "乙", 6000])
    final.save(final_path)
    final.close()

    summary = create_review_workbook(
        str(base_path),
        str(final_path),
        str(review_path),
    )

    review = load_workbook(review_path, data_only=False)
    payroll = review["工资核算"]
    audit = review["修改记录"]
    assert payroll["C2"].fill.patternType is None
    assert payroll["C2"].comment is not None
    assert "位置：工资核算!C2" in payroll["C2"].comment.text
    assert "原值：8000" in payroll["C2"].comment.text
    assert "新值：9000" in payroll["C2"].comment.text
    assert payroll["A3"].comment is not None
    assert "本行新增" in payroll["A3"].comment.text
    assert "位置：工资核算!A3" in payroll["A3"].comment.text
    assert audit.max_row == 3
    assert [cell.value for cell in audit[1]] == [
        "序号", "工作表", "人员标识", "姓名", "变更类型", "字段", "原值", "新值", "单元格位置", "变更说明",
    ]
    assert [cell.value for cell in audit[2]] == [
        1, "工资核算", "E001", "甲", "字段修改", "月基本薪资", 8000, 9000, "C2", "工资核算!C2：甲的“月基本薪资”由“8000”改为“9000”",
    ]
    assert [cell.value for cell in audit[3]] == [
        2, "工资核算", "E002", "乙", "新增行", "整行", None, None, "A3", "工资核算!A3：新增人员/记录“乙”",
    ]
    assert audit.auto_filter.ref == "A1:J3"
    assert summary == {"change_count": 2, "audit_row_count": 2}
    review.close()
