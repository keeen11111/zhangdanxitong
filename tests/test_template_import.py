"""主工资模板导入的回归测试。"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.entity_engine import export_to_template, import_file_to_entities, import_template_to_entities


def test_template_import_preserves_identity_and_salary_fields(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工资核算"
    headers = [
        "新工号",
        "姓名",
        "新公司",
        "新成本中心/部门",
        "新岗位",
        "是否试用期",
        "转正日期",
        "交通通讯补贴标准",
        "月基本薪资",
        "月绩效标准",
        "绩效奖金",
        "身份证号码",
        "工资银行卡账号",
        "开户银行",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=2, column=column, value=header)

    values = [
        "E001",
        "测试员工",
        "测试公司",
        "财务部",
        "专员",
        "N",
        "2026-01-01",
        300,
        8000,
        2000,
        0,
        "110101199001011234",
        "6222000000000000",
        "测试银行",
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(row=4, column=column, value=value)

    path = tmp_path / "工资模板.xlsx"
    workbook.save(path)

    result = import_template_to_entities(str(path))

    profile = result["employee_profile"][0]
    assert profile["是否试用期"] == "N"
    assert profile["转正日期"] == "2026-01-01"
    assert profile["身份证号"] == "110101199001011234"
    assert profile["银行卡号"] == "6222000000000000"
    assert profile["开户银行"] == "测试银行"

    salary = result["salary_detail"][0]
    assert salary["工号"] == "E001"
    assert salary["月基本薪资"] == 8000
    assert salary["月绩效标准"] == 2000
    assert salary["绩效奖金"] == 0
    assert salary["交通通讯补贴标准"] == 300


def test_bonus_import_uses_scalar_when_headers_repeat(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "奖金-7月"
    sheet.append(["姓名", "工号", "部门", "部门", "岗位", "绩效奖金(月)"])
    sheet.append(["张三", "E001", "财务部", None, "专员", 500])
    path = tmp_path / "薪资变更.xlsx"
    workbook.save(path)

    result = import_file_to_entities(str(path), path.name)

    profile = next(row for row in result["employee_profile"] if row.get("姓名") == "张三")
    assert profile["部门"] == "财务部"
    assert "dtype:" not in profile["部门"]


def test_export_inserts_employee_before_total_and_extends_formula_ranges(tmp_path: Path) -> None:
    workbook = Workbook()
    payroll = workbook.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    payroll["Z2"] = "月基本薪资"
    payroll["AA2"] = "月绩效标准"
    payroll["X2"] = "月度合计"

    for row, employee_id, name in [(4, "E001", "甲"), (5, "E002", "乙")]:
        payroll.cell(row=row, column=1, value=employee_id)
        payroll.cell(row=row, column=2, value=name)
        payroll.cell(row=row, column=24, value=1100 if row == 4 else 2200)
    payroll["Z6"] = "=SUM(Z4:Z5)"

    summary = workbook.create_sheet("汇总")
    summary["A1"] = "=SUM('工资核算'!Z4:Z5)"
    summary["B1"] = "=VLOOKUP(A1,[1]Sheet1!$A:$B,2,0)"
    summary["C1"] = "#N/A"

    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook.save(template_path)

    entities = {
        "employee_profile": [
            {"工号": "E001", "姓名": "甲"},
            {"工号": "E002", "姓名": "乙"},
            {"工号": "E003", "姓名": "丙"},
        ],
        "salary_detail": [
            {"工号": "E001", "月基本薪资": 1000, "月绩效标准": 100},
            {"工号": "E002", "月基本薪资": 2000, "月绩效标准": 200},
            {"工号": "E003", "月基本薪资": 3000, "月绩效标准": 300},
        ],
    }

    export_to_template(str(template_path), entities, str(output_path))

    exported = load_workbook(output_path, data_only=False)
    payroll_out = exported["工资核算"]
    assert payroll_out["A6"].value == "E003"
    assert payroll_out["X6"].value == "=SUM(Z6:AA6)"
    assert payroll_out["Y6"].value == '=IF(OR(X6="",X6=0,Z6=""),"",Z6/X6)'
    assert payroll_out["Z7"].value == "=SUM(Z4:Z6)"
    assert exported["汇总"]["A1"].value == "=SUM('工资核算'!Z4:Z6)"
    assert exported["汇总"]["B1"].value == "=VLOOKUP(A1,[1]Sheet1!$A:$B,2,0)"
    assert exported["汇总"]["C1"].value is None


def test_imports_social_billing_contributions_for_payroll_export(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["鹤安长泰202606月账单"])
    sheet.append([None, "姓名", "身份证号码", None, None, None, None, "单位申报基数", None, None, None, None, "社会保险", None, None, None, None, None, None, None, None, None, None, "住房公积金"])
    sheet.append([None] * 31)
    sheet.append([None] * 12 + ["公司16%", "个人8%", None, "公司0.5%", "个人0.5%", None, "公司0.2%", "公司9.8%", "个人2%+3", None, None, None, "公司12%", "个人12%"])
    sheet.append([
        1, "张三", "110101199001011234", None, None, None, None, 7162,
        7162, 7162, 7162, 7162, 1145.92, 572.96, 1718.88, 35.81, 35.81,
        71.62, 14.32, 701.88, 146.24, 848.12, 2652.94, 7162, 860, 860,
        1720, 4372.94,
    ])
    path = tmp_path / "鹤安长泰2026.06.xlsx"
    workbook.save(path)

    result = import_file_to_entities(str(path), path.name)

    social = result["social_security"]
    assert len(social) == 1
    assert social[0]["个人养老"] == 572.96
    assert social[0]["个人医疗"] == 146.24
    assert social[0]["个人公积金"] == 860
    assert social[0]["个人合计"] == 1615.01
    assert social[0]["公司养老"] == 1145.92
    assert social[0]["公司医疗"] == 701.88
    assert social[0]["公司公积金"] == 860
    assert social[0]["公司合计"] == 2757.93

    template = Workbook()
    payroll = template.active
    payroll.title = "工资核算"
    payroll["A2"] = "新工号"
    payroll["B2"] = "姓名"
    template_path = tmp_path / "工资核算模板.xlsx"
    output_path = tmp_path / "工资核算总表.xlsx"
    template.save(template_path)

    export_to_template(
        str(template_path),
        {
            "employee_profile": [{
                "工号": "E001",
                "姓名": "张三",
                "身份证号": "110101199001011234",
            }],
            "social_security": social,
        },
        str(output_path),
    )

    exported = load_workbook(output_path, data_only=False)["工资核算"]
    assert exported["BC4"].value == 1615.01
    assert exported["BD4"].value == 572.96
    assert exported["BF4"].value == 860
    assert exported["BH4"].value == 2757.93
    assert exported["BI4"].value == 1145.92
    assert exported["BK4"].value == 860


def test_imports_unknown_vendor_social_billing_with_two_row_grouped_headers(tmp_path: Path) -> None:
    """Social bills must be recognized by their headers, not by vendor/file name."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "缴费明细"
    sheet.append(["2026年6月缴费清单"])
    sheet.append([
        "序号", "员工姓名", "证件号码", "参保单位",
        "养老保险", None, "失业保险", None, "工伤保险",
        "医疗保险", None, "住房公积金", None,
    ])
    sheet.append([
        None, None, None, None,
        "单位缴纳", "个人缴纳", "单位缴纳", "个人缴纳", "单位缴纳",
        "单位缴纳", "个人缴纳", "单位缴纳", "个人缴纳",
    ])
    sheet.append([
        1, "李四", "110101199001011235", "测试公司",
        1200, 600, 40, 40, 12, 720, 150, 500, 500,
    ])
    sheet.append([
        None, None, "——", None,
        1200, 600, 40, 40, 12, 720, 150, 500, 500,
    ])
    path = tmp_path / "供应商A_202606缴费明细.xlsx"
    workbook.save(path)

    result = import_file_to_entities(str(path), path.name)

    social = result["social_security"]
    assert len(social) == 1
    assert social[0]["姓名"] == "李四"
    assert social[0]["缴费公司"] == "测试公司"
    assert social[0]["公司养老"] == 1200
    assert social[0]["个人养老"] == 600
    assert social[0]["公司失业"] == 40
    assert social[0]["个人失业"] == 40
    assert social[0]["公司工伤"] == 12
    assert social[0]["公司医疗"] == 720
    assert social[0]["个人医疗"] == 150
    assert social[0]["公司公积金"] == 500
    assert social[0]["个人公积金"] == 500
