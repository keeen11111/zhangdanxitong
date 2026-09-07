"""Personnel-only export using the same project isolation and download contracts."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from core._new_hire_sync import _layout
from core._structured_hire import _structured_people
from backend.workbook_calculation import recalculate_personnel_workbook, preserve_review_formula_caches
from core.sheet_mapper import apply_semantic_sheet_updates
from core.unified_integration import write_manual_issues_workbook


def _export_personnel_updates(
    master_path: Path,
    source_paths: list[Path],
    source_names: dict[str, str],
    output_dir: Path,
    salary_month: str,
    matching_policy: dict[str, Any] | None,
) -> dict[str, Any]:
    # Reuse the existing change ledger, including per-cell source provenance.
    from backend.routers.financial_workbooks import _create_review_workbook

    result_id = uuid4().hex[:12]
    filename = f"人员信息更新_{result_id}.xlsx"
    review_filename = f"人员信息更新_修改记录_{result_id}.xlsx"
    output_path = output_dir / filename
    result = apply_semantic_sheet_updates(
        master_path, source_paths, output_path, salary_month=salary_month,
        source_names=source_names, matching_policy=matching_policy,
    )
    if result.get('structured_hire') and not result['issues']:
        result['calculation'] = recalculate_personnel_workbook(master_path, output_path, result.get('row_insertions', []))
    _create_review_workbook(output_path, output_dir / review_filename, result["updates"])
    if result.get('calculation'):
        preserve_review_formula_caches(output_path, output_dir / review_filename)
    issues_filename = None
    if result["issues"]:
        issues_filename = f"人员信息更新_未完成事项_{result_id}.xlsx"
        write_manual_issues_workbook(str(output_dir / issues_filename), result["issues"])
    book = load_workbook(output_path)
    try:
        layout = _layout(book["工资核算"]) if "工资核算" in book.sheetnames else None
        output_people = len(layout.rows) if layout else len(_structured_people(book))
        target_sheets = book.sheetnames
    finally:
        book.close()
    coverage = result["personnel_coverage"]
    validation = {
        "scope": "personnel_updates", "personnel_coverage": coverage,
        "blocker_count": len(result["issues"]), "warning_count": 0,
        "blockers": result['issues'], "warnings": [],
        "can_export": not result['issues'], "calculation": result.get('calculation'),
        "message": f"人员信息及薪资标准已完整处理 {coverage['processed_person_count']} 人，未完成 {coverage['unprocessed_person_count']} 人。此结果不代表整月工资核算完成。",
    }
    return {
        **result, "filename": filename, "review_filename": review_filename,
        "issues_filename": issues_filename, "issue_count": len(result["issues"]),
        "source_file_count": len(source_paths), "source_person_count": coverage["input_person_count"],
        "input_person_count": coverage["input_person_count"],
        "output_person_count": output_people,
        "manual_only_person_count": coverage["unprocessed_person_count"],
        "coverage_gap_count": coverage["unprocessed_person_count"],
        "target_sheets": target_sheets, "validation": validation,
        "completed_at": datetime.now().isoformat(),
        "log": [
            "识别为客户单发入职通知，按工号同步相关人员明细表页。",
            f"来源 {coverage['input_person_count']} 人，完整自动处理 {coverage['processed_person_count']} 人，未完成 {coverage['unprocessed_person_count']} 人。",
            f"已记录 {result['auto_update_count']} 处单元格变更。",
        ],
    }
