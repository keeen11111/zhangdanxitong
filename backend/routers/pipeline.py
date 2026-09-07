"""数据处理流水线路由。

6个功能模块的API：
    A. 原始数据库（base_db）：从模板提取，可编辑
    B. 数据源头表（source_data）：从源文件提取，只读
    C. 匹配引擎（compute_diff）：A与B匹配计算差异
    D. 待处理数据（diff）：差异审核，确认/修改/忽略
    E. 最终数据库（final_db）：A+D确认结果拼接
    F. 模板导出（export）：E按模板渲染导出
"""
import json
import hashlib
import os
import re
import time
from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote

import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile as FastAPIUploadFile
from fastapi.responses import StreamingResponse
from fastapi.encoders import jsonable_encoder
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from backend.auth import get_current_user
from backend.database import DATA_DIR, EXPORT_DIR, UPLOAD_DIR, get_db
from backend.models import User, Project, UploadFile, EntitySnapshot
from core.entity_engine import (
    import_file_to_entities,
    import_template_to_entities,
    export_to_template,
    get_entity_overview,
)
from core.unified_integration import (
    _payroll_person_count,
    _verify_unchanged_formula_signature,
    combine_entity_sources,
    create_review_workbook,
    discover_people_in_workbook,
    export_unified_workbook,
    MANUAL_ISSUES_SHEET,
    MANUAL_REVIEW_OUTCOMES,
    write_manual_issues_workbook,
    remove_manual_issues_sheet,
)
from core.sheet_mapper import apply_semantic_sheet_updates
from core._new_hire_sync import _only_hire_notices
from backend.personnel_integration import _export_personnel_updates
from core.diff_engine import compute_diff, apply_confirmed_diff
from core.validator import DataValidator

router = APIRouter(prefix="/api/pipeline", tags=["pipeline"])

SESSION_DIR = os.path.join(DATA_DIR, "sessions")
BUILTIN_TEMPLATE_PATH = os.path.join(DATA_DIR, "templates", "大药房工资核算总表-v2.xlsx")
os.makedirs(SESSION_DIR, exist_ok=True)

PRESERVED_LAYOUT_SHEETS = {"工资核算"}
_PAYROLL_MASTER_FILE_TYPES = {"template", "financial_master"}
_PAYROLL_SOURCE_FILE_TYPES = {"source", "financial_source"}
_EXTERNAL_WORKBOOK_REFERENCE = re.compile(
    r"(?:'[^']*\[[^\]]+\][^']*'|\[[^\]]+\][^!]*?)!"
)
_EXCEL_ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
}


# ============================================================
# 工具函数
# ============================================================
def _load_project_or_404(project_id: str, user: User, db: Session) -> Project:
    p = db.query(Project).filter(Project.id == project_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="项目不存在")
    if p.owner.tenant_id != user.tenant_id:
        raise HTTPException(status_code=403, detail="无权访问")
    return p


def _build_export_preview(
    file_path: str,
    max_rows: int = 20,
    max_columns: int = 16,
    preferred_sheet: str | None = None,
) -> dict[str, Any]:
    """Return a small, display-safe preview of a generated Excel workbook."""
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        if preferred_sheet and preferred_sheet in workbook.sheetnames:
            sheet = workbook[preferred_sheet]
        else:
            sheet = workbook["工资核算"] if "工资核算" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows: list[list[Any]] = []
        for values in sheet.iter_rows(max_row=max_rows, max_col=max_columns, values_only=True):
            row = [value.isoformat() if isinstance(value, datetime) else value for value in values]
            while row and row[-1] in (None, ""):
                row.pop()
            if any(value not in (None, "") for value in row):
                rows.append(row)
        return {"sheet_name": sheet.title, "rows": rows}
    finally:
        workbook.close()


def _manual_location_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _attach_manual_issue_locations(workbook_path: str, issues: list[dict[str, Any]]) -> None:
    """Attach a safe Sheet/cell address for editable manual-review items."""
    if not issues or not os.path.exists(workbook_path):
        return
    employee_headers = {"工号", "新工号", "员工工号", "员工编号", "人员编码", "工牌号"}
    name_headers = {"姓名", "人员姓名", "员工姓名", "名字"}
    locations: dict[tuple[str, str, str], list[tuple[str, str]]] = {}
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            for header_row in range(1, min(sheet.max_row, 30) + 1):
                headers = {
                    column: _manual_location_text(sheet.cell(header_row, column).value)
                    for column in range(1, sheet.max_column + 1)
                }
                id_column = next((column for column, header in headers.items() if header in employee_headers), None)
                name_column = next((column for column, header in headers.items() if header in name_headers), None)
                field_columns = {
                    header: column for column, header in headers.items()
                    if header and header not in employee_headers and header not in name_headers
                }
                if not field_columns or (id_column is None and name_column is None):
                    continue
                for row in range(header_row + 1, sheet.max_row + 1):
                    employee_id = _manual_location_text(sheet.cell(row, id_column).value) if id_column else ""
                    person_name = _manual_location_text(sheet.cell(row, name_column).value) if name_column else ""
                    if not employee_id and not person_name:
                        continue
                    for field, column in field_columns.items():
                        # Read-only worksheets represent an empty cell as an
                        # EmptyCell without a coordinate.  The target address
                        # is defined by the row and column we already know.
                        coordinate = f"{openpyxl.utils.get_column_letter(column)}{row}"
                        if employee_id:
                            locations.setdefault(("employee", employee_id, field), []).append((sheet.title, coordinate))
                        if person_name:
                            locations.setdefault(("name", person_name, field), []).append((sheet.title, coordinate))
    finally:
        workbook.close()

    for issue in issues:
        field = _manual_location_text(issue.get("target_field"))
        employee_id = _manual_location_text(issue.get("employee_id"))
        person_name = _manual_location_text(issue.get("person_name"))
        candidates = locations.get(("employee", employee_id, field), []) if employee_id else []
        if not candidates and person_name and person_name != "全表":
            candidates = locations.get(("name", person_name, field), [])
        if len(candidates) == 1:
            sheet_name, cell = candidates[0]
            issue["target_sheet"] = sheet_name
            issue["target_cell"] = cell
            issue["target_location"] = f"{sheet_name}!{cell}"
        else:
            issue["target_location"] = "未能自动定位"


def _normalize_export_filename(filename: str) -> str:
    """Normalize a user-provided workbook filename without allowing path traversal."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", filename).strip().strip(". ")
    stem, suffix = os.path.splitext(cleaned)
    if suffix.lower() in {".xlsx", ".xls"}:
        cleaned = stem.strip().rstrip(". ")
    if not cleaned:
        raise ValueError("文件名不能为空")
    if cleaned.upper() in {"CON", "PRN", "AUX", "NUL", *{f"COM{i}" for i in range(1, 10)}, *{f"LPT{i}" for i in range(1, 10)}}:
        cleaned = f"{cleaned}_工资总表"
    return f"{cleaned[:120]}.xlsx"


def _unique_sheet_title(workbook: openpyxl.Workbook, preferred: str) -> str:
    """Return a valid, deterministic Excel sheet title that cannot overwrite data."""
    cleaned = re.sub(r"[\\/*?:\[\]]", "-", preferred).strip() or "来源数据"
    base = cleaned[:31]
    if base not in workbook.sheetnames:
        return base
    counter = 2
    while True:
        suffix = f"-{counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
        counter += 1


def _archived_source_sheet_title(
    workbook: openpyxl.Workbook,
    source_path: str,
    sheet_title: str,
) -> str:
    stem = os.path.splitext(os.path.basename(source_path))[0]
    preferred = f"来源-{stem}"
    if sheet_title not in stem:
        preferred = f"{preferred}-{sheet_title}"
    return _unique_sheet_title(workbook, preferred)


def _sheet_content_signature(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    cached_sheet: openpyxl.worksheet.worksheet.Worksheet | None = None,
) -> dict[str, Any]:
    """Capture source content that must survive the merge."""
    def normalized_value(value: Any) -> Any:
        # Excel serializes empty strings as blank cells when a workbook is
        # saved, so they must not make an otherwise lossless copy fail.
        if isinstance(value, str) and value == "":
            return None
        if hasattr(value, "text") and hasattr(value, "ref"):
            return (value.__class__.__name__, value.text, value.ref)
        return value

    cells = []
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            data_type = cell.data_type
            if cached_sheet is not None and _is_external_workbook_formula(value):
                cached_cell = cached_sheet[cell.coordinate]
                value = _cached_external_formula_value(cell, cached_cell)
                data_type = cached_cell.data_type
            normalized = normalized_value(value)
            if normalized is not None:
                cells.append((cell.coordinate, data_type, normalized))
    return {
        "cells": tuple(cells),
        "images": len(sheet._images),
        "charts": len(sheet._charts),
        "merged_ranges": tuple(sorted(str(item) for item in sheet.merged_cells.ranges)),
    }


def _formula_text(value: Any) -> str | None:
    if hasattr(value, "text") and isinstance(value.text, str):
        return value.text
    return value if isinstance(value, str) and value.startswith("=") else None


def _is_external_workbook_formula(value: Any) -> bool:
    formula = _formula_text(value)
    return bool(formula and _EXTERNAL_WORKBOOK_REFERENCE.search(formula))


def _cached_external_formula_value(
    source_cell: openpyxl.cell.cell.Cell,
    cached_cell: openpyxl.cell.cell.Cell,
) -> Any:
    cached_value = cached_cell.value
    if (
        cached_value is None
        or cached_cell.data_type == "e"
        or cached_value in _EXCEL_ERROR_VALUES
    ):
        raise ValueError(
            "来源工作表存在外部公式但没有可用缓存结果："
            f"{source_cell.parent.title}!{source_cell.coordinate}"
        )
    return cached_value


def _copy_source_sheet_content(
    source_sheet: openpyxl.worksheet.worksheet.Worksheet,
    target_sheet: openpyxl.worksheet.worksheet.Worksheet,
    cached_source_sheet: openpyxl.worksheet.worksheet.Worksheet | None = None,
) -> None:
    target_sheet.freeze_panes = source_sheet.freeze_panes
    target_sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
    target_sheet.auto_filter.ref = source_sheet.auto_filter.ref
    target_sheet.sheet_state = source_sheet.sheet_state
    target_sheet.sheet_properties = copy(source_sheet.sheet_properties)
    target_sheet.sheet_format = copy(source_sheet.sheet_format)
    target_sheet.page_margins = copy(source_sheet.page_margins)
    target_sheet.page_setup = copy(source_sheet.page_setup)
    target_sheet.print_options = copy(source_sheet.print_options)

    for source_row in source_sheet.iter_rows():
        for source_cell in source_row:
            copied_value = source_cell.value
            if _is_external_workbook_formula(copied_value):
                if cached_source_sheet is None:
                    raise ValueError(
                        "复制来源工作表时缺少外部公式缓存工作表："
                        f"{source_sheet.title}!{source_cell.coordinate}"
                    )
                cached_cell = cached_source_sheet.cell(source_cell.row, source_cell.column)
                copied_value = _cached_external_formula_value(source_cell, cached_cell)
            target_cell = target_sheet.cell(source_cell.row, source_cell.column, copied_value)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.alignment:
                target_cell.alignment = copy(source_cell.alignment)
            if source_cell.protection:
                target_cell.protection = copy(source_cell.protection)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))
    for key, dimension in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[key] = copy(dimension)
        target_sheet.column_dimensions[key].worksheet = target_sheet
    for key, dimension in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[key] = copy(dimension)
        target_sheet.row_dimensions[key].worksheet = target_sheet
    for source_image in source_sheet._images:
        image = OpenpyxlImage(BytesIO(source_image._data()))
        image.width = source_image.width
        image.height = source_image.height
        target_sheet.add_image(image, copy(source_image.anchor))
    for source_chart in source_sheet._charts:
        target_sheet.add_chart(copy(source_chart), copy(source_chart.anchor))


def _copy_source_workbooks(
    output_path: str,
    source_paths: list[str],
    source_names: dict[str, str] | None = None,
) -> list[str]:
    """Copy every source sheet without silently skipping or overwriting content."""
    target = openpyxl.load_workbook(output_path, data_only=False)
    copied_sheets: list[str] = []
    copied_canonical_titles: set[str] = set()
    expected_signatures: dict[str, dict[str, Any]] = {}
    original_names = {
        os.path.normcase(os.path.abspath(path)): name
        for path, name in (source_names or {}).items()
    }
    try:
        for source_path in source_paths:
            source_label = original_names.get(
                os.path.normcase(os.path.abspath(source_path)),
                source_path,
            )
            source = openpyxl.load_workbook(source_path, data_only=False)
            cached_source = openpyxl.load_workbook(source_path, data_only=True, keep_links=False)
            try:
                for source_sheet in source.worksheets:
                    cached_source_sheet = cached_source[source_sheet.title]
                    if source_sheet.title in PRESERVED_LAYOUT_SHEETS:
                        destination_title = _archived_source_sheet_title(
                            target,
                            source_label,
                            source_sheet.title,
                        )
                        issue_index = (
                            target.index(target["待人工处理"])
                            if "待人工处理" in target.sheetnames
                            else len(target.sheetnames)
                        )
                        target_sheet = target.create_sheet(destination_title, issue_index)
                    elif (
                        source_sheet.title in target.sheetnames
                        and source_sheet.title not in copied_canonical_titles
                    ):
                        current_sheet = target[source_sheet.title]
                        sheet_index = target.index(current_sheet)
                        target.remove(current_sheet)
                        target_sheet = target.create_sheet(source_sheet.title, sheet_index)
                        destination_title = source_sheet.title
                        copied_canonical_titles.add(source_sheet.title)
                    elif source_sheet.title not in target.sheetnames:
                        issue_index = (
                            target.index(target["待人工处理"])
                            if "待人工处理" in target.sheetnames
                            else len(target.sheetnames)
                        )
                        target_sheet = target.create_sheet(source_sheet.title, issue_index)
                        destination_title = source_sheet.title
                        copied_canonical_titles.add(source_sheet.title)
                    else:
                        destination_title = _archived_source_sheet_title(
                            target,
                            source_label,
                            source_sheet.title,
                        )
                        issue_index = (
                            target.index(target["待人工处理"])
                            if "待人工处理" in target.sheetnames
                            else len(target.sheetnames)
                        )
                        target_sheet = target.create_sheet(destination_title, issue_index)
                    _copy_source_sheet_content(source_sheet, target_sheet, cached_source_sheet)
                    expected_signatures[destination_title] = _sheet_content_signature(
                        source_sheet,
                        cached_source_sheet,
                    )
                    copied_sheets.append(destination_title)
            finally:
                source.close()
                cached_source.close()
        if hasattr(target, "calculation"):
            target.calculation.fullCalcOnLoad = True
            target.calculation.forceFullCalc = True
            target.calculation.calcMode = "auto"
        target.save(output_path)
    finally:
        target.close()

    verification = openpyxl.load_workbook(output_path, data_only=False)
    try:
        for sheet_title, expected in expected_signatures.items():
            if sheet_title not in verification.sheetnames:
                raise ValueError(f"来源工作表未写入输出文件：{sheet_title}")
            actual = _sheet_content_signature(verification[sheet_title])
            if actual != expected:
                raise ValueError(f"来源工作表数据完整性校验失败：{sheet_title}")
    finally:
        verification.close()
    return copied_sheets


_SEMANTIC_PURPOSE_RULES: tuple[tuple[tuple[str, ...], str], ...] = (
    (("工资核算",), "payroll_main"),
    (("工资汇总",), "salary_summary"),
    (("工资条",), "payslip"),
    (("人员异动", "入离职", "转岗", "调岗", "转正"), "personnel_changes"),
    (("考勤", "迟到", "早退", "旷工", "值班", "出勤"), "attendance"),
    (("社保", "公积金", "台账"), "social_security"),
    (("个税", "税务", "扣税"), "tax"),
    (("规则", "政策", "备忘"), "rules"),
    (("审批", "请款"), "approval"),
    (("年终奖",), "year_end_bonus"),
    (("奖金", "绩效", "提成", "佣金"), "bonus"),
    (("调差", "补发", "补扣"), "salary_adjustment"),
    (("福利", "高温", "防暑", "降温", "通讯费", "津贴", "补贴"), "allowance"),
)


def _sheet_text_tokens(sheet: Any) -> set[str]:
    tokens: set[str] = set()
    for row_number in range(1, min(sheet.max_row, 30) + 1):
        for column in range(1, min(sheet.max_column, 80) + 1):
            value = sheet.cell(row_number, column).value
            if value in (None, "") or not isinstance(value, str):
                continue
            normalized = re.sub(r"[\s_\-（）()/:：]+", "", value)
            if len(normalized) >= 2:
                tokens.add(normalized)
    return tokens


def _sheet_semantic_purpose(sheet: Any) -> str:
    text = re.sub(r"[\s_\-（）()/:：]+", "", str(sheet.title))
    header_text = "".join(_sheet_text_tokens(sheet))
    combined = f"{text}{header_text}"
    for keywords, purpose in _SEMANTIC_PURPOSE_RULES:
        if any(keyword in combined for keyword in keywords):
            return purpose
    return "other"


def _sheet_semantically_matches(source: Any, target: Any) -> bool:
    source_purpose = _sheet_semantic_purpose(source)
    target_purpose = _sheet_semantic_purpose(target)
    if source_purpose != "other" or target_purpose != "other":
        return source_purpose == target_purpose
    source_tokens = _sheet_text_tokens(source)
    target_tokens = _sheet_text_tokens(target)
    shared = source_tokens & target_tokens
    return len(shared) >= 2 and len(shared) / max(1, len(source_tokens)) >= 0.6


def _find_semantic_target(source: Any, targets: list[Any]) -> Any | None:
    exact = [
        target for target in targets
        if _sheet_semantic_purpose(source) != "other"
        and _sheet_semantic_purpose(source) == _sheet_semantic_purpose(target)
    ]
    if exact:
        return exact[0]
    if _sheet_semantic_purpose(source) in {
        "personnel_changes",
        "attendance",
        "social_security",
        "tax",
        "year_end_bonus",
        "bonus",
        "salary_adjustment",
        "allowance",
    }:
        return next(
            (target for target in targets if _sheet_semantic_purpose(target) == "payroll_main"),
            None,
        )
    return next((target for target in targets if _sheet_semantically_matches(source, target)), None)


def _copy_novel_source_sheets(
    output_path: str,
    source_paths: list[str],
    source_names: dict[str, str] | None = None,
) -> dict[str, list[Any]]:
    """Copy only source sheets with no semantic equivalent in the master.

    Recognized business sheets are already parsed into canonical entities and
    therefore update the existing master sheets. Unknown sheets are preserved
    as genuinely new business material so no information is silently dropped.
    """
    target = openpyxl.load_workbook(output_path, data_only=False)
    copied_sheets: list[str] = []
    mapped_sheets: list[dict[str, str]] = []
    try:
        for source_path in source_paths:
            source = openpyxl.load_workbook(source_path, data_only=False)
            cached_source = openpyxl.load_workbook(source_path, data_only=True, keep_links=False)
            try:
                for source_sheet in source.worksheets:
                    matching_target = _find_semantic_target(source_sheet, list(target.worksheets))
                    if matching_target is not None:
                        mapped_sheets.append({
                            "source_sheet": source_sheet.title,
                            "target_sheet": matching_target.title,
                            "purpose": _sheet_semantic_purpose(source_sheet),
                        })
                        continue
                    destination_title = _unique_sheet_title(target, source_sheet.title)
                    target_sheet = target.create_sheet(destination_title)
                    cached_sheet = cached_source[source_sheet.title]
                    _copy_source_sheet_content(source_sheet, target_sheet, cached_sheet)
                    copied_sheets.append(destination_title)
            finally:
                source.close()
                cached_source.close()
        target.save(output_path)
    finally:
        target.close()
    return {"copied_sheets": copied_sheets, "mapped_sheets": mapped_sheets}


def _apply_semantic_source_updates(
    output_path: str,
    source_paths: list[str],
    salary_month: str | None = None,
    source_names: dict[str, str] | None = None,
    sheet_overrides: dict[tuple[str, str], str] | None = None,
) -> dict[str, Any]:
    """Apply every uniquely understood source Sheet without adding new Sheets."""
    return apply_semantic_sheet_updates(
        output_path, source_paths, output_path,
        salary_month=salary_month,
        source_names=source_names,
        sheet_overrides=sheet_overrides,
    )


_SEMANTIC_REVIEW_ISSUE_TYPES = {
    "empty_source_sheet",
    "unprofiled_source_sheet",
    "ambiguous_sheet",
    "unmatched_sheet",
    "insufficient_topic_evidence",
    "missing_business_key",
    "source_formula_conflict",
    "formula_target_conflict",
    "duplicate_source_record",
    "ambiguous_record",
    "unknown_record",
    "conflicting_value",
}


def _semantic_issue_key(issue: dict[str, Any]) -> str:
    """Identify a source-workbook task even after the output workbook is regenerated."""
    identifying_fields = {
        field: issue.get(field)
        for field in (
            "issue_type", "source_files", "source_sheets", "source_row", "source_rows",
            "source_cell", "target_sheet", "target_cell", "business_key", "candidate_values",
        )
        if issue.get(field) not in (None, "", [], {})
    }
    encoded = json.dumps(identifying_fields, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _freeze_external_workbook_formulas(
    output_path: str,
    fallback_source_paths: str | list[str] | None = None,
) -> int:
    """Replace every external-workbook formula with its cached displayed value."""
    workbook = openpyxl.load_workbook(output_path, data_only=False, keep_links=False)
    cached_workbook = openpyxl.load_workbook(output_path, data_only=True, keep_links=False)
    if isinstance(fallback_source_paths, str):
        fallback_source_paths = [fallback_source_paths]
    fallback_workbooks = [
        (
            openpyxl.load_workbook(path, data_only=False, keep_links=False),
            openpyxl.load_workbook(path, data_only=True, keep_links=False),
        )
        for path in (fallback_source_paths or [])
    ]
    frozen_count = 0
    try:
        for sheet in workbook.worksheets:
            cached_sheet = cached_workbook[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if not _is_external_workbook_formula(cell.value):
                        continue
                    cached_cell = cached_sheet[cell.coordinate]
                    if cached_cell.value is None and fallback_workbooks:
                        matching_cached_cell = None
                        for fallback_workbook, fallback_cached_workbook in fallback_workbooks:
                            if sheet.title not in fallback_workbook.sheetnames:
                                continue
                            fallback_cell = fallback_workbook[sheet.title][cell.coordinate]
                            if fallback_cell.value == cell.value:
                                matching_cached_cell = fallback_cached_workbook[sheet.title][cell.coordinate]
                                break
                        if matching_cached_cell is None:
                            raise ValueError(
                                "目标外部公式与权威来源公式不一致，拒绝猜测缓存结果："
                                f"{sheet.title}!{cell.coordinate}"
                            )
                        cached_cell = matching_cached_cell
                    cell.value = _cached_external_formula_value(cell, cached_cell)
                    frozen_count += 1
        workbook.save(output_path)
    finally:
        workbook.close()
        cached_workbook.close()
        for fallback_workbook, fallback_cached_workbook in fallback_workbooks:
            fallback_workbook.close()
            fallback_cached_workbook.close()
    return frozen_count


def _clear_builtin_layout_source_data(output_path: str) -> list[str]:
    """Remove historical input values while retaining the built-in workbook layout.

    The bundled workbook is a presentation/calculation template, not project data.
    Keeping its old input rows would allow formulas to pull values from a previous
    company when a newly uploaded source uses a different sheet name.
    """
    workbook = openpyxl.load_workbook(output_path, data_only=False)
    cleared_sheets: list[str] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.title in PRESERVED_LAYOUT_SHEETS or sheet.title == "待人工处理":
                continue
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.__class__.__name__ == "MergedCell":
                        continue
                    cell.value = None
                    cell.comment = None
                    cell.hyperlink = None
            sheet._images = []
            sheet._charts = []
            cleared_sheets.append(sheet.title)
        workbook.save(output_path)
        return cleared_sheets
    finally:
        workbook.close()


_LEDGER_SOCIAL_FIELDS = {
    "缴费公司-与劳动合同公司主体一致": "缴费公司",
    "缴费公司": "缴费公司",
    "单位养老16%": "公司养老",
    "公司养老16%": "公司养老",
    "个人养老8%": "个人养老",
    "单位失业0.5%": "公司失业",
    "公司失业0.5%": "公司失业",
    "个人失业0.5%": "个人失业",
    "单位工伤0.3%、0.4%": "公司工伤",
    "公司工伤0.3%、0.4%": "公司工伤",
    "单位医疗9.8%": "公司医疗",
    "公司医疗9.8%": "公司医疗",
    "个人医疗2%+3": "个人医疗",
    "单位住房12%": "公司公积金",
    "公司住房12%": "公司公积金",
    "个人住房12%": "个人公积金",
}


def _ledger_header_field(value: Any) -> str | None:
    normalized = str(value or "").replace("\n", "").replace(" ", "").strip()
    return _LEDGER_SOCIAL_FIELDS.get(normalized)


def _write_social_contributions_to_ledger(
    output_path: str,
    entities: dict[str, list[dict]],
) -> int:
    """Write recognized social-insurance bill amounts into the output ledger."""
    social_rows = entities.get("social_security", [])
    profiles = entities.get("employee_profile", [])
    profiles_by_person_key = {
        str(profile.get("_person_key") or "").strip(): profile
        for profile in profiles
        if str(profile.get("_person_key") or "").strip()
    }
    profiles_by_id_card = {
        str(profile.get("身份证号") or "").strip().upper(): profile
        for profile in profiles
        if str(profile.get("身份证号") or "").strip()
    }
    profiles_by_name: dict[str, list[dict[str, Any]]] = {}
    for profile in profiles:
        name = str(profile.get("姓名") or "").strip()
        if name:
            profiles_by_name.setdefault(name, []).append(profile)

    rows_by_employee_id: dict[str, list[tuple[dict[str, Any], dict[str, Any] | None]]] = {}
    rows_by_name: dict[str, list[tuple[dict[str, Any], dict[str, Any] | None]]] = {}
    for record in social_rows:
        profile = profiles_by_person_key.get(str(record.get("_person_key") or "").strip())
        if profile is None:
            profile = profiles_by_id_card.get(
                str(record.get("身份证号") or "").strip().upper()
            )
        if profile is None:
            matching_profiles = profiles_by_name.get(str(record.get("姓名") or "").strip(), [])
            if len(matching_profiles) == 1:
                profile = matching_profiles[0]
        indexed = (record, profile)
        employee_id = str((profile or {}).get("工号") or record.get("工号") or "").strip()
        name = str((profile or {}).get("姓名") or record.get("姓名") or "").strip()
        if employee_id:
            rows_by_employee_id.setdefault(employee_id.upper(), []).append(indexed)
        if name:
            rows_by_name.setdefault(name, []).append(indexed)

    workbook = openpyxl.load_workbook(output_path, data_only=False)
    try:
        if "台账" not in workbook.sheetnames:
            return 0
        ledger = workbook["台账"]
        header_row = next(
            (
                row_index
                for row_index in range(1, min(ledger.max_row, 10) + 1)
                if (
                    any(str(cell.value or "").strip() == "姓名" for cell in ledger[row_index])
                    and any(_ledger_header_field(cell.value) for cell in ledger[row_index])
                )
            ),
            None,
        )
        if header_row is None:
            return 0

        name_column = next(
            cell.column
            for cell in ledger[header_row]
            if str(cell.value or "").strip() == "姓名"
        )
        field_columns = {
            field: cell.column
            for cell in ledger[header_row]
            if (field := _ledger_header_field(cell.value))
        }
        updated_rows = 0
        for row_index in range(header_row + 1, ledger.max_row + 1):
            employee_id = str(ledger.cell(row_index, 1).value or "").strip().upper()
            name = str(ledger.cell(row_index, name_column).value or "").strip()
            matching_rows = rows_by_employee_id.get(employee_id, []) if employee_id else []
            if not matching_rows:
                matching_rows = rows_by_name.get(name, [])
            if len(matching_rows) != 1:
                continue
            social_row, profile = matching_rows[0]
            changed = False
            if profile is not None:
                profile_employee_id = profile.get("工号")
                profile_name = profile.get("姓名")
                if profile_employee_id not in (None, "") and ledger.cell(row_index, 1).value != profile_employee_id:
                    ledger.cell(row_index, 1).value = profile_employee_id
                    changed = True
                if profile_name not in (None, "") and ledger.cell(row_index, name_column).value != profile_name:
                    ledger.cell(row_index, name_column).value = profile_name
                    changed = True
            for field, column in field_columns.items():
                value = social_row.get(field)
                if value not in (None, ""):
                    cell = ledger.cell(row_index, column)
                    if cell.value != value:
                        cell.value = value
                        changed = True
            if changed:
                updated_rows += 1
        if updated_rows:
            workbook.save(output_path)
        return updated_rows
    finally:
        workbook.close()


def _replace_issue_source_filename(output_path: str, stored_name: str, original_name: str) -> None:
    """Keep the manual-review sheet understandable by replacing storage hashes."""
    workbook = openpyxl.load_workbook(output_path, data_only=False)
    try:
        if "待人工处理" not in workbook.sheetnames:
            return
        sheet = workbook["待人工处理"]
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=6)
            if isinstance(cell.value, str):
                cell.value = cell.value.replace(stored_name, original_name)
        workbook.save(output_path)
    finally:
        workbook.close()


_CURRENT_ROSTER_FACT_ENTITIES = {
    "salary_detail",
    "attendance_record",
    "social_security",
    "tax_record",
}


def _record_confirms_current_roster(entity_id: str, record: dict[str, Any]) -> bool:
    """Use recognized current-period facts, never filenames, to confirm a person."""
    if entity_id not in _CURRENT_ROSTER_FACT_ENTITIES:
        return False
    identity_and_metadata = {
        "工号",
        "姓名",
        "身份证号",
        "公司",
        "部门",
        "缴费公司",
        "_source_file",
        "_source_sheet",
    }
    return any(
        key not in identity_and_metadata
        and not key.startswith("_")
        and value not in (None, "")
        for key, value in record.items()
    )


def _normalize_salary_month(year: str, month: str) -> str:
    return f"{int(year):04d}.{int(month):02d}"


def _extract_master_salary_month(filename: str) -> str | None:
    """Extract only an explicitly labelled accounting month from a master filename."""
    text = str(filename or "")
    match = re.search(
        r"所属月\s*[:：]?\s*(20\d{2})\s*[年./_-]?\s*(1[0-2]|0?[1-9])(?:\s*月)?",
        text,
    )
    if not match:
        return None
    return _normalize_salary_month(match.group(1), match.group(2))


def _resolve_effective_salary_month(project_month: str, master_filename: str) -> str:
    """Prefer the master's explicit accounting month without blocking mismatched project metadata."""
    master_month = _extract_master_salary_month(master_filename)
    return master_month or project_month


def _update_duty_roster_counts(output_path: str, source_paths: list[str]) -> int:
    """Replace duty-day inputs from explicit duty roster sheets while preserving formulas."""
    duty_counts: dict[str, int] = {}
    for source_path in source_paths:
        workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if "值班" not in sheet.title:
                    continue
                header_row = None
                name_column = None
                for row_number in range(1, min(sheet.max_row, 30) + 1):
                    for column in range(1, sheet.max_column + 1):
                        value = str(sheet.cell(row=row_number, column=column).value or "").strip()
                        if value in {"值班人员", "值班人员姓名"}:
                            header_row = row_number
                            name_column = column
                            break
                    if header_row is not None:
                        break
                if header_row is None or name_column is None:
                    continue
                for row_number in range(header_row + 1, sheet.max_row + 1):
                    name = str(sheet.cell(row=row_number, column=name_column).value or "").strip()
                    if not name or name in {"合计", "总计"}:
                        continue
                    duty_counts[name] = duty_counts.get(name, 0) + 1
        finally:
            workbook.close()

    if not duty_counts:
        return 0

    workbook = openpyxl.load_workbook(output_path, data_only=False)
    try:
        if "配送员值班费" not in workbook.sheetnames:
            return 0
        sheet = workbook["配送员值班费"]
        matched_people = 0
        for row_number in range(2, sheet.max_row + 1):
            name = str(sheet.cell(row=row_number, column=2).value or "").strip()
            if not name or name in {"合计", "总计"}:
                continue
            cell = sheet.cell(row=row_number, column=4)
            new_value = duty_counts.get(name)
            if name in duty_counts:
                matched_people += 1
            if cell.value != new_value:
                cell.value = new_value
        if matched_people:
            workbook.save(output_path)
        return matched_people
    finally:
        workbook.close()


def _standard_manual_follow_up_issues(master_filename: str) -> list[dict[str, Any]]:
    """Make the non-automated boundary explicit instead of silently omitting it."""
    source_files = [master_filename] if master_filename else []
    return [
        {
            "issue_type": "manual_formula_review",
            "person_name": "全表",
            "target_field": "公式及当期参数",
            "message": "系统保留了总表原有公式，未执行本月临时公式修改；请按当月操作手册人工修改并复核。",
            "source_files": source_files,
            "source_sheets": [],
        },
        {
            "issue_type": "manual_special_adjustment",
            "person_name": "全表",
            "target_field": "其他累计调差",
            "message": "其他累计调差及特殊补发补扣未自动计算，请人工清空、重填并复核。",
            "source_files": source_files,
            "source_sheets": ["其他调差累计"],
        },
        {
            "issue_type": "manual_person_exception",
            "person_name": "全表",
            "target_field": "新人、入离职及指定人员特殊事项",
            "message": "新人特殊公式、入离职金额和指定人员例外不自动处理，请逐项人工确认。",
            "source_files": source_files,
            "source_sheets": ["人员异动"],
        },
    ]


def _assign_manual_issue_ids(issues: list[dict[str, Any]]) -> None:
    seen: set[str] = set()
    for index, issue in enumerate(issues):
        base = str(issue.get("issue_id") or f"issue-{index}").strip() or f"issue-{index}"
        issue_id = base
        if issue_id in seen:
            issue_id = f"{base}-{index}"
            suffix = 2
            while issue_id in seen:
                issue_id = f"{base}-{index}-{suffix}"
                suffix += 1
        issue["issue_id"] = issue_id
        seen.add(issue_id)
        issue.setdefault("status", "pending")


_LEGACY_DEPARTMENT_CHANGE = re.compile(
    r"部门从总表值“(?P<current>[^”]*)”变更为“(?P<proposed>[^”]*)”"
)


def _hydrate_legacy_department_transfer_issues(final_data: dict[str, Any]) -> None:
    """Restore selectable fields omitted by exports created before manual confirmation."""
    profiles_by_name: dict[str, list[dict[str, Any]]] = {}
    for profile in final_data.get("entities", {}).get("employee_profile", []):
        name = str(profile.get("姓名") or "").strip()
        if name:
            profiles_by_name.setdefault(name, []).append(profile)

    for issue in final_data.get("issues", []):
        if issue.get("issue_type") != "department_transfer":
            continue
        name = str(issue.get("person_name") or "").strip()
        matches = profiles_by_name.get(name, [])
        if not issue.get("employee_id") and len(matches) == 1:
            employee_id = str(matches[0].get("工号") or "").strip()
            if employee_id:
                issue["employee_id"] = employee_id

        change = _LEGACY_DEPARTMENT_CHANGE.search(str(issue.get("message") or ""))
        if not change:
            continue
        issue.setdefault("current_value", change.group("current"))
        proposed = change.group("proposed").strip()
        if not issue.get("proposed_value") and proposed and "、" not in proposed:
            issue["proposed_value"] = proposed


def _confirm_department_transfers(
    final_data: dict[str, Any],
    issue_ids: list[str],
) -> int:
    """Apply only explicitly confirmed, unambiguous department transfers."""
    selected_ids = {issue_id.strip() for issue_id in issue_ids if issue_id.strip()}
    if not selected_ids:
        raise ValueError("请至少选择一条部门变更")

    _hydrate_legacy_department_transfer_issues(final_data)
    issues = final_data.get("issues", [])
    _assign_manual_issue_ids(issues)
    selected_issues = [issue for issue in issues if issue.get("issue_id") in selected_ids]
    if len(selected_issues) != len(selected_ids):
        raise ValueError("存在无效的待处理事项")

    profiles = final_data.get("entities", {}).get("employee_profile", [])
    updated = 0
    for issue in selected_issues:
        if issue.get("issue_type") != "department_transfer":
            raise ValueError("仅允许确认部门变更，其他待人工处理事项必须保持待处理")
        employee_id = str(issue.get("employee_id") or "").strip()
        new_department = issue.get("proposed_value")
        if not employee_id or new_department in (None, ""):
            raise ValueError("部门变更缺少工号或明确的新部门，不能自动更新")
        matches = [
            profile for profile in profiles
            if str(profile.get("工号") or "").strip() == employee_id
        ]
        if len(matches) != 1:
            raise ValueError(f"工号 {employee_id} 未能唯一匹配总表人员，不能自动更新")
        matches[0]["部门"] = new_department
        issue["status"] = "confirmed"
        updated += 1
    return updated


_GUIDED_RESOLUTION_FIELDS = {"公司", "部门", "岗位", "工作地点", "本月状态", "离职日期"}
_MANUAL_REVIEW_FORBIDDEN_FIELDS = {"工号", "姓名", "身份证号", "_person_key"}
_MANUAL_REVIEW_REQUIRED_HEADERS = {"事项编号", "处理结果", "处理值", "处理备注"}


def _resolve_manual_issue(
    final_data: dict[str, Any],
    issue_id: str,
    action: str,
) -> str:
    """Apply a guided choice for a uniquely actionable manual issue."""
    _hydrate_legacy_department_transfer_issues(final_data)
    issues = final_data.get("issues", [])
    _assign_manual_issue_ids(issues)
    issue = next((item for item in issues if item.get("issue_id") == issue_id), None)
    if issue is None:
        raise ValueError("待处理事项不存在")
    if issue.get("status") == "confirmed":
        return str(issue.get("proposed_value") or issue.get("current_value") or "")
    if action not in {"apply_proposed", "keep_current"}:
        raise ValueError("当前事项不支持该处理方式")
    if action == "keep_current":
        issue["status"] = "confirmed"
        issue["resolution"] = action
        return str(issue.get("current_value") or "")

    field = str(issue.get("target_field") or "").strip()
    proposed_value = issue.get("proposed_value")
    employee_id = str(issue.get("employee_id") or "").strip()
    if field not in _GUIDED_RESOLUTION_FIELDS:
        raise ValueError("该事项没有可安全自动写入的标准字段，请复制修改内容处理")
    if not employee_id or proposed_value in (None, ""):
        raise ValueError("该事项缺少唯一工号或推荐值，请人工核对")
    profiles = [
        profile for profile in final_data.get("entities", {}).get("employee_profile", [])
        if str(profile.get("工号") or "").strip() == employee_id
    ]
    if len(profiles) != 1:
        raise ValueError("无法唯一匹配员工，不能自动写入")
    profiles[0][field] = proposed_value
    issue["status"] = "confirmed"
    issue["resolution"] = action
    return str(proposed_value)


def _manual_review_headers(sheet: Any) -> dict[str, int]:
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value not in (None, "")
    }
    missing = _MANUAL_REVIEW_REQUIRED_HEADERS - set(headers)
    if missing:
        raise ValueError(f"人工处理表缺少必要列：{'、'.join(sorted(missing))}")
    return headers


def _manual_review_value(value: Any, issue: dict[str, Any]) -> Any:
    if value in (None, ""):
        raise ValueError(f"事项 {issue.get('issue_id')} 选择更新到总表时必须填写处理值")
    reference = issue.get("proposed_value", issue.get("current_value"))
    if isinstance(reference, bool):
        if isinstance(value, str) and value.strip() in {"是", "true", "True", "1"}:
            return True
        if isinstance(value, str) and value.strip() in {"否", "false", "False", "0"}:
            return False
    if isinstance(reference, int) and not isinstance(reference, bool) and isinstance(value, str):
        return int(float(value.replace(",", "").strip()))
    if isinstance(reference, float) and isinstance(value, str):
        return float(value.replace(",", "").strip())
    return value.strip() if isinstance(value, str) else value


def _manual_review_target_record(final_data: dict[str, Any], issue: dict[str, Any]) -> dict[str, Any]:
    field = str(issue.get("target_field") or "").strip()
    employee_id = str(issue.get("employee_id") or "").strip()
    person_name = str(issue.get("person_name") or "").strip()
    if not employee_id and (not person_name or person_name == "全表"):
        raise ValueError(f"事项 {issue.get('issue_id')} 缺少可唯一定位的人员，不能回写总表")
    if not field or field in _MANUAL_REVIEW_FORBIDDEN_FIELDS:
        raise ValueError(f"事项 {issue.get('issue_id')} 的目标字段不能安全回写")

    matches: list[dict[str, Any]] = []
    for entity_name, records in final_data.get("entities", {}).items():
        if not isinstance(records, list):
            continue
        for record in records:
            record_employee_id = str(record.get("工号") or "").strip()
            record_name = str(record.get("姓名") or "").strip()
            if employee_id and record_employee_id != employee_id:
                continue
            if not employee_id and record_name != person_name:
                continue
            if field in record or (entity_name == "employee_profile" and field in _GUIDED_RESOLUTION_FIELDS):
                matches.append(record)
    if len(matches) != 1:
        raise ValueError(
            f"事项 {issue.get('issue_id')} 无法唯一定位“{field}”的总表数据，不能自动回写"
        )
    return matches[0]


def _can_update_manual_issue_online(final_data: dict[str, Any], issue: dict[str, Any]) -> bool:
    """Expose only items the same write path can uniquely and safely update."""
    try:
        _manual_review_target_record(final_data, issue)
    except ValueError:
        return False
    return True


def _apply_manual_review_decisions(
    final_data: dict[str, Any],
    decisions: list[dict[str, Any]],
    resolution: str = "online_review",
) -> int:
    """Validate every decision before applying the batch to the final data."""
    _assign_manual_issue_ids(final_data.get("issues", []))
    issues_by_id = {
        str(issue.get("issue_id")): issue
        for issue in final_data.get("issues", [])
    }
    prepared: list[tuple[dict[str, Any], str, dict[str, Any] | None, Any, str]] = []
    seen_issue_ids: set[str] = set()

    for decision in decisions:
        issue_id = str(decision.get("issue_id") or "").strip()
        outcome = str(decision.get("outcome") or "待处理").strip()
        note = str(decision.get("note") or "").strip()
        if not issue_id:
            raise ValueError("人工处理事项缺少编号")
        if issue_id in seen_issue_ids:
            raise ValueError(f"人工处理事项重复填写：{issue_id}")
        seen_issue_ids.add(issue_id)
        if outcome not in MANUAL_REVIEW_OUTCOMES:
            raise ValueError(f"事项 {issue_id} 的处理结果无效")
        issue = issues_by_id.get(issue_id)
        if issue is None:
            raise ValueError(f"人工处理事项包含当前批次不存在的事项：{issue_id}")
        if issue.get("status") == "confirmed":
            continue

        target_record: dict[str, Any] | None = None
        value: Any = None
        if outcome == "更新到总表":
            target_record = _manual_review_target_record(final_data, issue)
            value = _manual_review_value(decision.get("value"), issue)
        prepared.append((issue, outcome, target_record, value, note))

    for issue, outcome, target_record, value, note in prepared:
        if outcome == "更新到总表":
            assert target_record is not None
            target_record[str(issue["target_field"]).strip()] = value
            issue["resolution"] = resolution
        else:
            issue["resolution"] = "keep_current"
        issue["status"] = "confirmed"
        issue["resolution_note"] = note
    return len(prepared)


def _apply_manual_review_workbook(final_data: dict[str, Any], content: bytes) -> int:
    """Apply only explicit, safely addressable workbook decisions to the final data set."""
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("无法读取人工处理表，请上传系统下载后填写的 .xlsx 文件") from exc
    try:
        if MANUAL_ISSUES_SHEET not in workbook.sheetnames:
            raise ValueError("人工处理表中未找到“待人工处理”工作表")
        sheet = workbook[MANUAL_ISSUES_SHEET]
        headers = _manual_review_headers(sheet)
        issue_rows: dict[str, tuple[str, Any, str]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            issue_id = str(row[headers["事项编号"] - 1] or "").strip()
            outcome = str(row[headers["处理结果"] - 1] or "待处理").strip()
            if not issue_id or outcome == "待处理":
                continue
            if outcome not in MANUAL_REVIEW_OUTCOMES:
                raise ValueError(f"事项 {issue_id} 的处理结果无效")
            if issue_id in issue_rows:
                raise ValueError(f"人工处理表中重复填写了事项 {issue_id}")
            issue_rows[issue_id] = (
                outcome,
                row[headers["处理值"] - 1],
                str(row[headers["处理备注"] - 1] or "").strip(),
            )
    finally:
        workbook.close()

    return _apply_manual_review_decisions(
        final_data,
        [
            {
                "issue_id": issue_id,
                "outcome": outcome,
                "value": value,
                "note": note,
            }
            for issue_id, (outcome, value, note) in issue_rows.items()
        ],
        resolution="imported_workbook",
    )


def _record_identity_tokens(record: dict[str, Any]) -> set[tuple[str, str]]:
    tokens: set[tuple[str, str]] = set()
    for field in ("工号", "身份证号", "姓名"):
        value = str(record.get(field) or "").strip()
        if value:
            tokens.add((field, value.upper() if field != "姓名" else value))
    return tokens


def _find_template_files(project_id: str, db: Session) -> list[tuple[UploadFile, str]]:
    """返回项目中所有仍存在的模板，按上传时间从早到晚排列。"""
    files = db.query(UploadFile).filter(
        UploadFile.project_id == project_id,
        UploadFile.file_type == "template",
    ).order_by(UploadFile.created_at.asc()).all()
    return [
        (file, path)
        for file in files
        if os.path.exists(path := os.path.join(UPLOAD_DIR, file.stored_path))
    ]


def _find_template_file(project_id: str, db: Session) -> str | None:
    """从项目全部文件中自动选择最适合承载最终输出的工作簿。"""
    selected = _select_template_file(project_id, db)
    return selected[1] if selected else None


def _choose_template_candidate(
    files: list[UploadFile],
    compatible_file_ids: set[str] | None = None,
) -> UploadFile | None:
    """选择总表底板；旧分类结果只作弱参考，文件名和结构优先。"""
    if compatible_file_ids is not None:
        files = [file for file in files if str(file.id) in compatible_file_ids]

    def score(file: UploadFile) -> tuple[int, int, int, str]:
        name = os.path.splitext(str(file.original_name))[0].replace(" ", "").lower()
        keyword_score = 0
        if "主模板" in name or "导出模板" in name:
            keyword_score = 500
        elif "总表" in name:
            keyword_score = 400
        elif "工资核算模板" in name:
            keyword_score = 350
        elif str(getattr(file, "file_type", "")) == "template":
            keyword_score = 10
        return (
            keyword_score,
            int(getattr(file, "col_count", 0) or 0),
            int(getattr(file, "row_count", 0) or 0),
            str(getattr(file, "created_at", "") or ""),
        )

    return max(files, key=score) if files else None


def _supports_unified_export(file_path: str) -> bool:
    """A unified export template must contain the sheet written by the exporter."""
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
    except Exception:
        return False
    try:
        return "工资核算" in workbook.sheetnames
    finally:
        workbook.close()


def _resolve_layout_template_path(uploaded_template_path: str) -> str:
    """The explicitly uploaded master is both the data baseline and output layout."""
    return uploaded_template_path


def _select_template_file(project_id: str, db: Session) -> tuple[UploadFile | None, str] | None:
    files = db.query(UploadFile).filter(
        UploadFile.project_id == project_id,
    ).order_by(UploadFile.created_at.asc()).all()
    existing = [
        (file, path)
        for file in files
        if os.path.exists(path := os.path.join(UPLOAD_DIR, file.stored_path))
    ]
    explicit_masters = [
        (file, path)
        for file, path in existing
        if str(getattr(file, "file_type", "")) in _PAYROLL_MASTER_FILE_TYPES
        and _supports_unified_export(path)
    ]
    return explicit_masters[0] if len(explicit_masters) == 1 else None


def _resolve_export_template(
    project_id: str,
    template_file_id: str | None,
    db: Session,
) -> tuple[UploadFile, str] | None:
    if template_file_id:
        file = db.query(UploadFile).filter(
            UploadFile.project_id == project_id,
            UploadFile.id == template_file_id,
        ).first()
        if file:
            path = os.path.join(UPLOAD_DIR, file.stored_path)
            if os.path.exists(path) and _supports_unified_export(path):
                return file, path
    return _select_template_file(project_id, db)


def _session_path(project_id: str, suffix: str) -> str:
    return os.path.join(SESSION_DIR, f"{project_id}_{suffix}.json")


def _manual_sheet_overrides(final_data: dict[str, Any]) -> dict[tuple[str, str], str]:
    """Return the user's explicit source-Sheet-to-master-Sheet choices."""
    overrides: dict[tuple[str, str], str] = {}
    for item in final_data.get("sheet_mappings", []):
        if not isinstance(item, dict):
            continue
        source_file = str(item.get("source_file") or "").strip()
        source_sheet = str(item.get("source_sheet") or "").strip()
        target_sheet = str(item.get("target_sheet") or "").strip()
        if source_file and source_sheet and target_sheet:
            overrides[(source_file, source_sheet)] = target_sheet
    return overrides


def _workbook_sheet_names(workbook_path: str) -> list[str]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


_INTEGRATION_PROGRESS_STAGES = (
    ("base", "加载原始数据库"),
    ("source", "解析来源文件"),
    ("match", "构建完整人员并集"),
    ("export", "生成可下载 Excel"),
)


def _default_integration_progress_stages() -> list[dict[str, str]]:
    return [
        {"key": key, "label": label, "status": "pending"}
        for key, label in _INTEGRATION_PROGRESS_STAGES
    ]


def _save_integration_progress(
    project_id: str,
    *,
    status: str,
    stages: list[dict[str, Any]],
    detail: str | None = None,
) -> dict[str, Any]:
    """Persist live stage state without putting it in export-result metadata."""
    progress = {
        "status": status,
        "stages": stages,
        "detail": detail,
        "updated_at": datetime.now().isoformat(),
    }
    _save_json(project_id, "integration_progress", progress)
    return progress


def _load_integration_progress(project_id: str) -> dict[str, Any]:
    return _load_json(project_id, "integration_progress") or {
        "status": "idle",
        "stages": _default_integration_progress_stages(),
        "detail": None,
        "updated_at": "",
    }


def _project_source_signature(files: list[Any]) -> str:
    """Build a stable fingerprint for every workbook that affects an export."""
    records = []
    for file in files:
        file_type = str(getattr(file, "file_type", ""))
        if file_type not in _PAYROLL_MASTER_FILE_TYPES | _PAYROLL_SOURCE_FILE_TYPES:
            continue
        created_at = getattr(file, "created_at", None)
        records.append({
            "id": str(getattr(file, "id", "")),
            "original_name": str(getattr(file, "original_name", "")),
            "stored_path": str(getattr(file, "stored_path", "")),
            "file_type": str(getattr(file, "file_type", "")),
            "row_count": int(getattr(file, "row_count", 0) or 0),
            "col_count": int(getattr(file, "col_count", 0) or 0),
            "created_at": created_at.isoformat() if hasattr(created_at, "isoformat") else str(created_at or ""),
        })
    payload = json.dumps(
        sorted(records, key=lambda item: item["id"]),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _current_source_signature(project_id: str, db: Session) -> tuple[str, int]:
    files = db.query(UploadFile).filter(
        UploadFile.project_id == project_id,
        UploadFile.file_type.in_([*_PAYROLL_MASTER_FILE_TYPES, *_PAYROLL_SOURCE_FILE_TYPES]),
    ).all()
    return _project_source_signature(files), len(files)


def _load_current_export_meta(project_id: str, db: Session) -> dict[str, Any]:
    """Load export metadata only when it still matches the uploaded workbooks."""
    meta_path = _session_path(project_id, "export_meta")
    if not os.path.exists(meta_path):
        raise HTTPException(status_code=400, detail="尚未生成导出文件")
    with open(meta_path, "r", encoding="utf-8") as fp:
        meta = json.load(fp)
    if meta.get("status") in {"processing", "blocked", "stale"} or not meta.get("filename"):
        raise HTTPException(status_code=400, detail="当前批次尚未生成可用结果")
    current_signature, _ = _current_source_signature(project_id, db)
    if not meta.get("source_signature") or meta["source_signature"] != current_signature:
        raise HTTPException(status_code=409, detail="来源文件已变化，请重新生成工资总表")
    return meta


def _ensure_manual_issues_export(project_id: str, meta: dict[str, Any]) -> tuple[str, str]:
    """Return a fresh standalone workbook containing only still-pending issues."""
    filename = meta.get("issues_filename")
    relative_path = meta.get("issues_path")
    if not filename or not relative_path:
        export_stem = os.path.splitext(str(meta["filename"]))[0]
        filename = f"待人工处理_{export_stem}.xlsx"
        relative_path = os.path.join(project_id, filename)
        meta["issues_filename"] = filename
        meta["issues_path"] = relative_path
    absolute_path = os.path.join(EXPORT_DIR, relative_path)
    _assign_manual_issue_ids(meta.get("issues", []))
    pending_issues = [
        issue for issue in meta.get("issues", [])
        if issue.get("status") != "confirmed"
    ]
    os.makedirs(os.path.dirname(absolute_path), exist_ok=True)
    write_manual_issues_workbook(absolute_path, pending_issues)
    meta["issue_count"] = len(pending_issues)
    with open(_session_path(project_id, "export_meta"), "w", encoding="utf-8") as fp:
        json.dump(meta, fp, ensure_ascii=False)
    return filename, absolute_path


_REVIEW_EXPORT_FORMAT_VERSION = 4


def _ensure_review_export(
    project_id: str,
    meta: dict[str, Any],
    db: Session,
) -> tuple[str, str]:
    """Return a review copy whose values match the formal workbook exactly."""
    filename = meta.get("review_filename")
    relative_path = meta.get("review_path")
    if not filename or not relative_path:
        raise HTTPException(status_code=404, detail="当前批次没有修改稿")
    absolute_path = os.path.join(EXPORT_DIR, relative_path)
    formal_path = os.path.join(EXPORT_DIR, str(meta.get("path", "")))
    needs_rebuild = (
        not os.path.exists(absolute_path)
        or meta.get("review_format_version") != _REVIEW_EXPORT_FORMAT_VERSION
    )
    if needs_rebuild:
        if not os.path.exists(formal_path):
            raise HTTPException(status_code=404, detail="正式稿文件已丢失，无法重建修改稿")
        base_path = meta.get("review_base_path")
        if not base_path or not os.path.exists(base_path):
            selected_template = _select_template_file(project_id, db)
            if not selected_template:
                raise HTTPException(status_code=404, detail="总表底板文件已丢失，无法重建修改稿")
            _, base_path = selected_template
        os.makedirs(os.path.dirname(absolute_path), exist_ok=True)
        create_review_workbook(base_path, formal_path, absolute_path)
        meta["review_base_path"] = base_path
        meta["review_format_version"] = _REVIEW_EXPORT_FORMAT_VERSION
        with open(_session_path(project_id, "export_meta"), "w", encoding="utf-8") as fp:
            json.dump(meta, fp, ensure_ascii=False)
    return filename, absolute_path


def _save_json(project_id: str, suffix: str, data: dict):
    with open(_session_path(project_id, suffix), "w", encoding="utf-8") as fp:
        json.dump(data, fp, ensure_ascii=False, default=str)


def _load_json(project_id: str, suffix: str) -> dict | None:
    p = _session_path(project_id, suffix)
    if not os.path.exists(p):
        return None
    with open(p, "r", encoding="utf-8") as fp:
        return json.load(fp)


def _validate_final_entities(entities: dict[str, list[dict]]) -> dict[str, Any]:
    """将最终实体拼成员工校验视图，返回可供 API 和前端使用的阻断结果。"""
    salary_by_id = {
        str(row.get("工号", "")).strip(): row
        for row in entities.get("salary_detail", [])
        if str(row.get("工号", "")).strip()
    }
    audit_rows: list[dict[str, Any]] = []
    for profile in entities.get("employee_profile", []):
        employee_id = str(profile.get("工号", "")).strip()
        salary = salary_by_id.get(employee_id, {})
        audit_rows.append({
            "工号": employee_id,
            "姓名": profile.get("姓名", ""),
            "身份证号": profile.get("身份证号", ""),
            "基本工资": salary.get("月基本薪资", profile.get("月基本薪资", "")),
            "转正日期": profile.get("转正日期", ""),
            "试用期状态": profile.get("是否试用期", profile.get("试用期状态", "")),
            "车贴标准": salary.get("交通通讯补贴标准", ""),
        })

    frame = pd.DataFrame(
        audit_rows,
        columns=["工号", "姓名", "身份证号", "基本工资", "转正日期", "试用期状态", "车贴标准"],
    )
    validator = DataValidator()
    audited = validator.audit_employee_profiles(frame)
    metrics = validator.get_completeness_metrics(audited)

    blockers: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for row_index, row in audited.iterrows():
        item = {
            "row_index": int(row_index),
            "employee_id": str(row.get("工号", "")),
            "employee_name": str(row.get("姓名", "")),
            "status": str(row.get("诊断状态", "")),
            "missing_fields": list(row.get("缺失字段清单", [])),
        }
        if item["status"] in {
            DataValidator.STATUS_MISSING_ID_CARD,
            DataValidator.STATUS_MISSING_BASE_SALARY,
            DataValidator.STATUS_MISSING_OTHER,
        }:
            blockers.append(item)
        elif item["status"] in {
            DataValidator.STATUS_MISSING_CONFIRM_DATE,
            DataValidator.STATUS_WARNING_ONLY,
        }:
            warnings.append(item)

    return {"metrics": metrics, "blockers": blockers, "warnings": warnings}


def _ensure_exportable(entities: dict[str, list[dict]]) -> dict[str, Any]:
    """校验最终库并在存在阻断项时拒绝导出。"""
    validation = _validate_final_entities(entities)
    blocker_count = int(validation["metrics"].get("high_risk_count", 0))
    if blocker_count > 0:
        raise HTTPException(
            status_code=400,
            detail=f"最终数据库存在 {blocker_count} 条阻断项，请补齐工号、姓名、身份证号和基本工资后再导出",
        )
    return validation


def _reuse_cached_base_db(existing: dict[str, Any] | None, refresh: bool) -> bool:
    """仅在未请求刷新时复用旧的原始库会话。"""
    return bool(existing) and not refresh


def _content_disposition(filename: str) -> str:
    """生成兼容中文文件名且可由 Starlette 写入的下载响应头。"""
    encoded = quote(filename, safe="._-")
    return f'attachment; filename="payroll-result.xlsx"; filename*=UTF-8\'\'{encoded}'


def _confirm_all_pending_diff(diff: dict[str, Any]) -> int:
    """确定性确认全部待处理差异，并保留已忽略或已确认状态。"""
    confirmed_now = 0
    summary = diff.setdefault("summary", {})
    for diff_type, raw_value in diff.items():
        if diff_type == "summary":
            continue
        if isinstance(raw_value, list):
            items = raw_value
        elif isinstance(raw_value, dict) and isinstance(raw_value.get("items"), list):
            items = raw_value["items"]
        else:
            continue

        for item in items:
            if str(item.get("status", "pending")) == "pending":
                item["status"] = "confirmed"
                confirmed_now += 1
        summary[f"{diff_type}_confirmed"] = sum(
            1 for item in items if item.get("status") == "confirmed"
        )
    return confirmed_now


# ============================================================
# 模块A：原始数据库
# ============================================================
class BaseDbResponse(BaseModel):
    entities: dict[str, list[dict]]
    overview: list[dict]
    source: str  # "template" / "snapshot_clone"
    created_at: str


@router.post("/{project_id}/base-db/load", response_model=BaseDbResponse)
def load_base_db(project_id: str, refresh: bool = False,
                 user: User = Depends(get_current_user),
                 db: Session = Depends(get_db)):
    """模块A：加载原始数据库。

    优先级：已有base_db session > 克隆最近快照 > 从模板提取。
    """
    p = _load_project_or_404(project_id, user, db)

    # 1. 检查是否已有base_db session
    existing = _load_json(p.id, "base_db")
    if _reuse_cached_base_db(existing, refresh):
        return BaseDbResponse(
            entities=existing.get("entities", {}),
            overview=get_entity_overview(existing.get("entities", {})),
            source=existing.get("source", "unknown"),
            created_at=existing.get("created_at", ""),
        )

    # 2. 尝试克隆最近快照
    latest_snap = db.query(EntitySnapshot).filter(
        EntitySnapshot.project_id == p.id
    ).order_by(EntitySnapshot.salary_month.desc()).first()

    if latest_snap and not refresh:
        snap_data = json.loads(latest_snap.data_json)
        entities = snap_data.get("entities", {})
        # 清空月度动态字段
        for eid in ["attendance_record", "tax_record"]:
            if eid in entities:
                entities[eid] = []
        for rec in entities.get("employee_profile", []):
            rec["本月状态"] = "在职"
            rec.pop("离职日期", None)
        data = {
            "entities": entities,
            "source": f"snapshot_clone:{latest_snap.salary_month}",
            "created_at": datetime.now().isoformat(),
        }
        _save_json(p.id, "base_db", data)
        return BaseDbResponse(
            entities=entities,
            overview=get_entity_overview(entities),
            source=data["source"],
            created_at=data["created_at"],
        )

    # 3. 从全部模板提取并按实体合并
    templates = _find_template_files(p.id, db)
    if not templates:
        raise HTTPException(status_code=400, detail="请先上传导出模板（首次使用）或保存历史快照")

    entities: dict[str, list[dict]] = {}
    tpl_logs: list[str] = []
    for template, template_path in templates:
        tpl_logs.append(f"处理模板：{template.original_name}")
        tpl_result = import_template_to_entities(template_path)
        tpl_logs.extend(tpl_result.pop("_log", []))
        for entity_id, records in tpl_result.items():
            entities.setdefault(entity_id, []).extend(records)

    source = "templates" if len(templates) > 1 else "template"
    data = {
        "entities": entities,
        "source": source,
        "created_at": datetime.now().isoformat(),
        "logs": tpl_logs,
    }
    _save_json(p.id, "base_db", data)
    return BaseDbResponse(
        entities=entities,
        overview=get_entity_overview(entities),
        source=source,
        created_at=data["created_at"],
    )


@router.get("/{project_id}/base-db", response_model=BaseDbResponse)
def get_base_db(project_id: str, user: User = Depends(get_current_user),
                db: Session = Depends(get_db)):
    """模块A：获取已生成的原始数据库，用于恢复页面阶段状态。"""
    p = _load_project_or_404(project_id, user, db)
    data = _load_json(p.id, "base_db")
    if not data:
        raise HTTPException(status_code=400, detail="尚未加载原始数据库")
    entities = data.get("entities", {})
    return BaseDbResponse(
        entities=entities,
        overview=get_entity_overview(entities),
        source=data.get("source", "unknown"),
        created_at=data.get("created_at", ""),
    )


@router.put("/{project_id}/base-db", response_model=BaseDbResponse)
def update_base_db(project_id: str, payload: dict,
                   user: User = Depends(get_current_user),
                   db: Session = Depends(get_db)):
    """模块A：手动编辑原始数据库。"""
    p = _load_project_or_404(project_id, user, db)
    data = {
        "entities": payload,
        "source": "manual_edit",
        "created_at": datetime.now().isoformat(),
    }
    _save_json(p.id, "base_db", data)
    return BaseDbResponse(
        entities=payload,
        overview=get_entity_overview(payload),
        source="manual_edit",
        created_at=data["created_at"],
    )


# ============================================================
# 模块B：数据源头表
# ============================================================
class SourceDataResponse(BaseModel):
    entities: dict[str, list[dict]]
    overview: list[dict]
    logs: list[str]
    created_at: str


@router.post("/{project_id}/source-data/extract", response_model=SourceDataResponse)
def extract_source_data(project_id: str, user: User = Depends(get_current_user),
                        db: Session = Depends(get_db)):
    """模块B：从已上传的源文件提取数据源头表。

    将所有非模板文件解析为实体数据，不做合并去重（保留原始结构）。
    """
    p = _load_project_or_404(project_id, user, db)
    files = db.query(UploadFile).filter(
        UploadFile.project_id == p.id,
        UploadFile.file_type != "template",
    ).all()

    if not files:
        raise HTTPException(status_code=400, detail="请先上传源数据文件")

    all_logs: list[str] = []
    collected: dict[str, list[dict]] = {}

    for f in files:
        fp = os.path.join(UPLOAD_DIR, f.stored_path)
        if not os.path.exists(fp):
            all_logs.append(f"文件丢失: {f.original_name}")
            continue

        all_logs.append(f"处理源文件: {f.original_name}")
        try:
            result = import_file_to_entities(fp, f.original_name)
            for eid, records in result.items():
                if eid == "_log":
                    all_logs.extend(records)
                    continue
                collected.setdefault(eid, []).extend(records)
            discovered = discover_people_in_workbook(fp, f.original_name)
            collected.setdefault("employee_profile", []).extend(discovered)
            all_logs.append(f"  人员兜底扫描: 识别 {len(discovered)} 条姓名记录")
        except Exception as e:
            all_logs.append(f"解析失败: {e}")

    data = {
        "entities": collected,
        "logs": all_logs,
        "created_at": datetime.now().isoformat(),
    }
    _save_json(p.id, "source_data", data)

    return SourceDataResponse(
        entities=collected,
        overview=get_entity_overview(collected),
        logs=all_logs,
        created_at=data["created_at"],
    )


@router.get("/{project_id}/source-data", response_model=SourceDataResponse)
def get_source_data(project_id: str, user: User = Depends(get_current_user),
                    db: Session = Depends(get_db)):
    """模块B：获取已提取的数据源头表。"""
    p = _load_project_or_404(project_id, user, db)
    data = _load_json(p.id, "source_data")
    if not data:
        raise HTTPException(status_code=400, detail="尚未提取数据源头表")
    return SourceDataResponse(
        entities=data.get("entities", {}),
        overview=get_entity_overview(data.get("entities", {})),
        logs=data.get("logs", []),
        created_at=data.get("created_at", ""),
    )


# ============================================================
# 模块C+D：匹配引擎 + 差异审核
# ============================================================
class DiffResponse(BaseModel):
    diff: dict[str, Any]
    summary: dict[str, Any]
    created_at: str


@router.post("/{project_id}/diff/compute", response_model=DiffResponse)
def compute_pipeline_diff(project_id: str, user: User = Depends(get_current_user),
                          db: Session = Depends(get_db)):
    """模块C：运行匹配引擎，计算原始数据库与数据源头表的差异。"""
    p = _load_project_or_404(project_id, user, db)

    base_data = _load_json(p.id, "base_db")
    source_data = _load_json(p.id, "source_data")

    if not base_data:
        raise HTTPException(status_code=400, detail="请先加载原始数据库（模块A）")
    if not source_data:
        raise HTTPException(status_code=400, detail="请先提取数据源头表（模块B）")

    diff = compute_diff(
        base_data.get("entities", {}),
        source_data.get("entities", {}),
    )

    result = {
        "diff": diff,
        "summary": diff.get("summary", {}),
        "created_at": datetime.now().isoformat(),
    }
    _save_json(p.id, "diff", result)
    return DiffResponse(**result)


@router.get("/{project_id}/diff", response_model=DiffResponse)
def get_diff(project_id: str, user: User = Depends(get_current_user),
             db: Session = Depends(get_db)):
    """模块D：获取当前差异。"""
    p = _load_project_or_404(project_id, user, db)
    data = _load_json(p.id, "diff")
    if not data:
        raise HTTPException(status_code=400, detail="尚未计算差异")
    return DiffResponse(**data)


class DiffUpdateIn(BaseModel):
    """更新差异项状态。"""
    diff_type: str  # hire / depart / transfer / salary_change / attendance / social / tax
    item_index: int
    status: str  # confirmed / ignored / pending
    updates: dict | None = None  # 修改差异项内容


@router.post("/{project_id}/diff/update", response_model=DiffResponse)
def update_diff_item(project_id: str, payload: DiffUpdateIn,
                     user: User = Depends(get_current_user),
                     db: Session = Depends(get_db)):
    """模块D：更新单个差异项的状态（确认/忽略/修改）。"""
    p = _load_project_or_404(project_id, user, db)
    data = _load_json(p.id, "diff")
    if not data:
        raise HTTPException(status_code=400, detail="尚未计算差异")

    diff = data["diff"]
    diff_list = diff.get(payload.diff_type, [])

    if payload.item_index < 0 or payload.item_index >= len(diff_list):
        raise HTTPException(status_code=400, detail="差异项索引无效")

    item = diff_list[payload.item_index]
    item["status"] = payload.status
    if payload.updates:
        item.update(payload.updates)

    _save_json(p.id, "diff", data)

    # 重新计算汇总
    summary = diff.get("summary", {})
    confirmed = sum(1 for v in diff_list if v.get("status") == "confirmed")
    summary[f"{payload.diff_type}_confirmed"] = confirmed

    return DiffResponse(diff=diff, summary=summary, created_at=data["created_at"])


class BatchUpdateIn(BaseModel):
    """批量更新差异项状态。"""
    diff_type: str
    indices: list[int]
    status: str


@router.post("/{project_id}/diff/batch-update", response_model=DiffResponse)
def batch_update_diff(project_id: str, payload: BatchUpdateIn,
                      user: User = Depends(get_current_user),
                      db: Session = Depends(get_db)):
    """模块D：批量确认/忽略差异项。"""
    p = _load_project_or_404(project_id, user, db)
    data = _load_json(p.id, "diff")
    if not data:
        raise HTTPException(status_code=400, detail="尚未计算差异")

    diff = data["diff"]
    diff_list = diff.get(payload.diff_type, [])

    for idx in payload.indices:
        if 0 <= idx < len(diff_list):
            diff_list[idx]["status"] = payload.status

    _save_json(p.id, "diff", data)
    return DiffResponse(diff=diff, summary=diff.get("summary", {}), created_at=data["created_at"])


# ============================================================
# 模块E：最终数据库
# ============================================================
class FinalDbResponse(BaseModel):
    entities: dict[str, list[dict]]
    overview: list[dict]
    logs: list[str]
    validation: dict[str, Any]
    created_at: str


@router.post("/{project_id}/final-db/build", response_model=FinalDbResponse)
def build_final_db(project_id: str, user: User = Depends(get_current_user),
                   db: Session = Depends(get_db)):
    """模块E：将原始数据库+已确认差异拼接为最终数据库。"""
    p = _load_project_or_404(project_id, user, db)

    base_data = _load_json(p.id, "base_db")
    diff_data = _load_json(p.id, "diff")

    if not base_data:
        raise HTTPException(status_code=400, detail="请先加载原始数据库（模块A）")
    if not diff_data:
        raise HTTPException(status_code=400, detail="请先计算差异（模块C）")

    final_db = apply_confirmed_diff(
        base_data.get("entities", {}),
        diff_data.get("diff", {}),
    )

    # 统计日志
    diff = diff_data.get("diff", {})
    summary = diff.get("summary", {})
    logs = [
        f"原始数据库: {sum(len(v) for v in base_data.get('entities', {}).values())} 条记录",
        f"确认入职: {summary.get('hire_confirmed', 0)} 人",
        f"确认离职: {summary.get('depart_confirmed', 0)} 人",
        f"确认调岗: {summary.get('transfer_confirmed', 0)} 人",
        f"确认薪资变动: {summary.get('salary_change_confirmed', 0)} 项",
        f"最终数据库: {sum(len(v) for v in final_db.values())} 条记录",
    ]

    result = {
        "entities": final_db,
        "logs": logs,
        "validation": _validate_final_entities(final_db),
        "created_at": datetime.now().isoformat(),
    }
    _save_json(p.id, "final_db", result)

    # 同时同步到entities session（兼容导出接口）
    _save_json(p.id, "entities", {
        "entities": final_db,
        "imported_at": datetime.now().isoformat(),
        "logs": logs,
    })

    return FinalDbResponse(
        entities=final_db,
        overview=get_entity_overview(final_db),
        logs=logs,
        validation=result["validation"],
        created_at=result["created_at"],
    )


@router.get("/{project_id}/final-db", response_model=FinalDbResponse)
def get_final_db(project_id: str, user: User = Depends(get_current_user),
                 db: Session = Depends(get_db)):
    """模块E：获取最终数据库。"""
    p = _load_project_or_404(project_id, user, db)
    data = _load_json(p.id, "final_db")
    if not data:
        raise HTTPException(status_code=400, detail="尚未生成最终数据库")
    return FinalDbResponse(
        entities=data.get("entities", {}),
        overview=get_entity_overview(data.get("entities", {})),
        logs=data.get("logs", []),
        validation=data.get("validation") or _validate_final_entities(data.get("entities", {})),
        created_at=data.get("created_at", ""),
    )


# ============================================================
# 模块F：模板导出（复用entities路由的导出逻辑）
# ============================================================
class ExportResult(BaseModel):
    filename: str
    log: list[str]
    issue_count: int = 0
    source_person_count: int = 0
    output_person_count: int = 0
    input_person_count: int = 0
    manual_only_person_count: int = 0
    coverage_gap_count: int = 0
    source_file_count: int = 0
    issues: list[dict[str, Any]] = Field(default_factory=list)
    validation: dict[str, Any] = Field(default_factory=dict)
    completed_at: str = ""
    issues_filename: str | None = None
    review_filename: str | None = None
    duration_seconds: float = 0
    stage_durations: dict[str, float] = Field(default_factory=dict)
    target_sheets: list[str] = Field(default_factory=list)


class ExportRenameIn(BaseModel):
    filename: str


class IntegrationResult(BaseModel):
    status: str
    filename: str | None
    stages: list[dict[str, Any]]
    overview: list[dict]
    validation: dict[str, Any]
    log: list[str]
    completed_at: str
    issue_count: int = 0
    source_person_count: int = 0
    output_person_count: int = 0
    input_person_count: int = 0
    manual_only_person_count: int = 0
    coverage_gap_count: int = 0
    source_file_count: int = 0
    issues: list[dict[str, Any]] = Field(default_factory=list)
    issues_filename: str | None = None
    review_filename: str | None = None
    duration_seconds: float = 0
    stage_durations: dict[str, float] = Field(default_factory=dict)


class IntegrationProgress(BaseModel):
    status: str
    stages: list[dict[str, Any]] = Field(default_factory=list)
    detail: str | None = None
    updated_at: str = ""


class MatchingSourceRuleIn(BaseModel):
    """A scoped source-sheet/target-field hint for one project period."""

    company: str | None = Field(default=None, max_length=200)
    month: str | None = Field(default=None, max_length=20)
    source_sheet: str | None = Field(default=None, max_length=120)
    target_sheet: str | None = Field(default=None, max_length=120)
    target_field: str | None = Field(default=None, max_length=120)
    action: str = Field(default="review", pattern="^(allow|review|ignore)$")


class MatchingPolicyIn(BaseModel):
    """Safe, tenant-isolated matching controls; no database schema change."""

    auto_match_threshold: float = Field(default=0.85, ge=0.85, le=0.99)
    field_weights: dict[str, float] = Field(default_factory=lambda: {
        "姓名": 0.50,
        "公司": 0.25,
        "部门": 0.15,
        "岗位": 0.10,
    })
    source_rules: list[MatchingSourceRuleIn] = Field(default_factory=list, max_length=200)


class MatchingPolicyOut(MatchingPolicyIn):
    project_id: str
    company: str
    salary_month: str
    updated_at: str


def _default_matching_policy(project: Project) -> dict[str, Any]:
    return {
        "project_id": str(project.id),
        "company": str(project.name),
        "salary_month": str(project.salary_month),
        "auto_match_threshold": 0.85,
        "field_weights": {"姓名": 0.50, "公司": 0.25, "部门": 0.15, "岗位": 0.10},
        "source_rules": [],
        "updated_at": "",
    }


@router.get("/{project_id}/matching-policy", response_model=MatchingPolicyOut)
def get_matching_policy(project_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    project = _load_project_or_404(project_id, user, db)
    policy = _load_json(project.id, "matching_policy") or _default_matching_policy(project)
    policy["project_id"] = str(project.id)
    policy["company"] = str(project.name)
    policy["salary_month"] = str(project.salary_month)
    return MatchingPolicyOut.model_validate(policy)


@router.put("/{project_id}/matching-policy", response_model=MatchingPolicyOut)
def update_matching_policy(
    project_id: str,
    payload: MatchingPolicyIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = _load_project_or_404(project_id, user, db)
    allowed_fields = {"姓名", "公司", "部门", "岗位"}
    unknown_fields = set(payload.field_weights) - allowed_fields
    if unknown_fields:
        raise HTTPException(status_code=422, detail=f"不支持的匹配字段：{'、'.join(sorted(unknown_fields))}")
    if any(value < 0 for value in payload.field_weights.values()):
        raise HTTPException(status_code=422, detail="匹配字段权重不能为负数")
    total = sum(payload.field_weights.values())
    if total <= 0:
        raise HTTPException(status_code=422, detail="至少需要一个大于 0 的匹配字段权重")
    policy = {
        **payload.model_dump(),
        "project_id": str(project.id),
        "company": str(project.name),
        "salary_month": str(project.salary_month),
        "updated_at": datetime.now().isoformat(),
    }
    _save_json(project.id, "matching_policy", policy)
    return MatchingPolicyOut.model_validate(policy)


def _load_template_base_entities(
    template_file: UploadFile | None,
    template_path: str,
) -> tuple[dict[str, list[dict]], list[str]]:
    """Load uploaded master data; built-in workbooks supply layout only."""
    if template_file is None:
        return {}, [f"内置模板仅作为输出版式：{os.path.basename(template_path)}"]

    template_result = import_template_to_entities(template_path)
    base_entities: dict[str, list[dict]] = {}
    logs: list[str] = []
    for entity_id, records in template_result.items():
        if entity_id == "_log":
            logs.extend(records)
            continue
        for record in records:
            record["_source_priority"] = 0
            record["_source_file"] = template_file.original_name
            if entity_id == "employee_profile":
                record["_roster_authoritative"] = True
        base_entities.setdefault(entity_id, []).extend(records)
    return base_entities, logs


def _ensure_recognized_people(prepared: dict[str, Any]) -> None:
    """Prevent unrelated or unrecognized workbooks from producing a blank/wrong ledger."""
    if int(prepared.get("source_person_count", 0) or 0) <= 0:
        raise HTTPException(
            status_code=400,
            detail="未识别到有效人员，请确认素材中包含姓名，并至少包含工号或身份证号等可核对字段",
        )


@router.post("/{project_id}/integrate", response_model=IntegrationResult)
def integrate_pipeline(project_id: str, user: User = Depends(get_current_user),
                       db: Session = Depends(get_db)):
    """用 Python 合并全部来源人员，异常留空后生成唯一总表。"""
    integration_started = time.perf_counter()
    p = _load_project_or_404(project_id, user, db)
    _save_json(p.id, "export_meta", {"status": "processing"})
    stages = _default_integration_progress_stages()
    _save_integration_progress(p.id, status="processing", stages=stages, detail="等待开始")
    logs: list[str] = []
    stage_durations: dict[str, float] = {}

    stage_started = time.perf_counter()
    stages[0]["status"] = "running"
    _save_integration_progress(p.id, status="processing", stages=stages, detail=stages[0]["label"])
    selected_template = _select_template_file(p.id, db)
    if not selected_template:
        _save_integration_progress(p.id, status="blocked", stages=stages, detail="缺少可用总表底板")
        raise HTTPException(status_code=400, detail="请先在“总表”上传区上传且仅上传一份包含“工资核算”Sheet 的总表")
    template_file, template_path = selected_template
    effective_salary_month = _resolve_effective_salary_month(
        p.salary_month,
        template_file.original_name,
    )
    stage_durations["base"] = round(time.perf_counter() - stage_started, 2)
    stages[0]["status"] = "completed"
    _save_integration_progress(p.id, status="processing", stages=stages, detail="已加载总表底板")
    logs.append(
        f"自动选择总表底板：{template_file.original_name if template_file else os.path.basename(template_path)}"
    )

    source_entities: dict[str, list[dict]] = {}
    failed_source_files: list[str] = []
    stage_started = time.perf_counter()
    stages[1]["status"] = "running"
    _save_integration_progress(p.id, status="processing", stages=stages, detail=stages[1]["label"])
    source_query = db.query(UploadFile).filter(
        UploadFile.project_id == p.id,
        UploadFile.file_type.in_([*_PAYROLL_SOURCE_FILE_TYPES]),
    )
    source_files = source_query.order_by(UploadFile.created_at.asc()).all()
    notice_paths = [Path(UPLOAD_DIR) / file.stored_path for file in source_files]
    if _only_hire_notices(notice_paths):
        try:
            result = _export_personnel_updates(
                Path(template_path), notice_paths,
                {str(path): file.original_name for path, file in zip(notice_paths, source_files)},
                Path(EXPORT_DIR) / p.id, effective_salary_month,
                _load_json(p.id, "matching_policy"),
            )
            result["source_signature"], _ = _current_source_signature(p.id, db)
            result["path"] = os.path.join(p.id, result["filename"])
            result["review_path"] = os.path.join(p.id, result["review_filename"])
            result["review_base_path"] = template_path
            result["review_format_version"] = _REVIEW_EXPORT_FORMAT_VERSION
            if result["issues_filename"]:
                result["issues_path"] = os.path.join(p.id, result["issues_filename"])
            result["duration_seconds"] = round(time.perf_counter() - integration_started, 2)
            _save_json(p.id, "export_meta", jsonable_encoder(result))
            for stage in stages:
                stage["status"] = "completed"
            status = "completed_with_issues" if result["issues"] else "completed"
            _save_integration_progress(p.id, status=status, stages=stages, detail="人员信息同步已结束")
            return IntegrationResult(**{
                **result, "status": status, "stages": stages, "overview": [],
            })
        except Exception:
            _save_integration_progress(p.id, status="failed", stages=stages, detail="人员信息同步失败")
            raise
    base_entities, template_logs = _load_template_base_entities(template_file, template_path)
    logs.extend(template_logs)
    for source_file in source_files:
        source_path = os.path.join(UPLOAD_DIR, source_file.stored_path)
        if not os.path.exists(source_path):
            logs.append(f"文件丢失：{source_file.original_name}")
            continue
        try:
            parsed = import_file_to_entities(source_path, source_file.original_name)
            confirmed_tokens: set[tuple[str, str]] = set()
            for entity_id, records in parsed.items():
                if entity_id == "_log":
                    logs.extend(records)
                else:
                    for record in records:
                        if _record_confirms_current_roster(entity_id, record):
                            record["_roster_confirmation"] = True
                            confirmed_tokens.update(_record_identity_tokens(record))
                    source_entities.setdefault(entity_id, []).extend(records)
            parsed_profiles = parsed.get("employee_profile", [])
            discovered = (
                []
                if parsed_profiles
                else discover_people_in_workbook(source_path, source_file.original_name)
            )
            for record in discovered:
                if _record_identity_tokens(record) & confirmed_tokens:
                    record["_roster_confirmation"] = True
            source_entities.setdefault("employee_profile", []).extend(discovered)
            if parsed_profiles:
                logs.append(
                    f"来源文件：{source_file.original_name}，已结构化识别 {len(parsed_profiles)} 条人员记录"
                )
            else:
                logs.append(f"来源文件：{source_file.original_name}，兜底识别 {len(discovered)} 条人员记录")
        except Exception as exc:
            logs.append(f"解析失败：{source_file.original_name}：{exc}")
            failed_source_files.append(source_file.original_name)
    stage_durations["source"] = round(time.perf_counter() - stage_started, 2)
    stages[1]["status"] = "completed"
    _save_integration_progress(p.id, status="processing", stages=stages, detail="已完成来源文件解析")
    logs.append(f"参与人员汇总的其他文件：{len(source_files)} 个")
    if failed_source_files:
        detail = "、".join(failed_source_files[:5])
        if len(failed_source_files) > 5:
            detail += f" 等 {len(failed_source_files)} 个文件"
        _save_json(
            p.id,
            "export_meta",
            {"status": "blocked", "detail": f"以下变更文件解析失败：{detail}"},
        )
        _save_integration_progress(p.id, status="blocked", stages=stages, detail=f"以下变更文件解析失败：{detail}")
        raise HTTPException(
            status_code=400,
            detail=f"以下变更文件解析失败，已停止更新以避免漏人：{detail}",
        )

    stage_started = time.perf_counter()
    stages[2]["status"] = "running"
    _save_integration_progress(p.id, status="processing", stages=stages, detail=stages[2]["label"])
    matching_policy = _load_json(p.id, "matching_policy") or _default_matching_policy(p)
    prepared = combine_entity_sources(
        base_entities,
        source_entities,
        salary_month=effective_salary_month,
        matching_policy=matching_policy,
    )
    try:
        _ensure_recognized_people(prepared)
    except HTTPException as exc:
        _save_json(p.id, "export_meta", {"status": "blocked", "detail": exc.detail})
        _save_integration_progress(p.id, status="blocked", stages=stages, detail=str(exc.detail))
        raise
    final_entities = prepared["entities"]
    validation = _validate_final_entities(final_entities)
    validation["preparation_issue_count"] = len(prepared["issues"])
    source_signature, _ = _current_source_signature(p.id, db)
    final_issues = prepared["issues"] + _standard_manual_follow_up_issues(
        template_file.original_name
    )
    _assign_manual_issue_ids(final_issues)
    final_payload = {
        "entities": final_entities,
        "issues": final_issues,
        "coverage": {
            "input_person_count": prepared["input_person_count"],
            "output_person_count": prepared["output_person_count"],
            "manual_only_person_count": prepared["manual_only_person_count"],
            "coverage_gap_count": prepared["coverage_gap_count"],
        },
        "stage_durations": stage_durations,
        "validation": validation,
        "template_file_id": template_file.id if template_file else None,
        "template_name": template_file.original_name if template_file else os.path.basename(template_path),
        "salary_month": effective_salary_month,
        "matching_policy": matching_policy,
        "source_signature": source_signature,
        "created_at": datetime.now().isoformat(),
    }
    _save_json(p.id, "final_db", final_payload)
    _save_json(p.id, "entities", {
        "entities": final_entities,
        "imported_at": datetime.now().isoformat(),
        "logs": logs,
    })
    stages[2]["status"] = "completed"
    stage_durations["match"] = round(time.perf_counter() - stage_started, 2)
    _save_integration_progress(p.id, status="processing", stages=stages, detail="已完成人员匹配与校验")
    logs.append(
        f"人员覆盖校验：输入候选 {prepared['input_person_count']} 人，"
        f"自动写入 {prepared['output_person_count']} 人，"
        f"仅人工复核 {prepared['manual_only_person_count']} 人"
    )

    stage_started = time.perf_counter()
    stages[3]["status"] = "running"
    _save_integration_progress(p.id, status="processing", stages=stages, detail=stages[3]["label"])
    try:
        exported = export_pipeline(project_id, user=user, db=db)
    except HTTPException as exc:
        _save_integration_progress(p.id, status="failed", stages=stages, detail=str(exc.detail))
        raise
    except Exception:
        _save_integration_progress(p.id, status="failed", stages=stages, detail="生成 Excel 时发生错误")
        raise
    stage_durations["export"] = round(time.perf_counter() - stage_started, 2)
    stages[3]["status"] = "completed"
    logs.extend(exported.log)
    completed_status = "completed_with_issues" if exported.issue_count else "completed"
    _save_integration_progress(
        p.id,
        status=completed_status,
        stages=stages,
        detail="已生成最新成果，可在成果页查收",
    )
    return IntegrationResult(
        status=completed_status,
        filename=exported.filename,
        stages=stages,
        overview=get_entity_overview(final_entities),
        validation=validation,
        log=logs,
        completed_at=datetime.now().isoformat(),
        issue_count=exported.issue_count,
        source_person_count=exported.source_person_count,
        output_person_count=exported.output_person_count,
        input_person_count=exported.input_person_count,
        manual_only_person_count=exported.manual_only_person_count,
        coverage_gap_count=exported.coverage_gap_count,
        source_file_count=exported.source_file_count,
        issues=exported.issues,
        issues_filename=exported.issues_filename,
        review_filename=exported.review_filename,
        duration_seconds=round(time.perf_counter() - integration_started, 2),
        stage_durations=stage_durations,
    )


@router.get("/{project_id}/integrate/progress", response_model=IntegrationProgress)
def get_integration_progress(
    project_id: str,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Return live integration progress for an authorized project."""
    project = _load_project_or_404(project_id, user, db)
    return IntegrationProgress(**_load_integration_progress(project.id))


@router.post("/{project_id}/export", response_model=ExportResult)
def export_pipeline(project_id: str, user: User = Depends(get_current_user),
                    db: Session = Depends(get_db)):
    """模块F：按模板导出最终数据库。"""
    export_started = time.perf_counter()
    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先生成最终数据库（模块E）")
    current_signature, _ = _current_source_signature(p.id, db)
    if final_data.get("source_signature") != current_signature:
        raise HTTPException(status_code=409, detail="来源文件已变化，请重新生成工资总表")
    _hydrate_legacy_department_transfer_issues(final_data)
    _assign_manual_issue_ids(final_data.get("issues", []))

    selected_template = _resolve_export_template(
        p.id,
        final_data.get("template_file_id"),
        db,
    )
    if not selected_template:
        raise HTTPException(status_code=400, detail="请先上传且仅上传一份明确指定的总表")
    _, uploaded_template_path = selected_template
    template_path = _resolve_layout_template_path(uploaded_template_path)

    os.makedirs(os.path.join(EXPORT_DIR, p.id), exist_ok=True)
    export_month = str(final_data.get("salary_month") or p.salary_month)
    timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
    out_name = f"工资核算_{export_month}_正式稿_{timestamp}.xlsx"
    out_path = os.path.join(EXPORT_DIR, p.id, out_name)
    issues_name = f"待人工处理_{export_month}_{timestamp}.xlsx"
    issues_path = os.path.join(EXPORT_DIR, p.id, issues_name)
    review_name = f"工资核算_{export_month}_修改稿_{timestamp}.xlsx"
    review_path = os.path.join(EXPORT_DIR, p.id, review_name)

    try:
        result = export_unified_workbook(
            template_path=template_path,
            entities=final_data["entities"],
            output_path=out_path,
            additional_issues=[
                issue for issue in final_data.get("issues", [])
                if issue.get("status") != "confirmed"
            ],
            coverage=final_data.get("coverage", {}),
            verify_formulas=False,
        )
        if os.path.normcase(os.path.abspath(template_path)) == os.path.normcase(
            os.path.abspath(BUILTIN_TEMPLATE_PATH)
        ):
            cleared_sheets = _clear_builtin_layout_source_data(out_path)
            result["log"].append(
                f"已隔离内置模板历史来源数据: {len(cleared_sheets)} 个工作表"
            )
        _replace_issue_source_filename(
            out_path,
            os.path.basename(template_path),
            os.path.basename(template_path),
        )
        source_paths = []
        source_names: dict[str, str] = {}
        for file in db.query(UploadFile).filter(
            UploadFile.project_id == p.id,
            UploadFile.file_type.in_([*_PAYROLL_SOURCE_FILE_TYPES]),
        ).order_by(UploadFile.created_at.asc()).all():
            path = os.path.join(UPLOAD_DIR, file.stored_path)
            if not os.path.exists(path):
                continue
            is_selected_template = bool(selected_template[0]) and str(file.id) == str(selected_template[0].id)
            if is_selected_template and template_path == uploaded_template_path:
                continue
            source_paths.append(path)
            source_names[path] = file.original_name
        # Every source Sheet is compared to every existing master Sheet. An
        # unmatched Sheet is an auditable manual exception, never a new tab.
        sheet_decisions = _apply_semantic_source_updates(
            out_path,
            source_paths,
            salary_month=str(final_data.get("salary_month") or export_month),
            source_names=source_names,
            sheet_overrides=_manual_sheet_overrides(final_data),
        )
        ignored_semantic_issues = set(final_data.get("ignored_semantic_issue_keys", []))
        result["issues"].extend(
            issue for issue in sheet_decisions["issues"]
            if _semantic_issue_key(issue) not in ignored_semantic_issues
        )
        _assign_manual_issue_ids(result["issues"])
        result["issue_count"] = len(result["issues"])
        formula_count = None
        if _payroll_person_count(template_path) == result["output_person_count"]:
            # Validate the master formulas before the intentional external-link
            # freeze converts those formulas into cached values.
            formula_count = _verify_unchanged_formula_signature(template_path, out_path)
            result["log"].append(
                f"来源整合后的公式完整性校验通过：原有 {formula_count} 个公式全部保留"
            )
        frozen_external_formula_count = _freeze_external_workbook_formulas(
            out_path,
            fallback_source_paths=[template_path, *source_paths],
        )
        if frozen_external_formula_count:
            result["log"].append(
                f"已将外部工作簿公式冻结为缓存结果：{frozen_external_formula_count} 个单元格"
            )
        duty_rows_updated = _update_duty_roster_counts(out_path, source_paths)
        ledger_rows_updated = _write_social_contributions_to_ledger(
            out_path,
            final_data["entities"],
        )
        remove_manual_issues_sheet(out_path)
        _attach_manual_issue_locations(out_path, result["issues"])
        write_manual_issues_workbook(issues_path, result["issues"])
        if template_path != uploaded_template_path:
            result["log"].append(f"输出格式底板：{os.path.basename(template_path)}")
        review_summary = create_review_workbook(
            template_path,
            out_path,
            review_path,
            issues=result["issues"],
        )
        result["log"].append(
            f"来源 Sheet 语义归并：{len(sheet_decisions['matches'])} 个，"
            f"自动覆盖 {sheet_decisions['auto_update_count']} 个单元格，"
            f"复杂异常 {len(sheet_decisions['issues'])} 项"
        )
        result["log"].append(
            f"已生成修改稿：标记 {review_summary['change_count']} 项更新内容"
        )
        if duty_rows_updated:
            result["log"].append(f"已按值班名单更新配送员值班天数: {duty_rows_updated} 人")
        if ledger_rows_updated:
            result["log"].append(f"已更新台账社保公积金: {ledger_rows_updated} 人")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出失败: {e}")

    export_meta = {
        "path": os.path.join(p.id, out_name),
        "filename": out_name,
        "issue_count": result["issue_count"],
        "source_person_count": result["source_person_count"],
        "output_person_count": result["output_person_count"],
        "input_person_count": result["input_person_count"],
        "manual_only_person_count": result["manual_only_person_count"],
        "coverage_gap_count": result["coverage_gap_count"],
        "source_file_count": 0,
        "issues": result["issues"],
        "validation": final_data.get("validation", {}),
        "completed_at": datetime.now().isoformat(),
        "issues_path": os.path.join(p.id, issues_name),
        "issues_filename": issues_name,
        "review_path": os.path.join(p.id, review_name),
        "review_filename": review_name,
        "review_base_path": template_path,
        "review_format_version": _REVIEW_EXPORT_FORMAT_VERSION,
        "duration_seconds": round(time.perf_counter() - export_started, 2),
        "stage_durations": final_data.get("stage_durations", {}),
        "target_sheets": _workbook_sheet_names(out_path),
    }
    export_meta["source_signature"], uploaded_file_count = _current_source_signature(p.id, db)
    # The selected master is a baseline, not a source/change workbook.
    export_meta["source_file_count"] = max(0, uploaded_file_count - 1)
    with open(_session_path(p.id, "export_meta"), "w", encoding="utf-8") as fp:
        json.dump(export_meta, fp, ensure_ascii=False)

    return ExportResult(
        filename=out_name,
        log=result["log"],
        issue_count=result["issue_count"],
        source_person_count=result["source_person_count"],
        output_person_count=result["output_person_count"],
        input_person_count=result["input_person_count"],
        manual_only_person_count=result["manual_only_person_count"],
        coverage_gap_count=result["coverage_gap_count"],
        source_file_count=export_meta["source_file_count"],
        issues=result["issues"],
        validation=export_meta["validation"],
        completed_at=export_meta["completed_at"],
        issues_filename=issues_name,
        review_filename=review_name,
        duration_seconds=export_meta["duration_seconds"],
    )


class DepartmentTransferConfirmIn(BaseModel):
    issue_ids: list[str]


class ManualIssueResolveIn(BaseModel):
    issue_id: str = Field(min_length=1, max_length=120)
    action: str = Field(pattern="^(apply_proposed|keep_current)$")


class ManualReviewDecisionIn(BaseModel):
    issue_id: str = Field(min_length=1, max_length=120)
    outcome: str = Field(pattern="^(更新到总表|保留总表)$")
    value: str | int | float | bool | None = None
    note: str = Field(default="", max_length=1000)
    remember: bool = False


class ManualReviewApplyIn(BaseModel):
    decisions: list[ManualReviewDecisionIn] = Field(min_length=1, max_length=1000)


class ManualSheetMappingIn(BaseModel):
    issue_id: str = Field(min_length=1, max_length=120)
    target_sheet: str = Field(min_length=1, max_length=120)


class ManualSheetMappingsIn(BaseModel):
    mappings: list[ManualSheetMappingIn] = Field(min_length=1, max_length=500)


class ManualSemanticIssueIgnoreIn(BaseModel):
    issue_id: str = Field(min_length=1, max_length=120)


@router.post(
    "/{project_id}/manual-review/import",
    response_model=ExportResult,
)
def import_manual_review_workbook(
    project_id: str,
    file: FastAPIUploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Import explicit manual decisions and regenerate the formal and review workbooks."""
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="请上传系统下载后填写的 .xlsx 人工处理表")
    content = file.file.read()
    if not content or len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="人工处理表不能为空且不能超过 10 MB")

    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先完成数据整合")
    try:
        updated = _apply_manual_review_workbook(final_data, content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not updated:
        raise HTTPException(status_code=400, detail="人工处理表中没有可更新到总表的已处理事项")

    _save_json(p.id, "final_db", final_data)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, f"已导入人工处理表并更新 {updated} 项；未处理事项保留在最新清单中")
    return result


@router.post(
    "/{project_id}/manual-review/apply",
    response_model=ExportResult,
)
def apply_online_manual_review(
    project_id: str,
    payload: ManualReviewApplyIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Apply webpage decisions in one validated batch and regenerate exports."""
    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先完成数据整合")
    decisions = [decision.model_dump() for decision in payload.decisions]
    try:
        updated = _apply_manual_review_decisions(
            final_data,
            decisions,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not updated:
        raise HTTPException(status_code=400, detail="所选事项已处理或没有可更新内容")

    _save_json(p.id, "final_db", final_data)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, f"已在网页完成 {updated} 项人工处理，并生成最新总表")
    remembered = 0
    issues_by_id = {str(issue.get("issue_id")): issue for issue in final_data.get("issues", [])}
    for decision in decisions:
        if not decision.get("remember"):
            continue
        issue = issues_by_id.get(str(decision.get("issue_id")))
        if not issue:
            continue
        try:
            memory_item = dict(issue)
            memory_item["source_sheet"] = (issue.get("source_sheets") or [None])[0]
            experience_store.record_agent_decision(
                tenant_id=str(user.tenant_id),
                project_id=str(p.id),
                created_by=str(user.id),
                diff_type=str(issue.get("issue_type") or "manual_review"),
                item=memory_item,
                match_fields=["issue_type", "target_field", "source_sheet"],
                decision=str(decision.get("outcome") or ""),
                updates={str(issue.get("target_field") or ""): decision.get("value")},
                note=str(decision.get("note") or ""),
                period_key=str(final_data.get("salary_month") or p.salary_month),
                confidence=1.0,
            )
            remembered += 1
        except ValueError:
            continue
    if remembered:
        result.log.insert(1, f"已将 {remembered} 项确认沉淀为本公司候选经验规则")
    return result


def _save_manual_sheet_mappings(
    project_id: str,
    mappings: list[ManualSheetMappingIn],
    user: User,
    db: Session,
) -> tuple[Project, dict[str, Any], int]:
    """Validate and persist Sheet choices together before regenerating exports."""
    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先完成数据整合")
    meta = _load_current_export_meta(p.id, db)
    _assign_manual_issue_ids(meta.get("issues", []))
    selected_template = _resolve_export_template(p.id, final_data.get("template_file_id"), db)
    if not selected_template:
        raise HTTPException(status_code=400, detail="找不到当前总表，无法确认目标工作表")
    _, template_path = selected_template
    target_sheets = _workbook_sheet_names(_resolve_layout_template_path(template_path))
    issues_by_id = {str(item.get("issue_id")): item for item in meta.get("issues", [])}
    seen_sources: set[tuple[str, str]] = set()
    validated: list[tuple[str, str, str]] = []
    for mapping in mappings:
        issue = issues_by_id.get(mapping.issue_id.strip())
        if not issue:
            raise HTTPException(status_code=404, detail="有待确认任务已不存在，请刷新页面")
        if str(issue.get("issue_type") or "") not in {
            "ambiguous_sheet", "unmatched_sheet", "insufficient_topic_evidence",
        }:
            raise HTTPException(status_code=400, detail="所选任务不需要选择目标工作表")
        source_files = issue.get("source_files") or []
        source_sheets = issue.get("source_sheets") or []
        if len(source_files) != 1 or len(source_sheets) != 1:
            raise HTTPException(status_code=400, detail="有任务缺少唯一来源，不能安全保存映射")
        if mapping.target_sheet not in target_sheets:
            raise HTTPException(status_code=400, detail=f"所选目标工作表“{mapping.target_sheet}”不在当前总表中")
        source_file, source_sheet = str(source_files[0]), str(source_sheets[0])
        if (source_file, source_sheet) in seen_sources:
            raise HTTPException(status_code=400, detail="同一来源工作表不能重复选择")
        seen_sources.add((source_file, source_sheet))
        validated.append((source_file, source_sheet, mapping.target_sheet))

    remaining_mappings = [
        mapping for mapping in final_data.get("sheet_mappings", [])
        if isinstance(mapping, dict)
        and (str(mapping.get("source_file") or ""), str(mapping.get("source_sheet") or "")) not in seen_sources
    ]
    remaining_mappings.extend({"source_file": source_file, "source_sheet": source_sheet, "target_sheet": target_sheet}
                              for source_file, source_sheet, target_sheet in validated)
    final_data["sheet_mappings"] = remaining_mappings
    _save_json(p.id, "final_db", final_data)
    return p, final_data, len(validated)


@router.post(
    "/{project_id}/manual-review/sheet-mappings",
    response_model=ExportResult,
)
def apply_manual_sheet_mapping(
    project_id: str,
    payload: ManualSheetMappingIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Persist one Sheet destination choice and regenerate the workbook."""
    p, _, updated = _save_manual_sheet_mappings(project_id, [payload], user, db)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, f"已确认 1 个工作表更新位置，并生成最新总表")
    return result


@router.post(
    "/{project_id}/manual-review/sheet-mappings/batch",
    response_model=ExportResult,
)
def apply_manual_sheet_mappings(
    project_id: str,
    payload: ManualSheetMappingsIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Persist all checked Sheet destination choices in one export operation."""
    p, _, updated = _save_manual_sheet_mappings(project_id, payload.mappings, user, db)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, f"已批量确认 {updated} 个工作表更新位置，并生成最新总表")
    return result


@router.post(
    "/{project_id}/manual-review/semantic-issues/ignore",
    response_model=ExportResult,
)
def ignore_semantic_manual_issue(
    project_id: str,
    payload: ManualSemanticIssueIgnoreIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Mark a source-workbook task as intentionally excluded for this batch."""
    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先完成数据整合")
    meta = _load_current_export_meta(p.id, db)
    _assign_manual_issue_ids(meta.get("issues", []))
    issue = next(
        (item for item in meta.get("issues", []) if str(item.get("issue_id")) == payload.issue_id.strip()),
        None,
    )
    if not issue:
        raise HTTPException(status_code=404, detail="该待确认任务已不存在，请刷新页面")
    if str(issue.get("issue_type") or "") not in _SEMANTIC_REVIEW_ISSUE_TYPES:
        raise HTTPException(status_code=400, detail="该任务不能通过“本次不更新”处理")
    ignored = set(final_data.get("ignored_semantic_issue_keys", []))
    ignored.add(_semantic_issue_key(issue))
    final_data["ignored_semantic_issue_keys"] = sorted(ignored)
    _save_json(p.id, "final_db", final_data)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, "已标记 1 项为本次不更新，并生成最新总表")
    return result


@router.post(
    "/{project_id}/manual-review/department-transfers/confirm",
    response_model=ExportResult,
)
def confirm_department_transfers(
    project_id: str,
    payload: DepartmentTransferConfirmIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Apply explicitly selected department transfers and regenerate the workbook."""
    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先完成数据整合")
    try:
        updated = _confirm_department_transfers(final_data, payload.issue_ids)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    _save_json(p.id, "final_db", final_data)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, f"已确认并更新 {updated} 条部门变更；其他待人工处理事项保持不变")
    return result


@router.post(
    "/{project_id}/manual-review/issue/resolve",
    response_model=ExportResult,
)
def resolve_manual_issue(
    project_id: str,
    payload: ManualIssueResolveIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Apply a guided choice and regenerate both workbook outputs."""
    p = _load_project_or_404(project_id, user, db)
    final_data = _load_json(p.id, "final_db")
    if not final_data:
        raise HTTPException(status_code=400, detail="请先完成数据整合")
    try:
        resolved_value = _resolve_manual_issue(
            final_data,
            payload.issue_id.strip(),
            payload.action,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    _save_json(p.id, "final_db", final_data)
    result = export_pipeline(project_id, user=user, db=db)
    result.log.insert(0, f"已按引导选项处理 {payload.issue_id}：{resolved_value}")
    return result


@router.get("/{project_id}/export/latest", response_model=ExportResult)
def get_latest_pipeline_export(project_id: str, user: User = Depends(get_current_user),
                               db: Session = Depends(get_db)):
    """模块F：获取最近一次已生成的导出结果，用于页面刷新后回显。"""
    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    final_data = _load_json(p.id, "final_db") or {}
    _hydrate_legacy_department_transfer_issues(final_data)
    current_issues = meta.get("issues", final_data.get("issues", []))
    _assign_manual_issue_ids(current_issues)
    export_path = os.path.join(EXPORT_DIR, meta.get("path", ""))
    if not os.path.exists(export_path):
        raise HTTPException(status_code=404, detail="导出文件已丢失")
    # Location data is attached while generating the export.  Re-scanning a
    # large workbook here makes this lightweight result query exceed the
    # frontend request timeout even though the export already exists.
    for issue in current_issues:
        issue["can_update_online"] = _can_update_manual_issue_online(final_data, issue)
    visible_issues = [issue for issue in current_issues if issue.get("status") != "confirmed"]
    issues_filename, _ = _ensure_manual_issues_export(p.id, meta)
    return ExportResult(
        filename=meta["filename"],
        log=[],
        issue_count=len(visible_issues),
        source_person_count=int(meta.get("source_person_count", 0)),
        output_person_count=int(meta.get("output_person_count", 0)),
        input_person_count=int(meta.get("input_person_count", meta.get("source_person_count", 0))),
        manual_only_person_count=int(meta.get("manual_only_person_count", 0)),
        coverage_gap_count=int(meta.get("coverage_gap_count", 0)),
        source_file_count=int(meta.get("source_file_count", 0)),
        issues=visible_issues,
        validation=meta.get("validation", {}),
        completed_at=meta.get("completed_at", ""),
        issues_filename=issues_filename,
        review_filename=meta.get("review_filename"),
        duration_seconds=float(meta.get("duration_seconds", 0) or 0),
        stage_durations=meta.get("stage_durations", {}),
        target_sheets=list(meta.get("target_sheets", [])),
    )


@router.patch("/{project_id}/export/filename", response_model=ExportResult)
def rename_pipeline_export(
    project_id: str,
    payload: ExportRenameIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Rename the latest generated workbook inside its project export directory."""
    p = _load_project_or_404(project_id, user, db)
    meta_path = _session_path(p.id, "export_meta")
    meta = _load_current_export_meta(p.id, db)
    try:
        filename = _normalize_export_filename(payload.filename)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    current_path = os.path.join(EXPORT_DIR, meta.get("path", ""))
    if not meta.get("filename") or not os.path.isfile(current_path):
        raise HTTPException(status_code=404, detail="导出文件已丢失")
    project_export_dir = os.path.join(EXPORT_DIR, p.id)
    new_path = os.path.join(project_export_dir, filename)
    if os.path.abspath(new_path) != os.path.abspath(current_path) and os.path.exists(new_path):
        raise HTTPException(status_code=409, detail="同名文件已存在，请换一个文件名")
    if os.path.abspath(new_path) != os.path.abspath(current_path):
        os.replace(current_path, new_path)
    meta["filename"] = filename
    meta["path"] = os.path.join(p.id, filename)
    with open(meta_path, "w", encoding="utf-8") as fp:
        json.dump(meta, fp, ensure_ascii=False)
    return ExportResult(
        filename=filename,
        log=[],
        issue_count=int(meta.get("issue_count", 0)),
        source_person_count=int(meta.get("source_person_count", 0)),
        output_person_count=int(meta.get("output_person_count", 0)),
        input_person_count=int(meta.get("input_person_count", meta.get("source_person_count", 0))),
        manual_only_person_count=int(meta.get("manual_only_person_count", 0)),
        coverage_gap_count=int(meta.get("coverage_gap_count", 0)),
        source_file_count=int(meta.get("source_file_count", 0)),
        issues=meta.get("issues", []),
        validation=meta.get("validation", {}),
        completed_at=meta.get("completed_at", ""),
        issues_filename=meta.get("issues_filename"),
        review_filename=meta.get("review_filename"),
        duration_seconds=float(meta.get("duration_seconds", 0) or 0),
        stage_durations=meta.get("stage_durations", {}),
    )


@router.get("/{project_id}/export/preview")
def preview_pipeline_export(project_id: str, user: User = Depends(get_current_user),
                            db: Session = Depends(get_db)):
    """Preview the generated workbook in the app without downloading it."""
    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    export_path = os.path.join(EXPORT_DIR, meta.get("path", ""))
    if not meta.get("filename") or not os.path.exists(export_path):
        raise HTTPException(status_code=404, detail="导出文件已丢失")

    return {"filename": meta["filename"], **_build_export_preview(export_path)}


@router.get("/{project_id}/export/download")
def download_pipeline_export(project_id: str, user: User = Depends(get_current_user),
                              db: Session = Depends(get_db)):
    """模块F：下载导出文件。"""
    import io
    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    fp_full = os.path.join(EXPORT_DIR, meta["path"])
    if not os.path.exists(fp_full):
        raise HTTPException(status_code=404, detail="导出文件已丢失")
    with open(fp_full, "rb") as f:
        data = f.read()
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(meta["filename"])},
    )


@router.get("/{project_id}/export/issues/preview")
def preview_manual_issues_export(project_id: str, user: User = Depends(get_current_user),
                                 db: Session = Depends(get_db)):
    """Preview the separate workbook that lists manual follow-up items."""
    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    filename, path = _ensure_manual_issues_export(p.id, meta)
    return {"filename": filename, **_build_export_preview(path)}


@router.get("/{project_id}/export/issues/download")
def download_manual_issues_export(project_id: str, user: User = Depends(get_current_user),
                                  db: Session = Depends(get_db)):
    """Download the separate workbook that lists manual follow-up items."""
    import io

    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    filename, path = _ensure_manual_issues_export(p.id, meta)
    with open(path, "rb") as fp:
        data = fp.read()
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


@router.get("/{project_id}/export/review/preview")
def preview_review_export(project_id: str, user: User = Depends(get_current_user),
                          db: Session = Depends(get_db)):
    """Preview the highlighted modification draft."""
    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    filename, path = _ensure_review_export(p.id, meta, db)
    return {
        "filename": filename,
        **_build_export_preview(path, preferred_sheet="工资核算"),
    }


@router.get("/{project_id}/export/review/download")
def download_review_export(project_id: str, user: User = Depends(get_current_user),
                           db: Session = Depends(get_db)):
    """Download the highlighted modification draft."""
    import io

    p = _load_project_or_404(project_id, user, db)
    meta = _load_current_export_meta(p.id, db)
    filename, path = _ensure_review_export(p.id, meta, db)
    with open(path, "rb") as fp:
        data = fp.read()
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )
