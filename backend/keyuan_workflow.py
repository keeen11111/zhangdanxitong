"""Bounded execution support for the verified 科园 2026.07 payroll batch."""
from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
import re
import shutil
from typing import Any, Mapping

import openpyxl


_MASTER_PATTERN = re.compile(
    r"^202607（所属月202606\).*北京科园.*鹤安.*大药房工资核算总表.*\.xlsx$"
)
_OUTPUT_TEMPLATE_PATTERN = re.compile(
    r"^202608（所属月202607）.*北京科园.*鹤安.*大药房工资核算总表.*\.xlsx$"
)
_SALARY_PATTERN = re.compile(r"^薪资数据-科园-7月薪资.*\.xlsx$")
_SOCIAL_NAMES = ("益药科园2026.07.xlsx", "鹤安长泰2026.07.xlsx")
_TAX_NAMES = (
    "202608_税款计算_工资薪金所得-益药科园.xls",
    "202608_税款计算_劳务报酬-益药科园.xls",
    "202608_税款计算_工资薪金所得-鹤安长泰.xls",
    "202608_税款计算_劳务报酬-鹤安长泰.xls",
)


@dataclass(frozen=True)
class KeyuanBatch:
    master_filename: str
    salary_source_name: str
    salary_source: Path
    social_sources: tuple[tuple[str, Path], ...]
    tax_sources: tuple[tuple[str, Path], ...]
    payroll_period: str = "2026.07"
    payment_period: str = "2026.08"
    output_filename: str = "202608（所属月202607）-北京科园-鹤安-大药房工资核算总表-v2.xlsx"

    def matches_project_month(self, project_month: str) -> bool:
        normalized = str(project_month or "").strip().replace("-", ".")
        return normalized == self.payroll_period


def detect_keyuan_batch(
    master_filename: str,
    source_paths: Mapping[str, str | Path] | None,
) -> KeyuanBatch | None:
    """Return the verified batch contract only for the complete known input set."""
    if not _MASTER_PATTERN.search(Path(str(master_filename or "")).name):
        return None
    sources = {
        Path(str(name)).name: Path(path)
        for name, path in (source_paths or {}).items()
    }
    salary_names = [name for name in sources if _SALARY_PATTERN.search(name)]
    if len(salary_names) != 1:
        return None
    if any(name not in sources for name in (*_SOCIAL_NAMES, *_TAX_NAMES)):
        return None
    required_names = {salary_names[0], *_SOCIAL_NAMES, *_TAX_NAMES}
    if any(not sources[name].is_file() for name in required_names):
        return None
    tax_like_names = {name for name in sources if "税款计算" in name}
    if tax_like_names != set(_TAX_NAMES):
        return None
    return KeyuanBatch(
        master_filename=Path(str(master_filename)).name,
        salary_source_name=salary_names[0],
        salary_source=sources[salary_names[0]],
        social_sources=tuple((name, sources[name]) for name in _SOCIAL_NAMES),
        tax_sources=tuple((name, sources[name]) for name in _TAX_NAMES),
    )


class KeyuanWorkflowError(RuntimeError):
    """A safe, resumable failure from the fixed workbook executor."""


def execute_keyuan_batch(
    batch: KeyuanBatch,
    *,
    master_path: Path,
    source_paths: Mapping[str, str | Path],
    output_path: Path,
    work_root: Path,
) -> dict[str, Any]:
    """Run the audited pure-Python workbook pass against only this run's inputs."""
    master_path = Path(master_path)
    output_path = Path(output_path)
    if output_path.resolve() in {master_path.resolve(), batch.salary_source.resolve()}:
        raise KeyuanWorkflowError("科园执行器只能写入独立结果草稿")

    job_root = Path(work_root) / f"keyuan-{output_path.stem}"
    sample_dir = job_root / "inputs"
    tax_dir = sample_dir / "回复：202608科园、鹤安应发工资的全部附件20260811"
    sample_dir.mkdir(parents=True, exist_ok=True)
    tax_dir.mkdir(parents=True, exist_ok=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    input_files = [(batch.master_filename, master_path)]
    # When the user uploads the next-month reference workbook as a source,
    # process against it so the exported workbook retains its full layout,
    # formulas, historical rows, and print settings.  The selected original
    # master remains untouched and is still recorded as the batch master.
    reference_candidates = [
        (Path(str(name)).name, Path(path))
        for name, path in (source_paths or {}).items()
        if _OUTPUT_TEMPLATE_PATTERN.search(Path(str(name)).name)
    ]
    if reference_candidates:
        input_files.append(reference_candidates[0])
    input_files.extend((batch.salary_source_name, batch.salary_source) for _ in [0])
    input_files.extend(batch.social_sources)
    input_files.extend(batch.tax_sources)
    for name, source in input_files:
        source = Path(source)
        if not source.is_file():
            raise KeyuanWorkflowError(f"输入文件不存在：{name}")
        destination_dir = tax_dir if name in dict(batch.tax_sources) else sample_dir
        shutil.copy2(source, destination_dir / name)

    staged_social = {name: sample_dir / name for name, _source in batch.social_sources}
    staged_tax = {name: tax_dir / name for name, _source in batch.tax_sources}
    processing_master = sample_dir / batch.master_filename
    if reference_candidates and (sample_dir / reference_candidates[0][0]).is_file():
        processing_master = sample_dir / reference_candidates[0][0]
    from scripts.keyuan_python_executor import process_keyuan_python
    try:
        result = process_keyuan_python(
            processing_master,
            sample_dir / batch.salary_source_name,
            staged_social,
            staged_tax,
            output_path,
        )
    except Exception as exc:
        if isinstance(exc, KeyuanWorkflowError):
            raise
        raise KeyuanWorkflowError(f"科园纯Python执行失败，未发布结果：{exc}") from exc
    try:
        workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=False)
        required_sheets = {"工资核算", "台账", "个税导出核对", "OA请款及审批"}
        missing = required_sheets.difference(workbook.sheetnames)
        formula_errors = [
            f"{sheet.title}!{cell.coordinate}"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and "#REF!" in cell.value
        ]
        structural = {
            "sheet_names": not missing,
            "payroll_columns": workbook["工资核算"].max_column == 99,
            "payroll_header": workbook["工资核算"]["A2"].value == "班制",
            "formula_references": not formula_errors,
        }
        workbook.close()
    except Exception as exc:
        raise KeyuanWorkflowError("生成的科园工作簿无法重新打开，未发布结果") from exc
    if not all(structural.values()):
        raise KeyuanWorkflowError(f"科园结果结构校验失败：{structural}")
    changes = result.get("changes") if isinstance(result, dict) else []
    issues = result.get("issues") if isinstance(result, dict) else []
    return {
        "status": "passed" if not issues else "needs_review",
        "output_path": str(output_path),
        "download_filename": batch.output_filename,
        "change_count": len(changes) if isinstance(changes, list) else 0,
        "changes": changes if isinstance(changes, list) else [],
        "issues": issues if isinstance(issues, list) else [],
        "structural": structural,
    }
