"""Bounded execution of supported natural-language workbook instructions."""
from __future__ import annotations

from collections import Counter
import hashlib
import json
from pathlib import Path
import re
from typing import Any, Mapping

import openpyxl

from backend.roster_sync import RosterSyncError, sync_master_to_source_roster


_QUOTED_TEXT = re.compile(r'["“]([^"”]+)["”]')
_TARGET_SHEET = re.compile(r"(?:表页|工作表|表)[—\-：:]\s*([^，,。；;\n]+)")
_TARGET_COLUMN = re.compile(r"([A-Za-z]{1,3})\s*列")


def _compact(value: Any) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", str(value or "")).lower()


def _is_total(value: str) -> bool:
    return _compact(value) in {"合计", "总计", "小计"}


def parse_supported_workbook_rule(instruction: str) -> dict[str, str] | None:
    """Parse one of the deliberately narrow supported workbook instructions.

    This parser is intentionally not a general language-to-Excel executor. A
    rule is accepted only when its source, target and operation boundaries are
    explicit enough for a deterministic, fail-closed update.
    """
    text = str(instruction or "")
    if "值班" in text and re.search(r"出现\s*(?:的\s*)?次数", text):
        sheet_match = _TARGET_SHEET.search(text)
        column_match = _TARGET_COLUMN.search(text)
        if sheet_match is None or column_match is None:
            return None
        target_sheet = re.split(r"\s*(?:填|将|把|清除|根据)", sheet_match.group(1).strip(), maxsplit=1)[0].strip()
        target_column = column_match.group(1).upper()
        if not target_sheet or "清除" not in text or "填" not in text:
            return None
        quoted = [value.strip() for value in _QUOTED_TEXT.findall(text) if value.strip()]
        source_hint = next((value for value in quoted if "值班" not in value and len(_compact(value)) >= 4), "")
        source_sheet = next((value for value in quoted if _compact(value) == "值班"), "值班")
        return {
            "kind": "duty_roster_name_count",
            "target_sheet": target_sheet,
            "target_column": target_column,
            "source_sheet": source_sheet,
            "source_hint": source_hint,
        }

    # Match the actual source-sheet token after phrases such as
    # ``变更表格里面的``.  The old expression accepted only ``表`` and
    # therefore captured the connective text (``格里面的派遣``) as part of
    # the sheet name, making an otherwise safe deterministic rule fail to
    # resolve its source and fall back to the model's manual write loop.
    roster_match = re.search(
        r"(?:只|仅)(?:需要|要|需)?(?:处理|保留|写入)\s*"
        r"(?:来源|变更)?(?:文件|表格|表)?(?:里|中|内|的|\s)*"
        r"([A-Za-z0-9_\-\u4e00-\u9fff]{1,20})\s*(?:sheet|表页|工作表)",
        text,
        flags=re.IGNORECASE,
    )
    has_roster_scope = bool(re.search(r"不保留|不显示|删除|最终名单|名单.*(?:为准|一致)", text))
    if roster_match is None or not has_roster_scope:
        return None
    period_match = re.search(r"(?<!\d)(20\d{2})[.\-/年](0?[1-9]|1[0-2])(?:月)?", text)
    source_sheet = roster_match.group(1).strip()
    # The capture can begin before a connective when the wording contains
    # ``表格里面的``/``文件里的``.  Strip only these explicit prefixes so a
    # real sheet name such as ``派遣`` is resolved without guessing.
    source_sheet = re.split(
        r"表格里面的|文件里面的|表格里的|文件里的|其中的|里面的|里的|中的",
        source_sheet,
    )[-1].strip()
    source_sheet = re.sub(r"^(?:面的|的)", "", source_sheet).strip()
    rule = {
        "kind": "source_sheet_roster_sync",
        "target_sheet": "明细",
        "source_sheet": source_sheet,
        "source_hint": "",
    }
    if period_match:
        rule["target_period"] = f"{period_match.group(1)}-{int(period_match.group(2)):02d}"
    return rule


def _find_name_column(sheet: openpyxl.worksheet.worksheet.Worksheet, *, source: bool) -> tuple[int, int] | None:
    for row_number in range(1, min(sheet.max_row, 30) + 1):
        for column in range(1, min(sheet.max_column, 100) + 1):
            header = _compact(sheet.cell(row_number, column).value)
            if source:
                if "值班人员" in header:
                    return row_number, column
            elif "姓名" in header or "值班人员" in header:
                return row_number, column
    return None


def _resolve_source(
    source_paths: Mapping[str, str | Path], *, source_hint: str, source_sheet: str,
) -> tuple[str, Path] | None:
    candidates: list[tuple[int, str, Path]] = []
    hint = _compact(source_hint).lstrip("0123456789")
    for filename, raw_path in source_paths.items():
        path = Path(raw_path)
        if not path.is_file() or path.suffix.lower() != ".xlsx":
            continue
        filename_key = _compact(Path(str(filename)).name)
        score = 0
        if hint:
            if hint in filename_key or filename_key in hint:
                score = 2
            else:
                continue
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            has_sheet = any(
                _compact(name) == _compact(source_sheet)
                or _compact(name).startswith(_compact(source_sheet))
                for name in workbook.sheetnames
            )
            workbook.close()
        except (OSError, ValueError, KeyError):
            continue
        if has_sheet:
            candidates.append((score, str(filename), path))
    if not candidates:
        return None
    candidates.sort(key=lambda item: (-item[0], Path(item[1]).name))
    best_score = candidates[0][0]
    best = [item for item in candidates if item[0] == best_score]
    if len(best) != 1:
        return None
    _, filename, path = best[0]
    return filename, path


def _rule_id(rule: Mapping[str, str], source_name: str) -> str:
    content = json.dumps({**rule, "source_name": source_name}, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(content.encode("utf-8")).hexdigest()[:24]


def apply_supported_workbook_rules(
    *, instruction: str, draft_path: Path, source_paths: Mapping[str, str | Path],
    previous_results: list[dict[str, Any]] | None = None, target_period: str | None = None,
) -> dict[str, list[dict[str, Any]]]:
    """Apply an explicit supported instruction to the run's draft only.

    A missing source, sheet, or name column produces an actionable review item
    and leaves the draft unchanged.  Source paths come solely from the active
    run's allow-list; callers cannot supply a model-selected external path.
    """
    rule = parse_supported_workbook_rule(instruction)
    if rule is None:
        return {"results": [], "updates": []}
    source = _resolve_source(source_paths, source_hint=rule["source_hint"], source_sheet=rule["source_sheet"])
    if source is None:
        return {"results": [{
            "status": "needs_review", "kind": rule["kind"],
            "detail": f"未能从本次上传文件中唯一定位包含“{rule['source_sheet']}”表页的来源文件，未写入总表。",
        }], "updates": []}
    source_name, source_path = source
    rule_id = _rule_id(rule, source_name)
    if any(item.get("rule_id") == rule_id and item.get("status") == "applied" for item in previous_results or []):
        return {"results": [{"rule_id": rule_id, "status": "skipped", "kind": rule["kind"], "detail": "该自然语言规则已在当前草稿执行。"}], "updates": []}
    if not draft_path.is_file():
        return {"results": [{
            "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
            "detail": "当前运行尚未生成独立草稿，未写入原始总表。",
        }], "updates": []}

    if rule["kind"] == "source_sheet_roster_sync":
        effective_period = str(rule.get("target_period") or target_period or "").strip()
        try:
            result = sync_master_to_source_roster(
                master_path=draft_path,
                source_path=source_path,
                output_path=draft_path,
                source_sheet=rule["source_sheet"],
                target_sheet=rule["target_sheet"],
                target_period=effective_period,
            )
        except RosterSyncError as exc:
            return {"results": [{
                "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
                "detail": f"按来源名单同步未执行：{exc}",
                "source_file": source_name, "source_sheet": rule["source_sheet"],
            }], "updates": []}
        return {"results": [{
            "rule_id": rule_id,
            "status": "applied",
            "kind": rule["kind"],
            "detail": (
                f"已仅按“{rule['source_sheet']}”表页同步“{rule['target_sheet']}”："
                f"最终 {result['employee_count']} 人，删除非名单人员 {len(result['removed_names'])} 人，"
                "并同步重算明细派生金额及付款通知书。"
            ),
            "source_file": source_name,
            **{key: result[key] for key in (
                "source_sheet", "target_sheet", "target_period", "employee_count",
                "retained_names", "removed_names", "unmapped_source_headers", "change_count",
            )},
        }], "updates": result["changes"]}

    source_book = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    try:
        source_sheet_name = next(
            (
                name for name in source_book.sheetnames
                if _compact(name) == _compact(rule["source_sheet"])
                or _compact(name).startswith(_compact(rule["source_sheet"]))
            ),
            None,
        )
        assert source_sheet_name is not None
        source_sheet = source_book[source_sheet_name]
        source_header = _find_name_column(source_sheet, source=True)
        if source_header is None:
            return {"results": [{
                "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
                "detail": "来源“值班”表页未找到“值班人员”姓名列，未写入总表。",
            }], "updates": []}
        source_header_row, source_name_column = source_header
        counts: Counter[str] = Counter()
        for row_number in range(source_header_row + 1, source_sheet.max_row + 1):
            name = str(source_sheet.cell(row_number, source_name_column).value or "").strip()
            if name and not _is_total(name):
                counts[name] += 1
    finally:
        source_book.close()
    if not counts:
        return {"results": [{
            "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
            "detail": "来源“值班”表页没有可统计的值班人员，未清空总表原值。",
        }], "updates": []}

    workbook = openpyxl.load_workbook(draft_path, data_only=False)
    try:
        if rule["target_sheet"] not in workbook.sheetnames:
            return {"results": [{
                "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
                "detail": f"总表未找到“{rule['target_sheet']}”表页，未写入。",
            }], "updates": []}
        target_sheet = workbook[rule["target_sheet"]]
        target_header = _find_name_column(target_sheet, source=False)
        if target_header is None:
            return {"results": [{
                "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
                "detail": f"“{rule['target_sheet']}”未找到姓名列，未写入。",
            }], "updates": []}
        target_header_row, target_name_column = target_header
        target_column = openpyxl.utils.column_index_from_string(rule["target_column"])
        target_rows = [
            (row_number, str(target_sheet.cell(row_number, target_name_column).value or "").strip())
            for row_number in range(target_header_row + 1, target_sheet.max_row + 1)
            if str(target_sheet.cell(row_number, target_name_column).value or "").strip()
            and not _is_total(str(target_sheet.cell(row_number, target_name_column).value or "").strip())
        ]
        formula_rows = [row for row, _ in target_rows if target_sheet.cell(row, target_column).data_type == "f"]
        if formula_rows:
            return {"results": [{
                "rule_id": rule_id, "status": "needs_review", "kind": rule["kind"],
                "detail": f"目标 {rule['target_column']} 列包含公式（如第 {formula_rows[0]} 行），为保护公式未写入。",
            }], "updates": []}
        updates: list[dict[str, Any]] = []
        matched_people = 0
        cleared_people = 0
        for row_number, name in target_rows:
            cell = target_sheet.cell(row_number, target_column)
            old_value = cell.value
            new_value = counts.get(name)
            if old_value not in (None, ""):
                cleared_people += 1
            if new_value is not None:
                matched_people += 1
            if old_value != new_value:
                cell.value = new_value
                updates.append({
                    "rule_id": rule_id, "rule": "自然语言指令：按值班人员出现次数填写值班天数",
                    "target_sheet": rule["target_sheet"], "target_cell": cell.coordinate,
                    "person_name": name, "old_value": old_value, "new_value": new_value,
                    "source_file": source_name, "source_sheet": source_sheet_name,
                    "source_row": "值班人员姓名出现次数",
                })
        workbook.save(draft_path)
    finally:
        workbook.close()
    unmatched_names = [name for _, name in target_rows if name not in counts]
    return {"results": [{
        "rule_id": rule_id, "status": "applied", "kind": rule["kind"],
        "detail": f"已清空并重填 {rule['target_sheet']}!{rule['target_column']} 列：匹配 {matched_people} 人，未匹配 {len(unmatched_names)} 人保持为空。",
        "source_file": source_name, "source_sheet": source_sheet_name,
        "matched_people": matched_people, "cleared_people": cleared_people,
        "unmatched_names": unmatched_names,
    }], "updates": updates}
