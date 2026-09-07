"""Regression for dual identifiers, subtotal blocks and array-formula templates."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

from core.sheet_mapper import apply_semantic_sheet_updates


def structured_files(tmp_path: Path) -> tuple[Path, Path, Path]:
    master, source, output = (tmp_path / name for name in ('master.xlsx', 'notice.xlsx', 'out.xlsx'))
    book = Workbook()
    pay = book.active
    pay.title = '工资核算'
    pay.append(['人员编码', '姓名', '新工号', '所属公司', '新公司', '成本中心/部门',
                '新成本中心/部门', '岗位', '新岗位', '入职日期', '转正日期',
                '月度合计(12月薪)', '固浮比', '月基本薪资', '月绩效标准',
                '交通通讯补贴标准', '餐补标准', '防暑降温费标准', '过节费', '独生子女补贴', '实发'])
    pay.append(['编码', None, '工号'])
    pay.append(['009', '原员工', '009', '甲', '甲公司', '旧部门', '新部门', '原岗位', '原岗位',
                datetime(2020, 1, 1), None, 300, .8, '=L3*M3', '=L3-N3'])
    pay['U3'] = ArrayFormula(ref='U3', text='=MAX(N3*{1;2},0)')
    pay['U5'] = '=SUM(U3:U4)'
    pay.auto_filter.ref = 'A1:U5'
    pay.conditional_formatting.add('N3:U4', CellIsRule(operator='lessThan', formula=['0']))
    for title, rows in {
        '人员异动': [['入职情况：'], ['序号', '姓 名', '员工工号', '部门', '组别', '职务', '人员类别', '入职日期', '试用结束日期'],
                   [1, '原员工', '009', '甲', '旧部门'], [], ['离职情况：'], ['序号', '姓 名', '员工工号']],
        '考勤': [['集团编号', '姓名', '成本中心/部门', '入职日期', '餐补标准随工资发', '出勤天数'],
               ['009', '原员工', '旧部门', None, 20, 21], [], [None, None, None, None, None, '=SUM(F2:F3)']],
        '绩效考核': [['人员编码', '姓名', '所属公司', '成本中心/部门', '岗位', '考核成绩'], ['009', '原员工', '甲', '旧部门', '原岗位', 1], [], ['人员编码', '姓名', '其他奖金']],
        '台账': [['列号'], ['集团编号', '姓名', '缴费公司-与劳动合同公司主体一致', '个人养老8%', '个人住房12%'],
               ['009', '原员工', '甲', 24, 36], [None, None, '=SUBTOTAL(3,C3:C3)', '=SUM(D3:D3)', '=SUM(E3:E3)']],
        '防暑降温费': [['集团编号', '姓名', '成本中心/部门', '出勤天数', '班制', '防暑降温费标准', '防暑降温费', '防暑降温费'],
                    ['009', '原员工', '旧部门', 21, None, 600, 0, 0], [], [None, '总经理']],
        '个税申报': [['工号', '姓名', '*证照类型', '*证照号码', '*本期收入'], [None, '原员工', '身份证', 'private', '=工资核算!U3'],
                   [None, '=SUBTOTAL(3,B2:B2)', None, '甲', '=SUM(E2:E2)'], [None, '其他公司员工', None, None, 100], [None, '=SUBTOTAL(3,B4:B4)', None, '乙']],
        '个税导出核对': [['工号', '姓名', '证件类型', '证件号码', '应补(退)税额'], [None, '原员工', '身份证', 'private', 10], [], [None, '=SUBTOTAL(3,B2:B3)']],
        '甲报盘': [[], [None, '收款户名', '收款账号', '金额', '银行'], [None, '原员工', 'private', '=工资核算!U3', 'bank']],
        '乙报盘': [[], [None, '收款户名', '收款账号', '金额', '银行'], [None, '其他公司员工', 'private', 100, 'bank']],
        '汇总': [['核对'], ['=台账!D4'], ['=个税申报!E3']],
    }.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book['人员异动'].merge_cells('A1:B1')
    book['人员异动'].merge_cells('A5:B5')
    dv = DataValidation(type='list', formula1='"身份证,护照"')
    book['个税申报'].add_data_validation(dv)
    dv.add('C2:C5')
    book.save(master)
    notice = Workbook()
    notice.active.append(['入职'])
    notice.active.append(['姓名', '新工号', '新公司', '新成本中心/部门', '入职时间', '试用期结束时间',
                          '岗位', '基本工资', '绩效工资', '车贴', '饭贴', '高温费', '高温费别再',
                          '过节费', '独生子女费', '社保基数', '公积金基数'])
    notice.active.append(['新员工', '010', '甲公司', '新部门', datetime(2026, 5, 12), datetime(2026, 11, 11),
                          '新岗位', 200, 50, 30, 40, 600, '6-9月发放，每月150元', 0, 0, 250, 250])
    notice.save(source)
    return master, source, output


def test_structured_master_preserves_blocks_and_all_notice_fields(tmp_path: Path) -> None:
    master, source, output = structured_files(tmp_path)
    result = apply_semantic_sheet_updates(master, [source], output, salary_month='2026-05')
    assert result['issues'] == []
    assert result['personnel_coverage']['match_rate'] == 1
    book = load_workbook(output)
    pay = book['工资核算']
    assert pay['C4'].value == '010'
    assert pay['C4'].number_format == '@'
    assert pay['L4'].value == '=200+50'
    assert pay['N4'].value == '=L4*M4'
    assert isinstance(pay['U4'].value, ArrayFormula)
    assert pay['U4'].value.ref == 'U4'
    assert pay['U3'].value.text == '=MAX(N3*{1;2},0)'
    assert book['台账']['A4'].value == '010'
    assert book['台账']['D4'].value is None
    assert book['台账']['D5'].value == '=SUM(D3:D4)'
    assert book['汇总']['A2'].value == '=台账!D5'
    assert book['个税申报']['B3'].value == '新员工'
    assert book['个税申报']['B4'].value == '=SUBTOTAL(3,B2:B3)'
    assert book['个税申报']['C3'].value is None
    assert book['个税申报']['D3'].value is None
    assert str(book['个税申报'].data_validations.dataValidation[0].sqref) == 'C2:C6'
    assert book['人员异动']['B4'].value == '新员工'
    assert book['人员异动']['A5'].value == '离职情况：'
    assert book['考勤']['F3'].value is None
    assert book['绩效考核']['F3'].value is None
    assert book['甲报盘']['B4'].value == '新员工'
    assert book['甲报盘']['C4'].value is None
    assert book['乙报盘']['B4'].value is None
    assert len(result['matches']) == 9
    assert result['personnel_coverage']['matched_field_count'] == 17
    second = apply_semantic_sheet_updates(output, [source], tmp_path / 'again.xlsx', salary_month='2026-05')
    assert second['issues'] == []
    assert second['auto_update_count'] == 0


def test_conflicting_structured_identity_writes_nothing(tmp_path: Path) -> None:
    master, source, output = structured_files(tmp_path)
    book = load_workbook(master)
    book['考勤']['A3'] = '010'
    book['考勤']['B3'] = '不同人'
    book.save(master)
    result = apply_semantic_sheet_updates(master, [source], output)
    assert result['auto_update_count'] == 0
    assert result['personnel_coverage']['match_rate'] == 0
