from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from core._personnel_formula_cache import refresh_personnel_formula_caches
from core.sheet_mapper import apply_semantic_sheet_updates
from backend.tests.core.test_structured_hire import structured_files


@pytest.mark.parametrize(('formula', 'expected'), [
    ('=2500+1500', 4000),
    ('=ROUND(4000*0.625,-1)', 2500),
    ('=IF(COUNT(B1:B4)<4,"",9999)', None),
    ('=IFERROR(VLOOKUP("missing",A1:B4,2,0),"missing")', 'missing'),
    ('=VLOOKUP("甲",A1:B4,2,0)', 100),
    ('=SUMIF(A1:A4,"甲",B1:B4)', 300),
    ('=COUNTIF(A1:A4,"甲")', 2),
    ('=SUBTOTAL(3,A1:A4)', 3),
    ('=SUBTOTAL(9,B1:B4)', 600),
    ('=ROUND(MAX(10000*{0.03;0.1;0.2}-{0;2520;16920},0),2)', 300),
    ('=ROUND(-2.675,2)', -2.68),
    ('=IF(1,42,1/0)', 42),
])
def test_payroll_formula_results(formula: str, expected: object, tmp_path: Path) -> None:
    master, output = tmp_path / 'base.xlsx', tmp_path / 'out.xlsx'
    book = Workbook()
    book.active.append(['甲', 100])
    book.active.append(['乙', 300])
    book.active.append(['甲', 200])
    book.save(master)
    book.active['D1'] = formula
    book.save(output)
    refresh_personnel_formula_caches(master, output, [])
    assert load_workbook(output, data_only=True).active['D1'].value == expected


def test_new_payroll_error_stops_success(tmp_path: Path) -> None:
    master, output = tmp_path / 'base.xlsx', tmp_path / 'out.xlsx'
    book = Workbook()
    book.save(master)
    book.active['A1'] = '=1/0'
    book.save(output)
    with pytest.raises(ValueError, match='新的公式错误'):
        refresh_personnel_formula_caches(master, output, [])


def test_calculated_standards_and_empty_monthly_pay_survive_export(tmp_path: Path) -> None:
    master, source, output = structured_files(tmp_path)
    result = apply_semantic_sheet_updates(master, [source], output)
    refresh_personnel_formula_caches(master, output, result['row_insertions'])
    book = load_workbook(output, data_only=True)
    assert book['工资核算']['N4'].value == 200
    assert book['工资核算']['O4'].value == 50
    assert book['工资核算']['U4'].value is None
    assert book['台账']['D5'].value == 24
    assert book['个税申报']['B4'].value == 2
