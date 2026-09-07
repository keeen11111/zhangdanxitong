"""Internal, deterministic handling of explicit standalone onboarding notices.

Personnel standards may fan out to several ledgers. Transactional source sheets
still use the ordinary one-topic mapper. No employee data is inferred from a
neighbouring employee, and all proposals for a person are checked before writing.
"""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import math
from pathlib import Path
import re
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from core.schema_config import ENTITY_BY_ID


_FIELDS = {
    *(field.name for field in ENTITY_BY_ID["employee_profile"].fields),
    "月基本薪资", "月绩效标准", "交通通讯补贴标准", "餐补标准",
    "防暑降温费标准", "过节费", "独生子女补贴", "单位申报基数",
    "公积金基数", "高温费备注",
}
_IDENTITY = {"工号", "姓名", "身份证号"}
_TEXT_FIELDS = _IDENTITY | {"银行卡号"}
_TOTALS = {"合计", "小计", "总计", "汇总"}
_HIRE_TITLE = re.compile(r"^(?:(?:20\d{2}年)?\d{1,2}月)?(?:入职|新增人员|新入职人员|入职人员|入职通知|新增员工)$")
_REFERENCE = re.compile(r"^(\$?[A-Z]{1,3})(\$?)([1-9]\d*)(?::(\$?[A-Z]{1,3})(\$?)([1-9]\d*))?$")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normal(value: Any) -> str:
    return re.sub(r"[\s（）()_：:]", "", _text(value))


_ALIASES: dict[str, str] = {}
_FIELD_TYPES: dict[str, str] = {"公积金基数": "amount"}
for _entity in ENTITY_BY_ID.values():
    for _field in _entity.fields:
        if _field.name in _FIELDS:
            _FIELD_TYPES[_field.name] = _field.dtype
            for _alias in (_field.name, *_field.aliases):
                _ALIASES.setdefault(_normal(_alias), _field.name)
_ALIASES.update({"高温费别再": "高温费备注", "高温费备注": "高温费备注", "公积金基数": "公积金基数"})


def _cell_value(cell: Any, field: str) -> Any:
    value = cell.value
    if field in _TEXT_FIELDS:
        text = _text(value)
        if isinstance(value, (int, float)) and re.fullmatch(r"0+", cell.number_format):
            text = text.zfill(len(cell.number_format))
        return text
    return value


def _validated_value(cell: Any, field: str) -> Any:
    """Normalize only explicit source values; invalid inputs never become zero."""
    value = _cell_value(cell, field)
    if value in (None, ""):
        return value
    if cell.data_type == "e" or isinstance(cell.value, bool):
        raise ValueError("字段含Excel错误值或布尔值")
    if field in _TEXT_FIELDS and isinstance(cell.value, (int, float)):
        number = cell.value
        if not math.isfinite(number) or number < 0 or int(number) != number or len(str(int(number))) > 15:
            raise ValueError("数值型身份标识可能已丢失精度，请使用原始文本标识")
    dtype = _FIELD_TYPES.get(field)
    if dtype == "amount":
        text = _text(value).removeprefix("￥").removeprefix("¥").strip()
        if "," in text and not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):
            raise ValueError("金额的千位分隔符不合法")
        try:
            number = Decimal(text.replace(",", ""))
            if not number.is_finite() or not math.isfinite(float(number)):
                raise ValueError("金额必须为有限数值")
        except (InvalidOperation, OverflowError) as error:
            raise ValueError("金额不是有效数值") from error
        return float(number)
    if dtype == "date":
        if isinstance(value, datetime) and value.tzinfo is None:
            return value
        if isinstance(value, date) and not isinstance(value, datetime):
            return datetime.combine(value, datetime.min.time())
        if isinstance(value, str):
            for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
                try:
                    return datetime.strptime(value.strip(), pattern)
                except ValueError:
                    continue
        raise ValueError("日期不是有效年月日")
    if field == "姓名" and not isinstance(value, str):
        raise ValueError("姓名必须为文本")
    return value


def _formula(cell: Any) -> bool:
    return cell.data_type == "f" or (cell.value is not None and not isinstance(cell.value, (str, int, float, bool, date)))


@dataclass
class _Layout:
    sheet: Any
    header: int
    columns: dict[str, int]
    rows: list[int]
    insert_at: int


def _layout(sheet: Any) -> _Layout | None:
    candidates: list[tuple[int, dict[str, int]]] = []
    for row in range(1, min(sheet.max_row, 30) + 1):
        columns: dict[str, int] = {}
        duplicate = False
        for cell in sheet[row]:
            field = _ALIASES.get(_normal(cell.value))
            if field:
                if field in columns:
                    duplicate = True
                columns[field] = cell.column
        if not duplicate and "姓名" in columns and ("工号" in columns or "身份证号" in columns):
            candidates.append((row, columns))
    if len(candidates) != 1:
        return None
    header, columns = candidates[0]
    rows = []
    insert_at = header + 1
    formula_identity = False
    for row in range(header + 1, sheet.max_row + 1):
        cells = [sheet.cell(row, column) for column in columns.values()]
        if any(_text(cell.value) in _TOTALS for cell in sheet[row]):
            insert_at = row
            break
        if any(_normal(cell.value) in _ALIASES for cell in cells):
            break
        name = sheet.cell(row, columns["姓名"])
        key = sheet.cell(row, columns.get("工号", columns.get("身份证号")))
        formula_identity |= _formula(key) or _formula(name)
        if (_text(name.value) or _text(key.value)) and not _formula(key) and not _formula(name):
            rows.append(row)
            insert_at = row + 1
    if formula_identity:
        return None
    return _Layout(sheet, header, columns, rows, insert_at)


def _issue(kind: str, message: str, record: dict[str, Any], **extra: Any) -> dict[str, Any]:
    return {
        "issue_type": kind, "message": message,
        "source_files": [record["source_file"]], "source_sheets": [record["source_sheet"]],
        "source_row": record.get("source_row"), "person_name": record.get("values", {}).get("姓名"),
        "employee_id": record.get("values", {}).get("工号"),
        **extra,
    }


def _read_notice(sheet: Any, filename: str) -> tuple[bool, list[dict[str, Any]], list[dict[str, Any]]]:
    labels = [sheet.title, *(_text(cell.value) for row in sheet.iter_rows(max_row=min(2, sheet.max_row)) for cell in row)]
    if not any(_HIRE_TITLE.fullmatch(_normal(label)) for label in labels):
        return False, [], []
    source = {"source_file": filename, "source_sheet": sheet.title, "values": {}}
    layout = _layout(sheet)
    if layout is None:
        return True, [], [_issue("hire_missing_identity", "入职通知需要唯一表头及姓名、工号或身份证号。", source)]
    records, issues = [], []
    for row in range(layout.header + 1, sheet.max_row + 1):
        values = {field: _cell_value(sheet.cell(row, column), field) for field, column in layout.columns.items()}
        if not any(value not in (None, "") for value in values.values()):
            continue
        record = {**source, "source_row": row, "values": values}
        if not _text(values.get("姓名")) or not _text(values.get("工号")):
            issues.append(_issue("hire_missing_identity", "入职记录缺少姓名或工号，未猜测人员身份。", record))
            continue
        if any(_formula(sheet.cell(row, column)) for column in layout.columns.values()):
            issues.append(_issue("hire_source_formula", "入职通知包含公式，需要提供明确的人员信息。", record))
            continue
        unknown = [cell.coordinate for cell in sheet[row]
                   if cell.value not in (None, "") and cell.column not in layout.columns.values()]
        if unknown:
            issues.append(_issue("hire_unknown_field", "入职通知存在尚未识别的非空字段，未静默丢弃。", record, source_cells=unknown))
            continue
        invalid_fields = []
        for field, column in layout.columns.items():
            try:
                values[field] = _validated_value(sheet.cell(row, column), field)
            except ValueError:
                invalid_fields.append(field)
        if invalid_fields:
            issues.append(_issue(
                "hire_invalid_value", "入职记录含无效金额、日期或身份标识，未写入该人员。", record,
                fields=invalid_fields,
                source_cells=[sheet.cell(row, layout.columns[field]).coordinate for field in invalid_fields],
            ))
            continue
        records.append(record)
    if not records and not issues:
        issues.append(_issue("empty_source_sheet", "入职通知没有人员记录。", source))
    return True, records, issues


def _only_hire_notices(paths: list[Path]) -> bool:
    """Route only wholly explicit notices; mixed monthly packages keep their pipeline."""
    if not paths:
        return False
    for path in paths:
        if not path.is_file():
            return False
        try:
            book = load_workbook(path, data_only=False)
        except (OSError, ValueError, BadZipFile):
            return False
        try:
            if not all(_read_notice(sheet, path.name)[0] for sheet in book):
                return False
        finally:
            book.close()
    return True


def _shift_reference(token: str, formula_sheet: str, target: str, start: int, expand_boundary: bool = True) -> str:
    qualifier, separator, address = token.rpartition("!")
    if separator:
        if qualifier.strip("'").replace("''", "'") != target:
            return token
    else:
        address = token
        if formula_sheet != target:
            return token
    match = _REFERENCE.fullmatch(address)
    if not match:
        return token
    col, dollar, row, end_col, end_dollar, end_row = match.groups()
    first = int(row)
    last = int(end_row) if end_row else None
    # A data range ending immediately above the inserted row must include it.
    new_first = first + (first >= start)
    result = f"{col}{dollar}{new_first}"
    if last is not None:
        new_last = last + (last >= start or (expand_boundary and first < start and last == start - 1))
        result += f":{end_col}{end_dollar}{new_last}"
    return f"{qualifier}!{result}" if separator else result


def _shift_formula(formula: str, formula_sheet: str, target: str, start: int, expand_boundary: bool = True) -> str:
    tokens = Tokenizer(formula).items
    return "=" + "".join(
        _shift_reference(token.value, formula_sheet, target, start, expand_boundary)
        if token.type == "OPERAND" and token.subtype == "RANGE" else token.value
        for token in tokens
    )


def _audit(cell: Any, value: Any, record: dict[str, Any], updates: list[dict[str, Any]]) -> None:
    def display(item: Any) -> Any:
        return item.text if isinstance(item, ArrayFormula) else item
    if isinstance(cell.value, ArrayFormula) and isinstance(value, ArrayFormula):
        if cell.value.text == value.text and cell.value.ref == value.ref:
            return
    if _text(cell.value) == _text(value):
        return
    updates.append({
        "target_sheet": cell.parent.title, "target_cell": cell.coordinate,
        "old_value": display(cell.value), "new_value": display(value),
        **{key: record[key] for key in ("source_file", "source_sheet", "source_row")},
    })
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "f"


def _insert_person(master: Any, layout: _Layout, record: dict[str, Any], updates: list[dict[str, Any]]) -> int:
    sheet, row = layout.sheet, layout.insert_at
    template = layout.rows[-1] if layout.rows else None
    originals = list(sheet[template]) if template else []
    # Retarget existing references as Excel would, including cross-sheet totals.
    for other in master:
        for cells in other:
            for cell in cells:
                if isinstance(cell.value, str) and cell.data_type == "f":
                    expand = other.title != sheet.title or cell.row >= row
                    _audit(cell, _shift_formula(cell.value, other.title, sheet.title, row, expand), record, updates)
    for defined in master.defined_names.values():
        if defined.attr_text and "!" in defined.attr_text:
            defined.attr_text = _shift_reference(defined.attr_text, sheet.title, sheet.title, row)
    for other in master:
        for defined in other.defined_names.values():
            if defined.attr_text:
                defined.attr_text = _shift_reference(defined.attr_text, other.title, sheet.title, row)
    if sheet.print_area:
        sheet.print_area = _shift_formula("=" + sheet.print_area, sheet.title, sheet.title, row)[1:]
    if sheet.freeze_panes:
        sheet.freeze_panes = _shift_reference(str(sheet.freeze_panes), sheet.title, sheet.title, row)
    # Move audited addresses along with their cells before recording new values.
    for update in updates:
        if update["target_sheet"] == sheet.title:
            cell = sheet[update["target_cell"]]
            if cell.row >= row:
                update["target_cell"] = f"{get_column_letter(cell.column)}{cell.row + 1}"
    sheet.insert_rows(row)
    for index in sorted(list(sheet.row_dimensions), reverse=True):
        if index >= row:
            dimension = copy(sheet.row_dimensions[index])
            dimension.index = index + 1
            sheet.row_dimensions[index + 1] = dimension
    if template:
        sheet.row_dimensions[row] = copy(sheet.row_dimensions[template])
        sheet.row_dimensions[row].index = row
    for src in originals:
        dst = sheet.cell(row, src.column)
        dst._style = copy(src._style)
        if isinstance(src.value, str) and src.data_type == "f":
            _audit(dst, Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate), record, updates)
    for merged in sheet.merged_cells.ranges:
        if merged.min_row >= row:
            merged.shift(row_shift=1)
    return row


def _insertion_error(master: Any, layout: _Layout) -> str | None:
    sheet, row = layout.sheet, layout.insert_at
    if row <= layout.header:
        return "无法确定人员明细的新增行位置。"
    # Unsupported structures must never be silently corrupted by openpyxl.
    if sheet.tables or sheet.data_validations.count or len(sheet.conditional_formatting) or sheet._charts or sheet._images or sheet.auto_filter.ref:
        return "目标表页含表格对象、数据验证或图表，尚不支持可靠地扩展其人员行。"
    if any(merged.min_row <= row <= merged.max_row for merged in sheet.merged_cells.ranges):
        return "新增位置涉及合并单元格，无法安全插行。"
    if any(_formula(cell) and not isinstance(cell.value, str) for other in master for cells in other for cell in cells):
        return "工作簿含数组公式，无法安全自动调整新增行的引用。"
    return None


def _apply_hire_notices(master: Any, sources: list[tuple[Path, str, Any]], policy_action: Any,
                        salary_month: str | None = None) -> dict[str, Any]:
    consumed: set[tuple[str, str]] = set()
    records, issues, updates, matches = [], [], [], []
    mixed_batch = False
    for path, filename, book in sources:
        for sheet in book:
            handled, found, errors = _read_notice(sheet, filename)
            if handled:
                consumed.add((str(path.resolve()), sheet.title))
                records.extend(found)
                issues.extend(errors)
            else:
                mixed_batch = True
    # Monthly mixed packages retain the existing conflict/priority resolver.
    # Never mutate hire cells before another source's proposals are checked.
    if mixed_batch:
        consumed.clear()
        records.clear()
        issues.clear()
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[_text(record["values"]["工号"]).upper()].append(record)
    ids_by_name: dict[str, set[str]] = defaultdict(set)
    for employee_id, group in grouped.items():
        for record in group:
            ids_by_name[_text(record["values"]["姓名"])].add(employee_id)
    invalid_rows = [issue for issue in issues if issue.get("source_row") is not None]
    invalid_ids = {_text(issue.get("employee_id")).upper() for issue in invalid_rows if _text(issue.get("employee_id"))}
    invalid_names = {_text(issue.get("person_name")) for issue in invalid_rows
                     if not _text(issue.get("employee_id")) and _text(issue.get("person_name"))}
    processed = 0
    matched_fields = input_fields = 0
    structured = False
    for employee_id, group in grouped.items():
        record = group[0]
        if employee_id in invalid_ids or any(_text(item["values"]["姓名"]) in invalid_names for item in group):
            continue
        if any(len(ids_by_name[_text(item["values"]["姓名"])]) > 1 for item in group):
            issues.append(_issue("hire_identity_conflict", "入职通知中同一姓名对应多个工号，未猜测人员关系。", record))
            continue
        merged: dict[str, Any] = {}
        field_sources: dict[str, dict[str, Any]] = {}
        conflicts = set()
        for item in group:
            for field, value in item["values"].items():
                if value in (None, ""):
                    continue
                if field in merged and _text(merged[field]) != _text(value):
                    conflicts.add(field)
                merged[field] = value
                field_sources[field] = item
        if conflicts:
            issues.append(_issue("hire_conflicting_source", "多份入职通知对同一人员提供了不同值。", record, fields=sorted(conflicts)))
            continue
        record = {**record, "values": merged}
        input_fields += len(merged)
        from core._structured_hire import _apply_structured_person
        handled, structured_errors, covered = _apply_structured_person(
            master, record, field_sources, policy_action, salary_month, updates, matches,
        )
        if handled:
            structured = True
            issues.extend(structured_errors)
            if not structured_errors:
                processed += 1
                matched_fields += len(covered)
            continue
        layouts = []
        plans: list[tuple[_Layout, int | None, dict[str, int]]] = []
        errors = []
        for sheet in master:
            layout = _layout(sheet)
            if layout:
                layouts.append(layout)
                continue
            for cells in sheet.iter_rows(max_row=min(sheet.max_row, 30)):
                fields = {_ALIASES.get(_normal(cell.value)) for cell in cells} - {None}
                if {"工号", "姓名"} <= fields and (fields & merged.keys()) - _IDENTITY:
                    errors.append(_issue("hire_unsupported_layout", "相关人员表页的表头、区块或身份公式无法唯一解析。", record, target_sheet=sheet.title))
                    break
        for layout in layouts:
            columns = layout.columns
            hits = []
            for row in layout.rows:
                old_id = _text(_cell_value(layout.sheet.cell(row, columns["工号"]), "工号")).upper() if "工号" in columns else ""
                old_name = _text(layout.sheet.cell(row, columns["姓名"]).value)
                old_card = _text(_cell_value(layout.sheet.cell(row, columns["身份证号"]), "身份证号")).upper() if "身份证号" in columns else ""
                new_card = _text(merged.get("身份证号")).upper()
                if old_id == employee_id or old_name == _text(merged["姓名"]) or (old_card and old_card == new_card):
                    if (old_id and old_id != employee_id) or (old_name and old_name != _text(merged["姓名"])) or (new_card and old_card and new_card != old_card):
                        errors.append(_issue("hire_identity_conflict", "总表的工号与姓名对应关系冲突，未写入该人员。", record, target_sheet=layout.sheet.title))
                    else:
                        hits.append(row)
            if len(hits) > 1:
                errors.append(_issue("hire_identity_conflict", "同一人员在目标表页有多条记录。", record, target_sheet=layout.sheet.title))
            mapped = {field: column for field, column in columns.items() if field in merged and merged[field] not in (None, "")}
            if not (mapped.keys() - _IDENTITY) or "工号" not in columns:
                continue
            row = hits[0] if len(hits) == 1 else None
            for field, column in mapped.items():
                if policy_action(record["source_sheet"], field) in {"ignore", "review"}:
                    errors.append(_issue("configured_field_review", "匹配策略要求暂停该入职字段。", record, target_field=field))
                if row and _formula(layout.sheet.cell(row, column)):
                    errors.append(_issue("formula_target_conflict", "人员字段目标为公式，未覆盖。", record, target_sheet=layout.sheet.title, target_field=field))
            if row is None and (reason := _insertion_error(master, layout)):
                errors.append(_issue("hire_unsupported_layout", reason, record, target_sheet=layout.sheet.title))
            plans.append((layout, row, mapped))
        covered_fields = {field for _, _, mapped in plans for field in mapped}
        unmapped = set(merged) - covered_fields
        if plans and unmapped:
            errors.append(_issue("hire_unmapped_fields", "入职通知的部分字段没有可靠的总表目标，未报告完全匹配。", record, fields=sorted(unmapped)))
        if errors:
            issues.extend(errors)
            continue
        if not plans:
            issues.append(_issue("hire_unmatched_target", "未找到有明确工号及相关人员字段的目标明细表页。", record))
            continue
        for layout, row, mapped in plans:
            is_new = row is None
            if row is None:
                row = _insert_person(master, layout, record, updates)
            for field, column in mapped.items():
                cell = layout.sheet.cell(row, column)
                _audit(cell, merged[field], field_sources[field], updates)
                if is_new and field in _TEXT_FIELDS:
                    cell.number_format = "@"
                elif isinstance(merged[field], (date, datetime)) and cell.number_format == "General":
                    cell.number_format = "yyyy-mm-dd"
            matches.append({"source_file": record["source_file"], "source_sheet": record["source_sheet"], "target_sheet": layout.sheet.title, "confidence": 1.0, "match_basis": "explicit_hire_employee_id"})
        processed += 1
        matched_fields += len(covered_fields)
    identified_names = set(ids_by_name) | {
        _text(issue.get("person_name")) for issue in invalid_rows if _text(issue.get("employee_id"))
    }
    unidentified_rows = {
        (tuple(issue["source_files"]), tuple(issue["source_sheets"]), issue["source_row"])
        for issue in invalid_rows if not _text(issue.get("employee_id")) and not _text(issue.get("person_name"))
    }
    total = len(set(grouped) | invalid_ids) + len(invalid_names - identified_names) + len(unidentified_rows)
    return {
        "consumed": consumed, "issues": issues, "updates": updates, "matches": matches,
        "structured_hire": structured,
        "personnel_coverage": {"input_person_count": total, "processed_person_count": processed,
                               "unprocessed_person_count": total - processed,
                               "input_field_count": input_fields, "matched_field_count": matched_fields,
                               "match_rate": processed / total if total else None},
    }
