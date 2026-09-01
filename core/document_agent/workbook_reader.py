"""Bounded read-only workbook observations shared by planning and execution."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .orchestrator import ToolExecutionError, ToolRegistry


def bounded_range(area: str, limit: int = 1000) -> tuple[int, int, int, int]:
    try:
        left, top, right, bottom = range_boundaries(area)
        if (None in (left, top, right, bottom) or not 1 <= left <= right <= 16384
                or not 1 <= top <= bottom <= 1048576 or (right-left+1)*(bottom-top+1) > limit):
            raise ValueError()
        return left, top, right, bottom
    except (TypeError, ValueError) as exc:
        raise ToolExecutionError("范围必须是有效的有限单元格区域，每次最多1000格") from exc


class WorkbookReader:
    def __init__(self, master: Path, sources: dict[str, Path], *, master_name: str | None = None,
                 target: Callable[[], Path] | None = None):
        self.master = master
        self.sources = sources
        self.master_name = master_name or master.name
        self.target = target or (lambda: master)

    def _path(self, filename: str) -> Path:
        if filename == self.master_name:
            return self.target()
        if filename not in self.sources:
            raise ToolExecutionError("文件不在当前任务只读白名单内")
        return self.sources[filename]

    def _read(self, path: Path, sheet: str, area: str) -> dict[str, Any]:
        left, top, right, bottom = bounded_range(area)
        books = []
        try:
            books = [openpyxl.load_workbook(path, read_only=True, data_only=False)]
            books.append(openpyxl.load_workbook(path, read_only=True, data_only=True))
            if sheet not in books[0]:
                raise ToolExecutionError("工作表不存在")
            rows, cached = [[list(row) for row in book[sheet].iter_rows(
                min_row=top, max_row=bottom, min_col=left, max_col=right, values_only=True,
            )] for book in books]
            missing = [f"{get_column_letter(left+c)}{top+r}" for r, row in enumerate(rows)
                       for c, value in enumerate(row)
                       if isinstance(value, str) and value.startswith("=") and cached[r][c] is None]
            return {"sheet": sheet, "range": area, "rows": rows, "cached_rows": cached,
                    "missing_formula_caches": missing}
        finally:
            for book in books:
                book.close()

    def read_range(self, sheet: str, range: str) -> dict[str, Any]:
        return self._read(self.target(), sheet, range)

    def read_source_range(self, filename: str, sheet: str, range: str) -> dict[str, Any]:
        if filename not in self.sources:
            raise ToolExecutionError("来源文件不在当前任务白名单内")
        return self._read(self.sources[filename], sheet, range)

    def search_cells(self, filename: str, sheet: str, query: str, exact: bool) -> dict[str, Any]:
        if not isinstance(query, str) or not 1 <= len(query) <= 200 or not isinstance(exact, bool):
            raise ToolExecutionError("搜索条件不合法")
        workbook = openpyxl.load_workbook(self._path(filename), read_only=True, data_only=False)
        try:
            if sheet and sheet not in workbook:
                raise ToolExecutionError("工作表不存在")
            matches = []
            complete, inspected = True, 0
            for ws in ([workbook[sheet]] if sheet else workbook.worksheets):
                # Formatting noise must not explode a lookup into billions of cells.
                columns = min(ws.max_column or 1, 256)
                rows = min(ws.max_row or 1, 20000, max(0, (1_000_000-inspected)//columns))
                complete = complete and columns == ws.max_column and rows == ws.max_row
                for row in ws.iter_rows(max_row=rows, max_col=columns):
                    for cell in row:
                        value = cell.value
                        if value is not None and (str(value).strip() == query.strip() if exact else query in str(value)):
                            matches.append({"sheet": ws.title, "cell": cell.coordinate, "value": value})
                            if len(matches) >= 100:
                                return {"matches": matches, "complete": False, "limit": 100}
                inspected += rows*columns
                if inspected >= 1_000_000:
                    complete = False
                    break
            return {"matches": matches, "complete": complete}
        finally:
            workbook.close()

    def register(self, registry: ToolRegistry) -> list[dict[str, Any]]:
        definitions = [
            ("read_range", "读取当前总表/草稿范围及公式缓存，最多1000格。", {"sheet": "string", "range": "string"}),
            ("read_source_range", "读取白名单来源范围及公式缓存，最多1000格。", {"filename": "string", "sheet": "string", "range": "string"}),
            ("search_cells", "在总表或来源中定位姓名/表头；sheet为空搜索各表。完整性不足时缩小范围，不可认定无匹配。", {"filename": "string", "sheet": "string", "query": "string", "exact": "boolean"}),
        ]
        schemas = []
        for name, description, fields in definitions:
            registry.register(name, getattr(self, name))
            schemas.append({"type": "function", "function": {"name": name, "description": description,
                "parameters": {"type": "object", "properties": {key: {"type": kind} for key, kind in fields.items()},
                               "required": list(fields), "additionalProperties": False}}})
        return schemas
