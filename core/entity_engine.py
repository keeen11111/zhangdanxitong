"""实体引擎：源数据 → 业务实体 → 模板导出。

导入流程：
    1. 读取源 Excel 文件
    2. 根据 schema_config 的映射规则，将源列名匹配到实体字段
    3. 按主键（工号/身份证号）去重合并，生成 5 个实体数据集

导出流程：
    1. 读取 Master 模板（保留公式与样式）
    2. 以 employee_profile 为基准，LEFT JOIN 其他实体
    3. 按 EXPORT_MAPPING 将字段写入对应列位
    4. 公式列保持模板原样，不覆盖
"""
from __future__ import annotations

import os
import copy
import re
from datetime import datetime, date
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.formula import ArrayFormula

from core.schema_config import (
    ALL_ENTITIES,
    ENTITY_BY_ID,
    EXPORT_MAPPING,
    EXPORT_BY_COLUMN,
    detect_file_type,
    EntityDef,
    EntityField,
)
from core.normalizer import DataNormalizer


_normalizer = DataNormalizer()


def _formula_text(value: Any) -> str | None:
    """Return formula text for normal and legacy array formulas."""
    if isinstance(value, str) and value.startswith("="):
        return value
    if isinstance(value, ArrayFormula) and isinstance(value.text, str):
        return value.text
    return None


def _translated_formula(value: Any, origin: str, target: str) -> str | ArrayFormula | None:
    """Translate a formula while preserving openpyxl's array-formula wrapper."""
    formula = _formula_text(value)
    if formula is None:
        return None
    translated = Translator(formula, origin=origin).translate_formula(target)
    if isinstance(value, ArrayFormula):
        return ArrayFormula(ref=target, text=translated)
    return translated

_IMPORT_FIELD_ALIASES: dict[str, list[str]] = {
    "工号": ["员工编号", "人员编码", "集团编码", "职工编号"],
    "姓名": ["人员姓名", "员工姓名", "业务员", "值班人员姓名"],
    "业绩奖金": ["本月实发合计", "门店业绩提成", "Q4奖金应发合计"],
    "DTP绩效奖金": ["奖金金额"],
    "其它工资调差": ["值班费合计"],
    "年终奖": ["年终奖纳税调整"],
}


def _field_candidates(field: EntityField) -> list[str]:
    """Return schema aliases plus conservative aliases used by uploaded ledgers."""
    return [field.name, *field.aliases, *_IMPORT_FIELD_ALIASES.get(field.name, [])]


# ============================================================
# 导入：源数据 → 实体
# ============================================================

def _normalize_header(val: Any) -> str:
    """规范化列名用于匹配：去空白、去换行、统一命名。"""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(val).strip()
    # 去除换行符和多余空格
    s = s.replace("\n", "").replace("\r", "").replace("  ", " ").strip()
    return s


def _find_header_row(df_raw: pd.DataFrame, entity: EntityDef, max_scan: int = 30) -> int:
    """嗅探真实表头所在行号（0-based）。

    策略：扫描前 max_scan 行，找到包含最多实体字段别名的行作为表头。
    如果主键字段是必需的，则表头行必须包含主键列，否则跳过该行。
    """
    # 找到主键字段定义
    pk_field = None
    for f in entity.fields:
        if f.name == entity.primary_key:
            pk_field = f
            break

    best_row = 0
    best_score = -1
    for row_idx in range(min(max_scan, len(df_raw))):
        row_values = df_raw.iloc[row_idx].tolist()
        row_headers = [_normalize_header(v) for v in row_values]

        # 人员业务表常只有姓名；允许姓名作为关联入口，后续再由花名册补齐工号。
        if pk_field and pk_field.required:
            pk_candidates = _field_candidates(pk_field)
            name_field = next((field for field in entity.fields if field.name == "姓名"), None)
            identity_candidates = pk_candidates + (_field_candidates(name_field) if name_field else [])
            has_identity = False
            for h in row_headers:
                if not h:
                    continue
                if h in identity_candidates:
                    has_identity = True
                    break
                # 模糊匹配主键
                for c in identity_candidates:
                    if len(c) >= 2 and c in h:
                        has_identity = True
                        break
                if has_identity:
                    break
            if not has_identity:
                continue

        score = 0
        for field in entity.fields:
            # 检查字段名本身和所有别名
            candidates = _field_candidates(field)
            for h in row_headers:
                if h in candidates:
                    score += 1
                    break
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def _match_columns(headers: list[str], entity: EntityDef) -> dict[str, str]:
    """匹配源列名 → 实体字段名。

    返回: {源列名: 实体字段名}
    """
    mapping: dict[str, str] = {}
    normalized_headers = {h: _normalize_header(h) for h in headers}

    for field in entity.fields:
        candidates = _field_candidates(field)
        matched = False
        for src_header, norm_header in normalized_headers.items():
            if matched:
                break
            if src_header in mapping:
                continue
            for candidate in candidates:
                if norm_header == candidate:
                    mapping[src_header] = field.name
                    matched = True
                    break
                # 模糊匹配：源列名包含候选名
                if (
                    len(candidate) >= 3
                    and candidate not in {"加班时长"}
                    and candidate in norm_header
                ):
                    mapping[src_header] = field.name
                    matched = True
                    break
    return mapping


def _clean_value(val: Any, dtype: str) -> Any:
    """按字段类型清洗值。"""
    if isinstance(val, pd.Series):
        candidates = [
            item
            for item in val.tolist()
            if item is not None and not (isinstance(item, float) and pd.isna(item))
        ]
        val = candidates[0] if candidates else None
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if dtype == "text":
        return _normalizer.clean_text(val)
    elif dtype == "amount":
        return _normalizer.clean_amount(val)
    elif dtype == "date":
        return _normalizer.clean_date(val)
    elif dtype == "int":
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return 0
    return val


def _parse_sheet_to_entity(
    df_raw: pd.DataFrame,
    entity: EntityDef,
    header_row: int | None = None,
    source_file: str = "",
    source_sheet: str = "",
) -> tuple[list[dict], dict[str, str]]:
    """将一个 Sheet 解析为实体记录列表。

    Args:
        df_raw: 原始 DataFrame（未处理表头）
        entity: 目标实体定义
        header_row: 指定表头行号（0-based），None 则自动嗅探
        source_file: 源文件名（用于溯源）
        source_sheet: 源 Sheet 名（用于溯源）

    Returns:
        (记录列表, 列名匹配结果)
    """
    if df_raw.empty:
        return [], {}

    # 确定表头行
    if header_row is None:
        header_row = _find_header_row(df_raw, entity)
    if header_row >= len(df_raw):
        return [], {}

    # 提取表头行作为列名
    headers = [_normalize_header(h) for h in df_raw.iloc[header_row].tolist()]

    # pandas 对重复列名执行 row.get() 时会返回 Series，随后被清洗成
    # “姓名...Name:...dtype:object”一类伪姓名。保留第一列原名，后续同名列
    # 加内部后缀，使实体字段稳定匹配到靠前的主数据列。
    header_counts: dict[str, int] = {}
    unique_headers: list[str] = []
    for header in headers:
        count = header_counts.get(header, 0) + 1
        header_counts[header] = count
        unique_headers.append(header if count == 1 else f"{header}__{count}")

    # 数据行（表头行之后）
    data_df = df_raw.iloc[header_row + 1:].copy()
    data_df.columns = unique_headers
    # 去除全空行
    data_df = data_df.dropna(how="all")

    # 匹配列名
    col_mapping = _match_columns(unique_headers, entity)

    effective_date_column = next(
        (
            header
            for header in unique_headers
            if _normalize_header(header).split("__", 1)[0]
            in {"生效日期", "变更日期", "异动日期", "调整日期", "调薪日期", "发生日期"}
        ),
        None,
    )

    if not col_mapping:
        return [], {}

    # 提取数据并清洗
    records: list[dict] = []
    for _, row in data_df.iterrows():
        record: dict[str, Any] = {}
        for src_col, field_name in col_mapping.items():
            # 找到字段定义
            field_def = None
            for f in entity.fields:
                if f.name == field_name:
                    field_def = f
                    break
            if field_def is None:
                continue
            raw_val = row.get(src_col)
            record[field_name] = _clean_value(raw_val, field_def.dtype)

        # 溯源信息
        record["_source_file"] = source_file
        record["_source_sheet"] = source_sheet
        if effective_date_column:
            effective_date = _clean_value(row.get(effective_date_column), "date")
            if effective_date:
                record["_effective_date"] = effective_date

        if entity.id == "attendance_record":
            overtime_parts = [
                record.pop("工作日加班时长", None),
                record.pop("公休日加班时长", None),
            ]
            explicit_total = record.get("普通加班时长")
            if explicit_total in (None, "") and any(
                value not in (None, "") for value in overtime_parts
            ):
                record["普通加班时长"] = sum(
                    float(value) for value in overtime_parts if value not in (None, "")
                )

        # 检查主键是否有值
        pk_val = record.get(entity.primary_key)
        name_val = record.get("姓名")
        if (pk_val and str(pk_val).strip()) or (name_val and str(name_val).strip()):
            records.append(record)

    return records, col_mapping


_SEMANTIC_IDENTITY_FIELDS = {"工号", "姓名", "身份证号"}


def _parse_sheet_semantically(
    df_raw: pd.DataFrame,
    source_file: str,
    source_sheet: str,
) -> dict[str, tuple[list[dict], dict[str, str]]]:
    """Infer applicable entities from column meaning, independent of names.

    One sheet may legitimately contain profile, salary, attendance and tax
    fields together.  Identity-only matches are retained only for the employee
    profile; transactional entities need at least one entity-specific field so
    a generic 姓名/工号 pair does not create false records everywhere.
    """
    parsed: dict[str, tuple[list[dict], dict[str, str]]] = {}
    for entity in ALL_ENTITIES:
        records, column_mapping = _parse_sheet_to_entity(
            df_raw,
            entity,
            None,
            source_file,
            source_sheet,
        )
        if not records:
            continue

        mapped_fields = set(column_mapping.values())
        business_fields = mapped_fields - _SEMANTIC_IDENTITY_FIELDS
        if entity.id == "employee_profile":
            has_identity_pair = (
                "姓名" in mapped_fields
                and bool({"工号", "身份证号"} & mapped_fields)
            )
            if not business_fields and not has_identity_pair:
                continue
        elif entity.id == "social_security":
            contribution_fields = business_fields - {"缴费公司"}
            if not contribution_fields:
                continue
        elif not business_fields:
            continue

        parsed[entity.id] = (records, column_mapping)
    return parsed


def _append_semantic_sheet_results(
    result: dict[str, list[dict]],
    log: list[str],
    df_raw: pd.DataFrame,
    source_file: str,
    source_sheet: str,
) -> None:
    parsed = _parse_sheet_semantically(df_raw, source_file, source_sheet)
    if not parsed:
        log.append(f"  Sheet「{source_sheet}」: 无可匹配数据")
        return
    for entity_id, (records, column_mapping) in parsed.items():
        result[entity_id].extend(records)
        entity = ENTITY_BY_ID[entity_id]
        log.append(
            f"  Sheet「{source_sheet}」→ 实体「{entity.name}」: "
            f"{len(records)} 条记录, 语义匹配 {len(column_mapping)} 列"
        )


_SOCIAL_NAME_HEADERS = ("姓名", "员工姓名", "人员姓名")
_SOCIAL_ID_HEADERS = ("身份证", "证件号码", "证件号", "身份证号码")


def _social_field_from_header(header: str) -> str | None:
    """Map a flattened social-billing header to one canonical entity field."""
    value = _normalize_header(header).replace(" ", "")
    if not value:
        return None
    if any(token in value for token in _SOCIAL_NAME_HEADERS):
        return "姓名"
    if any(token in value for token in _SOCIAL_ID_HEADERS):
        return "身份证号"
    if any(token in value for token in ("缴费公司", "缴费单位", "参保单位", "单位名称")):
        return "缴费公司"

    for keyword, field in (
        ("单位申报基数", "单位申报基数"),
        ("养老基数", "养老基数"),
        ("失业基数", "失业基数"),
        ("工伤基数", "工伤基数"),
        ("医疗基数", "医疗基数"),
    ):
        if keyword in value:
            return field

    if "补充医疗" in value or "商业保险" in value or "商保" in value:
        return "公司补充医疗"
    if "基数" in value or "合计" in value or "小计" in value:
        return None

    is_personal = any(token in value for token in ("个人", "职工", "员工"))
    is_company = any(token in value for token in ("单位", "公司", "企业"))
    if not (is_personal or is_company):
        return None
    party = "个人" if is_personal else "公司"
    if "养老" in value:
        return f"{party}养老"
    if "失业" in value:
        return f"{party}失业"
    if "工伤" in value:
        return f"{party}工伤"
    if "生育" in value:
        return f"{party}生育" if party == "公司" else None
    if "医疗" in value:
        return f"{party}医疗"
    if "公积金" in value or "住房" in value:
        return f"{party}公积金"
    # Some legacy bills omit the insurance category on the detail row, while
    # retaining the standard contribution rate.  Use this only as a fallback
    # after semantic labels have been exhausted.
    if party == "个人" and "8%" in value:
        return "个人养老"
    if party == "个人" and "2%" in value:
        return "个人医疗"
    if party == "个人" and "0.5%" in value:
        return "个人失业"
    if party == "公司" and "16%" in value:
        return "公司养老"
    if party == "公司" and "9.8%" in value:
        return "公司医疗"
    if party == "公司" and "0.5%" in value:
        return "公司失业"
    if party == "公司" and "0.2%" in value:
        return "公司工伤"
    return None


def _social_billing_layout(df_raw: pd.DataFrame) -> tuple[int, dict[int, str]] | None:
    """Find a social-billing header band and return its first data row/field columns.

    Supplier bills often use merged two-to-four-row headers.  This detection
    flattens the header band column by column and relies on field semantics,
    rather than a supplier name, sheet name, or fixed column positions.
    """
    if df_raw.empty:
        return None
    scan_limit = min(len(df_raw), 12)
    for primary_row in range(scan_limit):
        primary_headers = [_normalize_header(value) for value in df_raw.iloc[primary_row].tolist()]
        name_columns = [
            index for index, header in enumerate(primary_headers)
            if _social_field_from_header(header) == "姓名"
        ]
        id_columns = [
            index for index, header in enumerate(primary_headers)
            if _social_field_from_header(header) == "身份证号"
        ]
        if not name_columns or not id_columns:
            continue

        name_column = name_columns[0]
        id_column = id_columns[0]
        data_start = None
        for row_index in range(primary_row + 1, len(df_raw)):
            name = _clean_value(df_raw.iloc[row_index, name_column], "text")
            identity = _clean_value(df_raw.iloc[row_index, id_column], "text")
            if not name:
                continue
            if _social_field_from_header(str(name)) in {"姓名", "身份证号"}:
                continue
            if identity or name:
                data_start = row_index
                break
        if data_start is None or data_start == primary_row + 1:
            continue

        flattened_headers: list[str] = ["" for _ in df_raw.columns]
        for row_index in range(primary_row, data_start):
            row_labels = [_normalize_header(value) for value in df_raw.iloc[row_index].tolist()]
            is_detail_row = any(
                any(token in label for token in ("个人", "公司", "单位缴纳", "%"))
                for label in row_labels
            )
            carry = ""
            for column_index, label in enumerate(row_labels):
                if label:
                    carry = label
                # The primary row carries independent columns; only lower
                # group-header rows are horizontally expanded across merges.
                # A merged primary group (养老保险、社会保险、公积金) is also
                # expanded, while ordinary headers such as 医疗基数 are not.
                primary_group = any(token in carry for token in ("保险", "公积金", "社会保险", "住房"))
                if row_index == primary_row:
                    part = label or (carry if primary_group else "")
                else:
                    part = label if is_detail_row else carry
                if part and part not in flattened_headers[column_index]:
                    flattened_headers[column_index] += part

        field_columns = {
            column_index: field
            for column_index, header in enumerate(flattened_headers)
            if (field := _social_field_from_header(header))
        }
        contribution_count = sum(
            field.startswith(("个人", "公司"))
            for field in field_columns.values()
        )
        if (
            "姓名" in field_columns.values()
            and "身份证号" in field_columns.values()
            and contribution_count >= 2
        ):
            return data_start, field_columns
    return None


def _is_social_billing_sheet(df_raw: pd.DataFrame) -> bool:
    return _social_billing_layout(df_raw) is not None


def _parse_social_billing_sheet(
    df_raw: pd.DataFrame,
    source_file: str,
    source_sheet: str,
) -> tuple[list[dict], dict[str, str]]:
    """Parse any recognized grouped social-insurance bill into canonical fields."""
    layout = _social_billing_layout(df_raw)
    if layout is None:
        return [], {}
    data_start, field_columns = layout
    columns = {field: index for index, field in field_columns.items()}
    text_fields = {"姓名", "身份证号", "缴费公司"}
    records: list[dict[str, Any]] = []
    for _, row in df_raw.iloc[data_start:].dropna(how="all").iterrows():
        record = {
            field: _clean_value(row.iloc[index], "text" if field in text_fields else "amount")
            for field, index in columns.items()
        }
        # Settlement rows often retain a placeholder identity (for example
        # "——") together with contribution totals.  A ledger row is matched
        # by a real employee name, so rows without one must not be imported.
        if not record.get("姓名"):
            continue

        personal_amounts = [
            record[field]
            for field in ("个人养老", "个人医疗", "个人公积金", "个人失业")
            if record.get(field) is not None
        ]
        company_amounts = [
            record[field]
            for field in ("公司养老", "公司医疗", "公司公积金", "公司失业", "公司工伤", "公司生育", "公司补充医疗")
            if record.get(field) is not None
        ]
        record["个人合计"] = round(sum(personal_amounts), 2) if personal_amounts else None
        record["公司合计"] = round(sum(company_amounts), 2) if company_amounts else None
        record["_source_file"] = source_file
        record["_source_sheet"] = source_sheet
        records.append(record)
    return records, {field: str(index + 1) for field, index in columns.items()}


# ============================================================
# 分区表格解析（入离职表等）
# ============================================================

# 区块关键词 → 区块类型
BLOCK_KEYWORDS: list[tuple[str, str]] = [
    ("入职", "hire"),
    ("离职", "depart"),
    ("退休", "depart"),
    ("转正", "confirm"),
    ("调动", "transfer"),
    ("其他", "other"),
]


def _detect_block_type(text: str) -> str | None:
    """检测文本是否是区块标题，返回区块类型。"""
    s = text.strip()
    if not s:
        return None
    for keyword, block_type in BLOCK_KEYWORDS:
        if keyword in s:
            return block_type
    return None


def _is_block_title_row(df_raw: pd.DataFrame, row_idx: int) -> tuple[bool, str | None]:
    """判断某行是否是区块标题行。

    判断标准：第一列有值且是区块关键词，其余列基本为空。
    """
    if row_idx >= len(df_raw):
        return False, None
    row = df_raw.iloc[row_idx].tolist()
    first_val = row[0] if row else None
    if first_val is None or (isinstance(first_val, float) and pd.isna(first_val)):
        return False, None
    first_str = str(first_val).strip()
    if not first_str:
        return False, None

    block_type = _detect_block_type(first_str)
    if block_type is None:
        return False, None

    # 验证：其余列应基本为空（允许1个非空单元格作为容错）
    non_empty = sum(
        1 for v in row[1:]
        if v is not None and not (isinstance(v, float) and pd.isna(v)) and str(v).strip()
    )
    if non_empty <= 1:
        return True, block_type
    return False, None


def _parse_block_sheet(
    df_raw: pd.DataFrame,
    source_file: str,
    source_sheet: str,
    name_to_id: dict[str, str],
) -> tuple[dict[str, list[dict]], list[str]]:
    """解析分区表格（如入离职、转岗、转正表）。

    每个区块有独立的标题行和表头行，数据行到下一个区块标题或"无"结束。

    Args:
        df_raw: 原始 DataFrame
        source_file: 源文件名
        source_sheet: 源 Sheet 名
        name_to_id: 姓名→工号 映射（用于离职/转正等无工号区块的关联）

    Returns:
        ({实体ID: [记录]}, 日志列表)
    """
    result: dict[str, list[dict]] = {"employee_profile": [], "salary_detail": []}
    log: list[str] = []

    # 扫描所有行，找出区块边界
    blocks: list[tuple[int, str, int]] = []  # (标题行, 区块类型, 表头行)
    for i in range(len(df_raw)):
        is_title, btype = _is_block_title_row(df_raw, i)
        if is_title and btype:
            header_row = i + 1
            if header_row < len(df_raw):
                blocks.append((i, btype, header_row))

    if not blocks:
        log.append(f"  Sheet「{source_sheet}」: 未检测到分区标题")
        return result, log

    for idx, (title_row, block_type, header_row) in enumerate(blocks):
        # 区块结束行 = 下一个区块标题行 或 文件末尾
        end_row = blocks[idx + 1][0] if idx + 1 < len(blocks) else len(df_raw)

        # 提取表头
        headers = [_normalize_header(v) for v in df_raw.iloc[header_row].tolist()]

        # 提取数据行
        data_count = 0
        for row_idx in range(header_row + 1, end_row):
            row_values = df_raw.iloc[row_idx].tolist()
            first_cell = row_values[0] if row_values else None
            first_str = str(first_cell).strip() if first_cell is not None and not (
                isinstance(first_cell, float) and pd.isna(first_cell)
            ) else ""

            # 跳过"无"标记和空行
            if not first_str or first_str == "无":
                continue
            # 跳过嵌套的区块标题
            if _detect_block_type(first_str):
                continue

            # 构建行数据字典（表头→值）
            row_dict = {}
            for col_idx, h in enumerate(headers):
                if col_idx < len(row_values) and h:
                    row_dict[h] = row_values[col_idx]

            # 根据区块类型解析
            emp_rec, sal_rec = _parse_block_row(row_dict, block_type, name_to_id, source_file, source_sheet)
            if emp_rec:
                result["employee_profile"].append(emp_rec)
                if sal_rec:
                    result["salary_detail"].append(sal_rec)
                data_count += 1
                # 更新姓名→工号映射
                name = emp_rec.get("姓名")
                emp_id = emp_rec.get("工号")
                if name and emp_id:
                    name_to_id[name] = emp_id

        block_label = {"hire": "入职", "depart": "离职/退休", "confirm": "转正",
                       "transfer": "调岗", "other": "其他"}.get(block_type, block_type)
        log.append(f"  Sheet「{source_sheet}」区块「{block_label}」: {data_count} 条记录")

    return result, log


def _parse_block_row(
    row_dict: dict[str, Any],
    block_type: str,
    name_to_id: dict[str, str],
    source_file: str,
    source_sheet: str,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    """解析单行分区数据，返回 (员工档案记录, 薪资明细记录)。

    根据区块类型提取不同字段，并通过姓名映射补齐工号。
    """
    # 通用：尝试从行数据中获取姓名和工号
    name = _clean_value(row_dict.get("姓名"), "text")
    emp_id = _clean_value(row_dict.get("新工号") or row_dict.get("员工编号") or row_dict.get("集团编码"), "text")

    # 如果没有工号，尝试通过姓名映射查找
    if (not emp_id) and name and name in name_to_id:
        emp_id = name_to_id[name]

    if not name and not emp_id:
        return None, None

    emp_rec: dict[str, Any] = {
        "_source_file": source_file,
        "_source_sheet": source_sheet,
    }
    sal_rec: dict[str, Any] = {
        "_source_file": source_file,
        "_source_sheet": source_sheet,
    }

    if block_type == "hire":
        # 入职区：提取完整员工信息
        emp_rec["工号"] = emp_id
        emp_rec["姓名"] = name
        emp_rec["公司"] = _clean_value(row_dict.get("新公司"), "text")
        emp_rec["部门"] = _clean_value(row_dict.get("新成本中心/部门") or row_dict.get("成本中心/部门"), "text")
        emp_rec["入职日期"] = _clean_value(row_dict.get("入职时间"), "date")
        emp_rec["转正日期"] = _clean_value(row_dict.get("试用期结束时间"), "date")
        emp_rec["岗位"] = _clean_value(row_dict.get("岗位"), "text")
        emp_rec["开户银行"] = _clean_value(row_dict.get("银行"), "text")
        emp_rec["银行卡号"] = _clean_value(row_dict.get("银行卡号"), "text")
        emp_rec["本月状态"] = "入职"
        # 薪资标准
        sal_rec["工号"] = emp_id
        sal_rec["姓名"] = name
        sal_rec["月基本薪资"] = _clean_value(row_dict.get("基本工资"), "amount")
        sal_rec["月绩效标准"] = _clean_value(row_dict.get("绩效工资"), "amount")
        sal_rec["交通通讯补贴标准"] = _clean_value(row_dict.get("车贴"), "amount")
        sal_rec["餐补标准"] = _clean_value(row_dict.get("饭贴"), "amount")
        sal_rec["防暑降温费标准"] = _clean_value(row_dict.get("高温费"), "amount")
        sal_rec["过节费"] = _clean_value(row_dict.get("过节费"), "amount")
        sal_rec["独生子女补贴"] = _clean_value(row_dict.get("独生子女费"), "amount")

    elif block_type == "depart":
        # 离职区：只有姓名、最后工作日、离职原因
        emp_rec["姓名"] = name
        emp_rec["工号"] = emp_id
        emp_rec["离职日期"] = _clean_value(row_dict.get("最后工作日"), "date")
        emp_rec["本月状态"] = "离职"
        sal_rec = None  # 离职区无薪资数据

    elif block_type == "confirm":
        # 转正区：姓名、转正日期、备注
        emp_rec["姓名"] = name
        emp_rec["工号"] = emp_id
        emp_rec["转正日期"] = _clean_value(row_dict.get("转正日期"), "date")
        emp_rec["本月状态"] = "转正"
        sal_rec = None

    elif block_type == "transfer":
        # 调岗区：姓名、新工号、新公司、新部门、岗位、薪资标准
        emp_rec["工号"] = emp_id
        emp_rec["姓名"] = name
        emp_rec["公司"] = _clean_value(row_dict.get("新公司"), "text")
        emp_rec["部门"] = _clean_value(row_dict.get("新成本中心/部门") or row_dict.get("调入部门"), "text")
        emp_rec["岗位"] = _clean_value(row_dict.get("现职务") or row_dict.get("岗位"), "text")
        emp_rec["入职日期"] = _clean_value(row_dict.get("入职时间"), "date")
        emp_rec["本月状态"] = "调岗"
        # 薪资标准
        sal_rec["工号"] = emp_id
        sal_rec["姓名"] = name
        sal_rec["月基本薪资"] = _clean_value(row_dict.get("基本工资"), "amount")
        sal_rec["月绩效标准"] = _clean_value(row_dict.get("绩效工资"), "amount")
        sal_rec["交通通讯补贴标准"] = _clean_value(row_dict.get("车贴"), "amount")
        sal_rec["餐补标准"] = _clean_value(row_dict.get("饭贴"), "amount")
        sal_rec["防暑降温费标准"] = _clean_value(row_dict.get("高温费"), "amount")
        sal_rec["过节费"] = _clean_value(row_dict.get("过节费"), "amount")
        sal_rec["独生子女补贴"] = _clean_value(row_dict.get("独生子女费"), "amount")

    elif block_type == "other":
        # 其他区：姓名、备注
        emp_rec["姓名"] = name
        emp_rec["工号"] = emp_id
        emp_rec["本月状态"] = "其他"
        sal_rec = None

    return emp_rec, sal_rec


# ============================================================
# 奖金 Sheet 解析（同时提取员工档案 + 薪资明细）
# ============================================================

def _parse_bonus_sheet(
    df_raw: pd.DataFrame,
    source_file: str,
    source_sheet: str,
) -> tuple[list[dict], list[dict], dict[str, str]]:
    """解析奖金 Sheet，同时提取员工档案基础信息和薪资明细。

    奖金 Sheet 表头在第3行（0-based=2），包含：姓名、工号、部门、岗位、各类奖金。

    Returns:
        (员工档案记录列表, 薪资明细记录列表, 姓名→工号映射)
    """
    emp_records: list[dict] = []
    sal_records: list[dict] = []
    name_to_id: dict[str, str] = {}

    # 嗅探表头行（找包含"工号"的行）
    header_row = None
    for i in range(min(10, len(df_raw))):
        row_vals = [_normalize_header(v) for v in df_raw.iloc[i].tolist()]
        if any("工号" in h for h in row_vals):
            header_row = i
            break

    if header_row is None:
        return emp_records, sal_records, name_to_id

    headers = [_normalize_header(v) for v in df_raw.iloc[header_row].tolist()]
    data_df = df_raw.iloc[header_row + 1:].copy()
    data_df.columns = headers
    data_df = data_df.dropna(how="all")

    for _, row in data_df.iterrows():
        name = _clean_value(row.get("姓名"), "text")
        emp_id = _clean_value(row.get("工号"), "text")
        if not name and not emp_id:
            continue

        # 员工档案基础信息
        emp_rec: dict[str, Any] = {
            "工号": emp_id,
            "姓名": name,
            "部门": _clean_value(row.get("部门"), "text"),
            "岗位": _clean_value(row.get("岗位"), "text"),
            "本月状态": "在职",
            "_source_file": source_file,
            "_source_sheet": source_sheet,
        }
        emp_records.append(emp_rec)
        if name and emp_id:
            name_to_id[name] = emp_id

        # 薪资明细（奖金字段）
        sal_rec: dict[str, Any] = {
            "工号": emp_id,
            "姓名": name,
            "_source_file": source_file,
            "_source_sheet": source_sheet,
        }
        # 匹配奖金列
        for src_col in headers:
            if not src_col:
                continue
            # 绩效奖金
            if "绩效奖金" in src_col and "月" in src_col:
                sal_rec["绩效奖金"] = _clean_value(row.get(src_col), "amount")
            elif "提成" in src_col:
                sal_rec["业绩奖金"] = _clean_value(row.get(src_col), "amount")
            elif "计件" in src_col:
                sal_rec["DTP绩效奖金"] = _clean_value(row.get(src_col), "amount")
            elif "项目奖金" in src_col and "DTP" in src_col:
                sal_rec["DTP绩效奖金"] = _clean_value(row.get(src_col), "amount")
            elif "项目奖金" in src_col and "非DTP" in src_col:
                sal_rec["业绩奖金"] = _clean_value(row.get(src_col), "amount")
        sal_records.append(sal_rec)

    return emp_records, sal_records, name_to_id


# ============================================================
# 模板初始数据导入
# ============================================================

def import_template_to_entities(
    file_path: str,
) -> dict[str, list[dict]]:
    """从导出模板中提取初始数据，构成完整的数据库基础。

    模板中的关键 Sheet：
        - 工资核算：表头在第2行，数据从第4行起。包含完整员工档案。
        - 个税导出核对：表头在第1行，数据从第2行起。包含身份证号和个税数据。

    Returns:
        {实体ID: [记录字典]}
    """
    result: dict[str, list[dict]] = {
        "employee_profile": [],
        "salary_detail": [],
        "tax_record": [],
    }
    log: list[str] = []

    wb = openpyxl.load_workbook(file_path, data_only=True)

    # ---- 1. 工资核算 sheet → 员工档案 ----
    if "工资核算" in wb.sheetnames:
        ws = wb["工资核算"]
        # 表头在第2行（0-based: 1）
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=2, column=col_idx).value
            headers.append(_normalize_header(val) if val else "")

        profile_entity = ENTITY_BY_ID["employee_profile"]
        salary_entity = ENTITY_BY_ID["salary_detail"]
        profile_mapping = _match_columns(headers, profile_entity)
        salary_mapping = _match_columns(headers, salary_entity)
        profile_fields = {field.name: field for field in profile_entity.fields}
        salary_fields = {field.name: field for field in salary_entity.fields}

        # 数据从第4行起（跳过第3行占位符）
        emp_count = 0
        salary_count = 0
        for row_idx in range(4, ws.max_row + 1):
            emp_id = _clean_value(ws.cell(row=row_idx, column=1).value, "text")  # A: 新工号
            name = _clean_value(ws.cell(row=row_idx, column=2).value, "text")    # B: 姓名
            if not emp_id and not name:
                continue
            # 跳过占位行
            if name == "人员姓名":
                continue

            # 构建行数据字典
            row_dict = {}
            for col_idx, h in enumerate(headers, start=1):
                if h and col_idx <= ws.max_column:
                    row_dict[h] = ws.cell(row=row_idx, column=col_idx).value

            rec: dict[str, Any] = {
                "工号": emp_id,
                "姓名": name,
                "公司": _clean_value(row_dict.get("新公司"), "text"),
                "部门": _clean_value(row_dict.get("新成本中心/部门") or row_dict.get("成本中心/部门"), "text"),
                "岗位": _clean_value(row_dict.get("新岗位") or row_dict.get("岗位"), "text"),
                "人员类别": _clean_value(row_dict.get("人员类别"), "text"),
                "合同类别": _clean_value(row_dict.get("合同类别"), "text"),
                "本月状态": _clean_value(row_dict.get("本月状态"), "text") or "在职",
                "入职日期": _clean_value(row_dict.get("入职日期"), "date"),
                "工作地点": _clean_value(row_dict.get("工作地点"), "text"),
                "_source_file": os.path.basename(file_path),
                "_source_sheet": "工资核算",
            }

            # 使用统一实体字段映射补齐身份、试用期、银行卡等档案字段。
            for source_field, target_field in profile_mapping.items():
                field = profile_fields[target_field]
                cleaned = _clean_value(row_dict.get(source_field), field.dtype)
                if cleaned is not None:
                    rec[target_field] = cleaned

            # 从表头中匹配更多字段（餐补标准、班制等）
            for src_col, val in row_dict.items():
                if not src_col or val is None:
                    continue
                # 餐补标准
                if "餐补" in src_col and "标准" in src_col:
                    rec["餐补标准"] = _clean_value(val, "text")
                # 班制
                elif src_col == "班制" or "排班" in src_col:
                    rec["班制"] = _clean_value(val, "text")
                # 出勤天数
                elif "出勤天数" in src_col:
                    rec["出勤天数"] = _clean_value(val, "amount")

            result["employee_profile"].append(rec)
            emp_count += 1

            # 工资核算页同时是基本薪资与补贴标准的主数据来源。
            salary_rec: dict[str, Any] = {
                "工号": emp_id,
                "姓名": name,
                "_source_file": os.path.basename(file_path),
                "_source_sheet": "工资核算",
            }
            has_salary_value = False
            for source_field, target_field in salary_mapping.items():
                field = salary_fields[target_field]
                cleaned = _clean_value(row_dict.get(source_field), field.dtype)
                if cleaned is None:
                    continue
                salary_rec[target_field] = cleaned
                if target_field not in {"工号", "姓名"}:
                    has_salary_value = True
            if salary_rec.get("月基本薪资") in (None, ""):
                monthly_total = _clean_value(ws.cell(row=row_idx, column=24).value, "amount")
                fixed_ratio = _clean_value(ws.cell(row=row_idx, column=25).value, "amount")
                if monthly_total is not None and fixed_ratio is not None:
                    base_salary = round(float(monthly_total) * float(fixed_ratio), -1)
                    salary_rec["月基本薪资"] = base_salary
                    salary_rec["月绩效标准"] = float(monthly_total) - base_salary
                    has_salary_value = True
            if has_salary_value:
                result["salary_detail"].append(salary_rec)
                salary_count += 1

        log.append(f"  模板 Sheet「工资核算」→ 员工档案: {emp_count} 人")
        log.append(f"  模板 Sheet「工资核算」→ 薪资明细: {salary_count} 人")
    else:
        log.append("  模板中缺少「工资核算」Sheet")

    # ---- 2. 个税导出核对 sheet → 个税记录 + 身份证号回填员工档案 ----
    if "个税导出核对" in wb.sheetnames:
        ws = wb["个税导出核对"]
        # 表头在第1行
        headers_tax = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col_idx).value
            headers_tax.append(_normalize_header(val) if val else "")

        # 数据从第2行起
        tax_count = 0
        # 构建姓名→身份证号映射（用于回填员工档案）
        name_to_idcard: dict[str, str] = {}

        for row_idx in range(2, ws.max_row + 1):
            # 构建行数据字典
            row_dict = {}
            for col_idx, h in enumerate(headers_tax, start=1):
                if h and col_idx <= ws.max_column:
                    row_dict[h] = ws.cell(row=row_idx, column=col_idx).value

            name = _clean_value(row_dict.get("姓名"), "text")
            id_card = _clean_value(row_dict.get("证件号码"), "text")
            if not name and not id_card:
                continue

            # 记录姓名→身份证号映射
            if name and id_card:
                name_to_idcard[name] = id_card

            tax_rec: dict[str, Any] = {
                "姓名": name,
                "身份证号": id_card,
                "税款所属期起": _clean_value(row_dict.get("税款所属期起"), "date"),
                "税款所属期止": _clean_value(row_dict.get("税款所属期止"), "date"),
                "本期收入": _clean_value(row_dict.get("本期收入"), "amount"),
                "累计收入额": _clean_value(row_dict.get("累计收入额"), "amount"),
                "累计减除费用": _clean_value(row_dict.get("累计减除费用"), "amount"),
                "累计专项扣除": _clean_value(row_dict.get("累计专项扣除"), "amount"),
                "累计应纳税所得额": _clean_value(row_dict.get("累计应纳税所得额"), "amount"),
                "累计应扣缴税额": _clean_value(row_dict.get("累计应扣缴税额"), "amount"),
                "累计已预缴税额": _clean_value(row_dict.get("累计已预缴税额"), "amount"),
                "累计应补退税额": _clean_value(row_dict.get("累计应补(退)税额"), "amount"),
                "_source_file": os.path.basename(file_path),
                "_source_sheet": "个税导出核对",
            }
            result["tax_record"].append(tax_rec)
            tax_count += 1

        log.append(f"  模板 Sheet「个税导出核对」→ 个税记录: {tax_count} 条, 身份证号映射: {len(name_to_idcard)} 条")

        # 回填身份证号到员工档案
        backfilled = 0
        for emp_rec in result["employee_profile"]:
            name = emp_rec.get("姓名", "")
            if name and name in name_to_idcard and not emp_rec.get("身份证号"):
                emp_rec["身份证号"] = name_to_idcard[name]
                backfilled += 1
        if backfilled > 0:
            log.append(f"  模板身份证号回填员工档案: {backfilled} 人")
    else:
        log.append("  模板中缺少「个税导出核对」Sheet")

    wb.close()
    result["_log"] = log  # type: ignore
    return result


# ============================================================
# 主导入函数
# ============================================================

def import_file_to_entities(
    file_path: str,
    filename: str,
) -> dict[str, list[dict]]:
    """将一个源文件导入为实体数据。

    对薪资文件做特殊编排：先解析奖金sheet获取员工列表，
    再解析入离职sheet分区提取变动信息。

    Returns:
        {实体ID: [记录字典]}
    """
    file_type = detect_file_type(filename)
    log: list[str] = []

    # 初始化结果
    result: dict[str, list[dict]] = {
        "employee_profile": [],
        "salary_detail": [],
        "attendance_record": [],
        "social_security": [],
        "tax_record": [],
    }

    xls = pd.ExcelFile(file_path)
    social_sheet_names: set[str] = set()
    for sheet_name in xls.sheet_names:
        df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
        if not _is_social_billing_sheet(df_raw):
            continue
        records, col_map = _parse_social_billing_sheet(df_raw, filename, sheet_name)
        result["social_security"].extend(records)
        social_sheet_names.add(sheet_name)
        log.append(
            f"  Sheet「{sheet_name}」→ 社保公积金: {len(records)} 条记录，"
            f"语义匹配 {len(col_map)} 列"
        )

    # A supplier-neutral social bill may have an arbitrary filename.  Once all
    # sheets are recognized as social bills, do not route it through payroll
    # parsing merely because its filename does not contain "社保".
    if social_sheet_names and len(social_sheet_names) == len(xls.sheet_names):
        result["_log"] = log  # type: ignore
        return result

    # 薪资文件：特殊编排
    if file_type == "salary":
        # 1. 先解析"奖金"sheet → 员工档案 + 薪资明细 + 姓名映射
        name_to_id: dict[str, str] = {}
        for sheet_name in xls.sheet_names:
            if sheet_name in social_sheet_names:
                continue
            if "奖金" in sheet_name:
                df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
                emp_recs, sal_recs, n2i = _parse_bonus_sheet(df_raw, filename, sheet_name)
                result["employee_profile"].extend(emp_recs)
                result["salary_detail"].extend(sal_recs)
                name_to_id.update(n2i)
                log.append(f"  Sheet「{sheet_name}」→ 员工档案+薪资明细: {len(emp_recs)} 人, 姓名映射 {len(n2i)} 条")

        # 2. 解析"入离职"sheet → 分区提取
        for sheet_name in xls.sheet_names:
            if sheet_name in social_sheet_names:
                continue
            if any(kw in sheet_name for kw in ["入离职", "转岗", "转正", "离职"]):
                df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
                block_result, block_logs = _parse_block_sheet(df_raw, filename, sheet_name, name_to_id)
                for eid, recs in block_result.items():
                    result[eid].extend(recs)
                log.extend(block_logs)

        # 3. 解析"考勤"sheet
        for sheet_name in xls.sheet_names:
            if sheet_name in social_sheet_names:
                continue
            if "考勤" in sheet_name:
                df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
                entity = ENTITY_BY_ID["attendance_record"]
                records, col_map = _parse_sheet_to_entity(df_raw, entity, None, filename, sheet_name)
                result["attendance_record"].extend(records)
                log.append(f"  Sheet「{sheet_name}」→ 实体「{entity.name}」: {len(records)} 条记录, 匹配 {len(col_map)} 列")

                profile_entity = ENTITY_BY_ID["employee_profile"]
                profile_records, profile_map = _parse_sheet_to_entity(
                    df_raw, profile_entity, None, filename, sheet_name
                )
                result["employee_profile"].extend(profile_records)
                log.append(
                    f"  Sheet「{sheet_name}」→ 员工档案补全: {len(profile_records)} 条记录, "
                    f"匹配 {len(profile_map)} 列"
                )

                salary_entity = ENTITY_BY_ID["salary_detail"]
                salary_records, salary_map = _parse_sheet_to_entity(
                    df_raw, salary_entity, None, filename, sheet_name
                )
                result["salary_detail"].extend(salary_records)
                log.append(
                    f"  Sheet「{sheet_name}」→ 薪资标准补全: {len(salary_records)} 条记录, "
                    f"匹配 {len(salary_map)} 列"
                )

        # 4. 其他 sheet 按列语义推断实体，不依赖文件名或 Sheet 名。
        for sheet_name in xls.sheet_names:
            if sheet_name in social_sheet_names:
                continue
            if any(kw in sheet_name for kw in ["奖金", "入离职", "转岗", "转正", "离职", "考勤"]):
                continue  # 已处理
            df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            _append_semantic_sheet_results(result, log, df_raw, filename, sheet_name)

    else:
        # Even known filename categories are parsed by column semantics.  This
        # keeps renamed files, mixed workbooks and unfamiliar Sheet names safe.
        for sheet_name in xls.sheet_names:
            if sheet_name in social_sheet_names:
                continue
            df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            _append_semantic_sheet_results(result, log, df_raw, filename, sheet_name)

    result["_log"] = log  # type: ignore
    return result


def merge_entities(
    entities_map: dict[str, list[dict]],
) -> dict[str, list[dict]]:
    """合并多次导入的实体数据，按主键去重。

    对于同一主键的记录，后导入的覆盖先导入的非空字段。
    对于无主键但有姓名的记录（如离职区），尝试通过姓名匹配已有记录。
    """
    merged: dict[str, list[dict]] = {}

    for entity_id, records in entities_map.items():
        if entity_id == "_log":
            continue
        entity = ENTITY_BY_ID.get(entity_id)
        if entity is None:
            continue

        # 第一轮：按主键去重合并
        pk_index: dict[str, dict] = {}
        # 同时建立姓名→主键的映射（用于无主键记录的关联）
        name_to_pk: dict[str, str] = {}
        no_pk_records: list[dict] = []  # 无主键的记录

        for rec in records:
            pk_val = str(rec.get(entity.primary_key, "")).strip()
            name_val = str(rec.get("姓名", "")).strip()

            if pk_val:
                if pk_val in pk_index:
                    # 合并：后值覆盖前值（仅覆盖非空）
                    existing = pk_index[pk_val]
                    for k, v in rec.items():
                        if v is not None and v != "":
                            existing[k] = v
                else:
                    pk_index[pk_val] = dict(rec)
                # 更新姓名映射
                if name_val:
                    name_to_pk[name_val] = pk_val
            elif name_val:
                # 无主键但有姓名，暂存
                no_pk_records.append(rec)

        # 第二轮：无主键记录通过姓名匹配
        for rec in no_pk_records:
            name_val = str(rec.get("姓名", "")).strip()
            if name_val in name_to_pk:
                pk_val = name_to_pk[name_val]
                existing = pk_index[pk_val]
                for k, v in rec.items():
                    if v is not None and v != "":
                        existing[k] = v
            else:
                # 无法匹配，作为独立记录保留
                pk_index[f"_nomatch_{name_val}"] = dict(rec)

        merged[entity_id] = list(pk_index.values())

    return merged


# ============================================================
# 导出：实体 → 模板
# ============================================================

def _build_export_rows(
    entities: dict[str, list[dict]],
) -> list[dict[str, Any]]:
    """以 employee_profile 为基准，LEFT JOIN 其他实体，构建导出行。

    关联键：
        employee_profile.工号 ↔ salary_detail.工号 ↔ attendance_record.工号 ↔ tax_record.工号
        employee_profile.身份证号 ↔ social_security.身份证号
    """
    # 构建索引
    def _index_by_key(records: list[dict], key: str) -> dict[str, dict]:
        idx = {}
        for r in records:
            k = str(r.get("_person_key") or r.get(key, "")).strip()
            if k:
                idx[k] = r
        return idx

    emp_records = entities.get("employee_profile", [])
    salary_idx = _index_by_key(entities.get("salary_detail", []), "工号")
    attendance_idx = _index_by_key(entities.get("attendance_record", []), "工号")
    tax_idx = _index_by_key(entities.get("tax_record", []), "工号")
    social_idx = _index_by_key(entities.get("social_security", []), "身份证号")

    export_rows: list[dict[str, Any]] = []

    for emp in emp_records:
        row: dict[str, Any] = {}
        # 员工档案字段
        for k, v in emp.items():
            if not k.startswith("_"):
                row[f"employee_profile.{k}"] = v

        # 关联薪资明细
        person_key = str(emp.get("_person_key", "")).strip()
        emp_id = str(emp.get("工号", "")).strip()
        salary_key = person_key if person_key in salary_idx else emp_id
        if salary_key and salary_key in salary_idx:
            for k, v in salary_idx[salary_key].items():
                if not k.startswith("_"):
                    row[f"salary_detail.{k}"] = v

        # 关联考勤记录
        attendance_key = person_key if person_key in attendance_idx else emp_id
        if attendance_key and attendance_key in attendance_idx:
            for k, v in attendance_idx[attendance_key].items():
                if not k.startswith("_"):
                    row[f"attendance_record.{k}"] = v

        # 关联个税记录
        tax_key = person_key if person_key in tax_idx else emp_id
        if tax_key and tax_key in tax_idx:
            for k, v in tax_idx[tax_key].items():
                if not k.startswith("_"):
                    row[f"tax_record.{k}"] = v

        # 关联社保（按身份证号）
        id_card = str(emp.get("身份证号", "")).strip()
        social_key = person_key if person_key in social_idx else id_card
        if social_key and social_key in social_idx:
            for k, v in social_idx[social_key].items():
                if not k.startswith("_"):
                    row[f"social_security.{k}"] = v

        export_rows.append(row)

    return export_rows


def _write_formula_inputs(
    workbook: openpyxl.Workbook,
    export_rows: list[dict[str, Any]],
) -> int:
    """Refresh non-formula cells that feed preserved master formulas."""
    updated = 0
    payroll = workbook["工资核算"]
    for row_offset, row_data in enumerate(export_rows):
        row_number = 4 + row_offset
        base_salary = row_data.get("salary_detail.月基本薪资")
        performance_salary = row_data.get("salary_detail.月绩效标准")
        if base_salary not in (None, "") and performance_salary not in (None, ""):
            monthly_total = base_salary + performance_salary
            monthly_total_cell = payroll.cell(row=row_number, column=24)
            ratio_cell = payroll.cell(row=row_number, column=25)
            if _formula_text(monthly_total_cell.value) is None:
                monthly_total_cell.value = monthly_total
                updated += 1
            if _formula_text(ratio_cell.value) is None:
                ratio_cell.value = base_salary / monthly_total if monthly_total else None
                updated += 1

    if "考勤" not in workbook.sheetnames:
        return updated

    attendance = workbook["考勤"]
    source_columns = {
        23: "attendance_record.姓名",
        24: "attendance_record.实际出勤天数",
        25: "attendance_record.事假时长",
        26: "attendance_record.病假时长",
        28: "attendance_record.法定节日加班时长",
        34: "attendance_record.出差天数",
    }
    for row_number in range(2, attendance.max_row + 1):
        source_name = str(attendance.cell(row=row_number, column=23).value or "").strip()
        if not source_name:
            continue
        matched = next(
            (
                row_data
                for row_data in export_rows
                if str(row_data.get("attendance_record.姓名") or "").strip() == source_name
            ),
            None,
        )
        if matched is None:
            continue
        for column, lookup_key in source_columns.items():
            value = matched.get(lookup_key)
            if value in (None, ""):
                continue
            cell = attendance.cell(row=row_number, column=column)
            if _formula_text(cell.value) is not None:
                continue
            cell.value = value
            updated += 1

        ordinary_overtime = matched.get("attendance_record.普通加班时长")
        if ordinary_overtime not in (None, ""):
            total_cell = attendance.cell(row=row_number, column=27)
            if _formula_text(total_cell.value) is None:
                total_cell.value = ordinary_overtime
                updated += 1
            else:
                detail_cells = [
                    attendance.cell(row=row_number, column=31),
                    attendance.cell(row=row_number, column=32),
                ]
                if all(_formula_text(cell.value) is None for cell in detail_cells):
                    detail_cells[0].value = ordinary_overtime
                    detail_cells[1].value = 0
                    updated += 2

        travel_days = matched.get("attendance_record.出差天数")
        if travel_days not in (None, ""):
            travel_cell = attendance.cell(row=row_number, column=34)
            if _formula_text(travel_cell.value) is None:
                travel_cell.value = travel_days
                updated += 1
    return updated


def _refresh_business_sheet_inputs(
    workbook: openpyxl.Workbook,
    export_rows: list[dict[str, Any]],
) -> dict[str, int]:
    """Refresh auditable business-sheet inputs without replacing formulas.

    These worksheets are calculation views of the payroll roster.  Matching is
    deliberately limited to employee id, identity card, or a unique name.  A
    row that cannot be matched unambiguously is left unchanged for manual
    review.
    """
    rows_by_employee_id: dict[str, dict[str, Any]] = {}
    rows_by_id_card: dict[str, dict[str, Any]] = {}
    rows_by_name: dict[str, list[dict[str, Any]]] = {}
    for row_data in export_rows:
        employee_id = str(row_data.get("employee_profile.工号") or "").strip()
        id_card = str(row_data.get("employee_profile.身份证号") or "").strip().upper()
        name = str(row_data.get("employee_profile.姓名") or "").strip()
        if employee_id:
            rows_by_employee_id[employee_id.upper()] = row_data
        if id_card:
            rows_by_id_card[id_card] = row_data
        if name:
            rows_by_name.setdefault(name, []).append(row_data)

    def matched_row(sheet: openpyxl.worksheet.worksheet.Worksheet, row_number: int) -> dict[str, Any] | None:
        employee_id = str(sheet.cell(row=row_number, column=1).value or "").strip().upper()
        if employee_id and employee_id in rows_by_employee_id:
            return rows_by_employee_id[employee_id]
        id_card = str(sheet.cell(row=row_number, column=4).value or "").strip().upper()
        if id_card and id_card in rows_by_id_card:
            return rows_by_id_card[id_card]
        name = str(sheet.cell(row=row_number, column=2).value or "").strip()
        name_matches = rows_by_name.get(name, [])
        return name_matches[0] if len(name_matches) == 1 else None

    def refresh_rows(
        sheet_name: str,
        start_row: int,
        columns: dict[int, str],
    ) -> int:
        if sheet_name not in workbook.sheetnames:
            return 0
        sheet = workbook[sheet_name]
        changed = 0
        for row_number in range(start_row, sheet.max_row + 1):
            current_id = sheet.cell(row=row_number, column=1).value
            current_name = sheet.cell(row=row_number, column=2).value
            if _formula_text(current_id) is not None or _formula_text(current_name) is not None:
                continue
            if current_id in (None, "") and current_name in (None, ""):
                continue
            row_data = matched_row(sheet, row_number)
            if row_data is None:
                continue
            for column, lookup_key in columns.items():
                value = row_data.get(lookup_key)
                if value in (None, ""):
                    continue
                cell = sheet.cell(row=row_number, column=column)
                if _formula_text(cell.value) is not None:
                    continue
                if cell.value != value:
                    cell.value = value
                    changed += 1
        return changed

    updated: dict[str, int] = {}
    updated["考勤"] = refresh_rows(
        "考勤",
        2,
        {
            1: "employee_profile.工号",
            2: "employee_profile.姓名",
            3: "employee_profile.部门",
            4: "employee_profile.本月状态",
            5: "employee_profile.入职日期",
            6: "salary_detail.餐补标准",
        },
    )
    updated["防暑降温费"] = refresh_rows(
        "防暑降温费",
        2,
        {
            1: "employee_profile.工号",
            2: "employee_profile.姓名",
            3: "employee_profile.部门",
            4: "attendance_record.实际出勤天数",
            5: "attendance_record.班制",
        },
    )
    updated["个税申报"] = refresh_rows(
        "个税申报",
        2,
        {
            1: "employee_profile.工号",
            2: "employee_profile.姓名",
            4: "employee_profile.身份证号",
        },
    )
    updated["台账"] = refresh_rows(
        "台账",
        3,
        {
            1: "employee_profile.工号",
            2: "employee_profile.姓名",
        },
    )

    if "个税导出核对" in workbook.sheetnames:
        tax_sheet = workbook["个税导出核对"]
        tax_header_mapping = {
            "工号": "employee_profile.工号",
            "姓名": "employee_profile.姓名",
            "证件号码": "employee_profile.身份证号",
            "税款所属期起": "tax_record.税款所属期起",
            "税款所属期止": "tax_record.税款所属期止",
            "本期收入": "tax_record.本期收入",
            "累计收入额": "tax_record.累计收入额",
            "累计减除费用": "tax_record.累计减除费用",
            "累计专项扣除": "tax_record.累计专项扣除",
            "累计子女教育支出扣除": "tax_record.累计子女教育",
            "累计继续教育支出扣除": "tax_record.累计继续教育",
            "累计住房贷款利息支出扣除": "tax_record.累计住房贷款利息",
            "累计住房租金支出扣除": "tax_record.累计住房租金",
            "累计赡养老人支出扣除": "tax_record.累计赡养老人",
            "累计婴幼儿照护费用": "tax_record.累计婴幼儿照护",
            "累计应纳税所得额": "tax_record.累计应纳税所得额",
            "累计应扣缴税额": "tax_record.累计应扣缴税额",
            "累计已预缴税额": "tax_record.累计已预缴税额",
            "累计应补(退)税额": "tax_record.累计应补退税额",
        }
        columns = {
            column: tax_header_mapping[header]
            for column in range(1, tax_sheet.max_column + 1)
            if (header := str(tax_sheet.cell(row=1, column=column).value or "").strip())
            in tax_header_mapping
        }
        updated["个税导出核对"] = refresh_rows("个税导出核对", 2, columns)

    if "工资条" in workbook.sheetnames:
        payslip = workbook["工资条"]
        formula_rows = [
            row_number
            for row_number in range(2, payslip.max_row + 1)
            if any(
                "工资核算" in (formula or "")
                for cell in payslip[row_number]
                if (formula := _formula_text(cell.value)) is not None
            )
        ]
        changed = 0
        if len(formula_rows) == len(export_rows):
            for row_number, row_data in zip(formula_rows, export_rows):
                name = row_data.get("employee_profile.姓名")
                cell = payslip.cell(row=row_number, column=4)
                if name not in (None, "") and cell.value != name:
                    cell.value = name
                    changed += 1
        updated["工资条"] = changed

    return {sheet_name: count for sheet_name, count in updated.items() if count}


def export_to_template(
    template_path: str,
    entities: dict[str, list[dict]],
    output_path: str,
    freeze_formulas: bool = True,
) -> list[str]:
    """将实体数据按模板导出。

    Args:
        template_path: Master 模板路径
        entities: {实体ID: [记录列表]}
        output_path: 输出文件路径
        freeze_formulas: True=保留模板公式, False=公式列留空

    Returns:
        导出日志列表
    """
    log: list[str] = []

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"模板文件不存在: {template_path}")

    # 构建导出行（关联后的扁平数据）
    export_rows = _build_export_rows(entities)
    log.append(f"构建导出行: {len(export_rows)} 条员工记录")

    if not export_rows:
        log.append("无可导出的员工数据")
        return log

    # 加载模板（保留公式与样式）
    wb = openpyxl.load_workbook(template_path)
    if "工资核算" not in wb.sheetnames:
        raise ValueError("模板中缺少「工资核算」Sheet")

    ws = wb["工资核算"]

    # 确定数据起始行：第4行（第1行序号，第2行表头，第3行别名，第4行起数据）
    DATA_START_ROW = 4

    template_employee_end = DATA_START_ROW - 1
    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        employee_id = ws.cell(row=row_idx, column=1).value
        employee_name = ws.cell(row=row_idx, column=2).value
        if employee_id in (None, "") and employee_name in (None, ""):
            break
        template_employee_end = row_idx

    template_employee_count = max(0, template_employee_end - DATA_START_ROW + 1)
    target_employee_end = DATA_START_ROW + len(export_rows) - 1
    inserted_employee_rows: set[int] = set()
    formula_templates: dict[int, tuple[Any, str]] = {}
    for col_idx in range(1, ws.max_column + 1):
        for row_idx in range(template_employee_end, DATA_START_ROW - 1, -1):
            source_cell = ws.cell(row=row_idx, column=col_idx)
            if _formula_text(source_cell.value) is not None:
                formula_templates[col_idx] = (copy.copy(source_cell.value), source_cell.coordinate)
                break

    if len(export_rows) > template_employee_count:
        extra = len(export_rows) - template_employee_count
        insert_at = template_employee_end + 1
        ws.insert_rows(insert_at, amount=extra)
        inserted_employee_rows.update(range(insert_at, insert_at + extra))
        source_row = template_employee_end
        for target_row in range(insert_at, insert_at + extra):
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
            for col_idx in range(1, ws.max_column + 1):
                src_cell = ws.cell(row=source_row, column=col_idx)
                dst_cell = ws.cell(row=target_row, column=col_idx)
                if src_cell.has_style:
                    dst_cell.font = copy.copy(src_cell.font)
                    dst_cell.border = copy.copy(src_cell.border)
                    dst_cell.fill = copy.copy(src_cell.fill)
                    dst_cell.number_format = src_cell.number_format
                    dst_cell.alignment = copy.copy(src_cell.alignment)
                    dst_cell.protection = copy.copy(src_cell.protection)
                translated = _translated_formula(
                    src_cell.value,
                    src_cell.coordinate,
                    dst_cell.coordinate,
                )
                if translated is None:
                    col_letter = get_column_letter(col_idx)
                    mapping = EXPORT_BY_COLUMN.get(col_letter)
                    if mapping and mapping[2] and col_idx in formula_templates:
                        template_formula, template_coordinate = formula_templates[col_idx]
                        translated = _translated_formula(
                            template_formula,
                            template_coordinate,
                            dst_cell.coordinate,
                        )
                if translated is not None:
                    dst_cell.value = translated
        log.append(f"插入员工行: {extra} 行（已复制公式与样式）")
    elif len(export_rows) < template_employee_count:
        remove_count = template_employee_count - len(export_rows)
        ws.delete_rows(DATA_START_ROW + len(export_rows), amount=remove_count)
        log.append(f"删除多余员工行: {remove_count} 行")

    if target_employee_end != template_employee_end:
        range_end = re.compile(rf"(:\$?[A-Z]{{1,3}}\$?){template_employee_end}(?!\d)")
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    formula = _formula_text(cell.value)
                    if formula is None:
                        continue
                    references_payroll = (
                        (sheet is ws and cell.row > target_employee_end)
                        or "工资核算!" in formula
                        or "'工资核算'!" in formula
                    )
                    if references_payroll:
                        updated = range_end.sub(rf"\g<1>{target_employee_end}", formula)
                        cell.value = (
                            ArrayFormula(ref=cell.coordinate, text=updated)
                            if isinstance(cell.value, ArrayFormula)
                            else updated
                        )

    # 统计
    filled_cells = 0
    skipped_formula_cells = 0
    skipped_empty_cells = 0

    # 逐行逐列写入
    for row_offset, row_data in enumerate(export_rows):
        target_row = DATA_START_ROW + row_offset

        for col_letter, (entity_id, field_name, is_formula) in EXPORT_BY_COLUMN.items():
            col_idx = column_index_from_string(col_letter)

            if is_formula:
                # 公式列：保留模板原样（不覆盖）
                if freeze_formulas:
                    skipped_formula_cells += 1
                else:
                    # 清空公式
                    ws.cell(row=target_row, column=col_idx).value = None
                continue

            if entity_id is None or field_name is None:
                continue

            cell = ws.cell(row=target_row, column=col_idx)
            existing_value = cell.value

            # The uploaded master workbook is authoritative. Any formula in it
            # must keep driving the result, even when an update file supplies a
            # cached/scalar value for the same mapped field.
            if freeze_formulas and _formula_text(existing_value) is not None:
                skipped_formula_cells += 1
                continue

            lookup_key = f"{entity_id}.{field_name}"
            value = row_data.get(lookup_key)
            if value is not None and value != "":
                cell.value = value
                filled_cells += 1
                continue

            if isinstance(existing_value, str) and existing_value.startswith("="):
                # 模板内的计算公式本身就是有效规则；缺少替代值时应保留。
                cell.value = existing_value
                skipped_formula_cells += 1
            else:
                cell.value = None
                skipped_empty_cells += 1

        # 新增员工没有模板旧值，按基础薪资+月绩效生成月度标准与固浮比。
        if target_row in inserted_employee_rows:
            ws[f"X{target_row}"] = f"=SUM(Z{target_row}:AA{target_row})"
            ws[f"Y{target_row}"] = (
                f'=IF(OR(X{target_row}="",X{target_row}=0,Z{target_row}=""),"",'
                f"Z{target_row}/X{target_row})"
            )

    log.append(f"写入完成: {filled_cells} 个单元格有值, {skipped_empty_cells} 个为空, {skipped_formula_cells} 个公式列保留")

    literal_errors_cleared = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "e":
                    cell.value = None
                    literal_errors_cleared += 1
    if literal_errors_cleared:
        log.append(f"模板异常: {literal_errors_cleared} 个错误单元格已留空")

    refreshed_business_sheets = _refresh_business_sheet_inputs(wb, export_rows)
    for sheet_name, cell_count in refreshed_business_sheets.items():
        log.append(f"已刷新业务表「{sheet_name}」: {cell_count} 个输入单元格")

    formula_inputs_updated = _write_formula_inputs(wb, export_rows)
    if formula_inputs_updated:
        log.append(f"已刷新公式依赖输入: {formula_inputs_updated} 个单元格")

    # 保存
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    if hasattr(wb, "calculation"):
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    wb.save(output_path)
    log.append(f"导出完成: {output_path}")

    return log


# ============================================================
# 实体数据概览（用于前端展示）
# ============================================================

def get_entity_overview(entities: dict[str, list[dict]]) -> list[dict]:
    """返回各实体的概览信息。"""
    overview = []
    for entity in ALL_ENTITIES:
        records = entities.get(entity.id, [])
        field_names = [f.name for f in entity.fields]

        # 统计完整度
        total = len(records)
        if total > 0:
            filled_cells = 0
            total_cells = total * len(field_names)
            for rec in records:
                for fn in field_names:
                    v = rec.get(fn)
                    if v is not None and v != "":
                        filled_cells += 1
            completeness = round(filled_cells / total_cells * 100, 1) if total_cells else 100.0
        else:
            completeness = 0.0

        overview.append({
            "id": entity.id,
            "name": entity.name,
            "icon": entity.icon,
            "primary_key": entity.primary_key,
            "row_count": total,
            "field_count": len(field_names),
            "completeness_pct": completeness,
        })
    return overview
