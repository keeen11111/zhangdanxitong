"""Incremental formula caches for personnel-only changes, without desktop Excel.

The confirmed master's cached results remain authoritative when inputs did not
change. A small, explicit formula grammar recalculates affected values. Unsupported
expressions fail closed; formula text and workbook formatting are never rewritten.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
import fnmatch
import math
import operator
import re
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter
from openpyxl.utils.datetime import to_excel


class _ExcelError(Exception):
    pass


class _Parser:
    def __init__(self, formula: str):
        self.tokens = [t for t in Tokenizer(formula).items if t.type != 'WHITE-SPACE']
        self.pos = 0

    def take(self) -> Any:
        token = self.tokens[self.pos]
        self.pos += 1
        return token

    def expression(self, minimum: int = 0) -> Any:
        token = self.take()
        if token.type == 'OPERAND':
            left = ('literal', token.value, token.subtype)
        elif token.type == 'OPERATOR-PREFIX':
            left = ('unary', token.value, self.expression(6))
        elif token.type == 'PAREN' and token.subtype == 'OPEN':
            left = self.expression()
            assert self.take().subtype == 'CLOSE'
        elif token.type in {'FUNC', 'ARRAY'} and token.subtype == 'OPEN':
            items = []
            while self.pos < len(self.tokens) and self.tokens[self.pos].subtype != 'CLOSE':
                if self.tokens[self.pos].type == 'SEP':
                    items.append(('literal', '', 'EMPTY'))
                else:
                    items.append(self.expression())
                if self.tokens[self.pos].subtype != 'CLOSE':
                    assert self.take().type == 'SEP'
            self.take()
            left = ('call', token.value[:-1].upper() if token.type == 'FUNC' else 'ARRAY', items)
        else:
            raise ValueError(f'Unsupported Excel token: {token.type}/{token.subtype}')
        priorities = {'=': 1, '<>': 1, '<': 1, '>': 1, '<=': 1, '>=': 1, '&': 2, '+': 3, '-': 3, '*': 4, '/': 4, '^': 5}
        while self.pos < len(self.tokens):
            token = self.tokens[self.pos]
            if token.type == 'OPERATOR-POSTFIX':
                self.take()
                left = ('binary', '/', left, ('literal', '100', 'NUMBER'))
                continue
            precedence = priorities.get(token.value, 0) if token.type == 'OPERATOR-INFIX' else 0
            if not precedence or precedence < minimum:
                break
            self.take()
            left = ('binary', token.value, left, self.expression(precedence + 1))
        return left

    def parse(self) -> Any:
        result = self.expression()
        if self.pos != len(self.tokens):
            raise ValueError('Unsupported trailing Excel expression')
        return result


@dataclass
class _Ref:
    engine: Any
    sheet: str
    first_col: int
    first_row: int
    last_col: int
    last_row: int

    def at(self, row: int, column: int = 0) -> Any:
        return self.engine.cell(self.sheet, self.first_row + row, self.first_col + column)

    def all(self) -> Any:
        for r in range(self.last_row - self.first_row + 1):
            for c in range(self.last_col - self.first_col + 1):
                yield self.at(r, c)


def _scalar(value: Any) -> Any:
    return value.at(0) if isinstance(value, _Ref) else value


def _number(value: Any) -> float:
    value = _scalar(value)
    if value in (None, ''):
        return 0
    if isinstance(value, (date, datetime)):
        return to_excel(value)
    try:
        return float(value)
    except (ValueError, TypeError) as error:
        raise _ExcelError('#VALUE!') from error


def _equal(first: Any, second: Any) -> bool:
    if isinstance(first, str) and isinstance(second, str):
        return first.casefold() == second.casefold()
    return type(first) is type(second) and first == second or (
        isinstance(first, (float, int)) and isinstance(second, (float, int)) and first == second
    )


def _criterion(value: Any, criterion: Any) -> bool:
    if isinstance(criterion, str):
        match = re.fullmatch(r'(<=|>=|<>|=|<|>)(.*)', criterion)
        if match:
            op, rhs = match.groups()
            try:
                rhs = float(rhs)
                lhs = _number(value)
            except (ValueError, _ExcelError):
                lhs = str(value or '').casefold()
                rhs = str(rhs).casefold()
            return {'<': operator.lt, '>': operator.gt, '<=': operator.le, '>=': operator.ge,
                    '=': operator.eq, '<>': operator.ne}[op](lhs, rhs)
        if '*' in criterion or '?' in criterion:
            return fnmatch.fnmatchcase(str(value or '').casefold(), criterion.casefold())
    return _equal(value, criterion)


class _Engine:
    def __init__(self, master: Any, original_cache: Any, output: Any, insertions: list[dict]):
        self.master, self.original_cache, self.output = master, original_cache, output
        self.insertions = insertions
        self.memo: dict[tuple, tuple[Any, bool]] = {}
        self.frames: list[bool] = []
        self.active: set[tuple] = set()

    def origin(self, sheet: str, row: int) -> int | None:
        for insertion in reversed(self.insertions):
            if insertion['sheet'] != sheet:
                continue
            if row == insertion['row']:
                return None
            if row > insertion['row']:
                row -= 1
        return row

    def cell(self, sheet: str, row: int, column: int) -> Any:
        key = (sheet, row, column)
        if key not in self.memo:
            if key in self.active:
                raise ValueError('Circular reference in affected payroll formulas')
            self.active.add(key)
            cell = self.output[sheet].cell(row, column)
            origin = self.origin(sheet, row)
            old = self.master[sheet].cell(origin, column) if origin else None
            cached = self.original_cache[sheet].cell(origin, column) if origin else None
            formula = getattr(cell.value, 'text', cell.value) if cell.data_type == 'f' else None
            if formula is None:
                value = _ExcelError(cell.value) if cell.data_type == 'e' else cell.value
                changed = old is None or old.value != cell.value
            else:
                old_formula = getattr(old.value, 'text', old.value) if old else None
                dirty = old is None or old_formula != formula
                # External files are frozen at their confirmed cached results.
                # A new external expression has no authoritative cached value.
                if re.search(r'\[[^\]]+\].*!', formula):
                    if dirty:
                        raise ValueError('Cannot calculate a new external-workbook reference')
                    value = _ExcelError(cached.value) if cached.data_type == 'e' else cached.value
                    changed = False
                else:
                    self.frames.append(dirty)
                    try:
                        value = _scalar(self.evaluate(_Parser(formula).parse(), sheet))
                        if isinstance(value, list):
                            raise ValueError('Multi-cell array result is not supported for personnel updates')
                    except _ExcelError as error:
                        value = error
                    dirty = self.frames.pop()
                    if not dirty and cached is not None and cached.value is not None:
                        value = _ExcelError(cached.value) if cached.data_type == 'e' else cached.value
                    before = cached.value if cached is not None else None
                    current = str(value) if isinstance(value, _ExcelError) else value
                    changed = current != before
            self.memo[key] = value, changed
            self.active.remove(key)
        value, changed = self.memo[key]
        if self.frames:
            self.frames[-1] |= changed
        if isinstance(value, _ExcelError):
            raise value
        return value

    def reference(self, address: str, sheet: str) -> _Ref:
        qualifier, separator, ref = address.rpartition('!')
        if separator:
            sheet = qualifier.strip("'").replace("''", "'")
        else:
            ref = address
        if sheet not in self.output:
            raise _ExcelError('#REF!')
        try:
            c1, r1, c2, r2 = range_boundaries(ref)
        except ValueError as error:
            raise ValueError('Unsupported named reference in payroll template') from error
        return _Ref(self, sheet, c1 or 1, r1 or 1,
                    c2 or self.output[sheet].max_column, min(r2 or self.output[sheet].max_row, self.output[sheet].max_row))

    def evaluate(self, node: Any, sheet: str) -> Any:
        kind = node[0]
        if kind == 'literal':
            _, value, subtype = node
            if subtype == 'NUMBER':
                return float(value)
            if subtype == 'TEXT':
                return value[1:-1].replace('""', '"')
            if subtype == 'LOGICAL':
                return value.upper() == 'TRUE'
            if subtype == 'ERROR':
                raise _ExcelError(value)
            if subtype == 'EMPTY':
                return None
            return self.reference(value, sheet)
        if kind == 'unary':
            value = _number(self.evaluate(node[2], sheet))
            return -value if node[1] == '-' else value
        if kind == 'binary':
            _, op, left, right = node
            a, b = _scalar(self.evaluate(left, sheet)), _scalar(self.evaluate(right, sheet))
            def apply(x: Any, y: Any) -> Any:
                if isinstance(x, list) or isinstance(y, list):
                    return [apply(x[i] if isinstance(x, list) else x, y[i] if isinstance(y, list) else y)
                            for i in range(len(x) if isinstance(x, list) else len(y))]
                if op == '&':
                    return str(x or '') + str(y or '')
                if op in {'=', '<>'}:
                    eq = _equal(x, y) or x in (None, '') and y in (None, '')
                    return eq if op == '=' else not eq
                try:
                    return {'+': operator.add, '-': operator.sub, '*': operator.mul, '/': operator.truediv,
                            '^': operator.pow, '<': operator.lt, '>': operator.gt,
                            '<=': operator.le, '>=': operator.ge}[op](_number(x), _number(y))
                except ZeroDivisionError as error:
                    raise _ExcelError('#DIV/0!') from error
            return apply(a, b)
        name, nodes = node[1:]
        def ev(index: int) -> Any:
            return self.evaluate(nodes[index], sheet)
        if name == 'IF':
            return ev(1) if bool(_scalar(ev(0))) else ev(2) if len(nodes) > 2 else False
        if name == 'IFERROR':
            try:
                return _scalar(ev(0))
            except _ExcelError:
                return ev(1)
        if name == 'VLOOKUP':
            lookup, area, column = _scalar(ev(0)), ev(1), int(_number(ev(2))) - 1
            if not isinstance(area, _Ref) or not 0 <= column <= area.last_col - area.first_col:
                raise _ExcelError('#REF!')
            approximate = len(nodes) < 4 or bool(_scalar(ev(3)))
            candidate = None
            for row in range(area.last_row - area.first_row + 1):
                key = area.at(row)
                if _equal(lookup, key):
                    found = area.at(row, column)
                    return 0 if found is None else found
                if approximate and isinstance(key, type(lookup)) and key <= lookup:
                    candidate = row
            if candidate is not None:
                return area.at(candidate, column)
            raise _ExcelError('#N/A')
        if name in {'SUMIF', 'COUNTIF'}:
            criteria_range, criterion = ev(0), _scalar(ev(1))
            sum_range = ev(2) if len(nodes) > 2 else criteria_range
            result = 0
            for row in range(criteria_range.last_row - criteria_range.first_row + 1):
                if _criterion(criteria_range.at(row), criterion):
                    result += 1 if name == 'COUNTIF' else _number(sum_range.at(row))
            return result
        if name == 'ROUND':
            value, digits = _number(ev(0)), int(_number(ev(1)))
            return float(Decimal(str(value)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))
        if name == 'SUBTOTAL':
            function = int(_number(ev(0)))
            name = {3: 'COUNTA', 9: 'SUM', 103: 'COUNTA', 109: 'SUM'}.get(function)
            if name is None:
                raise ValueError('Unsupported subtotal function')
            nodes = nodes[1:]
        if name in {'ARRAY', 'SUM', 'MAX', 'COUNT', 'COUNTA'}:
            values = []
            for child in nodes:
                try:
                    value = self.evaluate(child, sheet)
                    if isinstance(value, _Ref):
                        for r in range(value.last_row - value.first_row + 1):
                            for c in range(value.last_col - value.first_col + 1):
                                try:
                                    values.append(value.at(r, c))
                                except _ExcelError:
                                    if name not in {'COUNT', 'COUNTA'}:
                                        raise
                                    if name == 'COUNTA':
                                        values.append('#ERROR')
                    elif isinstance(value, list):
                        values.extend(value)
                    else:
                        values.append(value)
                except _ExcelError:
                    if name != 'COUNT':
                        raise
            if name == 'ARRAY':
                return values
            if name == 'COUNTA':
                return sum(v is not None for v in values)
            numbers = [_number(v) for v in values if isinstance(v, (int, float, date)) and not isinstance(v, bool)]
            return len(numbers) if name == 'COUNT' else max(numbers, default=0) if name == 'MAX' else sum(numbers)
        raise ValueError(f'Unsupported payroll formula function: {name}')


def refresh_personnel_formula_caches(master_path: Path, output_path: Path, insertions: list[dict]) -> dict[str, Any]:
    master = load_workbook(master_path, data_only=False)
    cached = load_workbook(master_path, data_only=True)
    output = load_workbook(output_path, data_only=False)
    try:
        engine = _Engine(master, cached, output, insertions)
        formulas = {}
        new_errors = []
        for sheet in output:
            values = {}
            for row in sheet:
                for cell in row:
                    if cell.data_type != 'f':
                        continue
                    try:
                        value = engine.cell(sheet.title, cell.row, cell.column)
                    except _ExcelError as error:
                        value = error
                        origin = engine.origin(sheet.title, cell.row)
                        original = cached[sheet.title].cell(origin, cell.column) if origin else None
                        if original is None or original.data_type != 'e' or original.value != str(error):
                            new_errors.append(f'{sheet.title}!{cell.coordinate}: {error}')
                    if isinstance(value, (int, float)) and not math.isfinite(value):
                        raise ValueError('Non-finite payroll formula result')
                    values[cell.coordinate] = value
            formulas[sheet.title] = values
        if new_errors:
            raise ValueError('人员合并产生新的公式错误：' + '；'.join(new_errors[:5]))
        _write_caches(output_path, output.sheetnames, formulas)
        return {'status': 'completed', 'engine': 'incremental_personnel_calculation',
                'formula_count': sum(map(len, formulas.values())), 'new_formula_error_count': 0}
    finally:
        master.close()
        cached.close()
        output.close()


def _write_caches(path: Path, sheet_names: list[str], formulas: dict) -> None:
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    temporary = path.with_suffix('.calculation.tmp')
    with ZipFile(path) as source, ZipFile(temporary, 'w') as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            match = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', entry.filename)
            if match:
                values = formulas[sheet_names[int(match[1]) - 1]]
                root = ET.fromstring(payload)
                for cell in root.iter(ns + 'c'):
                    if cell.attrib['r'] not in values:
                        continue
                    value = values[cell.attrib['r']]
                    cached = cell.find(ns + 'v')
                    if cached is None:
                        cached = ET.SubElement(cell, ns + 'v')
                    if isinstance(value, _ExcelError):
                        cell.attrib['t'], cached.text = 'e', str(value)
                    elif value is None or isinstance(value, str):
                        cell.attrib['t'], cached.text = 'str', value or ''
                    else:
                        cell.attrib['t'] = 'b' if isinstance(value, bool) else 'n'
                        numeric = to_excel(value) if isinstance(value, (date, datetime)) else value
                        cached.text = str(int(numeric)) if isinstance(numeric, bool) else str(numeric)
                payload = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(entry, payload)
    temporary.replace(path)
