"""Workbook-wide, content-based sheet mapping for arbitrary Excel templates."""

from __future__ import annotations

from collections import defaultdict
from contextlib import ExitStack
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook

from core._new_hire_sync import _apply_hire_notices


_KEY_PRIORITY = ("工号", "身份证号", "单据号", "项目编号", "订单号")
_TOPIC_NEUTRAL_HEADERS = {
    *_KEY_PRIORITY,
    "金额", "含税金额", "不含税金额", "税额", "日期", "备注", "状态",
    "币种", "数量", "单价", "序号", "编号", "名称", "说明",
}
_HEADER_ALIASES = {
    "员工工号": "工号", "人员工号": "工号", "职工号": "工号", "新工号": "工号",
    "员工编号": "工号", "人员编号": "工号", "工牌号": "工号",
    "人员姓名": "姓名", "员工姓名": "姓名", "姓名": "姓名",
    "身份证号码": "身份证号", "证件号码": "身份证号",
    "变更后部门": "部门", "所属部门": "部门", "部门名称": "部门",
    "职位": "岗位", "职务": "岗位", "岗位名称": "岗位",
    "单据编号": "单据号", "凭证号": "单据号",
    "费用项目": "项目名称", "项目": "项目名称",
}
_PREFIXES = ("本月", "当月", "最新", "调整后", "变更后", "更新后", "员工", "人员")
_HEADER_ROW_SCAN_LIMIT = 30
_AUTO_MATCH_SCORE = 0.78
_MIN_MATCH_GAP = 0.12


@dataclass(frozen=True)
class SheetProfile:
    title: str
    header_row: int
    headers: dict[int, str]
    data_rows: tuple[int, ...]


@dataclass(frozen=True)
class SheetMatch:
    source: SheetProfile
    target: SheetProfile | None
    column_map: dict[int, int]
    score: float
    candidates: tuple[str, ...]
    reason: str | None = None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normal_header(value: Any) -> str:
    text = re.sub(r"[\s_\-—:：()（）\[\]【】]", "", _text(value)).lower()
    for prefix in _PREFIXES:
        if text.startswith(prefix) and len(text) > len(prefix):
            text = text[len(prefix):]
    return _HEADER_ALIASES.get(text, text)


def _header_score(source: str, target: str) -> float:
    if source == target:
        return 1.0
    if source in target or target in source:
        return 0.86
    return SequenceMatcher(a=source, b=target).ratio()


def _sheet_profile(sheet: Any) -> SheetProfile | None:
    best: tuple[int, dict[int, str], float] | None = None
    for row_number in range(1, min(sheet.max_row, _HEADER_ROW_SCAN_LIMIT) + 1):
        headers = {
            column: _normal_header(sheet.cell(row_number, column).value)
            for column in range(1, sheet.max_column + 1)
            if _normal_header(sheet.cell(row_number, column).value)
        }
        if not headers or len(set(headers.values())) != len(headers):
            continue
        following_values = sum(
            sheet.cell(next_row, column).value not in (None, "")
            for next_row in range(row_number + 1, min(sheet.max_row, row_number + 4) + 1)
            for column in headers
        )
        score = len(headers) * 10 + following_values
        if best is None or score > best[2]:
            best = (row_number, headers, score)
    if best is None:
        return None
    header_row, headers, _ = best
    data_rows = tuple(
        row_number
        for row_number in range(header_row + 1, sheet.max_row + 1)
        if any(sheet.cell(row_number, column).value not in (None, "") for column in headers)
    )
    return SheetProfile(sheet.title, header_row, headers, data_rows)


def _column_map(source: SheetProfile, target: SheetProfile) -> tuple[dict[int, int], float]:
    mapped: dict[int, int] = {}
    used_target_columns: set[int] = set()
    scores: list[float] = []
    for source_column, source_header in source.headers.items():
        ranked = sorted(
            (
                (_header_score(source_header, target_header), target_column)
                for target_column, target_header in target.headers.items()
                if target_column not in used_target_columns
            ),
            reverse=True,
        )
        if ranked and ranked[0][0] >= _AUTO_MATCH_SCORE:
            score, target_column = ranked[0]
            mapped[source_column] = target_column
            used_target_columns.add(target_column)
            scores.append(score)
    if not mapped:
        return {}, 0.0
    source_coverage = len(mapped) / len(source.headers)
    target_coverage = len(mapped) / len(target.headers)
    quality = sum(scores) / len(scores)
    return mapped, round(quality * 0.6 + source_coverage * 0.3 + target_coverage * 0.1, 4)


def _has_topic_evidence(source: SheetProfile, target: SheetProfile, column_map: dict[int, int]) -> bool:
    """Require a shared, non-generic field before treating two Sheets as one topic."""
    return any(
        source.headers[source_column] == target.headers[target_column]
        and source.headers[source_column] not in _TOPIC_NEUTRAL_HEADERS
        for source_column, target_column in column_map.items()
    )


def _match_sheets(master_workbook: Any, source_workbook: Any) -> list[SheetMatch]:
    target_profiles = [profile for sheet in master_workbook.worksheets if (profile := _sheet_profile(sheet))]
    matches: list[SheetMatch] = []
    for source_sheet in source_workbook.worksheets:
        source = _sheet_profile(source_sheet)
        if source is None or not source.data_rows:
            continue
        ranked: list[tuple[float, SheetProfile, dict[int, int]]] = []
        for target in target_profiles:
            column_map, score = _column_map(source, target)
            if column_map:
                ranked.append((score, target, column_map))
        ranked.sort(key=lambda item: item[0], reverse=True)
        candidates = tuple(item[1].title for item in ranked[:3])
        if not ranked or ranked[0][0] < _AUTO_MATCH_SCORE:
            matches.append(SheetMatch(source, None, {}, 0.0, candidates))
            continue
        best_score, target, column_map = ranked[0]
        is_unique = len(ranked) == 1 or best_score - ranked[1][0] >= _MIN_MATCH_GAP
        if not is_unique:
            matches.append(SheetMatch(source, None, {}, best_score, candidates))
        elif not _has_topic_evidence(source, target, column_map):
            matches.append(SheetMatch(
                source, None, {}, best_score, candidates,
                reason="insufficient_topic_evidence",
            ))
        else:
            matches.append(SheetMatch(source, target, column_map, best_score, candidates))
    return matches


def _match_sheets_with_overrides(
    master_workbook: Any,
    source_workbook: Any,
    source_filename: str,
    sheet_overrides: dict[tuple[str, str], str] | None,
) -> list[SheetMatch]:
    """Apply an explicit user choice before falling back to safe automatic matching."""
    automatic_matches = _match_sheets(master_workbook, source_workbook)
    if not sheet_overrides:
        return automatic_matches

    target_profiles = {
        profile.title: profile
        for sheet in master_workbook.worksheets
        if (profile := _sheet_profile(sheet))
    }
    overridden: list[SheetMatch] = []
    for match in automatic_matches:
        target_name = sheet_overrides.get((source_filename, match.source.title))
        target = target_profiles.get(target_name or "")
        if target is None:
            overridden.append(match)
            continue
        column_map, score = _column_map(match.source, target)
        # The user has confirmed the business topic; the remaining key and
        # formula guards below still decide whether any value can be written.
        overridden.append(SheetMatch(
            match.source,
            target,
            column_map,
            score,
            (target.title,),
        ))
    return overridden


def _issue(issue_type: str, source_file: Path, source_sheet: str, message: str, **extra: Any) -> dict[str, Any]:
    return {
        "issue_type": issue_type,
        "message": message,
        "source_files": [source_file.name],
        "source_sheets": [source_sheet],
        **extra,
    }


def _key_columns(source: SheetProfile, target: SheetProfile, column_map: dict[int, int]) -> tuple[int, int] | None:
    for key_name in _KEY_PRIORITY:
        for source_column, target_column in column_map.items():
            if source.headers[source_column] == key_name and target.headers[target_column] == key_name:
                return source_column, target_column
    return None


def _row_index(sheet: Any, profile: SheetProfile, key_column: int) -> dict[str, list[int]]:
    index: dict[str, list[int]] = defaultdict(list)
    for row_number in profile.data_rows:
        value = _text(sheet.cell(row_number, key_column).value)
        if value:
            index[value].append(row_number)
    return index


def _is_formula_cell(cell: Any) -> bool:
    """Recognize scalar, array, and other formula cells exposed by openpyxl."""
    value = cell.value
    return cell.data_type == "f" or (isinstance(value, str) and value.startswith("="))


def _normalized_path(value: str | Path) -> str:
    return str(Path(value).resolve()).casefold()


def _salary_month_markers(salary_month: str | None) -> tuple[str, ...]:
    """Return the unambiguous labels used for the configured payroll month."""
    if not salary_month:
        return ()
    match = re.search(r"(20\d{2})[.\-/年\s]*(0?[1-9]|1[0-2])", str(salary_month))
    if not match:
        return ()
    year, month = int(match.group(1)), int(match.group(2))
    return (
        f"{year}.{month:02d}", f"{year}-{month:02d}", f"{year}/{month:02d}",
        f"{year}年{month}月", f"{month}月",
    )


def _current_month_attendance_or_bonus_candidate(
    candidates: list[tuple[Any, dict[str, Any]]], salary_month: str | None,
) -> tuple[Any, dict[str, Any]] | None:
    """Choose one agreed current-month attendance/bonus value, never a guess."""
    markers = _salary_month_markers(salary_month)
    if not markers:
        return None
    current_month_candidates = [
        candidate
        for candidate in candidates
        if any(keyword in f"{candidate[1]['source_file']} {candidate[1]['source_sheet']}" for keyword in ("考勤", "奖金"))
        and any(marker in f"{candidate[1]['source_file']} {candidate[1]['source_sheet']}" for marker in markers)
    ]
    if len({_text(value) for value, _ in current_month_candidates}) != 1:
        return None
    return current_month_candidates[0]


def _configured_field_action(
    policy: dict[str, Any] | None,
    source_sheet: str,
    target_field: str,
    salary_month: str | None,
) -> str | None:
    if not isinstance(policy, dict):
        return None
    matches: list[tuple[int, str]] = []
    for rule in policy.get("source_rules") or []:
        if not isinstance(rule, dict):
            continue
        if rule.get("source_sheet") and str(rule["source_sheet"]).strip() != source_sheet:
            continue
        if rule.get("target_field") and str(rule["target_field"]).strip() != target_field:
            continue
        if rule.get("month") and salary_month and str(rule["month"]).strip() != salary_month:
            continue
        action = str(rule.get("action") or "").strip()
        if action in {"allow", "review", "ignore"}:
            specificity = sum(bool(rule.get(key)) for key in ("company", "month", "source_sheet", "target_field"))
            matches.append((specificity, action))
    return max(matches, key=lambda entry: entry[0])[1] if matches else None


def _policy_confirms_unique_target(
    policy: dict[str, Any] | None,
    source_sheet: SheetProfile,
    target: SheetProfile | None,
) -> bool:
    """Allow a unique low-topic match only when policy names a target field.

    Generic sheets that share only neutral headers must stay in review. A
    company-specific source rule can provide the missing business context, but
    it must still point at a field present in the sole candidate target.
    """
    if not isinstance(policy, dict) or target is None:
        return False
    target_fields = set(target.headers.values())
    for rule in policy.get("source_rules") or []:
        if not isinstance(rule, dict):
            continue
        source_name = str(rule.get("source_sheet") or "").strip()
        target_field = str(rule.get("target_field") or "").strip()
        action = str(rule.get("action") or "").strip()
        if source_name == source_sheet.title and target_field in target_fields and action in {"allow", "review", "ignore"}:
            return True
    return False


def apply_semantic_sheet_updates(
    master_path: str | Path,
    update_paths: Iterable[str | Path],
    output_path: str | Path,
    salary_month: str | None = None,
    source_names: dict[str, str] | None = None,
    sheet_overrides: dict[tuple[str, str], str] | None = None,
    matching_policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Apply uniquely understood source Sheet changes to an existing master workbook.

    The function intentionally never creates a new Sheet. Unmatched or ambiguous sources
    are returned as issues for the separate manual-review workflow.
    """
    master_path = Path(master_path)
    output_path = Path(output_path)
    master = load_workbook(master_path, data_only=False)
    issues: list[dict[str, Any]] = []
    proposals: dict[tuple[str, int, int], list[tuple[Any, dict[str, Any]]]] = defaultdict(list)
    matches_out: list[dict[str, Any]] = []
    updates: list[dict[str, Any]] = []
    source_display_names = {
        _normalized_path(path): name for path, name in (source_names or {}).items()
    }
    sources_stack = ExitStack()
    try:
        sources = []
        for update_path_raw in update_paths:
            update_path = Path(update_path_raw)
            filename = source_display_names.get(_normalized_path(update_path), update_path.name)
            source_book = load_workbook(update_path, read_only=False, data_only=False)
            sources_stack.callback(source_book.close)
            sources.append((update_path, filename, source_book))
        hires = _apply_hire_notices(
            master, sources,
            lambda source_sheet, field: _configured_field_action(matching_policy, source_sheet, field, salary_month),
            salary_month,
        )
        issues.extend(hires["issues"])
        updates.extend(hires["updates"])
        matches_out.extend(hires["matches"])
        for update_path, filename, source_book in sources:
            display_path = Path(filename)
            try:
                for source_sheet in source_book.worksheets:
                    if (str(update_path.resolve()), source_sheet.title) in hires["consumed"]:
                        continue
                    source_profile = _sheet_profile(source_sheet)
                    if source_profile is None:
                        issues.append(_issue(
                            "unprofiled_source_sheet", display_path, source_sheet.title,
                            "来源 Sheet 无法识别表头和数据区域，未自动写入。",
                        ))
                    elif not source_profile.data_rows:
                        issues.append(_issue(
                            "empty_source_sheet", display_path, source_sheet.title,
                            "来源 Sheet 没有可处理的数据记录，未自动写入。",
                        ))
                for match in _match_sheets_with_overrides(
                    master,
                    source_book,
                    display_path.name,
                    sheet_overrides,
                ):
                    if (str(update_path.resolve()), match.source.title) in hires["consumed"]:
                        continue
                    if (
                        match.target is None
                        and match.reason == "insufficient_topic_evidence"
                        and len(match.candidates) == 1
                    ):
                        candidate_target = next(
                            (profile for sheet in master.worksheets
                             if (profile := _sheet_profile(sheet))
                             and profile.title == match.candidates[0]),
                            None,
                        )
                        if _policy_confirms_unique_target(matching_policy, match.source, candidate_target):
                            column_map, score = _column_map(match.source, candidate_target)
                            match = SheetMatch(
                                match.source,
                                candidate_target,
                                column_map,
                                score,
                                match.candidates,
                            )
                    if match.target is None:
                        issue_type = match.reason or (
                            "ambiguous_sheet" if len(match.candidates) > 1 else "unmatched_sheet"
                        )
                        issues.append(_issue(
                            issue_type, display_path, match.source.title,
                            "来源 Sheet 缺少足以唯一确认主题的证据，未自动写入。"
                            if issue_type == "insufficient_topic_evidence"
                            else "来源 Sheet 无法唯一映射到总表主题，未自动写入。",
                            candidate_target_sheets=list(match.candidates), confidence=match.score,
                        ))
                        continue
                    source_sheet = source_book[match.source.title]
                    target_sheet = master[match.target.title]
                    key_columns = _key_columns(match.source, match.target, match.column_map)
                    if key_columns is None:
                        issues.append(_issue(
                            "missing_business_key", display_path, match.source.title,
                            "已识别同主题，但无法推断可唯一定位记录的业务主键，未自动写入。",
                            target_sheet=match.target.title, confidence=match.score,
                        ))
                        continue
                    source_key_column, target_key_column = key_columns
                    target_rows = _row_index(target_sheet, match.target, target_key_column)
                    source_rows_by_key: dict[str, list[int]] = defaultdict(list)
                    for source_row in match.source.data_rows:
                        source_key_cell = source_sheet.cell(source_row, source_key_column)
                        if _is_formula_cell(source_key_cell):
                            issues.append(_issue(
                                "source_formula_conflict", display_path, match.source.title,
                                "来源业务主键是公式，未自动使用或写入。",
                                source_row=source_row, target_sheet=match.target.title,
                                source_cell=source_key_cell.coordinate,
                            ))
                            continue
                        key = _text(source_key_cell.value)
                        if key:
                            source_rows_by_key[key].append(source_row)
                    duplicate_source_keys = {
                        key: rows
                        for key, rows in source_rows_by_key.items()
                        if len(rows) > 1
                    }
                    matches_out.append({
                        "source_file": display_path.name, "source_sheet": match.source.title,
                        "target_sheet": match.target.title, "confidence": match.score,
                    })
                    for source_row in match.source.data_rows:
                        key = _text(source_sheet.cell(source_row, source_key_column).value)
                        if not key:
                            issues.append(_issue(
                                "missing_business_key", display_path, match.source.title,
                                "来源记录缺少业务主键，未自动写入。",
                                source_row=source_row, target_sheet=match.target.title,
                            ))
                            continue
                        duplicate_rows = duplicate_source_keys.get(key)
                        if duplicate_rows:
                            if source_row == duplicate_rows[0]:
                                issues.append(_issue(
                                    "duplicate_source_record", display_path, match.source.title,
                                    "来源 Sheet 中同一业务主键出现多次，未自动覆盖。",
                                    source_rows=duplicate_rows, target_sheet=match.target.title,
                                    business_key=key,
                                ))
                            continue
                        rows = target_rows.get(key, [])
                        if len(rows) > 1:
                            issues.append(_issue(
                                "ambiguous_record", display_path, match.source.title,
                                "总表中存在多条相同业务主键，未自动覆盖。",
                                source_row=source_row, target_sheet=match.target.title, business_key=key,
                            ))
                            continue
                        if not rows:
                            issues.append(_issue(
                                "unknown_record", display_path, match.source.title,
                                "来源记录在总表中不存在。通用模式无法可靠识别插入位置，未自动新增记录。",
                                source_row=source_row, target_sheet=match.target.title, business_key=key,
                            ))
                            continue
                        target_row = rows[0]
                        for source_column, target_column in match.column_map.items():
                            source_cell = source_sheet.cell(source_row, source_column)
                            value = source_cell.value
                            if value in (None, ""):
                                continue
                            target_cell = target_sheet.cell(target_row, target_column)
                            target_field = match.target.headers[target_column]
                            configured_action = _configured_field_action(
                                matching_policy, match.source.title, target_field, salary_month
                            )
                            if configured_action in {"ignore", "review"}:
                                issue_type = "configured_field_ignored" if configured_action == "ignore" else "configured_field_review"
                                issues.append(_issue(
                                    issue_type, display_path, match.source.title,
                                    f"匹配策略已要求{('忽略' if configured_action == 'ignore' else '人工确认')}字段“{target_field}”，未自动写入。",
                                    source_row=source_row,
                                    target_sheet=match.target.title,
                                    target_cell=target_cell.coordinate,
                                    target_field=target_field,
                                    candidate_values=[value],
                                ))
                                continue
                            if _is_formula_cell(source_cell):
                                issues.append(_issue(
                                    "source_formula_conflict", display_path, match.source.title,
                                    "来源单元格含公式，未自动写入总表。",
                                    source_row=source_row, target_sheet=match.target.title,
                                    source_cell=source_cell.coordinate,
                                ))
                                continue
                            if _is_formula_cell(target_cell):
                                issues.append(_issue(
                                    "formula_target_conflict", display_path, match.source.title,
                                    "目标单元格含公式，未使用来源值覆盖。",
                                    source_row=source_row, target_sheet=match.target.title,
                                    target_cell=target_cell.coordinate,
                                ))
                                continue
                            if _text(target_cell.value) != _text(value):
                                proposals[(match.target.title, target_row, target_column)].append((value, {
                                    "source_file": display_path.name, "source_sheet": match.source.title,
                                    "source_row": source_row,
                                }))
            finally:
                source_book.close()

        auto_update_count = len(updates)
        for (sheet_name, row, column), candidates in proposals.items():
            values = {_text(value) for value, _ in candidates}
            preferred_candidate = _current_month_attendance_or_bonus_candidate(candidates, salary_month)
            if len(values) != 1 and preferred_candidate is None:
                issues.append({
                    "issue_type": "conflicting_value",
                    "message": "多个来源为同一总表单元格提供了不同值，未自动覆盖。",
                    "target_sheet": sheet_name,
                    "target_cell": master[sheet_name].cell(row, column).coordinate,
                    "candidate_values": [value for value, _ in candidates],
                    "source_files": sorted({meta["source_file"] for _, meta in candidates}),
                    "source_sheets": sorted({meta["source_sheet"] for _, meta in candidates}),
                    "source_rows": [meta["source_row"] for _, meta in candidates],
                })
                continue
            target_cell = master[sheet_name].cell(row, column)
            new_value, source = preferred_candidate or candidates[0]
            updates.append({
                "target_sheet": sheet_name,
                "target_cell": target_cell.coordinate,
                "old_value": target_cell.value,
                "new_value": new_value,
                "source_file": source["source_file"],
                "source_sheet": source["source_sheet"],
                "source_row": source["source_row"],
            })
            target_cell.value = new_value
            auto_update_count += 1
        output_path.parent.mkdir(parents=True, exist_ok=True)
        master.save(output_path)
    finally:
        sources_stack.close()
        master.close()
    return {
        "output_path": str(output_path), "auto_update_count": auto_update_count,
        "issues": issues, "matches": matches_out, "updates": updates,
        "personnel_coverage": hires["personnel_coverage"],
        "structured_hire": hires["structured_hire"],
        "row_insertions": getattr(master, '_personnel_row_insertions', []),
    }
