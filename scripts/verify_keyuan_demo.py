from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import openpyxl


def equivalent(left: object, right: object) -> bool:
    if left in (None, "") and right in (None, ""):
        return True
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-9, abs_tol=1e-6)
    return left == right


def formula_signature(value: object) -> tuple[object, ...] | None:
    if isinstance(value, str) and value.startswith("="):
        return ("formula", value)
    if hasattr(value, "text") and hasattr(value, "ref"):
        return ("array_formula", value.text, value.ref)
    return None


def main() -> int:
    generated = Path(sys.argv[1])
    reference = Path(sys.argv[2])
    report = Path(sys.argv[3])
    actual = openpyxl.load_workbook(generated, data_only=False, read_only=False)
    expected = openpyxl.load_workbook(reference, data_only=False, read_only=False)
    actual_values = openpyxl.load_workbook(generated, data_only=True, read_only=False)
    expected_values = openpyxl.load_workbook(reference, data_only=True, read_only=False)
    checks: dict[str, object] = {}
    try:
        checks["sheet_names"] = actual.sheetnames == expected.sheetnames
        checks["payroll_columns"] = actual["工资核算"].max_column == 99
        checks["payroll_first_header"] = actual["工资核算"]["A2"].value == "班制"
        checks["oa_month"] = actual["OA请款及审批"]["A8"].value == "7月"
        checks["oa_history"] = [actual["OA请款及审批"].cell(row, 1).value for row in range(9, 15)] == ["6月", "5月", "4月", "3月", "2月", "1月"]
        checks["heat_days"] = actual["防暑降温费"]["K1"].value == 23
        checks["adjustment"] = [actual["其他调差累计"].cell(2, col).value for col in range(1, 4)] == ["李楠", 775.98, "12小时工作日加班、12小时法假加班，一共775.98元"]
        checks["duty_counts"] = [actual["配送员值班费"].cell(row, 4).value for row in range(2, 7)] == [0, 0, 4, 4, 0]
        payroll = actual["工资核算"]
        names = {payroll.cell(row, 3).value: row for row in range(4, 44)}
        checks["people_count"] = len([name for name in names if name]) == 40
        bonus_examples = {"郭利娜": 4680.18, "李楠": 1500, "闫美华": 1340, "程春娜": 2700}
        checks["bonuses"] = {
            name: payroll.cell(names[name], 31).value == value
            for name, value in bonus_examples.items()
        }
        checks["au_av_formulas"] = all(
            payroll.cell(row, col).value == expected["工资核算"].cell(row, col).value
            for row in range(4, 44)
            for col in (47, 48)
        )
        checks["payroll_formula_integrity"] = all(
            formula_signature(actual["工资核算"].cell(row, col).value)
            == formula_signature(expected["工资核算"].cell(row, col).value)
            for row in range(1, 53)
            for col in range(1, 94)
            if formula_signature(expected["工资核算"].cell(row, col).value) is not None
        )
        formula_errors: list[str] = []
        for sheet in actual.worksheets:
            for row in sheet.iter_rows(max_col=min(sheet.max_column, 200)):
                for cell in row:
                    if isinstance(cell.value, str) and "#REF!" in cell.value:
                        formula_errors.append(f"{sheet.title}!{cell.coordinate}")
        checks["formula_ref_errors"] = formula_errors
        actual_errors = sum(
            1
            for sheet in actual_values.worksheets
            for row in sheet.iter_rows(max_col=min(sheet.max_column, 200))
            for cell in row
            if isinstance(cell.value, str) and cell.value.upper() in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
        )
        expected_errors = sum(
            1
            for sheet in expected_values.worksheets
            for row in sheet.iter_rows(max_col=min(sheet.max_column, 200))
            for cell in row
            if isinstance(cell.value, str) and cell.value.upper() in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
        )
        checks["no_additional_formula_errors"] = actual_errors <= expected_errors
        ledger = actual["台账"]
        expected_ledger = expected["台账"]
        checks["social_ledger"] = all(
            ledger.cell(row, col).value == expected_ledger.cell(row, col).value
            for row in range(3, 44)
            for col in range(4, 14)
        )
        tax = actual["个税导出核对"]
        expected_tax = expected["个税导出核对"]
        checks["tax_groups"] = all(
            (tax.cell(row, col).value in (None, "") and expected_tax.cell(row, col).value in (None, ""))
            or tax.cell(row, col).value == expected_tax.cell(row, col).value
            for row in list(range(2, 41)) + list(range(42, 48))
            for col in range(1, 40)
        )
        checks["tax_subtotal_formulas"] = {
            "row41": [tax.cell(41, col).value for col in range(1, 40)] == [expected_tax.cell(41, col).value for col in range(1, 40)],
            "row48": [tax.cell(48, col).value for col in range(1, 40)] == [expected_tax.cell(48, col).value for col in range(1, 40)],
        }
        passed = all(
            value is True or value == [] or (isinstance(value, dict) and all(value.values()))
            for value in checks.values()
        )
        payload = {"status": "passed" if passed else "failed", "checks": checks}
        report.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0 if passed else 1
    finally:
        actual.close()
        expected.close()
        actual_values.close()
        expected_values.close()


if __name__ == "__main__":
    raise SystemExit(main())
