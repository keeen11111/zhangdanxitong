from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


def normalize(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        return round(value, 8)
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if hasattr(value, "text") and hasattr(value, "ref"):
        return {"text": value.text, "ref": value.ref}
    return str(value)


def equal(left: Any, right: Any) -> bool:
    left = normalize(left)
    right = normalize(right)
    if left in (None, "") and right in (None, ""):
        return True
    if (
        isinstance(left, (int, float))
        and not isinstance(left, bool)
        and isinstance(right, (int, float))
        and not isinstance(right, bool)
    ):
        return math.isclose(float(left), float(right), rel_tol=1e-9, abs_tol=1e-6)
    return left == right


def color_signature(color: Any) -> tuple[Any, ...] | None:
    if color is None:
        return None
    return (color.type, color.rgb, color.indexed, color.auto, color.theme, color.tint)


def style_signature(cell: Any) -> tuple[Any, ...]:
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    protection = cell.protection
    return (
        font.name,
        font.sz,
        font.b,
        font.i,
        font.u,
        font.strike,
        color_signature(font.color),
        fill.fill_type,
        color_signature(fill.fgColor),
        color_signature(fill.bgColor),
        tuple(
            (side.style, color_signature(side.color))
            for side in (border.left, border.right, border.top, border.bottom, border.diagonal)
        ),
        alignment.horizontal,
        alignment.vertical,
        alignment.text_rotation,
        alignment.wrap_text,
        alignment.shrink_to_fit,
        alignment.indent,
        cell.number_format,
        protection.locked,
        protection.hidden,
    )


def dimension_signature(dimension: Any) -> tuple[Any, ...]:
    if dimension is None:
        return (False, 0, False, None, None, False)
    return (
        dimension.hidden,
        dimension.outlineLevel,
        dimension.collapsed,
        getattr(dimension, "width", None),
        getattr(dimension, "height", None),
        getattr(dimension, "bestFit", None),
    )


def formula_value(cell: Any) -> Any:
    if cell.data_type == "f" or hasattr(cell.value, "text"):
        return normalize(cell.value)
    return None


def append_sample(samples: list[dict[str, Any]], limit: int, **item: Any) -> None:
    if len(samples) < limit:
        samples.append(item)


def compare(actual_path: Path, expected_path: Path, sample_limit: int) -> dict[str, Any]:
    actual_formula = openpyxl.load_workbook(actual_path, data_only=False, read_only=False, keep_links=True)
    expected_formula = openpyxl.load_workbook(expected_path, data_only=False, read_only=False, keep_links=True)
    actual_values = openpyxl.load_workbook(actual_path, data_only=True, read_only=False, keep_links=True)
    expected_values = openpyxl.load_workbook(expected_path, data_only=True, read_only=False, keep_links=True)
    try:
        sheet_names_equal = actual_formula.sheetnames == expected_formula.sheetnames
        sheet_reports: dict[str, Any] = {}
        samples: dict[str, list[dict[str, Any]]] = {
            key: []
            for key in (
                "formula",
                "constant",
                "cached_value",
                "style",
                "merge",
                "row_dimension",
                "column_dimension",
                "formula_error",
            )
        }

        for sheet_name in sorted(set(actual_formula.sheetnames) | set(expected_formula.sheetnames)):
            if sheet_name not in actual_formula or sheet_name not in expected_formula:
                sheet_reports[sheet_name] = {
                    "missing_in": "actual" if sheet_name not in actual_formula else "expected"
                }
                continue

            actual_sheet = actual_formula[sheet_name]
            expected_sheet = expected_formula[sheet_name]
            actual_value_sheet = actual_values[sheet_name]
            expected_value_sheet = expected_values[sheet_name]
            counts: Counter[str] = Counter()
            coordinates = set(actual_sheet._cells) | set(expected_sheet._cells)

            for row, column in sorted(coordinates):
                actual_cell = actual_sheet._cells.get((row, column)) or actual_sheet.cell(row, column)
                expected_cell = expected_sheet._cells.get((row, column)) or expected_sheet.cell(row, column)
                coordinate = actual_cell.coordinate
                actual_formula_value = formula_value(actual_cell)
                expected_formula_value = formula_value(expected_cell)
                if not equal(actual_formula_value, expected_formula_value):
                    counts["formula"] += 1
                    append_sample(
                        samples["formula"],
                        sample_limit,
                        sheet=sheet_name,
                        cell=coordinate,
                        actual=actual_formula_value,
                        expected=expected_formula_value,
                    )

                actual_constant = None if actual_formula_value is not None else normalize(actual_cell.value)
                expected_constant = None if expected_formula_value is not None else normalize(expected_cell.value)
                if not equal(actual_constant, expected_constant):
                    counts["constant"] += 1
                    append_sample(
                        samples["constant"],
                        sample_limit,
                        sheet=sheet_name,
                        cell=coordinate,
                        actual=actual_constant,
                        expected=expected_constant,
                    )

                actual_cached = actual_value_sheet.cell(row, column).value
                expected_cached = expected_value_sheet.cell(row, column).value
                if not equal(actual_cached, expected_cached):
                    counts["cached_value"] += 1
                    append_sample(
                        samples["cached_value"],
                        sample_limit,
                        sheet=sheet_name,
                        cell=coordinate,
                        actual=normalize(actual_cached),
                        expected=normalize(expected_cached),
                    )

                if style_signature(actual_cell) != style_signature(expected_cell):
                    counts["style"] += 1
                    append_sample(
                        samples["style"],
                        sample_limit,
                        sheet=sheet_name,
                        cell=coordinate,
                    )

                if isinstance(actual_cached, str) and actual_cached.upper() in ERROR_VALUES:
                    counts["actual_formula_error"] += 1
                    append_sample(
                        samples["formula_error"],
                        sample_limit,
                        sheet=sheet_name,
                        cell=coordinate,
                        side="actual",
                        value=actual_cached,
                    )
                if isinstance(expected_cached, str) and expected_cached.upper() in ERROR_VALUES:
                    counts["expected_formula_error"] += 1

            actual_merges = {str(item) for item in actual_sheet.merged_cells.ranges}
            expected_merges = {str(item) for item in expected_sheet.merged_cells.ranges}
            merge_difference = sorted(actual_merges ^ expected_merges)
            counts["merge"] = len(merge_difference)
            for item in merge_difference:
                append_sample(
                    samples["merge"],
                    sample_limit,
                    sheet=sheet_name,
                    range=item,
                    actual=item in actual_merges,
                    expected=item in expected_merges,
                )

            for index in sorted(set(actual_sheet.row_dimensions) | set(expected_sheet.row_dimensions)):
                if dimension_signature(actual_sheet.row_dimensions.get(index)) != dimension_signature(
                    expected_sheet.row_dimensions.get(index)
                ):
                    counts["row_dimension"] += 1
                    append_sample(
                        samples["row_dimension"],
                        sample_limit,
                        sheet=sheet_name,
                        row=index,
                    )

            for index in sorted(set(actual_sheet.column_dimensions) | set(expected_sheet.column_dimensions)):
                if dimension_signature(actual_sheet.column_dimensions.get(index)) != dimension_signature(
                    expected_sheet.column_dimensions.get(index)
                ):
                    counts["column_dimension"] += 1
                    append_sample(
                        samples["column_dimension"],
                        sample_limit,
                        sheet=sheet_name,
                        column=index,
                    )

            sheet_reports[sheet_name] = {
                "actual_dimension": [actual_sheet.max_row, actual_sheet.max_column],
                "expected_dimension": [expected_sheet.max_row, expected_sheet.max_column],
                **counts,
            }

        totals = {
            key: sum(int(report.get(key, 0)) for report in sheet_reports.values())
            for key in (
                "formula",
                "constant",
                "cached_value",
                "style",
                "merge",
                "row_dimension",
                "column_dimension",
                "actual_formula_error",
                "expected_formula_error",
            )
        }
        return {
            "sheet_names_equal": sheet_names_equal,
            "actual_sheet_names": actual_formula.sheetnames,
            "expected_sheet_names": expected_formula.sheetnames,
            "totals": totals,
            "sheets": sheet_reports,
            "samples": samples,
        }
    finally:
        actual_formula.close()
        expected_formula.close()
        actual_values.close()
        expected_values.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("actual", type=Path)
    parser.add_argument("expected", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--sample-limit", type=int, default=100)
    args = parser.parse_args()
    report = compare(args.actual, args.expected, args.sample_limit)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload, encoding="utf-8")
    print(
        json.dumps(
            {
                "sheet_names_equal": report["sheet_names_equal"],
                "totals": report["totals"],
                "sheets_with_differences": {
                    name: data
                    for name, data in report["sheets"].items()
                    if any(
                        int(data.get(key, 0))
                        for key in (
                            "formula",
                            "constant",
                            "cached_value",
                            "style",
                            "merge",
                            "row_dimension",
                            "column_dimension",
                        )
                    )
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
