import openpyxl, sys, json
a=openpyxl.load_workbook(sys.argv[1],data_only=False,read_only=True)
b=openpyxl.load_workbook(sys.argv[2],data_only=False,read_only=True)
out={}
for sn,rows,cols in [('台账',range(3,44),range(4,14)),('个税导出核对',list(range(2,41))+list(range(42,48)),range(1,40))]:
 d=[]
 for r in rows:
  for c in cols:
   av=a[sn].cell(r,c).value; bv=b[sn].cell(r,c).value
   if av!=bv:d.append({'cell':a[sn].cell(r,c).coordinate,'actual':str(av),'expected':str(bv)})
 out[sn]={'count':len(d),'sample':d[:50]}
print(json.dumps(out,ensure_ascii=False,indent=2))
a.close();b.close()
